# SPDX-License-Identifier: AGPL-3.0-or-later
"""End to end: configurable products. Options and choices, rules on BOM lines at any level, constraints, the BOM of
one combination, the check of every combination, saved variants and machines built as a variant."""
import io

import openpyxl


def _bom(c, pn, desc, lines):
    bid = c.call("POST", "/api/boms", {"part_number": pn, "description": desc})["bom_id"]
    c.call("PUT", f"/api/boms/{bid}/lines", {"lines": lines})
    return bid


OPTIONS = [
    {"code": "motor", "name": "Motor power", "choices": [{"code": "5k5", "label": "5.5 kW"}, {"code": "7K5", "label": "7.5 kW"},
                                                        {"code": "11K", "label": "11 kW"}]},
    {"code": "VOLTAGE", "name": "Voltage", "choices": [{"code": "230", "label": "230 V"}, {"code": "415", "label": "415 V"}]},
    {"code": "SEAL", "name": "Seal", "choices": [{"code": "PTFE"}, {"code": "VITON"}]},
]


def test_configurable_product(kosha):
    admin = kosha.client()
    eng = kosha.add_user("cf-eng", "engineer")
    viewer = kosha.add_user("cf-view", "viewer")
    seal = _bom(admin, "CF-200-030", "SEAL KIT", [{"part_number": "CF-200-031", "description": "SEAL PTFE", "qty": 2},
                                                  {"part_number": "CF-200-032", "description": "SEAL VITON", "qty": 2},
                                                  {"part_number": "CF-200-033", "description": "SPRING", "qty": 1}])
    top = _bom(admin, "CF-200-000", "MIXER CONFIGURABLE", [
        {"part_number": "CF-200-010", "description": "FRAME", "qty": 1},
        {"part_number": "CF-200-021", "description": "MOTOR 5.5 KW", "qty": 1},
        {"part_number": "CF-200-022", "description": "MOTOR 7.5 KW", "qty": 1},
        {"part_number": "CF-200-023", "description": "MOTOR 11 KW", "qty": 1},
        {"part_number": "CF-200-030", "description": "SEAL KIT", "qty": 1},
        {"part_number": "CF-200-041", "description": "CABLE 230 V", "qty": 3},
        {"part_number": "CF-200-042", "description": "CABLE 415 V", "qty": 3},
        {"part_number": "CF-200-050", "description": "OLD GUARD", "qty": 1}])
    product = admin.call("GET", f"/api/boms/{top}")["bom"]["item_id"]
    seal_kit = admin.call("GET", f"/api/boms/{seal}")["bom"]["item_id"]
    d = eng.call("GET", f"/api/items/{product}/config?source=in_work")
    ids = {r["part_number"]: r["item_id"] for r in d["structure"]["rows"]}
    assert d["options"] == [] and len(d["structure"]["rows"]) == 11

    # options: at least two choices each, codes once
    eng.call("PUT", f"/api/items/{product}/config/options", {"options": [{"code": "X", "choices": [{"code": "A"}]}]}, expect=400)
    eng.call("PUT", f"/api/items/{product}/config/options", {"options": [OPTIONS[1], OPTIONS[1]]}, expect=400)
    viewer.call("PUT", f"/api/items/{product}/config/options", {"options": OPTIONS}, expect=403)
    d = eng.call("PUT", f"/api/items/{product}/config/options", {"options": OPTIONS})
    assert [o["code"] for o in d["options"]] == ["MOTOR", "VOLTAGE", "SEAL"]
    assert [c["code"] for c in d["options"][0]["choices"]] == ["5K5", "7K5", "11K"]

    # rules on lines, at the top and inside the seal kit; a rule that can't be read says why
    rule = lambda parent, pn, text: {"parent_item_id": parent, "child_item_id": ids[pn], "rule": text}  # noqa: E731
    r = eng.call("PUT", f"/api/items/{product}/config/rules", {"rules": [rule(product, "CF-200-021", "FRAME = X")]}, expect=400)
    assert "CF-200-021: There is no option FRAME" in r["error"]
    eng.call("PUT", f"/api/items/{product}/config/rules", {"rules": [rule(product, "CF-200-021", "MOTOR = 4K")]}, expect=400)
    d = eng.call("PUT", f"/api/items/{product}/config/rules", {"rules": [
        rule(product, "CF-200-021", "motor = 5k5"), rule(product, "CF-200-022", "MOTOR = 7K5"),
        rule(product, "CF-200-023", "MOTOR = 11K"), rule(product, "CF-200-041", "VOLTAGE = 230"),
        rule(product, "CF-200-042", "voltage in (415)"), rule(product, "CF-200-050", "MOTOR = 11K and VOLTAGE = 230"),
        rule(seal_kit, "CF-200-031", "SEAL = PTFE"), rule(seal_kit, "CF-200-032", "SEAL = VITON")]})
    by_pn = {x["part_number"]: x["rule"] for x in d["structure"]["rows"]}
    assert by_pn["CF-200-021"] == "MOTOR = 5K5" and by_pn["CF-200-042"] == "VOLTAGE in (415)" and by_pn["CF-200-010"] == ""
    assert by_pn["CF-200-031"] == "SEAL = PTFE"

    # an 11 kW motor needs 415 V
    eng.call("PUT", f"/api/items/{product}/config/constraints", {"constraints": [{"if": "MOTOR = 11K", "then": ""}]}, expect=400)
    eng.call("PUT", f"/api/items/{product}/config/constraints", {"constraints": [
        {"if": "MOTOR = 11K", "then": "VOLTAGE = 415", "message": "The 11 kW motor runs on 415 V only."}]})

    # configure one combination (anyone may look)
    viewer.call("POST", f"/api/items/{product}/configure", {"choice": {"MOTOR": "7K5"}, "source": "in_work"}, expect=400)
    viewer.call("POST", f"/api/items/{product}/configure", {"choice": {"MOTOR": "9K", "VOLTAGE": "415", "SEAL": "PTFE"},
                                                             "source": "in_work"}, expect=400)
    c = viewer.call("POST", f"/api/items/{product}/configure", {"choice": {"motor": "7k5", "VOLTAGE": "415", "SEAL": "VITON"},
                                                                 "source": "in_work"})
    assert c["allowed"] and c["summary"] == "Motor power 7.5 kW, Voltage 415 V, Seal VITON"
    assert [(x["level"], x["part_number"]) for x in c["rows"]] == [
        (1, "CF-200-010"), (1, "CF-200-022"), (1, "CF-200-030"), (2, "CF-200-032"), (2, "CF-200-033"), (1, "CF-200-042")]
    c = viewer.call("POST", f"/api/items/{product}/configure", {"choice": {"MOTOR": "11K", "VOLTAGE": "230", "SEAL": "PTFE"},
                                                                 "source": "in_work"})
    assert not c["allowed"] and c["blocked_by"] == ["The 11 kW motor runs on 415 V only."]

    # every combination: the old guard is never used (only with a blocked combination)
    chk = viewer.call("POST", f"/api/items/{product}/config/check", {"source": "in_work"})
    assert (chk["combinations"], chk["allowed"], chk["blocked"]) == (12, 10, 2) and chk["distinct_boms"] == 10
    assert [(i["severity"], i["part_number"]) for i in chk["issues"]] == [("warning", "CF-200-050")]
    eng.call("PUT", f"/api/items/{product}/config/options", {"options": OPTIONS + [
        {"code": "PAINT", "choices": [{"code": "BLUE"}, {"code": "GREY"}]}]})
    chk = viewer.call("POST", f"/api/items/{product}/config/check", {"source": "in_work"})
    assert sum(1 for i in chk["issues"] if i["severity"] == "info" and "PAINT" in i["message"]) == 2

    # an option the rules use cannot be taken away
    r = eng.call("PUT", f"/api/items/{product}/config/options", {"options": OPTIONS[:2]}, expect=409)
    assert "CF-200-031" in r["error"]

    # variants need the released BOM; a viewer may look but not save
    choice = {"MOTOR": "11K", "VOLTAGE": "415", "SEAL": "PTFE", "PAINT": "BLUE"}
    eng.call("POST", f"/api/items/{product}/variants", {"choice": choice}, expect=409)
    admin.call("POST", f"/api/boms/{seal}/release")
    admin.call("POST", f"/api/boms/{top}/release")
    viewer.call("POST", f"/api/items/{product}/variants", {"choice": choice}, expect=403)
    eng.call("POST", f"/api/items/{product}/variants", {"choice": {**choice, "VOLTAGE": "230"}}, expect=409)
    v = eng.call("POST", f"/api/items/{product}/variants", {"choice": choice, "customer": "Example Foods",
                                                             "order_ref": "SO-0099"})["variant"]
    assert v["number"] == "CF-200-000-V01" and v["revision"] == "A" and v["choice"] == choice
    assert {ln["part_number"] for ln in v["lines"]} == {"CF-200-010", "CF-200-023", "CF-200-030", "CF-200-031",
                                                         "CF-200-033", "CF-200-042"}
    v2 = eng.call("POST", f"/api/items/{product}/variants", {"choice": {**choice, "MOTOR": "5K5"}, "name": "Small"})["variant"]
    assert v2["number"] == "CF-200-000-V02" and v2["name"] == "Small"

    data, _ = viewer.call("GET", f"/api/variants/{v['id']}/export", raw=True)
    ws = openpyxl.load_workbook(io.BytesIO(data)).active
    assert ws["A1"].value.startswith("CF-200-000-V01") and ws.max_row == 5 + 6

    # a machine built as the variant keeps the variant's BOM
    m = eng.call("POST", "/api/machines", {"variant_id": v["id"], "serial": "CF26-0001", "customer": "Example Foods"})["machine"]
    assert m["variant"] == "CF-200-000-V01" and m["rows"] == 6 and m["item_id"] == product
    assert [x["serial"] for x in viewer.call("GET", f"/api/variants/{v['id']}")["variant"]["machines"]] == ["CF26-0001"]

    # withdrawing: approvers only, with a reason; then no new machines from it
    eng.call("POST", f"/api/variants/{v2['id']}/withdraw", {"reason": "Order cancelled"}, expect=403)
    admin.call("POST", f"/api/variants/{v2['id']}/withdraw", {"reason": ""}, expect=400)
    admin.call("POST", f"/api/variants/{v2['id']}/withdraw", {"reason": "Order cancelled"})
    eng.call("POST", "/api/machines", {"variant_id": v2["id"], "serial": "CF26-0002"}, expect=409)
    assert [x["number"] for x in viewer.call("GET", f"/api/variants?item_id={product}")["variants"]] == ["CF-200-000-V01"]
    assert len(viewer.call("GET", f"/api/variants?item_id={product}&withdrawn=1")["variants"]) == 2
    prods = viewer.call("GET", "/api/products")["products"]
    assert [(p["part_number"], p["options"], p["rules"], p["variants"]) for p in prods] == [("CF-200-000", 4, 8, 1)]
