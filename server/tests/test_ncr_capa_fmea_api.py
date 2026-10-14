# SPDX-License-Identifier: AGPL-3.0-or-later
"""End to end: nonconformance reports, corrective and preventive actions, and FMEA. Who may do what (engineers raise,
approvers or the quality role decide, four eyes on closing a CAPA), dispositions that raise a deviation, numbering,
photos, discussions, and the risk numbers of an FMEA."""
import io

import openpyxl


def _uid(client) -> int:
    return client.call("GET", "/api/auth/me")["user"]["id"]


def test_ncr_to_capa_to_change_request(kosha):
    admin = kosha.client()
    eng = kosha.add_user("q-eng", "engineer")
    insp = kosha.add_user("q-insp", "viewer")       # an inspector: viewer plus the quality role
    viewer = kosha.add_user("q-view", "viewer")
    appr = kosha.add_user("q-appr", "approver")
    admin.call("PATCH", f"/api/users/{_uid(insp)}", {"roles": ["quality"]})
    part = admin.call("POST", "/api/items", {"part_number": "QA-100-020", "description": "AGITATOR SHAFT"})["item"]

    # raising: engineer or quality; a viewer may only look
    viewer.call("POST", "/api/ncrs", {"title": "Bore too big"}, expect=403)
    assert "Say in a few words" in insp.call("POST", "/api/ncrs", {"title": " "}, expect=400)["error"]
    insp.call("POST", "/api/ncrs", {"title": "x", "found_at": "moon"}, expect=400)
    n = insp.call("POST", "/api/ncrs", {"title": "Bore Ø40 measured 40.08", "item_id": part["id"],
                                         "found_at": "incoming", "batch": "Lot 12", "qty": 8,
                                         "supplier": "Example Machining"})["ncr"]
    assert n["number"].startswith("NCR-") and n["number"].endswith("-0001") and n["state"] == "open"
    assert n["revision"] == part["revision"] and n["can"]["edit"] and n["can"]["disposition"]
    assert not eng.call("GET", f"/api/ncrs/{n['id']}")["ncr"]["can"]["disposition"]
    n2 = eng.call("POST", "/api/ncrs", {"title": "Raised twice"})["ncr"]
    assert n2["number"].endswith("-0002")
    viewer.call("POST", f"/api/ncrs/{n2['id']}/cancel", {"reason": "dup"}, expect=403)
    eng.call("POST", f"/api/ncrs/{n2['id']}/cancel", {"reason": ""}, expect=400)
    assert eng.call("POST", f"/api/ncrs/{n2['id']}/cancel", {"reason": "Raised twice"})["ncr"]["state"] == "cancelled"
    eng.call("PATCH", f"/api/ncrs/{n2['id']}", {"title": "Changed"}, expect=409)

    # a photo, then a discussion on the report
    r = insp.call("POST", f"/api/ncrs/{n['id']}/attachments", files={"file": ("bore.jpg", b"\xff\xd8photo")})
    att = r["ncr"]["attachments"][0]
    data, _ = viewer.call("GET", f"/api/ncrs/{n['id']}/attachments/{att['id']}", raw=True)
    assert data == b"\xff\xd8photo"
    viewer.call("DELETE", f"/api/ncrs/{n['id']}/attachments/{att['id']}", expect=403)
    th = eng.call("POST", "/api/discussions", {"entity_type": "ncr", "entity_id": n["id"], "title": "Gauge",
                                               "body": "Which gauge was used?"})
    assert th["threads"][0]["title"] == "Gauge"

    # disposition: approver or quality; use as is needs a reason and raises a draft deviation
    eng.call("POST", f"/api/ncrs/{n['id']}/disposition", {"disposition": "use_as_is", "note": "ok"}, expect=403)
    insp.call("POST", f"/api/ncrs/{n['id']}/disposition", {"disposition": "use_as_is"}, expect=400)
    insp.call("POST", f"/api/ncrs/{n['id']}/disposition", {"disposition": "melt"}, expect=400)
    d = insp.call("POST", f"/api/ncrs/{n['id']}/disposition", {"disposition": "use_as_is",
                                                                "note": "Fit still within the bearing's limits."})["ncr"]
    assert d["state"] == "dispositioned" and d["disposition_label"] == "Use as is" and d["deviation"]["state"] == "draft"
    dev = admin.call("GET", f"/api/deviations/{d['deviation']['id']}")
    assert d["number"] in dev["title"] and dev["qty_limit"] == 8 and dev["serials"] == "Lot 12"
    insp.call("POST", f"/api/ncrs/{n['id']}/disposition", {"disposition": "scrap"}, expect=409)
    insp.call("POST", f"/api/ncrs/{n['id']}/close", {"note": ""}, expect=400)

    # a corrective action from the report
    c = eng.call("POST", f"/api/ncrs/{n['id']}/capa", {"owner_id": _uid(eng), "due_date": "2099-01-31"})["capa"]
    assert c["number"].startswith("CAPA-") and [x["number"] for x in c["ncrs"]] == [n["number"]]
    assert c["item"]["part_number"] == "QA-100-020" and not c["next"]["ok"]
    assert c["next"]["why"] == "Write down the root cause first."
    eng.call("POST", f"/api/ncrs/{n['id']}/capa", {}, expect=409)        # already linked
    eng.call("POST", f"/api/capas/{c['id']}/advance", {}, expect=409)
    c = eng.call("POST", f"/api/capas/{c['id']}/advance", {"root_cause": "Reamer worn; no tool life limit."})["capa"]
    assert c["state"] == "actions" and c["next"]["why"] == "Add at least one action."
    eng.call("POST", f"/api/capas/{c['id']}/advance", {}, expect=409)
    eng.call("POST", f"/api/capas/{c['id']}/actions", {"text": "Tool life limit 400 parts"}, expect=400)
    eng.call("POST", f"/api/capas/{c['id']}/actions", {"text": "Tool life", "owner_id": _uid(eng)}, expect=400)
    c = eng.call("POST", f"/api/capas/{c['id']}/actions", {"text": "Tool life limit 400 parts on the reamer",
                                                            "owner_id": _uid(insp), "due_date": "2000-01-01"})["capa"]
    c = eng.call("POST", f"/api/capas/{c['id']}/actions", {"text": "Typo", "owner_id": _uid(eng),
                                                            "due_date": "2099-01-01"})["capa"]
    a1, a2 = c["actions"]
    assert a1["overdue"] and not a2["overdue"]
    c = eng.call("DELETE", f"/api/capas/{c['id']}/actions/{a2['id']}")["capa"]
    assert len(c["actions"]) == 1
    r = eng.call("POST", f"/api/capas/{c['id']}/advance", {}, expect=409)
    assert "not done yet" in r["error"]
    viewer.call("PATCH", f"/api/capas/{c['id']}/actions/{a1['id']}", {"note": "done"}, expect=403)
    c = insp.call("PATCH", f"/api/capas/{c['id']}/actions/{a1['id']}", {"note": "Limit set in the CAM sheet"})["capa"]
    assert c["actions"][0]["done_by_name"] and not c["actions"][0]["overdue"]

    # the design has to change too: a change request from the CAPA
    ecr = eng.call("POST", f"/api/capas/{c['id']}/change-request", {"title": "Bore tolerance H7 to G7"})["capa"]["ecr"]
    assert ecr["number"].startswith("ECR-")
    eng.call("POST", f"/api/capas/{c['id']}/change-request", {}, expect=409)

    # checking it worked: someone other than the owner, approver or quality
    c = eng.call("POST", f"/api/capas/{c['id']}/advance", {})["capa"]
    assert c["state"] == "verifying" and not c["can"]["actions"]
    eng.call("POST", f"/api/capas/{c['id']}/advance", {"effectiveness": "fine"}, expect=403)
    admin.call("PATCH", f"/api/capas/{c['id']}", {"owner_id": _uid(appr)})
    appr.call("POST", f"/api/capas/{c['id']}/advance", {"effectiveness": "fine"}, expect=403)   # four eyes
    insp.call("POST", f"/api/capas/{c['id']}/advance", {"effectiveness": ""}, expect=400)
    c = insp.call("POST", f"/api/capas/{c['id']}/advance", {"effectiveness": "Too early to tell", "effective": False})["capa"]
    assert c["state"] == "open" and c["effectiveness"] == "Too early to tell"
    c = appr.call("POST", f"/api/capas/{c['id']}/advance", {})["capa"]
    c = appr.call("POST", f"/api/capas/{c['id']}/advance", {})["capa"]
    c = insp.call("POST", f"/api/capas/{c['id']}/advance", {"effectiveness": "3 months, 600 shafts, none out"})["capa"]
    assert c["state"] == "closed" and c["verified_by_name"]
    eng.call("PATCH", f"/api/capas/{c['id']}", {"title": "x"}, expect=409)
    insp.call("POST", f"/api/capas/{c['id']}/attachments", files={"file": ("8d.pdf", b"%PDF")}, expect=409)

    # the report is closed with what was done
    n = insp.call("POST", f"/api/ncrs/{n['id']}/close", {"note": "Used under DEV; supplier informed."})["ncr"]
    assert n["state"] == "closed" and n["capa"]["number"] == c["number"]

    # registers and the overview
    assert [x["number"] for x in viewer.call("GET", "/api/ncrs?state=open")["ncrs"]] == []
    assert len(viewer.call("GET", "/api/ncrs")["ncrs"]) == 2
    assert viewer.call("GET", "/api/ncrs?q=Example%20Machining")["ncrs"][0]["capa_number"] == c["number"]
    assert viewer.call("GET", "/api/capas?state=closed")["capas"][0]["ncrs"] == 1
    ov = insp.call("GET", "/api/quality")
    assert (ov["ncr_open"], ov["capa_open"], ov["may_decide"]) == (0, 0, True)
    assert not viewer.call("GET", "/api/quality")["may_raise"]


def test_fmea_risk_numbers(kosha):
    admin = kosha.client()
    eng = kosha.add_user("f-eng", "engineer")
    viewer = kosha.add_user("f-view", "viewer")
    part = admin.call("POST", "/api/items", {"part_number": "QA-200-000", "description": "PUMP ASSEMBLY"})["item"]
    eng.call("POST", "/api/fmeas", {"kind": "design"}, expect=400)
    viewer.call("POST", "/api/fmeas", {"item_id": part["id"]}, expect=403)
    f = eng.call("POST", "/api/fmeas", {"item_id": part["id"], "kind": "design", "team": "asha, tom"})["fmea"]
    assert f["title"] == "Design FMEA QA-200-000" and f["rows"] == []

    rows = [
        {"function": "Seal the shaft", "failure_mode": "Seal leaks", "effect": "Product loss", "severity": 6,
         "cause": "Dry running", "occurrence": 5, "detection": 4},                                   # 120 > 100
        {"function": "Carry the load", "failure_mode": "Shaft breaks", "effect": "Injury", "severity": 10,
         "cause": "Fatigue", "occurrence": 2, "detection": 3, "action": "Raise the fillet radius",
         "due_date": "2000-01-01"},                                                                   # severity 10
        {"function": "Pump", "failure_mode": "Low flow", "effect": "Slow filling", "severity": 3, "cause": "Wear",
         "occurrence": 3, "detection": 3},                                                            # fine
    ]
    eng.call("PUT", f"/api/fmeas/{f['id']}", {"rows": [{**rows[0], "severity": 11}]}, expect=400)
    r = eng.call("PUT", f"/api/fmeas/{f['id']}", {"rows": [{**rows[0], "failure_mode": ""}]}, expect=400)
    assert r["error"].startswith("Row 1:")
    viewer.call("PUT", f"/api/fmeas/{f['id']}", {"rows": rows}, expect=403)
    f = eng.call("PUT", f"/api/fmeas/{f['id']}", {"rows": rows})["fmea"]
    assert [x["rpn"] for x in f["rows"]] == [120, 60, 27]
    assert [x["needs_action"] for x in f["rows"]] == [True, True, False]
    assert "No action planned yet." in f["rows"][0]["flags"] and f["rows"][1]["overdue"]
    assert f["summary"]["need_action"] == 2 and f["summary"]["high_severity"] == 1 and f["summary"]["max_rpn"] == 120

    # the actions taken bring the risk down; severity 10 stays flagged whatever the RPN
    rows[0].update(action="Dry-run protection", action_taken="Flow switch fitted", new_severity=6, new_occurrence=2,
                   new_detection=3)
    rows[1].update(action_taken="R3 fillet", new_severity=10, new_occurrence=1, new_detection=2)
    f = eng.call("PUT", f"/api/fmeas/{f['id']}", {"rows": rows})["fmea"]
    assert [x["new_rpn"] for x in f["rows"]] == [36, 20, None]
    assert [x["needs_action"] for x in f["rows"]] == [False, True, False]

    # the company's limit
    viewer.call("PUT", "/api/quality/rpn-limit", {"rpn_limit": 30}, expect=403)
    admin.call("PUT", "/api/quality/rpn-limit", {"rpn_limit": 0}, expect=400)
    admin.call("PUT", "/api/quality/rpn-limit", {"rpn_limit": 25})
    f = viewer.call("GET", f"/api/fmeas/{f['id']}")["fmea"]
    assert f["rpn_limit"] == 25 and [x["needs_action"] for x in f["rows"]] == [True, True, True]

    data, _ = viewer.call("GET", f"/api/fmeas/{f['id']}/export", raw=True)
    ws = openpyxl.load_workbook(io.BytesIO(data)).active
    assert ws["B5"].value == "Seal leaks" and ws["J5"].value == 120 and ws["R5"].value == 36
    lst = viewer.call("GET", f"/api/fmeas?item_id={part['id']}")["fmeas"]
    assert [(x["rows"], x["max_rpn"]) for x in lst] == [(3, 120)]
