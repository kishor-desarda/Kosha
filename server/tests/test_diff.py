# SPDX-License-Identifier: AGPL-3.0-or-later
"""Measured change list: every comparer finds the facts and says them in plain English."""
import io

import pytest

from kosha.cad import dwg
from kosha.cad import dxfwriter as W
from kosha.cad import mesh as M
from kosha.diff import Change, Comparer, DiffEngine, Side, engine, render
from kosha.diff.bom import compare_lines
from kosha.diff.drawing import DrawingComparer, DwgComparer, noun_for
from kosha.diff.model import ModelComparer
from kosha.diff.props import compare_props
from kosha.diff.text import TextComparer

SHAFT_L = (60, 80, 220, 80, 80)


def diff(old: bytes, new: bytes, name="part.dxf", **kw):
    return engine().compare(Side.of_bytes(name, old, **kw.get("a", {})), Side.of_bytes(name, new, **kw.get("b", {})))


def texts(d):
    return [c.text for c in d.changes]


# ------------------------------------------------------------------ drawings

def test_shaft_diameter_change_is_found_with_its_zone():
    a = W.shaft(SHAFT_L, (80, 95, 110, 95, 80), values={"PART_NO": "MX-100-020", "MATERIAL": "42CrMo4", "REV": "C"})
    b = W.shaft(SHAFT_L, (80, 95, 115, 95, 80), values={"PART_NO": "MX-100-020", "MATERIAL": "42CrMo4", "REV": "C"})
    d = diff(a, b)
    dims = [c for c in d.changes if c.kind == "dimension"]
    assert [c.text for c in dims] == ["Diameter Ø110 changed to Ø115"]
    c = dims[0]
    assert (c.before, c.after, c.severity_hint) == ("Ø110", "Ø115", "major")
    x0, y0, x1, y1 = c.zone                 # the Ø label sits above the middle step of the shaft
    assert x0 < 175 < x1 and y0 < 202 < y1
    assert "1 polyline changed on layer 0" in texts(d)       # the outline
    assert "1 polyline moved on layer 0" in texts(d)         # the keyway follows the larger diameter
    assert not any(c.kind == "attribute" for c in d.changes)


def test_dimension_entities_with_fits_read_as_diameter_or_bore():
    def drawing(shaft_d, bore_d):
        w = W.DxfWriter()
        w.layer("DIM", 3)
        w.rect(0, 0, 200, 100)
        # a real DIMENSION entity without a block: Kosha uses its measured value and the "<> h6" override
        w._ent("DIMENSION", "DIM", (2, "*D9"), (10, 20.0), (20, 50.0), (11, 60.0), (21, 80.0), (70, 3), (1, "<> h6"),
               (42, float(shaft_d)))
        w.text(150, 80, 3.5, f"%%c{bore_d} H7", "DIM")
        return w.tobytes()

    d = diff(drawing(110, 80), drawing(115, 85))
    assert sorted(texts(d)) == ["Bore Ø80 H7 changed to Ø85 H7", "Diameter Ø110 h6 changed to Ø115 h6"]
    assert noun_for("R5") == "Radius" and noun_for("M16") == "Thread" and noun_for("120") == "Dimension"


def test_flange_hole_count_8_to_12():
    a = W.flange(200, 160, 8, 18, 100, values={"PART_NO": "SC-300-021"})
    b = W.flange(200, 160, 12, 18, 100, values={"PART_NO": "SC-300-021"})
    d = diff(a, b)
    holes = [c for c in d.changes if c.kind == "holes"]
    assert len(holes) == 1
    h = holes[0]
    assert (h.before, h.after) == (8, 12) and h.detail["action"] == "added" and h.detail["count"] == 4
    assert h.text == "4 holes Ø10.8 added on layer 0 (8 → 12; 4 repositioned)"
    assert h.zone is not None
    assert not any(c.kind == "entities" and c.detail["entity"] == "circle" for c in d.changes)   # not counted twice


def test_holes_added_on_their_own_layer_and_resized():
    def plate(holes, d=10):
        w = W.DxfWriter()
        w.rect(0, 0, 300, 150)
        for k in range(holes):
            w.circle(30 + 40 * k, 75, d / 2, layer="HOLES")
        return w.tobytes()

    assert texts(diff(plate(4), plate(6))) == ["2 holes Ø10 added on layer HOLES (4 → 6)"]
    assert texts(diff(plate(4), plate(4, 12))) == ["4 holes changed from Ø10 to Ø12 on layer HOLES"]
    assert texts(diff(plate(4), plate(3))) == ["1 hole Ø10 removed on layer HOLES (4 → 3)"]


def test_title_block_material_change():
    a = W.shaft(SHAFT_L, (80, 95, 110, 95, 80), values={"PART_NO": "MX-100-020", "MATERIAL": "C45", "REV": "A"})
    b = W.shaft(SHAFT_L, (80, 95, 110, 95, 80), values={"PART_NO": "MX-100-020", "MATERIAL": "42CrMo4", "REV": "B"})
    d = diff(a, b)
    assert texts(d) == ["Material changed from C45 to 42CrMo4", "Revision changed from A to B"]
    mat, rev = d.changes
    assert mat.severity_hint == "major" and rev.severity_hint == "info"
    assert mat.where == "Title block" and mat.zone is not None


def test_moved_and_added_entities_and_notes_and_extents():
    def part(dx, note, extra):
        w = W.DxfWriter()
        w.rect(0, 0, 200, 100)
        w.rect(20 + dx, 20, 30, 30)
        for k in range(extra):
            w.line(150, 10 + 10 * k, 190, 10 + 10 * k)
        w.text(10, 120, 3.5, note, "NOTES")
        return w.tobytes()

    d = diff(part(0, "BREAK ALL EDGES", 0), part(15, "BREAK ALL EDGES 0.5X45", 2))
    t = texts(d)
    assert "1 polyline moved on layer 0" in t
    assert "2 lines added on layer 0" in t
    assert "Note “BREAK ALL EDGES” changed to “BREAK ALL EDGES 0.5X45”" in t
    moved = next(c for c in d.changes if c.detail.get("action") == "moved")
    assert moved.detail["offset"] == [15.0, 0.0]
    housing = diff(W.bearing_housing(80, 220), W.bearing_housing(80, 230))
    assert housing.changes[0].text == "Overall size changed from 342 × 167 to 352 × 167 mm"
    assert "Dimension 220 changed to 230" in texts(housing)


def test_identical_drawings_have_no_changes():
    a = W.flange(200, 160, 8, 18, 100)
    assert diff(a, a).changes == []
    assert diff(a, a.replace(b"EOF", b"EOF\n")).changes == []      # different bytes, same drawing


def test_dwg_without_converter_compares_the_stored_title_block(monkeypatch):
    def no_converter(_):
        raise dwg.DwgUnavailable("no converter")
    monkeypatch.setattr(dwg, "to_dxf", no_converter)
    a = Side.of_bytes("x.dwg", b"AC1032 old", meta={"attributes": {"MATERIAL": "S275JR", "REV": "A"}})
    b = Side.of_bytes("x.dwg", b"AC1032 new", meta={"attributes": {"MATERIAL": "S355JR", "REV": "B"}})
    d = engine().compare(a, b)
    assert texts(d) == ["Material changed from S275JR to S355JR", "Revision changed from A to B"]
    assert any("DWG converter" in n for n in d.notes)


# ------------------------------------------------------------------ 3D and CAD properties

def test_mesh_size_volume_and_area():
    a = M.write_stl(M.shaft_mesh(SHAFT_L, (80, 95, 110, 95, 80)))
    b = M.write_stl(M.shaft_mesh(SHAFT_L, (80, 95, 115, 95, 80)))
    d = diff(a, b, "shaft.stl")
    kinds = [c.kind for c in d.changes]
    assert kinds == ["model_size", "volume", "area"]
    assert d.changes[0].text == "Model size changed from 520 × 110 × 110 to 520 × 115 × 115 mm"
    vol = d.changes[1]
    assert vol.before < vol.after and vol.text.startswith("Volume changed from ") and "cm³ (+" in vol.text


def test_solidworks_metadata_custom_properties_mass_and_components():
    sa = M.stats(M.write_stl(M.box_mesh(100, 50, 20)))
    sb = M.stats(M.write_stl(M.box_mesh(100, 50, 25)))
    a = Side.of_bytes("MX-1.SLDASM", b"sw", meta={"mesh": sa, "properties": {"Material": "C45", "Mass": "186.0", "Finish": "Painted"},
                                                    "structure": [{"part_number": "MX-1-010", "qty": 2}, {"part_number": "MX-1-020", "qty": 1}]})
    b = Side.of_bytes("MX-1.SLDASM", b"sw2", meta={"mesh": sb, "properties": {"Material": "42CrMo4", "Mass": "190.5 kg", "Finish": "Painted",
                                                                                "Coating": "Zinc"},
                                                    "structure": [{"part_number": "MX-1-010", "qty": 4}, {"part_number": "MX-1-030", "qty": 1}]})
    t = texts(engine().compare(a, b))
    assert "Model size changed from 100 × 50 × 20 to 100 × 50 × 25 mm" in t
    assert "Custom property Material changed from C45 to 42CrMo4" in t
    assert "Mass changed from 186 kg to 190.5 kg (+2.4 %)" in t
    assert "Custom property Coating set to Zinc" in t
    assert "Component MX-1-030 added (qty 1)" in t and "Component MX-1-020 removed" in t
    assert "Quantity of MX-1-010 changed from 2 to 4" in t
    assert not any("Finish" in x for x in t)


# ------------------------------------------------------------------ BOMs

def test_bom_lines_added_removed_quantity_and_swap():
    old = [{"part_number": "MX-100-010", "description": "IMPELLER DISC", "qty": 1, "find_number": "1"},
           {"part_number": "MX-100-031", "description": "BEARING HOUSING NDE", "qty": 1, "find_number": "4"},
           {"part_number": "STD-BLT-M16X60", "description": "HEX BOLT M16 X 60", "qty": 8, "find_number": "8"},
           {"part_number": "STD-WSH-M16", "description": "WASHER M16", "qty": 8, "find_number": "9"}]
    new = [{"part_number": "MX-100-010", "description": "IMPELLER DISC", "qty": 1, "find_number": "1"},
           {"part_number": "MX-100-035", "description": "BEARING HOUSING NDE HEAVY", "qty": 1, "find_number": "4"},
           {"part_number": "STD-BLT-M16X60", "description": "HEX BOLT M16 X 60", "qty": 12, "find_number": "8"},
           {"part_number": "MX-100-060", "description": "GUARD", "qty": 1, "find_number": "10"}]
    t = [c.text for c in compare_lines(old, new)]
    assert t == ["MX-100-031 replaced by MX-100-035 at find no. 4",
                 "MX-100-060 Guard added to the BOM (qty 1)",
                 "STD-WSH-M16 Washer M16 removed from the BOM (was qty 8)",
                 "Quantity of STD-BLT-M16X60 Hex Bolt M16 X 60 changed from 8 to 12"]
    assert compare_lines(old, old) == []


def test_bom_comparer_runs_for_any_file_type_when_both_sides_have_a_bom():
    a = Side.of_bytes("asm.stl", b"x", bom_lines=[{"part_number": "A", "qty": 1}], bom_id=1)
    b = Side.of_bytes("asm.stl", b"x", bom_lines=[{"part_number": "A", "qty": 2}], bom_id=2)
    assert texts(engine().compare(a, b)) == ["Quantity of A changed from 1 to 2"]
    same = Side.of_bytes("asm.stl", b"x", bom_lines=[{"part_number": "A", "qty": 2}], bom_id=1)
    assert engine().compare(a, same).changes == []        # the same BOM record on both sides


# ------------------------------------------------------------------ text documents

def _docx(paragraphs):
    import docx
    d = docx.Document()
    for p in paragraphs:
        d.add_paragraph(p)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _xlsx(cells):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Calc"
    for ref, v in cells.items():
        ws[ref] = v
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pdf(lines):
    """A one-page PDF with a text layer, written by hand."""
    stream = "BT /F1 12 Tf 72 720 Td " + " ".join(f"({x}) Tj 0 -16 Td" for x in lines) + " ET"
    objs = ["<< /Type /Catalog /Pages 2 0 R >>", "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
            f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream",
            "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offsets = "%PDF-1.4\n", []
    for i, o in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{o}\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n" + "".join(f"{o:010d} 00000 n \n" for o in offsets)
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n"
    return out.encode("latin-1")


def test_word_paragraphs_changed_added_removed():
    a = _docx(["Scope", "The shaft is made of C45.", "Paint RAL 5010.", "Old remark"])
    b = _docx(["Scope", "The shaft is made of 42CrMo4.", "Paint RAL 5010.", "Check hub seat fit."])
    d = diff(a, b, "spec.docx")
    assert d.changes[0].text == "2 paragraphs changed"
    assert d.changes[1].text == "Paragraph 2 changed: “The shaft is made of C45.” → “The shaft is made of 42CrMo4.”"


def test_excel_cells_and_text_files_and_pdf():
    d = diff(_xlsx({"A1": "Torque", "B1": 120, "B2": 5}), _xlsx({"A1": "Torque", "B1": 150, "C3": "new"}), "calc.xlsx")
    assert texts(d) == ["1 cell changed, 1 cell added, 1 cell removed", "Cell Calc!B1 changed from 120 to 150",
                        "Cell Calc!B2 cleared (was 5)", "Cell Calc!C3 set to new"]
    t = texts(diff(b"line one\nline two\n", b"line one\nline 2\nline three\n", "notes.txt"))
    assert t[0] == "1 line changed, 1 line added"
    assert "Line 2 changed: “line two” → “line 2”" in t and "Line 3 added: “line three”" in t
    p = texts(diff(_pdf(["Hub seat 110 h6", "Material C45"]), _pdf(["Hub seat 115 h6", "Material C45"]), "drg.pdf"))
    assert p[0] == "1 line changed" and "Hub seat 110 h6" in p[1] and "on page 1" in p[1]


def test_text_comparer_is_listed_for_office_and_text_types():
    assert TextComparer in engine().comparers_for("docx") and TextComparer in engine().comparers_for("pdf")
    assert DrawingComparer in engine().comparers_for("dxf") and DwgComparer in engine().comparers_for("dwg")
    assert ModelComparer in engine().comparers_for("sldprt")


# ------------------------------------------------------------------ properties and the engine

def test_item_properties():
    t = [c.text for c in compare_props({"description": "SHAFT", "material": "C45", "weight_kg": 186.0, "uom": "Nos"},
                                       {"description": "AGITATOR SHAFT", "material": "42CrMo4", "weight_kg": 190.5, "uom": "Nos"})]
    assert t == ["Description changed from SHAFT to AGITATOR SHAFT", "Material changed from C45 to 42CrMo4",
                 "Weight changed from 186 kg to 190.5 kg"]


def test_engine_registers_by_type_guards_failures_and_drops_repeated_facts():
    class Broken(Comparer):
        name, label, exts = "broken", "Broken comparer", ("xyz",)

        def compare(self, a, b):
            raise RuntimeError("boom")

    class Facts(Comparer):
        name, label, exts = "facts", "Facts", ("xyz",)

        def compare(self, a, b):
            return [Change("attribute", "Title block", "C45", "42CrMo4", "major", None, {"label": "Material"}),
                    Change("note", "Layer 0", None, "NEW NOTE", "info")]

    class Props(Comparer):
        name, label, exts = "props", "Props", ("*",)

        def compare(self, a, b):
            return [Change("property", "Part", "c45", "42crmo4", "major", None, {"label": "material"}),
                    Change("property", "Part", "SHAFT", "MAIN SHAFT", "minor", None, {"label": "Description"})]

    e = DiffEngine()
    for cls in (Broken, Facts, Props):
        e.register(cls)
    assert e.comparers_for("xyz") == [Broken, Facts, Props] and e.comparers_for("abc") == [Props]
    d = e.compare(Side.of_bytes("a.xyz", b"1"), Side.of_bytes("a.xyz", b"2"))
    assert [c.text for c in d.changes] == ["Material changed from C45 to 42CrMo4", "Description changed from SHAFT to MAIN SHAFT",
                                           "Note “NEW NOTE” added"]
    assert d.comparers == ["facts", "props"]
    assert any("Broken comparer" in n and "RuntimeError" in n for n in d.notes)
    assert d.to_dict()["counts"] == {"major": 1, "minor": 1, "info": 1}


def test_render_covers_every_kind():
    assert render(Change("attribute", "Title block", None, "S355", "minor", None, {"label": "Material"})) == "Material set to S355"
    assert render(Change("attribute", "Title block", "S355", None, "minor", None, {"label": "Material"})) == "Material removed (was S355)"
    assert render(Change("dimension", "Layer DIM", None, "Ø20", "major")) == "Dimension Ø20 added"
    assert render(Change("entities", "Layer 0", 3, None, "major", None,
                         {"action": "removed", "count": 3, "entity": "arc", "layer": "0"})) == "3 arcs removed on layer 0"
    assert render(Change("holes", "Layer 0", 8, 8, "major", None,
                         {"action": "moved", "count": 8, "d": 18, "layer": "0"})) == "8 holes Ø18 moved on layer 0"
    assert render(Change("other", "Somewhere", "x", "y")) == "Somewhere: x → y"


@pytest.mark.parametrize("value", ["Ø110 h6", "Ø80H7", "2xØ18", "R5", "M16x1.5", "120 ±0.1", "45°"])
def test_measurement_texts_count_as_dimensions_wherever_they_are(value):
    from kosha.diff.drawing import MEASUREMENT
    assert MEASUREMENT.match(value)
