# SPDX-License-Identifier: AGPL-3.0-or-later
"""End to end: manufacturing BOMs and process plans, made from the engineering BOM, checked against it, flattened,
released, carried to the next revision and sent to the ERP with their operations."""
import io

import openpyxl


def _bom(c, pn, desc, lines, revision=None, item_id=None):
    body = {"item_id": item_id, "revision": revision} if item_id else {"part_number": pn, "description": desc}
    bid = c.call("POST", "/api/boms", body)["bom_id"]
    c.call("PUT", f"/api/boms/{bid}/lines", {"lines": lines})
    return bid


def _issues(d, severity=None):
    return [(i["part_number"], i["message"]) for i in d["issues"] if severity in (None, i["severity"])]


def _plan_lines(d):
    return [{"item_id": ln["item_id"], "qty": ln["qty"], "op_no": ln["op_no"], "note": ln["note"]} for ln in d["lines"]]


OPS = [{"op_no": 10, "work_centre": "Assembly bench", "description": "Build the drive unit", "setup_min": 15, "run_min": 40},
       {"op_no": 20, "work_centre": "Assembly bench", "description": "Fit the tank and the drive", "run_min": 60},
       {"op_no": 30, "work_centre": "Paint shop", "description": "Paint", "setup_min": 10, "run_min": 45,
        "instructions": "Two coats, RAL 5015."}]


def test_manufacturing_bom_from_design_to_erp(kosha):
    admin = kosha.client()
    eng = kosha.add_user("mb-eng", "engineer")
    viewer = kosha.add_user("mb-view", "viewer")
    sub = _bom(admin, "MB-100-020", "DRIVE UNIT", [{"part_number": "MB-100-021", "description": "GEARBOX", "qty": 1},
                                                    {"part_number": "MB-100-022", "description": "MOTOR", "qty": 1},
                                                    {"part_number": "STD-BLT-M12", "description": "HEX BOLT M12", "qty": 4}])
    top = _bom(admin, "MB-100-000", "MIXER 500 L", [
        {"part_number": "MB-100-010", "description": "TANK", "qty": 1},
        {"part_number": "MB-100-020", "description": "DRIVE UNIT", "qty": 1},
        {"part_number": "STD-BLT-M12", "description": "HEX BOLT M12", "qty": 8}])
    item = admin.call("GET", f"/api/boms/{top}")["bom"]["item_id"]
    drive = admin.call("GET", "/api/search?q=MB-100-020")["items"][0]["id"]

    # made from the engineering BOM; only engineers make it
    viewer.call("POST", "/api/mboms", {"item_id": item}, expect=403)
    r = eng.call("POST", "/api/mboms", {"item_id": item})
    assert r["created"] and r["carried_from"] is None
    mid = r["mbom_id"]
    assert eng.call("POST", "/api/mboms", {"item_id": item})["mbom_id"] == mid       # the same revision: the same one
    d = eng.call("GET", f"/api/mboms/{mid}")
    assert d["mbom"]["revision"] == "A" and d["mbom"]["state"] == "draft" and d["ebom"]["id"] == top
    assert {ln["part_number"]: ln["qty"] for ln in d["lines"]} == {"MB-100-010": 1, "MB-100-020": 1, "STD-BLT-M12": 8}
    assert not _issues(d, "error")
    assert any("No process plan yet" in m for _, m in _issues(d, "warning"))
    assert d["sub_assemblies"] == [{"item_id": drive, "part_number": "MB-100-020"}]

    # the process plan: operations, and each part at the operation that uses it
    eng.call("PUT", f"/api/mboms/{mid}/plan", {"ops": [{"op_no": 10, "description": ""}]}, expect=400)
    eng.call("PUT", f"/api/mboms/{mid}/plan", {"ops": [OPS[0], OPS[0]]}, expect=400)
    d = eng.call("PUT", f"/api/mboms/{mid}/plan", {"ops": OPS})
    assert [o["op_no"] for o in d["ops"]] == [10, 20, 30] and d["totals"] == {"setup_min": 25, "run_min": 145}
    assert sum(1 for _, m in _issues(d, "warning") if "Not issued at an operation" in m) == 3

    # build the drive unit in one go with the mixer: its parts come up, at the drive unit's operation
    lines = _plan_lines(d)
    for ln in lines:
        ln["op_no"] = 10 if ln["item_id"] == drive else 20
    eng.call("PUT", f"/api/mboms/{mid}/plan", {"lines": lines})
    d = eng.call("POST", f"/api/mboms/{mid}/flatten", {"item_id": drive})
    got = {(ln["part_number"], ln["op_no"]): ln["qty"] for ln in d["lines"]}
    assert got == {("MB-100-010", 20): 1, ("STD-BLT-M12", 20): 8, ("MB-100-021", 10): 1, ("MB-100-022", 10): 1,
                   ("STD-BLT-M12", 10): 4}
    assert [f["part_number"] for f in d["flattened"]] == ["MB-100-020"] and not _issues(d, "error")

    # the workshop adds paint; issuing too few bolts is an error, too many a warning
    lines = _plan_lines(d) + [{"part_number": "MB-PAINT-5015", "description": "PAINT RAL 5015", "qty": 2, "op_no": 30}]
    eng.call("PUT", f"/api/mboms/{mid}/plan", {"lines": lines + [{"part_number": "X", "qty": 1, "op_no": 40}]}, expect=400)
    d = eng.call("PUT", f"/api/mboms/{mid}/plan", {"lines": lines})
    assert ("MB-PAINT-5015", "Added for manufacturing (not in the engineering BOM).") in _issues(d, "info")
    assert {ln["part_number"]: ln["origin"] for ln in d["lines"]}["MB-PAINT-5015"] == "added"
    bolt = next(ln["item_id"] for ln in d["lines"] if ln["part_number"] == "STD-BLT-M12")
    short = [dict(ln, qty=2) if (ln["item_id"], ln["op_no"]) == (bolt, 20) else ln for ln in _plan_lines(d)]
    d2 = eng.call("PUT", f"/api/mboms/{mid}/plan", {"lines": short})
    assert ("STD-BLT-M12", "The engineering BOM needs 12; the manufacturing BOM issues only 6.") in _issues(d2, "error")
    eng.call("PUT", f"/api/mboms/{mid}/plan", {"lines": _plan_lines(d)})

    # an operation still in use cannot be removed on its own
    eng.call("PUT", f"/api/mboms/{mid}/plan", {"ops": OPS[:2]}, expect=409)

    # release: approvers only, and only after the engineering BOM
    eng.call("POST", f"/api/mboms/{mid}/release", expect=403)
    d = admin.call("GET", f"/api/mboms/{mid}")
    assert d["can_release"] == {"ok": False, "why": "Release the engineering BOM revision A first."}
    admin.call("POST", f"/api/mboms/{mid}/release", expect=409)
    admin.call("PUT", "/api/erp/settings", {"bom_source": "sideways"}, expect=400)
    admin.call("PUT", "/api/erp/settings", {"bom_source": "mbom"})
    admin.call("POST", f"/api/boms/{sub}/release")
    admin.call("POST", f"/api/boms/{top}/release")
    d = admin.call("POST", f"/api/mboms/{mid}/release")
    assert d["mbom"]["state"] == "released" and d["can_release"]["why"] == "Already released."
    eng.call("PUT", f"/api/mboms/{mid}/plan", {"ops": OPS}, expect=409)
    eng.call("POST", f"/api/mboms/{mid}/rebuild", expect=409)

    # the ERP gets the manufacturing BOM with its operations (dry run here), not the engineering one
    jobs = admin.call("GET", "/api/erp/jobs")["jobs"]
    kinds = {j["kind"] for j in jobs}
    assert "mbom" in kinds and not any(j["kind"] == "bom" and j["entity_id"] == top for j in jobs)
    assert any(j["kind"] == "bom" and j["entity_id"] == sub for j in jobs)     # nobody planned the drive unit
    admin.call("POST", "/api/erp/process")
    job = next(j for j in admin.call("GET", "/api/erp/jobs")["jobs"] if j["kind"] == "mbom")
    assert job["status"] == "dry_run" and "manufacturing BOM rev A" in job["label"]
    import json
    payload = json.loads(job["payload_json"])
    assert [o["op_no"] for o in payload["operations"]] == [10, 20, 30]
    assert {(i["item_code"], i.get("operation")) for i in payload["items"]} >= {("MB-PAINT-5015", 30), ("MB-100-021", 10)}

    # the Excel export has the parts by operation and the process plan
    data, _ = admin.call("GET", f"/api/mboms/{mid}/export", raw=True)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Manufacturing BOM", "Process plan"]
    plan = [r for r in wb["Process plan"].iter_rows(min_row=4, values_only=True)]
    assert plan[0][:3] == (10, "Assembly bench", "Build the drive unit") and "MB-100-021 × 1" in plan[0][5]
    assert plan[-1][2:5] == ("Total", 25, 145)

    # revision B of the design: a new manufacturing BOM carries the plan, the flattening and the paint forward
    _bom(admin, None, None, [{"part_number": "MB-100-010", "description": "TANK", "qty": 1},
                             {"part_number": "MB-100-020", "description": "DRIVE UNIT", "qty": 1},
                             {"part_number": "MB-100-030", "description": "SIGHT GLASS", "qty": 1},
                             {"part_number": "STD-BLT-M12", "description": "HEX BOLT M12", "qty": 8}],
         revision="B", item_id=item)
    assert admin.call("GET", f"/api/mboms/{mid}")["newer_ebom"] == "B"
    r = eng.call("POST", "/api/mboms", {"item_id": item})
    assert r["created"] and r["carried_from"] == "A"
    d = eng.call("GET", f"/api/mboms/{r['mbom_id']}")
    got = {(ln["part_number"], ln["op_no"]): ln["qty"] for ln in d["lines"]}
    assert got[("MB-100-021", 10)] == 1 and got[("MB-PAINT-5015", 30)] == 2 and got[("MB-100-030", None)] == 1
    assert [o["op_no"] for o in d["ops"]] == [10, 20, 30]
    assert ("MB-100-030", "Not issued at an operation. Choose the operation that uses it.") in _issues(d, "warning")
    assert [m["revision"] for m in d["revisions"]] == ["A", "B"]

    # back to one line for the drive unit, at the operation its parts used
    d = eng.call("POST", f"/api/mboms/{r['mbom_id']}/flatten", {"item_id": drive, "on": False})
    got = {(ln["part_number"], ln["op_no"]): ln["qty"] for ln in d["lines"]}
    assert got[("MB-100-020", 10)] == 1 and ("MB-100-021", 10) not in got and got[("MB-PAINT-5015", 30)] == 2
    assert not d["flattened"] and not _issues(d, "error")

    listing = viewer.call("GET", f"/api/mboms?item_id={item}")["mboms"]
    assert [(m["revision"], m["state"]) for m in listing] == [("B", "draft"), ("A", "released")]
    admin.call("PUT", "/api/erp/settings", {"bom_source": "ebom"})


def test_no_engineering_bom_no_manufacturing_bom(kosha):
    admin = kosha.client()
    item = admin.call("POST", "/api/items", {"part_number": "MB-900-000", "description": "LONE PART"})
    iid = (item.get("item") or item)["id"]
    r = admin.call("POST", "/api/mboms", {"item_id": iid}, expect=409)
    assert "no engineering BOM" in r["error"]
    admin.call("GET", "/api/mboms/999999", expect=404)
