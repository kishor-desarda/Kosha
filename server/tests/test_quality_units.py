# SPDX-License-Identifier: AGPL-3.0-or-later
"""Deviations, discussions and notifications against a real (temporary) database, without a server."""
import io
from datetime import date, datetime, timedelta, timezone

import openpyxl
import pytest

from kosha import config, db, deviations, discussions, notifications
from kosha.web import ApiError


@pytest.fixture
def q(tmp_path, monkeypatch):
    """A fresh database with users, an item, a released drawing and a hook that records add-in events."""
    for name, value in (("DATA_DIR", tmp_path), ("DB_PATH", tmp_path / "kosha.db"), ("BLOB_DIR", tmp_path / "blobs"),
                        ("WORK_DIR", tmp_path / "work")):
        monkeypatch.setattr(config, name, value)
    db.close_thread_connection()
    db.init()
    now = db.now()
    users = {}
    for username, role in (("eve", "engineer"), ("ann", "approver"), ("bob", "approver"), ("vic", "viewer"),
                           ("root", "admin"), ("gone", "approver")):
        users[username] = {"id": db.insert("users", {"username": username, "full_name": username.title(), "role": role,
                                                    "password_hash": "x", "created_at": now,
                                                    "active": 0 if username == "gone" else 1}),
                           "username": username, "full_name": username.title(), "role": role}
    folder = db.insert("folders", {"parent_id": None, "name": "Kosha", "created_at": now})
    item = db.insert("items", {"part_number": "MX-100-050", "description": "SHAFT KEY", "material": "C45", "revision": "B",
                               "created_at": now, "updated_at": now})
    doc = db.insert("documents", {"folder_id": folder, "filename": "MX-100-050 Shaft Key.dxf", "ext": "dxf",
                                  "doc_type": "drawing", "item_id": item, "part_number": "MX-100-050", "revision": "B",
                                  "state": "released", "latest_version": 1, "created_at": now, "updated_at": now})
    db.insert("versions", {"document_id": doc, "version": 1, "revision": "B", "blob_sha": "0" * 64, "size": 1,
                           "created_at": now})
    events = []
    from kosha.addins import registry
    monkeypatch.setattr(registry, "emit", lambda name, payload: events.append((name, payload)))
    yield {"u": users, "item": item, "doc": doc, "events": events}
    db.close_thread_connection()


def _raises(code_or_status, fn, *args, **kwargs):
    with pytest.raises(ApiError) as e:
        fn(*args, **kwargs)
    assert code_or_status in (e.value.code, e.value.status), (e.value.status, e.value.code, e.value.message)
    return e.value


def _ready(q, requester="eve", approvers=("ann",), **extra) -> int:
    u = q["u"]
    data = {"title": "42CrMo4 instead of C45", "departure": "material 42CrMo4 allowed", "serials": "batch 12",
            "affected": [{"item_id": q["item"]}], "approver_ids": [u[a]["id"] for a in approvers], **extra}
    return deviations.create(u[requester], data)


def _prefix() -> str:
    return f"DEV-{datetime.now(timezone.utc).year % 100:02d}-"


# ------------------------------------------------------------------ numbering

def test_numbers_run_per_year_and_are_never_reused(q):
    u = q["u"]
    ids = [_ready(q) for _ in range(3)]
    assert [deviations.dev_or_404(i)["number"] for i in ids] == [f"{_prefix()}{n:04d}" for n in (1, 2, 3)]
    deviations.cancel(u["eve"], ids[2])
    assert deviations.dev_or_404(_ready(q))["number"] == f"{_prefix()}0004"       # a cancelled number stays taken
    db.execute("DELETE FROM deviation_numbers")                                      # e.g. an old backup of the counter
    assert deviations.dev_or_404(_ready(q))["number"] == f"{_prefix()}0005"
    with pytest.raises(ApiError):                                                    # a failed create uses no number
        deviations.create(u["eve"], {"title": "x", "affected": [{"item_id": 99999}]})
    assert deviations.dev_or_404(_ready(q))["number"] == f"{_prefix()}0006"


# ------------------------------------------------------------------ lifecycle and roles

def test_roles_four_eyes_and_lifecycle(q):
    u, events = q["u"], q["events"]
    _raises(403, deviations.create, u["vic"], {"title": "Viewer tries"})
    dev = deviations.create(u["eve"], {"title": "Empty"})
    e = _raises("not_ready", deviations.submit, u["eve"], dev)
    for words in ("affected", "departure", "limit it", "approver"):
        assert words in e.message
    _raises("four_eyes", deviations.add_approver, u["eve"], dev, u["eve"]["id"])     # the requester is never an approver
    _raises(400, deviations.add_approver, u["eve"], dev, u["vic"]["id"])             # viewers can't approve
    _raises(400, deviations.add_approver, u["eve"], dev, u["gone"]["id"])            # inactive users neither

    dev = _ready(q, requester="ann", approvers=("bob",))                             # an approver raises one
    deviations.submit(u["ann"], dev)
    _raises(403, deviations.decide, u["ann"], dev, "approved")                       # ... and can't approve it
    _raises(403, deviations.decide, u["eve"], dev, "approved")
    _raises(400, deviations.decide, u["bob"], dev, "rejected", "")                   # a rejection needs a reason
    deviations.decide(u["bob"], dev, "rejected", "Say which heat treatment")
    assert deviations.dev_or_404(dev)["state"] == "rejected" and not events
    deviations.update(u["ann"], dev, {"departure": "material 42CrMo4 allowed, QT 800"})
    deviations.submit(u["ann"], dev)
    d = deviations.decide(u["bob"], dev, "approved", "ok")
    assert d["state"] == "active" and d["activated_at"]
    assert [n for n, _ in events] == ["deviation.approved"]
    assert events[0][1]["number"] == d["number"] and events[0][1]["affected"][0]["part_number"] == "MX-100-050"
    _raises("not_draft", deviations.update, u["ann"], dev, {"title": "late change"})
    _raises("closed", deviations.cancel, u["ann"], dev)                              # active ones are closed instead

    root_dev = _ready(q, requester="root", approvers=("ann",))                       # four-eyes binds admins too
    deviations.submit(u["root"], root_dev)
    _raises(403, deviations.decide, u["root"], root_dev, "approved")
    actions = deviations.detail(root_dev, u["root"])["actions"]
    assert not actions["decide"]["ok"] and "four-eyes" in actions["decide"]["why"]

    two = _ready(q, approvers=("ann", "bob"))                                        # every approver must agree
    deviations.submit(u["eve"], two)
    assert deviations.decide(u["ann"], two, "approved")["state"] == "submitted"
    assert deviations.decide(u["bob"], two, "approved")["state"] == "active"
    audit = {r["action"] for r in db.rows("SELECT action FROM audit WHERE entity = 'deviation' AND entity_id = ?", (two,))}
    assert {"deviation_create", "deviation_submit", "deviation_approved"} <= audit


def test_start_date_and_expiry(q, monkeypatch):
    u = q["u"]
    start = date.today() + timedelta(days=3)
    end = start + timedelta(days=7)
    dev = _ready(q, valid_from=start.isoformat(), valid_until=end.isoformat())
    deviations.submit(u["eve"], dev)
    assert deviations.decide(u["ann"], dev, "approved")["state"] == "approved"      # starts later
    assert deviations.sweep(date.today())["activated"] == []
    assert deviations.sweep(start)["activated"] == [deviations.dev_or_404(dev)["number"]]
    assert deviations.active_for("item", q["item"])["deviations"][0]["id"] == dev
    assert deviations.sweep(end)["expired"] == []                                    # the last day still counts
    res = deviations.sweep(end + timedelta(days=1))
    d = deviations.dev_or_404(dev)
    assert res["expired"] == [d["number"]] and d["state"] == "closed" and d["close_reason"] == "expired"
    assert not deviations.active_for("item", q["item"])["deviations"]
    assert any("has ended" in n["title"] for n in notifications.for_user(u["eve"]["id"]))
    assert deviations.sweep(end + timedelta(days=2)) == {"activated": [], "expired": [], "used_up": []}

    late = _ready(q, valid_until=(date.today() + timedelta(days=1)).isoformat())     # end date passes while waiting
    deviations.submit(u["eve"], late)
    monkeypatch.setattr(deviations, "today", lambda: date.today() + timedelta(days=5))
    _raises("expired", deviations.decide, u["ann"], late, "approved")


def test_quantity_usage_and_closing(q):
    u = q["u"]
    dev = _ready(q, qty_limit=10)
    deviations.submit(u["eve"], dev)
    deviations.decide(u["ann"], dev, "approved")
    _raises(403, deviations.record_usage, u["vic"], dev, 1)
    deviations.record_usage(u["eve"], dev, 4, "batch 12")
    assert "4 of 10 used" in deviations.scope_text(deviations.dev_or_404(dev))
    e = _raises("over_limit", deviations.record_usage, u["eve"], dev, 7)
    assert "Only 6 of 10" in e.message
    d = deviations.record_usage(u["eve"], dev, 6, "batch 12, rest")
    assert d["state"] == "closed" and d["close_reason"] == "used_up" and d["qty_used"] == 10

    dev2 = _ready(q, qty_limit=50)
    deviations.submit(u["eve"], dev2)
    deviations.decide(u["ann"], dev2, "approved")
    deviations.record_usage(u["eve"], dev2, 5)
    _raises(403, deviations.close, u["vic"], dev2, 12)
    d = deviations.close(u["eve"], dev2, 12, "Batch finished")
    assert d["state"] == "closed" and d["qty_used"] == 12 and d["close_reason"] == "closed"
    usage = db.rows("SELECT qty, reference FROM deviation_usage WHERE deviation_id = ? ORDER BY id", (dev2,))
    assert [x["qty"] for x in usage] == [5, 7] and "closing" in usage[1]["reference"]  # the difference is booked


def test_banner_text_and_bom_lines(q):
    u = q["u"]
    end = date.today() + timedelta(days=20)
    dev = _ready(q, qty_limit=20, valid_until=end.isoformat())
    deviations.submit(u["eve"], dev)
    d = deviations.decide(u["ann"], dev, "approved")
    assert d["number"] in deviations.banner_text(d)
    assert deviations.banner_text(d).startswith(f"Deviation {d['number']} active: material 42CrMo4 allowed for batch 12, ")
    assert "20 of 20 left" in deviations.banner_text(d) and "until " in deviations.banner_text(d)
    assert deviations.active_for("document", q["doc"])["deviations"][0]["banner"] == deviations.banner_text(d)
    now = db.now()
    asm = db.insert("items", {"part_number": "MX-100-000", "item_type": "assembly", "created_at": now, "updated_at": now})
    bom = db.insert("boms", {"item_id": asm, "revision": "A", "created_at": now, "updated_at": now})
    db.insert("bom_lines", {"bom_id": bom, "line_no": 1, "child_item_id": q["item"], "qty": 2})
    r = deviations.active_for("bom", bom)
    assert r["by_item"] == {str(q["item"]): [{"id": dev, "number": d["number"], "banner": deviations.banner_text(d)}]}


def test_make_permanent_and_links(q):
    u = q["u"]
    dev = _ready(q)
    _raises("not_allowed", deviations.make_permanent, u["eve"], dev)                 # drafts can't
    deviations.submit(u["eve"], dev)
    deviations.decide(u["ann"], dev, "approved")
    _raises(403, deviations.make_permanent, u["vic"], dev)
    res = deviations.make_permanent(u["eve"], dev)
    co = db.one("SELECT * FROM change_orders WHERE id = ?", (res["change_id"],))
    assert co["state"] == "draft" and co["requested_by"] == u["eve"]["id"] and "permanent" in co["title"]
    assert "42CrMo4" in co["description"]
    assert [r["document_id"] for r in db.rows("SELECT document_id FROM co_items WHERE co_id = ?", (co["id"],))] == [q["doc"]]
    assert [r["user_id"] for r in db.rows("SELECT user_id FROM co_approvers WHERE co_id = ?", (co["id"],))] == [u["ann"]["id"]]
    _raises("not_allowed", deviations.make_permanent, u["eve"], dev)                 # only once
    linked = deviations.active_for("change_order", co["id"])["deviations"]
    assert linked[0]["number"] == deviations.dev_or_404(dev)["number"] and linked[0]["relation"] == "made_permanent"

    other = db.insert("change_orders", {"number": "ECO-0099", "title": "older change", "created_at": db.now()})
    _raises(403, deviations.link_change, u["vic"], dev, other)
    deviations.link_change(u["eve"], dev, other)
    deviations.link_change(u["eve"], dev, other)                                     # twice is harmless
    assert [(c["number"], c["relation"]) for c in deviations.detail(dev, u["eve"])["changes"]] == [
        (co["number"], "made_permanent"), ("ECO-0099", "related")]
    deviations.unlink_change(u["eve"], dev, other)
    assert len(deviations.detail(dev, u["eve"])["changes"]) == 1


def test_editing_affected_and_approvers_only_in_draft(q):
    u = q["u"]
    dev = _ready(q, approvers=("ann", "bob"))
    aid = deviations.add_affected(u["eve"], dev, document_id=q["doc"])
    _raises("duplicate", deviations.add_affected, u["eve"], dev, document_id=q["doc"])
    _raises(400, deviations.add_affected, u["eve"], dev)                              # neither part nor file
    assert [a["label"] for a in deviations.affected(dev)] == ["MX-100-050 rev B", "MX-100-050 Shaft Key.dxf rev B"]
    _raises(403, deviations.remove_affected, u["vic"], dev, aid)
    deviations.remove_affected(u["eve"], dev, aid)
    deviations.remove_approver(u["eve"], dev, u["bob"]["id"])
    assert [a["user_id"] for a in deviations.detail(dev, u["eve"])["approvers"]] == [u["ann"]["id"]]
    deviations.update(u["ann"], dev, {"qty_limit": "12", "valid_until": ""})        # an approver may edit a draft
    _raises(400, deviations.update, u["eve"], dev, {"qty_limit": "-1"})
    _raises(400, deviations.update, u["eve"], dev, {"valid_from": "2030-01-02", "valid_until": "2030-01-01"})
    _raises(400, deviations.update, u["eve"], dev, {"valid_until": "30.11.2026"})
    deviations.submit(u["eve"], dev)
    _raises("not_draft", deviations.add_affected, u["eve"], dev, document_id=q["doc"])
    _raises("not_draft", deviations.remove_approver, u["eve"], dev, u["ann"]["id"])
    assert deviations.dev_or_404(dev)["qty_limit"] == 12


def test_attachments_are_stored_and_soft_deleted(q):
    u = q["u"]
    dev = _ready(q)
    aid = deviations.add_attachment(u["eve"], dev, "certificate.pdf", io.BytesIO(b"%PDF-1.4 certificate"))
    a = deviations.attachment_or_404(dev, aid)
    from kosha import store
    assert store.read_bytes(a["blob_sha"]) == b"%PDF-1.4 certificate"
    _raises(403, deviations.remove_attachment, u["vic"], dev, aid)
    _raises(403, deviations.add_attachment, u["vic"], dev, "x.txt", io.BytesIO(b"x"))
    deviations.remove_attachment(u["eve"], dev, aid)
    _raises(404, deviations.attachment_or_404, dev, aid)
    assert db.one("SELECT deleted, blob_sha FROM deviation_attachments WHERE id = ?", (aid,))["deleted"] == 1
    assert store.exists(a["blob_sha"])                                               # the file itself is kept


def test_register_export_filters(q):
    u = q["u"]
    active = _ready(q)
    deviations.submit(u["eve"], active)
    deviations.decide(u["ann"], active, "approved")
    _ready(q, title="Draft one")
    data, name = deviations.report_xlsx({"state": "active"})
    assert name.endswith(".xlsx")
    ws = openpyxl.load_workbook(io.BytesIO(data)).active
    assert ws["A3"].value == "Number" and ws["C3"].value == "Status"
    rows = [r for r in ws.iter_rows(min_row=4, values_only=True) if r[0]]
    assert len(rows) == 1 and rows[0][0] == deviations.dev_or_404(active)["number"] and rows[0][2] == "Active"
    assert "MX-100-050 rev B" in rows[0][3] and "Ann (approved)" in rows[0][13]
    all_rows = [r for r in openpyxl.load_workbook(io.BytesIO(deviations.report_xlsx({})[0])).active.iter_rows(
        min_row=4, values_only=True) if r[0]]
    assert len(all_rows) == 2
    assert len(deviations.listing({"item": "MX-100-050"})) == 2 and deviations.listing({"item": "NOPE"}) == []
    tomorrow = (date.today() + timedelta(days=1)).isoformat()
    assert deviations.listing({"date_from": tomorrow}) == [] and len(deviations.listing({"date_to": tomorrow})) == 2
    _raises(400, deviations.listing, {"state": "bogus"})


# ------------------------------------------------------------------ discussions

def test_mentions_notify_and_replies_reach_participants(q):
    u = q["u"]
    assert [x["username"] for x in discussions.mentioned_users("Hi @ann and @Bob. Also @nobody, @gone, mail a@b.c")] \
        == ["ann", "bob"]
    tid = discussions.start(u["eve"], "document", q["doc"], "@ann can you check the key width? @eve", "Key width")
    t = db.one("SELECT * FROM discussion_threads WHERE id = ?", (tid,))
    assert t["revision"] == "B"                                                      # documents: per revision
    ann = notifications.for_user(u["ann"]["id"])
    assert len(ann) == 1 and ann[0]["kind"] == "mention" and ann[0]["link"] == f"#/doc/{q['doc']}?thread={tid}"
    assert "Eve mentioned you on MX-100-050 Shaft Key.dxf rev B" in ann[0]["title"]
    assert not notifications.for_user(u["eve"]["id"])                                # mentioning yourself: no bell
    discussions.reply(u["ann"], tid, "Width is fine. @bob please confirm the tolerance.")
    assert notifications.for_user(u["eve"]["id"])[0]["kind"] == "discussion_reply"
    assert [n["kind"] for n in notifications.for_user(u["bob"]["id"])] == ["mention"]  # one notification, not two
    view = discussions.threads(u["vic"], "document", q["doc"])
    assert view["open"] == 1 and len(view["threads"][0]["comments"]) == 2
    assert view["threads"][0]["comments"][1]["mentions"] == ["bob"]
    _raises(400, discussions.start, u["eve"], "document", q["doc"], "x", revision="Q")
    _raises(400, discussions.start, u["eve"], "spaceship", 1, "x")
    _raises(404, discussions.start, u["eve"], "item", 9999, "x")
    _raises(400, discussions.start, u["eve"], "item", q["item"], "   ")


def test_edit_window_soft_delete_and_resolve(q, monkeypatch):
    u = q["u"]
    tid = discussions.start(u["eve"], "item", q["item"], "First text")
    cid = db.scalar("SELECT id FROM discussion_comments WHERE thread_id = ?", (tid,))
    _raises(403, discussions.edit, u["ann"], cid, "Not mine")
    _raises(403, discussions.remove, u["ann"], cid)
    _raises(403, discussions.remove, u["root"], cid)                                 # not even an admin
    discussions.edit(u["eve"], cid, "Second text, @ann")
    assert db.scalar("SELECT body FROM discussion_comment_edits WHERE comment_id = ?", (cid,)) == "First text"
    assert notifications.for_user(u["ann"]["id"])[0]["kind"] == "mention"            # new mentions on edit count
    later = datetime.now(timezone.utc) + timedelta(minutes=16)
    monkeypatch.setattr(discussions, "_utcnow", lambda: later)
    _raises("edit_window", discussions.edit, u["eve"], cid, "Third text")
    assert not discussions.threads(u["eve"], "item", q["item"])["threads"][0]["comments"][0]["can_edit"]["ok"]
    discussions.remove(u["eve"], cid)                                                # removing has no time limit
    c = discussions.threads(u["ann"], "item", q["item"])["threads"][0]["comments"][0]
    assert c["removed"] and c["body"] == "comment removed" and c["mentions"] == []
    assert db.scalar("SELECT body FROM discussion_comments WHERE id = ?", (cid,)) == "Second text, @ann"  # kept
    _raises("removed", discussions.remove, u["eve"], cid)
    _raises("removed", discussions.edit, u["eve"], cid, "bring it back")

    _raises(403, discussions.set_status, u["vic"], tid, True)                        # viewers only resolve their own
    discussions.set_status(u["ann"], tid, True)
    _raises("resolved", discussions.reply, u["vic"], tid, "late reply")
    discussions.set_status(u["ann"], tid, False)
    discussions.reply(u["vic"], tid, "late reply")
    own = discussions.start(u["vic"], "item", q["item"], "A viewer's question")
    discussions.set_status(u["vic"], own, True)
    actions = [r["action"] for r in db.rows("SELECT action FROM audit WHERE entity = 'discussion' ORDER BY id")]
    assert actions == ["discussion_start", "discussion_edit", "discussion_remove", "discussion_resolve",
                       "discussion_reopen", "discussion_reply", "discussion_start", "discussion_resolve"]


def test_threads_on_every_record_type(q):
    u = q["u"]
    dev = _ready(q)
    now = db.now()
    co = db.insert("change_orders", {"number": "ECO-0001", "title": "t", "created_at": now})
    bom = db.insert("boms", {"item_id": q["item"], "revision": "B", "created_at": now, "updated_at": now})
    for etype, eid, link in (("item", q["item"], "#/item/"), ("document", q["doc"], "#/doc/"), ("change_order", co, "#/change/"),
                             ("deviation", dev, "#/deviation/"), ("bom", bom, "#/bom/")):
        discussions.start(u["eve"], etype, eid, f"note on {etype}")
        v = discussions.threads(u["eve"], etype, eid)
        assert v["link"] == f"{link}{eid}" and v["threads"][0]["comments"][0]["body"] == f"note on {etype}"


# ------------------------------------------------------------------ notifications

def test_notifications_are_generic_and_private(q):
    u = q["u"]
    a = notifications.notify(u["ann"]["id"], "erp_failed", "Item MX-1 could not be sent", "#/erp", u["eve"]["id"])
    assert notifications.notify(u["gone"]["id"], "x", "inactive users get nothing") is None
    assert notifications.notify(None, "x", "nobody") is None
    b = notifications.notify(u["ann"]["id"], "x", "outside link dropped", "javascript:alert(1)")
    c = notifications.notify(u["ann"]["id"], "x", "protocol-relative dropped", "//example.com/x")
    links = {n["id"]: n["link"] for n in notifications.for_user(u["ann"]["id"])}
    assert links == {a: "#/erp", b: "", c: ""}
    assert notifications.unread_count(u["ann"]["id"]) == 3
    assert notifications.mark_read(u["bob"]["id"], [a]) == 0                          # not yours: nothing happens
    assert notifications.mark_read(u["ann"]["id"], [a]) == 1
    assert [n["id"] for n in notifications.for_user(u["ann"]["id"], unread_only=True)] == [c, b]
    assert notifications.mark_read(u["ann"]["id"]) == 2 and notifications.unread_count(u["ann"]["id"]) == 0
    assert notifications.notify_many([u["ann"]["id"], u["bob"]["id"], u["ann"]["id"]], "k", "t", skip=[u["bob"]["id"]]) \
        and notifications.unread_count(u["ann"]["id"]) == 1 and notifications.unread_count(u["bob"]["id"]) == 0
