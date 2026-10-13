# SPDX-License-Identifier: AGPL-3.0-or-later
"""Deviations, discussions and the notification bell end to end, on the Demo Engineering data."""
import io
import threading
from datetime import date, datetime, timedelta, timezone

import openpyxl
import pytest


@pytest.fixture(scope="module")
def k(demo_kosha):
    s = demo_kosha
    admin = s.client()
    users = {u["username"]: u for u in admin.call("GET", "/api/users")["users"]}
    find = s.client("asha").call
    key = find("GET", "/api/search?q=MX-100-050")["items"][0]
    shaft = [d for d in find("GET", "/api/search?q=MX-100-020")["documents"] if d["filename"].endswith(".dxf")][0]
    bom = [b for b in find("GET", "/api/boms")["boms"] if b["part_number"] == "MX-100-000"][0]
    return {"s": s, "admin": admin, "asha": s.client("asha"), "daniel": s.client("daniel"), "sofia": s.client("sofia"),
            "tom": s.client("tom"), "users": users, "key": key, "shaft": shaft, "bom": bom}


def _prefix():
    return f"DEV-{datetime.now(timezone.utc).year % 100:02d}-"


def test_demo_data_has_deviations_and_a_discussion(k):
    reg = k["tom"].call("GET", "/api/deviations")["deviations"]
    by_state = {d["state"]: d for d in reg}
    assert {"active", "submitted"} <= set(by_state)
    assert by_state["active"]["number"].startswith(_prefix()) and "batch 12" in by_state["active"]["banner"]
    sofia_home = k["sofia"].call("GET", "/api/deviations/home")
    assert [d["state"] for d in sofia_home["waiting_for_me"]] == ["submitted"]
    assert k["sofia"].call("GET", "/api/deviations/summary")["waiting_for_me"] == 1
    threads = k["tom"].call("GET", f"/api/discussions?entity_type=document&entity_id={k['shaft']['id']}")["threads"]
    assert threads and threads[0]["title"] == "Hub seat fretting"
    kinds = [n["kind"] for n in k["sofia"].call("GET", "/api/notifications")["notifications"]]
    assert "mention" in kinds and "deviation_approval" in kinds


def test_deviation_lifecycle_roles_and_banners(k):
    asha, daniel, sofia, tom, u = k["asha"], k["daniel"], k["sofia"], k["tom"], k["users"]
    body = {"title": "Bar stock 42CrMo4 for keys, batch 14", "departure": "material 42CrMo4 allowed instead of C45",
            "reason": "C45 out of stock", "serials": "batch 14", "qty_limit": 5,
            "valid_until": (date.today() + timedelta(days=10)).isoformat(),
            "affected": [{"item_id": k["key"]["id"]}, {"document_id": k["shaft"]["id"]}]}
    tom.call("POST", "/api/deviations", body, expect=403)                            # viewers don't raise deviations
    d = asha.call("POST", "/api/deviations", body)
    assert d["state"] == "draft" and d["number"].startswith(_prefix())
    assert not d["actions"]["submit"]["ok"] and "approver" in d["actions"]["submit"]["why"]
    r = asha.call("POST", f"/api/deviations/{d['id']}/approvers", {"user_id": u["asha"]["id"]}, expect=400)
    assert r["code"] == "four_eyes"
    asha.call("POST", f"/api/deviations/{d['id']}/submit", expect=409)
    asha.call("POST", f"/api/deviations/{d['id']}/approvers", {"user_id": u["daniel"]["id"]})
    d = asha.call("POST", f"/api/deviations/{d['id']}/submit")
    assert d["state"] == "submitted"
    asha.call("POST", f"/api/deviations/{d['id']}/decide", {"decision": "approved"}, expect=403)   # engineer: role
    sofia.call("POST", f"/api/deviations/{d['id']}/decide", {"decision": "approved"}, expect=403)  # not on the list
    assert daniel.call("GET", "/api/notifications?unread=1")["notifications"][0]["link"] == f"#/deviation/{d['id']}"
    d = daniel.call("POST", f"/api/deviations/{d['id']}/decide", {"decision": "approved", "comment": "OK for batch 14"})
    assert d["state"] == "active"
    assert any(n["kind"] == "deviation_approved" for n in asha.call("GET", "/api/notifications")["notifications"])

    item_banner = tom.call("GET", f"/api/deviations/active?entity_type=item&entity_id={k['key']['id']}")
    assert d["id"] in [x["id"] for x in item_banner["deviations"]]
    banner = [x for x in item_banner["deviations"] if x["id"] == d["id"]][0]["banner"]
    assert banner.startswith(f"Deviation {d['number']} active: material 42CrMo4 allowed instead of C45 for batch 14")
    doc_banner = tom.call("GET", f"/api/deviations/active?entity_type=document&entity_id={k['shaft']['id']}")
    assert d["id"] in [x["id"] for x in doc_banner["deviations"]]
    bom_banner = tom.call("GET", f"/api/deviations/active?entity_type=bom&entity_id={k['bom']['id']}")
    assert d["number"] in [x["number"] for x in bom_banner["by_item"][str(k["key"]["id"])]]
    assert str(k["shaft"]["item_id"]) in bom_banner["by_item"]                       # via the shaft drawing

    tom.call("POST", f"/api/deviations/{d['id']}/usage", {"qty": 1}, expect=403)
    r = asha.call("POST", f"/api/deviations/{d['id']}/usage", {"qty": 6}, expect=409)
    assert r["code"] == "over_limit"
    d = asha.call("POST", f"/api/deviations/{d['id']}/usage", {"qty": 2, "reference": "batch 14 / WO 77"})
    assert d["qty_used"] == 2 and d["state"] == "active"
    d = asha.call("POST", f"/api/deviations/{d['id']}/close", {"qty_used": 3, "note": "Batch done"})
    assert d["state"] == "closed" and d["qty_used"] == 3 and d["close_reason"] == "closed"
    after = tom.call("GET", f"/api/deviations/active?entity_type=item&entity_id={k['key']['id']}")
    assert d["id"] not in [x["id"] for x in after["deviations"]]

    perm = asha.call("POST", f"/api/deviations/{d['id']}/make-permanent")
    co = asha.call("GET", f"/api/changes/{perm['change_id']}")
    assert co["state"] == "draft" and d["number"] in co["title"]
    assert perm["changes"][0]["relation"] == "made_permanent"
    linked = tom.call("GET", f"/api/deviations/active?entity_type=change_order&entity_id={co['id']}")["deviations"]
    assert linked[0]["number"] == d["number"]
    asha.call("POST", f"/api/deviations/{d['id']}/make-permanent", expect=409)


def test_attachments_and_register_export(k):
    asha, tom, u = k["asha"], k["tom"], k["users"]
    d = asha.call("POST", "/api/deviations", {"title": "Photo test", "departure": "x", "serials": "SN 1",
                                              "affected": [{"item_id": k["key"]["id"]}], "approver_ids": [u["daniel"]["id"]]})
    d = asha.call("POST", f"/api/deviations/{d['id']}/attachments", files={"file": ("report.txt", b"hardness 32 HRC")})
    att = d["attachments"][0]
    assert att["filename"] == "report.txt" and att["size"] == 15 and "blob_sha" not in att
    raw, _ = tom.call("GET", f"/api/deviations/{d['id']}/attachments/{att['id']}", raw=True)
    assert raw == b"hardness 32 HRC"
    tom.call("DELETE", f"/api/deviations/{d['id']}/attachments/{att['id']}", expect=403)
    d = asha.call("DELETE", f"/api/deviations/{d['id']}/attachments/{att['id']}")
    assert d["attachments"] == []
    tom.call("GET", f"/api/deviations/{d['id']}/attachments/{att['id']}", expect=404)

    raw, ctype = tom.call("GET", "/api/deviations/export?state=active", raw=True)
    assert "spreadsheetml" in ctype
    ws = openpyxl.load_workbook(io.BytesIO(raw)).active
    rows = [r for r in ws.iter_rows(min_row=4, values_only=True) if r[0]]
    assert rows and all(r[2] == "Active" for r in rows)
    raw, _ = tom.call("GET", "/api/deviations/export?item=MX-100-050", raw=True)
    numbers = {r[0] for r in openpyxl.load_workbook(io.BytesIO(raw)).active.iter_rows(min_row=4, values_only=True) if r[0]}
    assert d["number"] in numbers
    tom.call("GET", "/api/deviations/export?state=nonsense", expect=400)
    tom.call("GET", "/api/deviations?date_from=not-a-date", expect=400)


def test_numbers_are_unique_under_parallel_requests(k):
    asha = k["asha"]
    out, errors = [], []

    def make(i):
        try:
            out.append(asha.call("POST", "/api/deviations", {"title": f"parallel {i}"})["number"])
        except AssertionError as e:   # pragma: no cover - reported below
            errors.append(str(e))

    threads = [threading.Thread(target=make, args=(i,)) for i in range(8)]
    for t in threads:
        t.start()
    for t in threads:
        t.join()
    assert not errors and len(set(out)) == 8
    nums = sorted(int(n.rsplit("-", 1)[1]) for n in out)
    assert nums == list(range(nums[0], nums[0] + 8))                                  # no gaps, no doubles


def test_discussion_mentions_bell_edit_and_soft_delete(k):
    asha, daniel, sofia, tom = k["asha"], k["daniel"], k["sofia"], k["tom"]
    before = tom.call("GET", "/api/notifications/count")["unread"]
    r = daniel.call("POST", "/api/discussions", {"entity_type": "item", "entity_id": k["key"]["id"], "title": "Key fit",
                                                 "body": "@tom can the workshop live with a P9 keyway? @nobody"})
    tid = r["thread_id"]
    thread = [t for t in r["threads"] if t["id"] == tid][0]
    first = thread["comments"][0]
    assert first["mentions"] == ["tom"] and first["can_edit"]["ok"]
    assert tom.call("GET", "/api/notifications/count")["unread"] == before + 1
    note = tom.call("GET", "/api/notifications?unread=1")["notifications"][0]
    assert note["kind"] == "mention" and note["link"] == f"#/item/{k['key']['id']}?thread={tid}"
    assert tom.call("POST", f"/api/notifications/{note['id']}/read")["unread"] == before
    asha.call("POST", f"/api/notifications/{note['id']}/read")                       # someone else's: no effect
    r = tom.call("POST", f"/api/discussions/{tid}/comments", {"body": "Yes, P9 is fine."})
    assert len([t for t in r["threads"] if t["id"] == tid][0]["comments"]) == 2
    assert daniel.call("GET", "/api/notifications?unread=1")["notifications"][0]["kind"] == "discussion_reply"

    tom.call("PATCH", f"/api/discussions/comments/{first['id']}", {"body": "changed by tom"}, expect=403)
    tom.call("DELETE", f"/api/discussions/comments/{first['id']}", expect=403)
    k["admin"].call("DELETE", f"/api/discussions/comments/{first['id']}", expect=403)
    r = daniel.call("PATCH", f"/api/discussions/comments/{first['id']}", {"body": "@tom @sofia P9 keyway okay?"})
    c = [t for t in r["threads"] if t["id"] == tid][0]["comments"][0]
    assert c["edited_at"] and c["body"].startswith("@tom @sofia")
    assert sofia.call("GET", "/api/notifications?unread=1")["notifications"][0]["kind"] == "mention"
    r = daniel.call("DELETE", f"/api/discussions/comments/{first['id']}")
    c = [t for t in r["threads"] if t["id"] == tid][0]["comments"][0]
    assert c["removed"] and c["body"] == "comment removed"

    other = asha.call("POST", "/api/discussions", {"entity_type": "deviation", "entity_id": 1, "body": "Who signs off?"})
    tom.call("POST", f"/api/discussions/{other['thread_id']}/resolve", expect=403)   # viewer, not the starter
    r = asha.call("POST", f"/api/discussions/{other['thread_id']}/resolve")
    assert [t for t in r["threads"] if t["id"] == other["thread_id"]][0]["status"] == "resolved"
    asha.call("POST", f"/api/discussions/{other['thread_id']}/comments", {"body": "late"}, expect=409)
    asha.call("POST", f"/api/discussions/{other['thread_id']}/reopen")
    asha.call("POST", f"/api/discussions/{other['thread_id']}/comments", {"body": "Daniel does."})
    audit = k["admin"].call("GET", f"/api/audit?entity=discussion&entity_id={tid}")["entries"]
    assert {e["action"] for e in audit} >= {"discussion_start", "discussion_reply", "discussion_edit", "discussion_remove"}

    assert asha.call("POST", "/api/notifications/read-all")["unread"] == 0
    asha.call("GET", "/api/discussions?entity_type=item", expect=400)
    asha.call("GET", "/api/discussions?entity_type=planet&entity_id=1", expect=400)
