# SPDX-License-Identifier: AGPL-3.0-or-later
"""Legacy migration batches end to end (specification 11): ZIP unpacking, the count check, quality queues, the
review grid and Excel round-trip, old paths and numbers in search, sample audit and sign-off, and Undo batch."""
import io
import time
import zipfile

import openpyxl

from kosha.cad import dxfwriter as W


def _tb(pn, d, m, r):
    return {"PART_NO": pn, "DESCRIPTION": d, "MATERIAL": m, "REV": r}


def _put(root, rel, data):
    p = root / rel
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_bytes(data)


def _zip(entries: dict) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in entries.items():
            z.writestr(name, data)
    return buf.getvalue()


def _wait(c, job_id):
    for _ in range(600):
        j = next(x for x in c.call("GET", "/api/import/jobs")["jobs"] if x["id"] == job_id)
        if j["status"] != "running":
            return j
        time.sleep(0.1)
    raise AssertionError("import did not finish")


def _start(c, src, target, **extra):
    r = c.call("POST", "/api/import/start", fields={"source": str(src), "target": target, "groups": "cad2d,pdf", **extra})
    return _wait(c, r["job"]["id"])


def test_batch_review_reconcile_and_sign_off(kosha, tmp_path):
    admin = kosha.client()
    eng = kosha.add_user("lb-eng", "engineer")
    src = tmp_path / "Old server"
    housing = W.bearing_housing(80, 220, values=_tb("LB-BH-01", "BEARING HOUSING", "FG 260", "Rev. 2"))
    _put(src, "2009/LB-BH-01 Bearing housing.dxf", housing)                                     # full title block
    _put(src, "2009/Copy of LB-BH-01 Bearing housing.dxf", housing)                             # identical copy
    _put(src, "2013/LB-SH-01 Shaft.dxf", W.shaft(values={"PART_NO": "LB-SH-01"}))              # number only
    _put(src, "2013/sketch.dxf", W.flange(120, 90, 4, 10, 50, values={}))                      # nothing known
    _put(src, "2013/sketch.dxf.bak", b"junk")
    _put(src, "2015/Pump order.zip", _zip({
        "Pump/LB-FL-150 Flange.dxf": W.flange(285, 240, 8, 22, 150, values=_tb("LB-FL-150", "FLANGE DN150", "S275JR", "0")),
        "Pump/Thumbs.db": b"junk", "Pump/readme.txt": b"not selected"}))
    _put(src, "2015/Bad.zip", _zip({"../outside.dxf": b"0\nEOF\n"}))

    j = _start(admin, src, "Legacy/Old server")
    assert j["status"] == "done" and j["imported"] == 4 and j["duplicates"] == 1, j

    # every file is accounted for: 7 on disk + 3 inside the good archive
    rec = eng.call("GET", f"/api/import/jobs/{j['id']}/reconcile")
    assert rec["found"] == 10 and rec["adds_up"], rec
    assert rec["by_status"] == {"imported": 4, "identical copies": 1, "set aside": 3, "archives unpacked": 1,
                                "errors": 1}, rec["by_status"]
    assert not rec["can_sign_off"] and any("read back" in b for b in rec["blockers"])

    # the archive's contents went to a folder named after it; the hostile one imported nothing
    data = eng.call("GET", f"/api/import/jobs/{j['id']}/records?status=error")["records"]
    assert "unsafe path" in data[0]["message"]
    flange = eng.call("GET", "/api/search?q=LB-FL-150")["documents"][0]
    assert flange["folder_name"] == "Pump"
    folders = {f["id"]: f for f in admin.call("GET", "/api/folders")["folders"]}
    assert folders[folders[flange["folder_id"]]["parent_id"]]["name"] == "Pump order"

    # quality queues with their reasons
    rows = {r["source_path"].replace("\\", "/").rsplit("/", 1)[-1]: r
            for r in eng.call("GET", f"/api/import/jobs/{j['id']}/records?status=imported")["records"]}
    assert rows["LB-BH-01 Bearing housing.dxf"]["queue"] == "auto" and rows["LB-BH-01 Bearing housing.dxf"]["score"] == 100
    sh = rows["LB-SH-01 Shaft.dxf"]
    assert sh["queue"] in ("review", "attention") and any("No material" in w for w in sh["reasons"])
    sk = rows["sketch.dxf"]
    assert sk["queue"] == "attention" and sk["reasons"][0].startswith("No part number")
    assert "Pump order.zip!/Pump/LB-FL-150 Flange.dxf" in rows["LB-FL-150 Flange.dxf"]["source_path"]
    counts = eng.call("GET", f"/api/import/jobs/{j['id']}/records")["counts"]
    assert counts["auto"] + counts["review"] + counts["attention"] == 4 and counts["confirmed"] == 0

    # old paths find the file (also the identical copy's path, which was not imported)
    assert [d["filename"] for d in eng.call("GET", "/api/search?q=Copy%20of%20LB-BH-01")["documents"]] == \
        ["LB-BH-01 Bearing housing.dxf"]

    # correct one row: the old number stays searchable
    fixed = eng.call("PATCH", f"/api/import/records/{sk['id']}",
                     {"part_number": "lb-fl-120", "description": "FLANGE DN50", "material": "S275JR", "confirm": True})["record"]
    assert (fixed["part_number"], fixed["material"]) == ("LB-FL-120", "S275JR") and fixed["confirmed_at"]
    doc = eng.call("GET", "/api/search?q=LB-FL-120")["documents"][0]
    assert doc["filename"] == "sketch.dxf" and doc["description"] == "FLANGE DN50"
    assert eng.call("GET", "/api/items?q=LB-FL-120")["items"][0]["material"] == "S275JR"
    eng.call("PATCH", f"/api/import/records/{sk['id']}", {"part_number": "bad number!"}, expect=400)
    eng.call("PATCH", f"/api/import/records/{sk['id']}", {"revision": "C"}, expect=400)

    # many rows at once: one field set across them
    n = eng.call("POST", f"/api/import/jobs/{j['id']}/records/confirm",
                 {"ids": [sh["id"]], "field": "material", "value": "42CrMo4"})["confirmed"]
    assert n == 1 and eng.call("GET", "/api/items?q=LB-SH-01")["items"][0]["material"] == "42CrMo4"
    eng.call("POST", f"/api/import/jobs/{j['id']}/records/confirm", {"ids": [sh["id"]], "field": "part_number",
                                                                     "value": "X"}, expect=400)

    # Excel round-trip: a bad row changes nothing; a good file applies
    data, _ = eng.call("GET", f"/api/import/jobs/{j['id']}/records.xlsx", raw=True)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.worksheets[0]
    head = [c.value for c in ws[1]]
    by_file = {row[head.index("File")].value: row for row in ws.iter_rows(min_row=2)}
    by_file["LB-FL-150 Flange.dxf"][head.index("Description")].value = "FLANGE  DN150 PN16"
    by_file["LB-FL-150 Flange.dxf"][head.index("Confirm")].value = "yes"
    by_file["LB-BH-01 Bearing housing.dxf"][head.index("Part number")].value = "no good!"
    buf = io.BytesIO()
    wb.save(buf)
    r = eng.call("POST", f"/api/import/jobs/{j['id']}/records.xlsx", files={"file": ("q.xlsx", buf.getvalue())})
    assert r["applied"] == 0 and r["problems"][0]["row"] >= 2 and "not a part number" in r["problems"][0]["message"]
    assert eng.call("GET", "/api/search?q=LB-FL-150")["documents"][0]["description"] == "FLANGE DN150"
    by_file["LB-BH-01 Bearing housing.dxf"][head.index("Part number")].value = "LB-BH-01"
    buf = io.BytesIO()
    wb.save(buf)
    r = eng.call("POST", f"/api/import/jobs/{j['id']}/records.xlsx", files={"file": ("q.xlsx", buf.getvalue())})
    assert (r["applied"], r["confirmed"], r["problems"]) == (1, 1, [])
    assert eng.call("GET", "/api/search?q=LB-FL-150")["documents"][0]["description"] == "FLANGE DN150 PN16"
    eng.call("POST", f"/api/import/jobs/{j['id']}/records.xlsx", files={"file": ("q.xlsx", b"not excel")}, expect=400)

    # read back, sample, sign off
    eng.call("POST", f"/api/import/jobs/{j['id']}/sign-off", expect=403)
    assert "cannot be signed off" in admin.call("POST", f"/api/import/jobs/{j['id']}/sign-off", expect=409)["error"]
    v = admin.call("POST", f"/api/import/jobs/{j['id']}/verify")
    assert (v["checked"], v["bad"]) == (4, [])
    sample = admin.call("POST", f"/api/import/jobs/{j['id']}/sample")["sample"]
    assert sample["size"] == 4 and sample["checked"] == 0
    pending = [r for r in eng.call("GET", f"/api/import/jobs/{j['id']}/records?status=imported")["records"]
               if r["sample_result"] == "pending"]
    for r in pending:
        eng.call("PUT", f"/api/import/records/{r['id']}/sample", {"result": "ok"})
    done = admin.call("POST", f"/api/import/jobs/{j['id']}/sign-off")
    assert done["signed_off_at"] and done["sample"]["accuracy"] == 100.0
    assert "signed off" in admin.call("POST", f"/api/import/jobs/{j['id']}/undo", expect=409)["error"]


def test_undo_batch_takes_back_only_what_nobody_used(kosha, tmp_path):
    admin = kosha.client()
    eng = kosha.add_user("ub-eng", "engineer")
    src = tmp_path / "Trial"
    _put(src, "UB-100-001 Plate.dxf", W.flange(100, 80, 4, 10, 40, values=_tb("UB-100-001", "PLATE", "S235", "A")))
    _put(src, "Sub/UB-100-002 Cover.dxf", W.flange(140, 110, 6, 10, 60, values=_tb("UB-100-002", "COVER", "S235", "A")))
    j = _start(admin, src, "Trial batch", state="in_work")
    assert j["imported"] == 2
    used = eng.call("GET", "/api/search?q=UB-100-002")["documents"][0]
    eng.call("POST", f"/api/documents/{used['id']}/checkout")

    plan = admin.call("GET", f"/api/import/jobs/{j['id']}/undo")
    assert [x["filename"] for x in plan["remove"]] == ["UB-100-001 Plate.dxf"]
    assert plan["keep"][0]["why"] == "it is checked out"
    eng.call("POST", f"/api/import/jobs/{j['id']}/undo", expect=403)
    r = admin.call("POST", f"/api/import/jobs/{j['id']}/undo")
    assert (r["removed"], r["parts_removed"], len(r["kept"])) == (1, 1, 1)
    found = eng.call("GET", "/api/search?q=UB-100-001")
    assert (found["documents"], found["items"]) == ([], [])
    assert eng.call("GET", "/api/search?q=UB-100-002")["documents"]                       # the used one stays
    assert "undone already" in admin.call("POST", f"/api/import/jobs/{j['id']}/undo", expect=409)["error"]
    eng.call("PATCH", f"/api/import/records/{plan['keep'][0]['record']}", {"description": "X"}, expect=409)

    # the file waits in the recycle bin; put back, it gets its part again
    gone = plan["remove"][0]["document_id"]
    admin.call("POST", f"/api/recycle-bin/{gone}/restore")
    assert eng.call("GET", "/api/items?q=UB-100-001")["items"][0]["description"] == "PLATE"
