# SPDX-License-Identifier: AGPL-3.0-or-later
"""Inputs library: reference material per part or per project.

Competitor benchmarks, customer requirements, supplier datasheets, test reports, photos, meeting notes,
standards: material that informs a design but is not itself a controlled design file. Inputs have no
revisions or release lifecycle, but they are kept safely, found quickly and linked to what they inform.

* ``InputLibrary``  the records: create, edit, tag, link, restrict, search, soft-delete.
* ``TextExtractor`` reads text from files (one reader per file type) and makes thumbnails. It runs in a
  separate process with a time limit, after the upload has been answered, so an upload never waits for it.
* ``SearchIndex``   a portable term table (no database-specific full-text engine): normalised lower-case
  terms, simple English plural stemming, and part-number-like tokens (``MX-100-020``, ``M16x60``,
  ``Ø110``) kept intact.

Files go through the content-addressed store, so an identical file is stored once; its text and thumbnail
are worked out once per fingerprint. Restricted inputs are filtered in every query (``visible_sql``),
never only on screen.
"""
from __future__ import annotations

import base64
import codecs
import html.parser
import io
import json
import logging
import math
import os
import re
import subprocess
import sys
import threading
import time
import unicodedata
import urllib.parse
import xml.parsers.expat
import zipfile
from collections import Counter
from dataclasses import dataclass
from datetime import datetime, timezone
from functools import lru_cache
from typing import BinaryIO, Callable, Iterable

from . import config, db, store
from .web import ROLE_LEVEL, ApiError, bad, forbidden, not_found

log = logging.getLogger("kosha.inputs")

# ------------------------------------------------------------------ settings --

SETTINGS_KEY = "inputs"
DEFAULT_KINDS = [
    {"key": "competitor_benchmark", "label": "Competitor benchmark"},
    {"key": "customer_requirement", "label": "Customer requirement"},
    {"key": "supplier_datasheet", "label": "Supplier datasheet"},
    {"key": "test_report", "label": "Test report"},
    {"key": "photo", "label": "Photo"},
    {"key": "meeting_note", "label": "Meeting note"},
    {"key": "standard", "label": "Standard"},
    {"key": "other", "label": "Other"},
]
DEFAULTS = {
    "kinds": DEFAULT_KINDS,
    "ocr": True,               # read text in photos and scanned PDFs (needs Tesseract)
    "ocr_lang": "eng",
    "ocr_max_pages": 10,       # scanned PDF pages read with OCR
    "max_mb": 100,             # larger files are stored but their text is not read
    "time_limit_s": 120,       # per file; the reader is stopped after this
    "max_chars": 1_000_000,    # text kept per file
}
LIMITS = {"ocr_max_pages": (0, 200), "max_mb": (1, 4096), "time_limit_s": (5, 3600), "max_chars": (10_000, 20_000_000)}

TYPE_GROUPS = {
    "pdf": ("PDF", ("pdf",)),
    "image": ("Images", ("png", "jpg", "jpeg", "gif", "bmp", "webp", "tif", "tiff", "heic")),
    "spreadsheet": ("Spreadsheets", ("xlsx", "xlsm", "xls", "csv", "tsv", "ods")),
    "document": ("Documents", ("docx", "doc", "odt", "txt", "rtf", "md")),
    "presentation": ("Presentations", ("pptx", "ppt", "odp")),
    "cad": ("CAD files", ("dwg", "dxf", "step", "stp", "igs", "iges", "stl", "sldprt", "sldasm", "slddrw", "x_t")),
    "email": ("E-mails", ("eml", "msg")),
}
TYPE_LABEL = {**{k: v[0] for k, v in TYPE_GROUPS.items()}, "other": "Other files", "link": "Web links", "note": "Notes"}
KNOWN_EXTS = tuple(e for _, exts in TYPE_GROUPS.values() for e in exts)
BROWSER_IMAGES = ("png", "jpg", "jpeg", "gif", "webp", "bmp")
INLINE_MIME = {"pdf": "application/pdf", "png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg",
               "gif": "image/gif", "webp": "image/webp", "bmp": "image/bmp", "svg": "image/svg+xml"}
STATUS_LABEL = {"pending": "Waiting to be read", "running": "Reading text…", "done": "Text read",
                "empty": "No text found", "skipped": "Text not read", "failed": "Could not read the text"}
CONFIDENTIALITY = ("normal", "restricted")
MAX_TAGS = 30


def type_group(ext: str) -> str:
    for key, (_, exts) in TYPE_GROUPS.items():
        if ext in exts:
            return key
    return "other"


def _today() -> str:
    return datetime.now(timezone.utc).date().isoformat()


def _ext(filename: str) -> str:
    name = filename.rsplit("/", 1)[-1]
    return name.rsplit(".", 1)[-1].lower()[:10] if "." in name else ""


def _clip(value, n: int) -> str:
    return str(value if value is not None else "").strip()[:n]


# ===================================================================== search index ==

STOPWORDS = frozenset("""a an and are as at be by for from has have in into is it its of on or that the this to was were
will with not no but if than then there these those which who what when where how all any can our we you your
""".split())

_W = r"[^\W_]+"
# joins kept inside one token: MX-100-020, DN150/PN16, 1.4401, 3,5, M16 x 60 (only between digits)
_JOIN = r"(?:[-\u2010\u2011\u2013/]|(?<=\d)[.,](?=\d)|(?<=\d)[ \t]*[xX\u00d7][ \t]*(?=\d))"
TOKEN_RE = re.compile(rf"\u2300?{_W}(?:{_JOIN}{_W})*")
_TRANS = str.maketrans({"\u2300": "ø", "\u00d7": "x", "\u2010": "-", "\u2011": "-", "\u2013": "-"})
_SEG_SPLIT = re.compile(r"[-/]")
_PIECES = re.compile(r"\d+|[^\W\d_]+")
_DECIMAL = re.compile(r"^\d+(?:[.,]\d+)+$")


def stem(word: str) -> str:
    """Very small English plural stemmer: pumps -> pump, batteries -> battery, boxes -> box."""
    if len(word) <= 3 or not word.isalpha():
        return word
    if word.endswith("ies") and len(word) > 4:
        return word[:-3] + "y"
    if word.endswith("sses"):
        return word[:-2]
    if word.endswith(("ches", "shes", "xes", "zes")):
        return word[:-2]
    if word.endswith(("ss", "us", "is")):
        return word
    if word.endswith("s"):
        return word[:-1]
    return word


def normalise(token: str) -> str:
    t = unicodedata.normalize("NFKC", token).translate(_TRANS).lower()
    if " " in t or "\t" in t:
        t = re.sub(r"[ \t]+", "", t)
    return stem(t) if t.isalpha() else t


def _keep(term: str) -> bool:
    return (len(term) >= 2 or term.isdigit()) and len(term) <= 64 and term not in STOPWORDS


def atoms(term: str) -> set[str]:
    """Searchable parts of a compound token: mx-100-020 -> mx, 100, 020; m16x60 -> 16, 60; ø110 -> 110."""
    out: set[str] = set()
    if _DECIMAL.match(term):
        return out
    for seg in _SEG_SPLIT.split(term):
        if not seg:
            continue
        if seg != term:
            out.add(stem(seg) if seg.isalpha() else seg)
        pieces = _PIECES.findall(seg)
        if len(pieces) > 1:
            out.update(stem(p) if p.isalpha() else p for p in pieces)
    return {a for a in out if a != term and (a.isdigit() or len(a) >= 2) and a not in STOPWORDS}


def tokens(text: str, limit: int | None = None) -> Iterable[str]:
    for m in TOKEN_RE.finditer(text, 0, limit or len(text)):
        t = normalise(m.group())
        if _keep(t):
            yield t


def term_matches(candidate: str, term: str) -> bool:
    return candidate == term or (len(term) >= 3 and candidate.startswith(term))


class SearchIndex:
    """Portable search index made of two plain term tables.

    ``input_index`` holds the terms of an input's own fields (title, tags, source, kind, comment, note, web
    link, file names, linked part numbers and descriptions, project codes, names and customers); it is
    rebuilt whenever one of them changes, which is cheap. ``input_text_terms`` holds the terms of the text
    read from a file, once per file fingerprint, so a long PDF is indexed once however many inputs hold it
    and editing a title never re-reads it. A search matches each query term in either table.
    """

    FIELD_WEIGHTS = {"title": 6.0, "tag": 4.0, "part": 4.0, "project": 4.0, "file": 3.0, "source": 2.5,
                     "project_name": 2.0, "kind": 1.5, "item_desc": 1.5, "customer": 1.5, "comment": 1.5,
                     "note": 1.2, "url": 1.0, "text": 1.0}
    MAX_TERMS = 20_000
    PREFIX_FACTOR = 0.6

    # -- building --
    def weigh(self, fields: Iterable[tuple[str, str]], limit: int | None = None) -> dict[str, float]:
        """Term weights: field weight x (1 + ln tf), capped; parts of compound tokens count half."""
        weights: dict[str, float] = {}
        for field, text in fields:
            if not text:
                continue
            base = self.FIELD_WEIGHTS.get(field, 1.0)
            counts: Counter = Counter(tokens(text, limit))
            parts: Counter = Counter()
            for t, n in counts.items():
                for a in atoms(t):
                    parts[a] += n
            for t, n in counts.items():
                weights[t] = weights.get(t, 0.0) + base * min(1 + math.log(n), 4.0)
            for t, n in parts.items():
                weights[t] = weights.get(t, 0.0) + 0.5 * base * min(1 + math.log(n), 4.0)
        if len(weights) > self.MAX_TERMS:
            weights = dict(sorted(weights.items(), key=lambda kv: -kv[1])[:self.MAX_TERMS])
        return weights

    @staticmethod
    def write_input(input_id: int, weights: dict[str, float]) -> None:
        with db.tx() as conn:
            conn.execute("DELETE FROM input_index WHERE input_id = ?", (input_id,))
            conn.executemany("INSERT INTO input_index (input_id, term, weight) VALUES (?, ?, ?)",
                             [(input_id, t, round(w, 3)) for t, w in weights.items()])

    @staticmethod
    def write_text(sha: str, weights: dict[str, float]) -> None:
        with db.tx() as conn:
            conn.execute("DELETE FROM input_text_terms WHERE sha = ?", (sha,))
            conn.executemany("INSERT INTO input_text_terms (sha, term, weight) VALUES (?, ?, ?)",
                             [(sha, t, round(w, 3)) for t, w in weights.items()])

    # -- querying --
    @staticmethod
    def query_terms(q: str) -> list[str]:
        seen: list[str] = []
        for m in TOKEN_RE.finditer(q or ""):
            t = normalise(m.group())
            if (len(t) >= 2 or t.isdigit()) and len(t) <= 64 and t not in seen:
                seen.append(t)
        useful = [t for t in seen if t not in STOPWORDS]
        return (useful or seen)[:12]

    def match_sql(self, terms: list[str], alias: str = "i") -> tuple[list[str], list, str, list]:
        """WHERE clauses (every term must match the input's fields or the text of one of its files) and a
        relevance score expression. Terms of three or more characters also match longer terms (prefix)."""
        where, wparams, scores, sparams = [], [], [], []
        own = f"FROM input_index x WHERE x.input_id = {alias}.id AND "
        text = (f"FROM input_files f JOIN input_text_terms x ON x.sha = f.blob_sha "
                f"WHERE f.input_id = {alias}.id AND f.removed = 0 AND ")
        for t in terms:
            cond, p = ("x.term LIKE ?", [t + "%"]) if len(t) >= 3 else ("x.term = ?", [t])
            where.append(f"(EXISTS (SELECT 1 {own}{cond}) OR EXISTS (SELECT 1 {text}{cond}))")
            wparams += p + p
            weight = f"COALESCE(SUM(x.weight * CASE WHEN x.term = ? THEN 1.0 ELSE {self.PREFIX_FACTOR} END), 0)"
            scores.append(f"(SELECT {weight} {own}{cond}) + (SELECT {weight} {text}{cond})")
            sparams += [t, *p, t, *p]
        return where, wparams, " + ".join(scores) or "0", sparams


    # -- highlighting --
    @staticmethod
    def _hit(token_text: str, terms: list[str]) -> bool:
        t = normalise(token_text)
        cands = [t, *atoms(t)]
        return any(term_matches(c, q) for c in cands for q in terms)

    def highlight(self, text: str, terms: list[str]) -> list[list]:
        """Split text into [segment, is_hit] pairs (rendered with <mark> by the client, never as HTML)."""
        out: list[list] = []
        pos = 0
        if terms:
            for m in TOKEN_RE.finditer(text):
                if self._hit(m.group(), terms):
                    if m.start() > pos:
                        out.append([text[pos:m.start()], False])
                    out.append([m.group(), True])
                    pos = m.end()
        if pos < len(text):
            out.append([text[pos:], False])
        return out

    @staticmethod
    def _anchors(term: str) -> list[str]:
        out = [term]
        if term.isalpha() and len(term) > 4:
            out.append(term[:-1])                      # battery -> batter (finds batteries)
        m = re.match(r"ø?([^\W\d_]*\d+|[^\W\d_]+)", term)
        if m and m.group(1) != term and len(m.group(1)) >= 2:
            out.append(m.group(1))                     # m16x60 -> m16 (finds "M16 x 60"), ø110 -> 110
        return out

    def snippet(self, text: str, terms: list[str], width: int = 240) -> list[list] | None:
        """The best window of ``text`` around the query terms, with the hits marked."""
        if not text or not terms:
            return None
        low = text.lower()
        cands: list[int] = []
        for t in terms:
            for a in self._anchors(t):
                start, found = 0, 0
                while found < 15:
                    p = low.find(a, start)
                    if p < 0:
                        break
                    cands.append(p)
                    start, found = p + 1, found + 1
                if found:
                    break
        best = None
        for p in sorted(set(cands))[:60]:
            a = max(0, p - width // 3)
            b = min(len(text), a + width)
            hits = []
            for m in TOKEN_RE.finditer(text, a, b):
                if self._hit(m.group(), terms):
                    hits.append((m.start(), m.end(), normalise(m.group())))
            if not hits:
                continue
            distinct = len({q for _, _, n in hits for q in terms if term_matches(n, q) or
                            any(term_matches(x, q) for x in atoms(n))})
            score = distinct * 100 + len(hits)
            if best is None or score > best[0]:
                best = (score, a, b, hits)
        if not best:
            return None
        _, a, b, hits = best
        while a > 0 and not text[a - 1].isspace() and a > best[1] - 20:
            a -= 1
        while b < len(text) and not text[b].isspace() and b < best[2] + 20:
            b += 1
        segs: list[list] = []
        pos = a
        for s, e, _ in hits:
            if s < pos:
                continue
            if s > pos:
                segs.append([text[pos:s], False])
            segs.append([text[s:e], True])
            pos = e
        if pos < b:
            segs.append([text[pos:b], False])
        for seg in segs:
            seg[0] = re.sub(r"\s+", " ", seg[0])
        if a > 0:
            segs.insert(0, ["… ", False])
        if b < len(text):
            segs.append([" …", False])
        return segs


# ===================================================================== text extraction ==

class ExtractError(Exception):
    """A file that cannot be read; the message is shown to people."""


class _Stop(Exception):
    pass


@dataclass
class Extracted:
    status: str                  # done | empty | skipped | failed
    text: str = ""
    method: str = ""
    note: str = ""
    pages: int | None = None
    file_type: str = ""
    thumb: tuple[str, bytes] | None = None
    thumb_original: bool = False

    def to_json(self) -> dict:
        d = {k: getattr(self, k) for k in ("status", "text", "method", "note", "pages", "file_type", "thumb_original")}
        d["thumb"] = [self.thumb[0], base64.b64encode(self.thumb[1]).decode()] if self.thumb else None
        return d

    @classmethod
    def from_json(cls, d: dict) -> "Extracted":
        thumb = (d["thumb"][0], base64.b64decode(d["thumb"][1])) if d.get("thumb") else None
        return cls(status=d.get("status", "failed"), text=d.get("text") or "", method=d.get("method") or "",
                   note=d.get("note") or "", pages=d.get("pages"), file_type=d.get("file_type") or "",
                   thumb=thumb, thumb_original=bool(d.get("thumb_original")))


def _refuse_dtd(*_a):
    raise ExtractError("The file contains an XML document type declaration, which Kosha does not read.")


def _xml_parser(namespaces: bool = True):
    """An expat parser that refuses document type declarations, so entity expansion attacks cannot work."""
    p = xml.parsers.expat.ParserCreate(namespace_separator="|") if namespaces else xml.parsers.expat.ParserCreate()
    p.StartDoctypeDeclHandler = _refuse_dtd
    p.EntityDeclHandler = _refuse_dtd
    return p


def _xml_text(stream: BinaryIO, *, text_tags: set | None = None, para_tags: Iterable = (), tab_tags: Iterable = (),
              break_tags: Iterable = (), space_tags: Iterable = (), limit: int = 1_000_000) -> str:
    """Stream text out of an XML part (Office Open XML, OpenDocument). Document type declarations are
    refused, so entity expansion attacks cannot work."""
    out: list[str] = []
    size = 0
    depth = 0
    para, tab, brk, space = set(para_tags), set(tab_tags), set(break_tags), set(space_tags)
    p = _xml_parser()

    def start(name, attrs):
        nonlocal depth
        local = name.rsplit("|", 1)[-1]
        if text_tags is not None and local in text_tags:
            depth += 1
        elif local in tab:
            out.append("\t")
        elif local in brk:
            out.append("\n")
        elif local in space:
            out.append(" ")

    def end(name):
        nonlocal depth
        local = name.rsplit("|", 1)[-1]
        if text_tags is not None and local in text_tags:
            depth -= 1
        if local in para:
            out.append("\n")

    def chars(data):
        nonlocal size
        if text_tags is None or depth > 0:
            out.append(data)
            size += len(data)
            if size > limit:
                raise _Stop

    p.StartElementHandler = start
    p.EndElementHandler = end
    p.CharacterDataHandler = chars
    try:
        p.ParseFile(stream)
    except _Stop:
        pass
    except xml.parsers.expat.ExpatError as e:
        raise ExtractError(f"The file is damaged ({e}).")
    return "".join(out)


def _decode(data: bytes) -> str:
    if data.startswith(codecs.BOM_UTF8):
        return data[3:].decode("utf-8", "replace")
    if data.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return data.decode("utf-16", "replace")
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("cp1252", "replace")


class _HtmlText(html.parser.HTMLParser):
    BLOCK = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "table", "section", "article"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.out: list[str] = []
        self.skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style", "head"):
            self.skip += 1
        elif tag in self.BLOCK:
            self.out.append("\n")

    def handle_endtag(self, tag):
        if tag in ("script", "style", "head") and self.skip:
            self.skip -= 1
        elif tag in self.BLOCK:
            self.out.append("\n")

    def handle_data(self, data):
        if not self.skip:
            self.out.append(data)


def html_to_text(markup: str) -> str:
    p = _HtmlText()
    p.feed(markup)
    p.close()
    return re.sub(r"\n\s*\n\s*\n+", "\n\n", "".join(p.out))


@lru_cache(maxsize=1)
def ocr_status() -> tuple[bool, str]:
    """(available, why not) for reading text in images."""
    try:
        import pytesseract  # optional
        from PIL import Image  # noqa: F401  optional
    except ImportError:
        return False, ("Reading text in images needs Tesseract and the Python packages pytesseract and Pillow on the "
                       "server. Everything else works without them.")
    try:
        pytesseract.get_tesseract_version()
    except Exception:
        return False, ("The Tesseract program was not found on the server. Install it (for example "
                       "'apt install tesseract-ocr') to read text in images and scanned PDFs.")
    return True, ""


class TextExtractor:
    """Reads the text of one file. One reader per file type, registered with ``TextExtractor.register``;
    an add-in or a later version can register more. ``run`` never raises: problems become a status and a
    note that people can read."""

    readers: dict[str, Callable[["TextExtractor", bytes, str], Extracted]] = {}
    ZIP_PART_LIMIT = 300 * 1024 * 1024       # uncompressed size of one part we are willing to read
    CAD_EXTS = ("dwg", "dxf", "step", "stp", "igs", "iges", "stl", "sldprt", "sldasm", "slddrw", "x_t")
    PREVIEW_EXTS = ("pdf", "dxf", "dwg", "stl", "sldprt", "sldasm", "slddrw")
    TEXT_EXTS = ("txt", "csv", "tsv", "md", "log", "json", "xml", "ini", "cfg", "yaml", "yml", "rtf")

    def __init__(self, options: dict | None = None):
        self.opt = {**DEFAULTS, **(options or {})}
        self.max_chars = int(self.opt["max_chars"])

    @classmethod
    def register(cls, file_type: str):
        def deco(fn):
            cls.readers[file_type] = fn
            return fn
        return deco

    # -- type detection (by content first, the extension only as a hint) --
    def detect(self, data: bytes, filename: str) -> str:
        ext = _ext(filename)
        head = data[:16]
        if head.startswith(b"%PDF") or b"%PDF-" in data[:1024]:
            return "pdf"
        if head.startswith((b"\x89PNG", b"\xff\xd8\xff", b"GIF87a", b"GIF89a", b"II*\x00", b"MM\x00*")) or \
                (head[:4] == b"RIFF" and data[8:12] == b"WEBP") or (head.startswith(b"BM") and ext == "bmp"):
            return "image"
        if ext in self.CAD_EXTS:
            return "cad"
        if head.startswith(b"PK\x03\x04"):
            try:
                with zipfile.ZipFile(io.BytesIO(data)) as zf:
                    names = set(zf.namelist())
                    if "word/document.xml" in names:
                        return "docx"
                    if "xl/workbook.xml" in names:
                        return "xlsx"
                    if "ppt/presentation.xml" in names:
                        return "pptx"
                    if "content.xml" in names and "mimetype" in names and \
                            zf.read("mimetype")[:60].startswith(b"application/vnd.oasis.opendocument"):
                        return "odf"
            except (zipfile.BadZipFile, OSError, KeyError):
                return "damaged"
            return "zip"
        if head.startswith(b"\xd0\xcf\x11\xe0"):
            return "ole"
        if ext in ("html", "htm"):
            return "html"
        if ext == "eml":
            return "email"
        sample = data[:8192]
        if ext in self.TEXT_EXTS or (sample and b"\x00" not in sample and self._mostly_text(sample)):
            return "text"
        return "binary"

    @staticmethod
    def _mostly_text(sample: bytes) -> bool:
        try:
            s = sample.decode("utf-8")
        except UnicodeDecodeError as e:
            if e.start < len(sample) - 4:      # not just a character cut at the end of the sample
                return False
            s = sample[:e.start].decode("utf-8", "ignore")
        printable = sum(1 for c in s if c.isprintable() or c in "\r\n\t")
        return bool(s) and printable / len(s) > 0.95

    # -- the whole job for one file --
    def run(self, data: bytes, filename: str) -> Extracted:
        try:
            file_type = self.detect(data, filename)
        except Exception:
            file_type = "binary"
        reader = self.readers.get(file_type)
        try:
            if reader is None:
                res = Extracted("skipped", note="Kosha can't read text from this type of file. Its title, tags, "
                                                "source and links are still searchable.")
            else:
                res = reader(self, data, filename)
        except ExtractError as e:
            res = Extracted("failed", note=str(e))
        except Exception as e:  # a damaged file must never stop the worker
            log.warning("Reading text from %s failed: %s", filename, e)
            res = Extracted("failed", note=f"The file could not be read ({e.__class__.__name__}). It is stored safely; "
                                           "its title, tags and links are still searchable.")
        res.file_type = file_type
        if res.text:
            res.text = self._tidy(res.text)[:self.max_chars]
            if res.status == "done" and not res.text.strip():
                res.status = "empty"
        elif res.status == "done":
            res.status = "empty"
        try:
            thumb = self.thumbnail(data, filename, file_type)
            if thumb == "original":
                res.thumb_original = True
            elif thumb:
                res.thumb = thumb
        except Exception as e:
            log.info("No thumbnail for %s: %s", filename, e)
        return res

    @staticmethod
    def _tidy(text: str) -> str:
        text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
        text = re.sub(r"[ \t\f\v]+\n", "\n", text)
        return re.sub(r"\n{4,}", "\n\n\n", text).strip()

    # -- helpers for readers --
    def _zip_part(self, zf: zipfile.ZipFile, name: str):
        info = zf.getinfo(name)
        if info.file_size > self.ZIP_PART_LIMIT:
            raise ExtractError(f"Part {name} is larger than {self.ZIP_PART_LIMIT // 2**20} MB once unpacked; not read.")
        return zf.open(info)

    def ocr_image(self, img) -> str:
        import pytesseract
        img = img.convert("L")
        w, h = img.size
        if max(w, h) > 3500:
            img.thumbnail((3500, 3500))
        elif max(w, h) < 1000:
            img = img.resize((w * 2, h * 2))
        timeout = max(5, min(int(self.opt["time_limit_s"]) - 2, 90))
        d = pytesseract.image_to_data(img, lang=self.opt.get("ocr_lang") or "eng", timeout=timeout,
                                      output_type=pytesseract.Output.DICT)
        lines: dict[tuple, list[str]] = {}
        for i, word in enumerate(d.get("text", [])):
            w_ = (word or "").strip()
            try:
                conf = float(d["conf"][i])
            except (TypeError, ValueError):
                conf = -1
            if not w_ or conf < 55 or not any(ch.isalnum() for ch in w_):
                continue
            lines.setdefault((d["block_num"][i], d["par_num"][i], d["line_num"][i]), []).append(w_)
        return "\n".join(" ".join(ws) for ws in lines.values())

    def ocr_allowed(self) -> tuple[bool, str]:
        if not self.opt.get("ocr"):
            return False, "Reading text in images (OCR) is switched off in the inputs settings."
        return ocr_status()

    # -- thumbnails (reusing the preview code of the Explorer for PDF and CAD files) --
    def thumbnail(self, data: bytes, filename: str, file_type: str):
        ext = _ext(filename)
        if file_type == "image":
            try:
                from PIL import Image, ImageOps
            except ImportError:
                return "original" if ext in BROWSER_IMAGES and len(data) <= 3_000_000 else None
            try:
                img = Image.open(io.BytesIO(data))
                img = ImageOps.exif_transpose(img)
                img.thumbnail((480, 480))
                buf = io.BytesIO()
                if img.mode in ("RGBA", "LA", "P"):
                    img.convert("RGBA").save(buf, format="PNG", optimize=True)
                    return ("png", buf.getvalue())
                img.convert("RGB").save(buf, format="JPEG", quality=82)
                return ("jpg", buf.getvalue())
            except Exception:
                return "original" if ext in BROWSER_IMAGES and len(data) <= 3_000_000 else None
        if file_type == "pdf" or ext in self.PREVIEW_EXTS:
            from .cad import process as cadproc
            r = cadproc.process(filename if file_type != "pdf" else "file.pdf", data)
            if r.get("preview") and r["preview"][0] in ("png", "svg", "jpg", "jpeg"):
                return r["preview"]
        return None


@TextExtractor.register("pdf")
def _read_pdf(self: TextExtractor, data: bytes, filename: str) -> Extracted:
    pages: list[str] = []
    n_pages = 0
    method = ""
    total = 0
    try:
        import pypdfium2 as pdfium
        try:
            pdf = pdfium.PdfDocument(data)
        except Exception as e:
            if "password" in str(e).lower():
                raise ExtractError("The PDF is protected with a password, so its text can't be read.")
            raise ExtractError(f"The PDF could not be opened ({e}).")
        try:
            n_pages = len(pdf)
            for i in range(n_pages):
                page = pdf[i]
                try:
                    tp = page.get_textpage()
                    try:
                        t = tp.get_text_range()
                    finally:
                        tp.close()
                finally:
                    page.close()
                pages.append(t)
                total += len(t)
                if total > self.max_chars:
                    break
        finally:
            pdf.close()
        method = "PDF text"
    except ImportError:
        try:
            from pypdf import PdfReader
        except ImportError:
            return Extracted("skipped", note="Reading PDF text needs pypdfium2 (or pypdf) on the server.")
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ExtractError("The PDF is protected with a password, so its text can't be read.")
        n_pages = len(reader.pages)
        for page in reader.pages:
            t = page.extract_text() or ""
            pages.append(t)
            total += len(t)
            if total > self.max_chars:
                break
        method = "PDF text"
    text = "\n\n".join(pages).strip()
    alnum = sum(1 for c in text if c.isalnum())
    if alnum >= max(20, 10 * min(n_pages, 5)):
        return Extracted("done", text, method, pages=n_pages)
    # no usable text layer: a scan. Read the first pages with OCR when allowed.
    ok, why = self.ocr_allowed()
    if not ok or int(self.opt["ocr_max_pages"]) <= 0:
        return Extracted("empty", text, method, note="This PDF has no text layer (probably a scan). " + (why or ""),
                         pages=n_pages)
    try:
        import pypdfium2 as pdfium
    except ImportError:
        return Extracted("empty", text, method, note="This PDF has no text layer; reading scans needs pypdfium2.",
                         pages=n_pages)
    pdf = pdfium.PdfDocument(data)
    out = []
    try:
        limit = min(len(pdf), int(self.opt["ocr_max_pages"]))
        for i in range(limit):
            page = pdf[i]
            try:
                img = page.render(scale=250 / 72).to_pil()
            finally:
                page.close()
            out.append(self.ocr_image(img))
    finally:
        pdf.close()
    note = ""
    if n_pages > int(self.opt["ocr_max_pages"]):
        note = f"Scanned PDF: the first {int(self.opt['ocr_max_pages'])} of {n_pages} pages were read with OCR."
    return Extracted("done", "\n\n".join(out), "OCR (scanned PDF)", note=note, pages=n_pages)


@TextExtractor.register("image")
def _read_image(self: TextExtractor, data: bytes, filename: str) -> Extracted:
    ok, why = self.ocr_allowed()
    if not ok:
        return Extracted("skipped", note=why)
    from PIL import Image, ImageOps
    try:
        img = ImageOps.exif_transpose(Image.open(io.BytesIO(data)))
    except Exception as e:
        raise ExtractError(f"The image could not be opened ({e.__class__.__name__}).")
    text = self.ocr_image(img)
    return Extracted("done", text, "OCR")


@TextExtractor.register("docx")
def _read_docx(self: TextExtractor, data: bytes, filename: str) -> Extracted:
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        names = zf.namelist()
        parts = ["word/document.xml"] + sorted(n for n in names if re.match(
            r"word/(header|footer|footnotes|endnotes|comments)\d*\.xml$", n))
        texts = []
        for n in parts:
            if n in names:
                with self._zip_part(zf, n) as f:
                    texts.append(_xml_text(f, text_tags={"t"}, para_tags={"p"}, tab_tags={"tab"},
                                           break_tags={"br", "cr"}, limit=self.max_chars))
    return Extracted("done", "\n".join(texts), "Word document")


def _numbered(names: Iterable[str], pattern: str) -> list[str]:
    found = [(int(m.group(1)), n) for n in names if (m := re.match(pattern, n))]
    return [n for _, n in sorted(found)]


@TextExtractor.register("pptx")
def _read_pptx(self: TextExtractor, data: bytes, filename: str) -> Extracted:
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        names = zf.namelist()
        slides = _numbered(names, r"ppt/slides/slide(\d+)\.xml$")
        notes = {int(m.group(1)): n for n in names if (m := re.match(r"ppt/notesSlides/notesSlide(\d+)\.xml$", n))}
        out = []
        for i, n in enumerate(slides, 1):
            with self._zip_part(zf, n) as f:
                out.append(f"Slide {i}\n" + _xml_text(f, text_tags={"t"}, para_tags={"p"}, break_tags={"br"},
                                                      limit=self.max_chars))
            if i in notes:
                with self._zip_part(zf, notes[i]) as f:
                    out.append(_xml_text(f, text_tags={"t"}, para_tags={"p"}, break_tags={"br"}, limit=self.max_chars))
    return Extracted("done", "\n".join(out), "PowerPoint slides", pages=len(slides))


@TextExtractor.register("xlsx")
def _read_xlsx(self: TextExtractor, data: bytes, filename: str) -> Extracted:
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        names = set(zf.namelist())
        shared: list[str] = []
        if "xl/sharedStrings.xml" in names:
            cur: list[str] | None = None
            in_t = 0
            p = _xml_parser()

            def s_start(name, attrs):
                nonlocal cur, in_t
                local = name.rsplit("|", 1)[-1]
                if local == "si":
                    cur = []
                elif local == "t":
                    in_t += 1

            def s_end(name):
                nonlocal cur, in_t
                local = name.rsplit("|", 1)[-1]
                if local == "si" and cur is not None:
                    shared.append("".join(cur))
                    cur = None
                elif local == "t":
                    in_t -= 1

            def s_chars(d):
                if in_t and cur is not None:
                    cur.append(d)

            p.StartElementHandler, p.EndElementHandler, p.CharacterDataHandler = s_start, s_end, s_chars
            with self._zip_part(zf, "xl/sharedStrings.xml") as f:
                p.ParseFile(f)
        # sheet names in workbook order
        sheet_names: dict[str, str] = {}
        if "xl/workbook.xml" in names and "xl/_rels/workbook.xml.rels" in names:
            rels = {}
            with self._zip_part(zf, "xl/_rels/workbook.xml.rels") as f:
                rp = _xml_parser(namespaces=False)
                rp.StartElementHandler = lambda n, a: rels.__setitem__(a.get("Id"), a.get("Target", "")) \
                    if n.endswith("Relationship") else None
                rp.ParseFile(f)
            with self._zip_part(zf, "xl/workbook.xml") as f:
                wp = _xml_parser()

                def w_start(name, attrs):
                    if name.rsplit("|", 1)[-1] == "sheet":
                        rid = next((v for k, v in attrs.items() if k.endswith("|id")), None)
                        target = rels.get(rid, "")
                        target = target.lstrip("/")
                        if not target.startswith("xl/"):
                            target = "xl/" + target
                        sheet_names[target] = attrs.get("name", "")
                wp.StartElementHandler = w_start
                wp.ParseFile(f)
        sheets = list(sheet_names) or _numbered(names, r"xl/worksheets/sheet(\d+)\.xml$")
        out: list[str] = []
        size = 0
        for sheet in sheets:
            if sheet not in names:
                continue
            out.append(f"Sheet: {sheet_names.get(sheet) or sheet.rsplit('/', 1)[-1][:-4]}")
            row: list[str] = []
            cell: dict = {}
            p = _xml_parser()

            def c_start(name, attrs):
                local = name.rsplit("|", 1)[-1]
                if local == "row":
                    row.clear()
                elif local == "c":
                    cell.clear()
                    cell.update(t=attrs.get("t", ""), v=[], capture=False)
                elif local in ("v", "t") and cell:
                    cell["capture"] = True

            def c_end(name):
                nonlocal size
                local = name.rsplit("|", 1)[-1]
                if local in ("v", "t") and cell:
                    cell["capture"] = False
                elif local == "c" and cell:
                    raw = "".join(cell["v"])
                    if cell["t"] == "s":
                        try:
                            raw = shared[int(raw)]
                        except (ValueError, IndexError):
                            raw = ""
                    elif cell["t"] == "b":
                        raw = "TRUE" if raw == "1" else "FALSE"
                    raw = raw.strip()
                    if raw:
                        row.append(raw)
                    cell.clear()
                elif local == "row" and row:
                    line = "\t".join(row)
                    out.append(line)
                    size += len(line)
                    if size > self.max_chars:
                        raise _Stop

            def c_chars(d):
                if cell and cell.get("capture"):
                    cell["v"].append(d)

            p.StartElementHandler, p.EndElementHandler, p.CharacterDataHandler = c_start, c_end, c_chars
            try:
                with self._zip_part(zf, sheet) as f:
                    p.ParseFile(f)
            except _Stop:
                break
    return Extracted("done", "\n".join(out), "Excel cells", pages=len(sheets))


@TextExtractor.register("odf")
def _read_odf(self: TextExtractor, data: bytes, filename: str) -> Extracted:
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        with self._zip_part(zf, "content.xml") as f:
            text = _xml_text(f, text_tags=None, para_tags={"p", "h"}, tab_tags={"tab"}, break_tags={"line-break"},
                             space_tags={"s"}, limit=self.max_chars)
    return Extracted("done", text, "OpenDocument")


@TextExtractor.register("text")
def _read_text(self: TextExtractor, data: bytes, filename: str) -> Extracted:
    text = _decode(data[: self.max_chars * 4])
    if _ext(filename) == "rtf" or text.startswith("{\\rtf"):
        text = re.sub(r"\\[a-z]+-?\d* ?|[{}]", " ", text)
        return Extracted("done", text, "RTF (plain text only)")
    return Extracted("done", text, "Plain text")


@TextExtractor.register("html")
def _read_html(self: TextExtractor, data: bytes, filename: str) -> Extracted:
    return Extracted("done", html_to_text(_decode(data[: self.max_chars * 4])), "Web page")


@TextExtractor.register("email")
def _read_email(self: TextExtractor, data: bytes, filename: str) -> Extracted:
    import email
    import email.policy
    msg = email.message_from_bytes(data, policy=email.policy.default)
    head = "\n".join(f"{k}: {msg.get(k)}" for k in ("Subject", "From", "To", "Date") if msg.get(k))
    body = ""
    part = msg.get_body(preferencelist=("plain", "html"))
    if part is not None:
        body = part.get_content()
        if part.get_content_type() == "text/html":
            body = html_to_text(body)
    return Extracted("done", head + "\n\n" + body, "E-mail")


@TextExtractor.register("ole")
def _read_ole(self: TextExtractor, data: bytes, filename: str) -> Extracted:
    return Extracted("skipped", note="Older Office files (.doc, .xls, .ppt) and Outlook .msg files can't be read yet. "
                                     "Save the file as .docx, .xlsx, .pptx or PDF to make its text searchable.")


@TextExtractor.register("cad")
def _read_cad(self: TextExtractor, data: bytes, filename: str) -> Extracted:
    return Extracted("skipped", note="CAD file: it is stored and previewed where possible; its text is not read.")


@TextExtractor.register("damaged")
def _read_damaged(self: TextExtractor, data: bytes, filename: str) -> Extracted:
    return Extracted("failed", note="The file looks damaged (it is not a complete Office or ZIP file), so its text "
                                    "can't be read. The file is stored exactly as it was uploaded.")


def run_isolated(path: str, filename: str, options: dict, timeout: float) -> Extracted:
    """Run the extractor in a separate process, so a slow or broken file is stopped after ``timeout`` seconds
    and can never take the server down with it."""
    env = dict(os.environ)
    env["PYTHONPATH"] = str(config.BASE_DIR) + (os.pathsep + env["PYTHONPATH"] if env.get("PYTHONPATH") else "")
    kwargs: dict = {}
    if os.name == "posix":
        kwargs["start_new_session"] = True
    elif hasattr(subprocess, "CREATE_NEW_PROCESS_GROUP"):
        kwargs["creationflags"] = subprocess.CREATE_NEW_PROCESS_GROUP
    proc = subprocess.Popen([sys.executable, "-m", "kosha.inputs"], stdin=subprocess.PIPE, stdout=subprocess.PIPE,
                            stderr=subprocess.PIPE, cwd=str(config.BASE_DIR), env=env, **kwargs)
    payload = json.dumps({"path": path, "filename": filename, "options": options}).encode()
    try:
        out, err = proc.communicate(payload, timeout=timeout)
    except subprocess.TimeoutExpired:
        try:
            if os.name == "posix":
                import signal
                os.killpg(proc.pid, signal.SIGKILL)
            else:
                proc.kill()
        except OSError:
            pass
        proc.communicate()
        return Extracted("failed", note=f"Reading the text took longer than {int(timeout)} seconds and was stopped. "
                                        "The file itself is stored safely.")
    if proc.returncode != 0:
        log.warning("Text reader for %s exited with %s: %s", filename, proc.returncode, err[-2000:].decode("utf-8", "replace"))
        return Extracted("failed", note=f"The text reader stopped unexpectedly (exit code {proc.returncode}). "
                                        "The file itself is stored safely.")
    try:
        line = out.decode("utf-8").strip().splitlines()[-1]
        return Extracted.from_json(json.loads(line))
    except (ValueError, IndexError):
        return Extracted("failed", note="The text reader gave no answer. The file itself is stored safely.")


def _child_main() -> int:
    req = json.loads(sys.stdin.buffer.read().decode("utf-8"))
    with open(req["path"], "rb") as f:
        data = f.read()
    res = TextExtractor(req.get("options")).run(data, req["filename"])
    sys.stdout.write("\n" + json.dumps(res.to_json()) + "\n")
    sys.stdout.flush()
    return 0


# ===================================================================== background worker ==

class ExtractionWorker:
    """Reads text and makes thumbnails in the background, one file fingerprint at a time.

    The queue is the ``input_content`` table itself, so it survives restarts: rows left ``running`` by a
    stopped server are put back to ``pending`` (or marked failed after three tries)."""

    MAX_ATTEMPTS = 3

    def __init__(self, library: "InputLibrary"):
        self.library = library
        self._thread: threading.Thread | None = None
        self._wake = threading.Event()
        self._lock = threading.Lock()

    def start(self) -> None:
        if os.environ.get("KOSHA_DISABLE_TEXT_WORKER") == "1":
            return
        with self._lock:
            if self._thread and self._thread.is_alive():
                return
            self._thread = threading.Thread(target=self._loop, name="kosha-inputs-text", daemon=True)
            self._thread.start()

    def wake(self) -> None:
        self.start()
        self._wake.set()

    def recover(self) -> None:
        with db.tx():
            db.execute("UPDATE input_content SET status = 'failed', note = ?, updated_at = ? "
                       "WHERE status = 'running' AND attempts >= ?",
                       ("Reading stopped several times (the server restarted). The file itself is stored safely.",
                        db.now(), self.MAX_ATTEMPTS))
            db.execute("UPDATE input_content SET status = 'pending', updated_at = ? WHERE status = 'running'", (db.now(),))

    def _loop(self) -> None:
        try:
            self.recover()
        except Exception:
            log.exception("Could not reset interrupted text reading")
        while True:
            try:
                busy = self.process_next()
            except Exception:
                log.exception("Text reading failed; trying the next file")
                busy = False
                time.sleep(2)
            if not busy:
                self._wake.wait(timeout=30)
                self._wake.clear()

    def claim(self) -> str | None:
        with db.tx():
            row = db.one("SELECT sha FROM input_content WHERE status = 'pending' ORDER BY updated_at, sha LIMIT 1")
            if not row:
                return None
            db.execute("UPDATE input_content SET status = 'running', attempts = attempts + 1, updated_at = ? "
                       "WHERE sha = ?", (db.now(), row["sha"]))
        return row["sha"]

    def process_next(self, isolate: bool = True) -> bool:
        sha = self.claim()
        if not sha:
            return False
        self.library.process_content(sha, isolate=isolate)
        return True

    def drain(self, isolate: bool = False, limit: int = 1000) -> int:
        """Process everything waiting (used by tests and the command line)."""
        n = 0
        while n < limit and self.process_next(isolate=isolate):
            n += 1
        return n


# ===================================================================== the library ==

@dataclass
class Upload:
    filename: str
    stream: BinaryIO


def _clean_filename(name: str) -> str:
    name = (name or "").replace("\\", "/").split("/")[-1]
    name = re.sub(r"[\x00-\x1f\x7f]", "", name).strip()
    return name[:255] or "file"


def title_from_filename(name: str) -> str:
    stem_ = name.rsplit(".", 1)[0] if "." in name else name
    return re.sub(r"\s+", " ", stem_.replace("_", " ")).strip()[:300] or name


def clean_tags(raw) -> list[str]:
    if raw is None:
        return []
    if isinstance(raw, str):
        s = raw.strip()
        if s.startswith("["):
            try:
                raw = json.loads(s)
            except ValueError:
                raw = s.strip("[]").split(",")
        else:
            raw = s.split(",")
    out, keys = [], set()
    for t in raw:
        t = re.sub(r"\s+", " ", str(t or "")).strip().lstrip("#").strip()[:40]
        if t and t.lower() not in keys:
            keys.add(t.lower())
            out.append(t)
    if len(out) > MAX_TAGS:
        raise bad(f"An input can have at most {MAX_TAGS} tags.")
    return out


def id_list(raw, what: str) -> list[int]:
    if raw is None or raw == "":
        return []
    if isinstance(raw, str):
        s = raw.strip()
        try:
            raw = json.loads(s) if s.startswith("[") else [x for x in s.split(",") if x.strip()]
        except ValueError:
            raise bad(f"{what} must be a list of numbers.")
    if isinstance(raw, (int, float)):
        raw = [raw]
    try:
        return list(dict.fromkeys(int(x) for x in raw))
    except (TypeError, ValueError):
        raise bad(f"{what} must be a list of numbers.")


def clean_url(raw) -> str:
    u = _clip(raw, 2000)
    if not u:
        return ""
    p = urllib.parse.urlparse(u)
    if p.scheme.lower() not in ("http", "https") or not p.netloc:
        raise bad("A web link must start with http:// or https://")
    return u


class InputLibrary:
    def __init__(self):
        self.index = SearchIndex()
        self.worker = ExtractionWorker(self)

    # ------------------------------------------------------------ settings --
    def settings(self) -> dict:
        raw = db.get_setting(SETTINGS_KEY) or {}
        out = {**DEFAULTS, **{k: v for k, v in raw.items() if k in DEFAULTS}}
        if not isinstance(out["kinds"], list) or not out["kinds"]:
            out["kinds"] = DEFAULT_KINDS
        return out

    def public_settings(self) -> dict:
        s = self.settings()
        ok, why = ocr_status()
        return {**s, "ocr_available": ok, "ocr_note": why, "types": TYPE_LABEL}

    def save_settings(self, user: dict, data: dict) -> dict:
        cur = self.settings()
        new = dict(cur)
        if "kinds" in data:
            kinds = []
            keys = set()
            for k in data["kinds"] or []:
                label = _clip(k.get("label") if isinstance(k, dict) else k, 60)
                if not label:
                    continue
                key = _clip(k.get("key") if isinstance(k, dict) else "", 40) or \
                    re.sub(r"[^a-z0-9]+", "_", label.lower()).strip("_")[:40] or "kind"
                base, i = key, 2
                while key in keys:
                    key, i = f"{base}_{i}", i + 1
                keys.add(key)
                kinds.append({"key": key, "label": label})
            if "other" not in keys:
                kinds.append({"key": "other", "label": "Other"})
                keys.add("other")
            used = db.rows("SELECT kind, COUNT(*) AS n FROM inputs GROUP BY kind")
            missing = [f"{r['kind']} ({r['n']})" for r in used if r["kind"] not in keys]
            if missing:
                raise bad("These kinds are still used, so they can't be removed: " + ", ".join(missing) +
                          ". Give those inputs another kind first.")
            new["kinds"] = kinds
        if "ocr" in data:
            v = data["ocr"]
            new["ocr"] = v.strip().lower() in ("1", "true", "yes", "on") if isinstance(v, str) else bool(v)
        if "ocr_lang" in data:
            lang = _clip(data["ocr_lang"], 40)
            if not re.match(r"^[a-z_]{3,8}(\+[a-z_]{3,8})*$", lang):
                raise bad("OCR language must be a Tesseract language code such as eng or eng+deu.")
            new["ocr_lang"] = lang
        for k, (lo, hi) in LIMITS.items():
            if k in data:
                try:
                    v = int(data[k])
                except (TypeError, ValueError):
                    raise bad(f"{k} must be a whole number.")
                if not lo <= v <= hi:
                    raise bad(f"{k} must be between {lo} and {hi}.")
                new[k] = v
        db.set_setting(SETTINGS_KEY, new)
        db.audit(user["id"], "inputs_settings", "settings", None, "")
        labels_changed = {k["key"]: k["label"] for k in new["kinds"]} != {k["key"]: k["label"] for k in cur["kinds"]}
        if labels_changed:
            self.reindex_where("SELECT id FROM inputs", ())
        return self.public_settings()

    def kind_labels(self) -> dict[str, str]:
        return {k["key"]: k["label"] for k in self.settings()["kinds"]}

    def default_kind(self, filename: str = "") -> str:
        keys = self.kind_labels()
        if filename and type_group(_ext(filename)) == "image" and "photo" in keys:
            return "photo"
        return "other" if "other" in keys else next(iter(keys))

    # -------------------------------------------------------------- access --
    @staticmethod
    def visible_sql(user: dict, alias: str = "i") -> tuple[str, list]:
        """SQL condition: inputs this user may see. Used by every query that touches inputs."""
        if user.get("role") == "admin":
            return "1 = 1", []
        return (f"({alias}.confidentiality = 'normal' OR {alias}.created_by = ? OR EXISTS (SELECT 1 FROM input_access ia "
                f"WHERE ia.input_id = {alias}.id AND (ia.user_id = ? OR ia.role = ?)))",
                [user["id"], user["id"], user.get("role", "")])

    @staticmethod
    def is_manager(user: dict, inp: dict) -> bool:
        return inp["created_by"] == user["id"] or ROLE_LEVEL.get(user.get("role"), 0) >= ROLE_LEVEL["approver"]

    def get(self, user: dict, input_id: int, include_deleted: bool = False) -> dict:
        vis, params = self.visible_sql(user)
        inp = db.one(f"SELECT i.* FROM inputs i WHERE i.id = ? AND {vis}", [input_id, *params])
        if not inp or (inp["deleted"] and not (include_deleted and self.is_manager(user, inp))):
            raise not_found("Input not found")      # the same answer whether it is missing or restricted
        return inp

    def _require_edit(self, user: dict) -> None:
        if ROLE_LEVEL.get(user.get("role"), 0) < ROLE_LEVEL["engineer"]:
            raise forbidden("Changing inputs needs the engineer role or higher.")

    def _require_manage(self, user: dict, inp: dict, what: str) -> None:
        if not self.is_manager(user, inp):
            raise forbidden(f"Only the person who added this input, an approver or an admin can {what}.")

    # ------------------------------------------------------------ cleaning --
    def _clean(self, data: dict, new: bool) -> dict:
        out: dict = {}
        if "title" in data or new:
            t = re.sub(r"\s+", " ", _clip(data.get("title"), 300))
            if not t:
                raise bad("Give the input a title.")
            out["title"] = t
        if "kind" in data or new:
            k = _clip(data.get("kind"), 40) or self.default_kind()
            if k not in self.kind_labels():
                raise bad(f"Unknown kind '{k}'. Choose one of: {', '.join(self.kind_labels().values())}.")
            out["kind"] = k
        if "source" in data:
            out["source"] = _clip(data.get("source"), 300)
        if "date" in data or "input_date" in data or new:
            from .projects import clean_date
            out["input_date"] = clean_date(data.get("date", data.get("input_date")), "Date") or (_today() if new else None)
        if "url" in data:
            out["url"] = clean_url(data.get("url"))
        if "note" in data:
            out["note"] = str(data.get("note") or "").strip()[:100_000]
        if "comment" in data:
            out["comment"] = _clip(data.get("comment"), 5000)
        if "confidentiality" in data or new:
            c = _clip(data.get("confidentiality"), 20).lower() or "normal"
            if c not in CONFIDENTIALITY:
                raise bad("Confidentiality must be normal or restricted.")
            out["confidentiality"] = c
        return out

    def _access(self, data: dict) -> tuple[list[int], list[str]] | None:
        if "access_users" not in data and "access_roles" not in data:
            return None
        users = id_list(data.get("access_users"), "access_users")
        roles_raw = data.get("access_roles") or []
        if isinstance(roles_raw, str):
            roles_raw = json.loads(roles_raw) if roles_raw.strip().startswith("[") else roles_raw.split(",")
        roles = list(dict.fromkeys(str(r).strip().lower() for r in roles_raw if str(r).strip()))
        for r in roles:
            if r not in ROLE_LEVEL:
                raise bad(f"Unknown role '{r}'.")
        for u in users:
            if not db.scalar("SELECT id FROM users WHERE id = ?", (u,)):
                raise bad(f"User {u} does not exist.")
        return users, roles

    def _check_links(self, item_ids: list[int], project_ids: list[int]) -> None:
        for i in item_ids:
            if not db.scalar("SELECT id FROM items WHERE id = ?", (i,)):
                raise bad(f"Part {i} does not exist.")
        for p in project_ids:
            if not db.scalar("SELECT id FROM projects WHERE id = ?", (p,)):
                raise bad(f"Project {p} does not exist.")

    # -------------------------------------------------------- write helpers --
    def _set_access(self, input_id: int, access: tuple[list[int], list[str]]) -> None:
        db.execute("DELETE FROM input_access WHERE input_id = ?", (input_id,))
        for u in access[0]:
            db.insert("input_access", {"input_id": input_id, "user_id": u, "role": None})
        for r in access[1]:
            db.insert("input_access", {"input_id": input_id, "user_id": None, "role": r})

    def _add_tags(self, input_id: int, tags: list[str]) -> int:
        n = 0
        for t in tags:
            if not db.scalar("SELECT 1 FROM input_tags WHERE input_id = ? AND tag_key = ?", (input_id, t.lower())):
                if db.scalar("SELECT COUNT(*) FROM input_tags WHERE input_id = ?", (input_id,)) >= MAX_TAGS:
                    raise bad(f"An input can have at most {MAX_TAGS} tags.")
                db.execute("INSERT INTO input_tags (input_id, tag, tag_key) VALUES (?, ?, ?)", (input_id, t, t.lower()))
                n += 1
        return n

    def _link(self, user: dict, input_id: int, item_ids: list[int], project_ids: list[int]) -> int:
        n = 0
        now = db.now()
        for i in item_ids:
            if not db.scalar("SELECT 1 FROM input_items WHERE input_id = ? AND item_id = ?", (input_id, i)):
                db.execute("INSERT INTO input_items (input_id, item_id, created_by, created_at) VALUES (?, ?, ?, ?)",
                           (input_id, i, user["id"], now))
                n += 1
        for p in project_ids:
            if not db.scalar("SELECT 1 FROM input_projects WHERE input_id = ? AND project_id = ?", (input_id, p)):
                db.execute("INSERT INTO input_projects (input_id, project_id, created_by, created_at) VALUES (?, ?, ?, ?)",
                           (input_id, p, user["id"], now))
                n += 1
        return n

    def _unlink(self, input_id: int, item_ids: list[int], project_ids: list[int]) -> int:
        n = 0
        for i in item_ids:
            n += db.connect().execute("DELETE FROM input_items WHERE input_id = ? AND item_id = ?", (input_id, i)).rowcount
        for p in project_ids:
            n += db.connect().execute("DELETE FROM input_projects WHERE input_id = ? AND project_id = ?",
                                      (input_id, p)).rowcount
        return n

    @staticmethod
    def store_upload(up: Upload) -> dict:
        """Write the bytes to the store first (flushed to disk) - before any database row points at them."""
        name = _clean_filename(up.filename)
        sha, size = store.put_stream(up.stream)
        if size == 0:
            raise bad(f"{name} is empty (0 bytes), so it was not added.")
        return {"filename": name, "ext": _ext(name), "sha": sha, "size": size}

    def _add_file_row(self, user: dict, input_id: int, f: dict) -> int:
        pos = (db.scalar("SELECT MAX(position) FROM input_files WHERE input_id = ?", (input_id,)) or 0) + 1
        db.insert("input_files", {"input_id": input_id, "position": pos, "filename": f["filename"], "ext": f["ext"],
                                  "blob_sha": f["sha"], "size": f["size"], "removed": 0, "created_by": user["id"],
                                  "created_at": db.now()})
        db.execute("INSERT INTO input_content (sha, status, updated_at) VALUES (?, 'pending', ?) "
                   "ON CONFLICT (sha) DO NOTHING", (f["sha"], db.now()))
        return pos

    def duplicates_of(self, user: dict, sha: str, except_input: int | None = None) -> list[dict]:
        """Other inputs (that this user may see) holding a file with the same fingerprint."""
        vis, params = self.visible_sql(user)
        return db.rows(f"SELECT DISTINCT i.id, i.title FROM input_files f JOIN inputs i ON i.id = f.input_id "
                       f"WHERE f.blob_sha = ? AND f.removed = 0 AND i.deleted = 0 AND i.id != ? AND {vis} "
                       f"ORDER BY i.id LIMIT 10", [sha, except_input or 0, *params])

    # ---------------------------------------------------------------- create --
    def create(self, user: dict, data: dict, files: list[dict] | None = None) -> int:
        """Create one input. ``files`` are already stored (see ``store_upload``)."""
        self._require_edit(user)
        files = files or []
        data = dict(data)
        if not _clip(data.get("title"), 300):
            if files:
                data["title"] = title_from_filename(files[0]["filename"])
            elif data.get("url"):
                data["title"] = urllib.parse.urlparse(clean_url(data["url"])).netloc
        if not _clip(data.get("kind"), 40):
            data["kind"] = self.default_kind(files[0]["filename"] if files else "")
        vals = self._clean(data, new=True)
        if not files and not vals.get("url") and not str(data.get("note") or "").strip():
            raise bad("Add a file, a web link or a note: an input needs at least one of them.")
        tags = clean_tags(data.get("tags"))
        access = self._access(data) or ([], [])
        item_ids = id_list(data.get("item_ids"), "item_ids")
        project_ids = id_list(data.get("project_ids"), "project_ids")
        self._check_links(item_ids, project_ids)
        now = db.now()
        with db.tx():
            iid = db.insert("inputs", {"source": "", "url": "", "note": "", "comment": "", **vals,
                                       "created_by": user["id"], "created_at": now, "updated_by": user["id"],
                                       "updated_at": now})
            if vals["confidentiality"] == "restricted":
                self._set_access(iid, access)
            self._add_tags(iid, tags)
            self._link(user, iid, item_ids, project_ids)
            for f in files:
                self._add_file_row(user, iid, f)
            db.audit(user["id"], "input_create", "input", iid, vals["title"][:200])
            self.reindex(iid)
        return iid

    def add_files(self, user: dict, input_id: int, files: list[dict]) -> list[int]:
        self._require_edit(user)
        self.get(user, input_id)
        with db.tx():
            positions = [self._add_file_row(user, input_id, f) for f in files]
            db.execute("UPDATE inputs SET updated_at = ?, updated_by = ? WHERE id = ?", (db.now(), user["id"], input_id))
            db.audit(user["id"], "input_add_files", "input", input_id, ", ".join(f["filename"] for f in files)[:500])
            self.reindex(input_id)
        return positions

    def set_file_removed(self, user: dict, input_id: int, n: int, removed: bool) -> None:
        self._require_edit(user)
        inp = self.get(user, input_id)
        f = db.one("SELECT * FROM input_files WHERE input_id = ? AND position = ?", (input_id, n))
        if not f:
            raise not_found("File not found")
        if not removed:
            self._require_manage(user, inp, "put back a removed file")
        elif not (inp["url"] or inp["note"]) and not db.scalar(
                "SELECT 1 FROM input_files WHERE input_id = ? AND removed = 0 AND id != ?", (input_id, f["id"])):
            raise bad("This is the only file of the input. Add another file first, or delete the whole input.")
        with db.tx():
            db.execute("UPDATE input_files SET removed = ? WHERE id = ?", (1 if removed else 0, f["id"]))
            db.execute("UPDATE inputs SET updated_at = ?, updated_by = ? WHERE id = ?", (db.now(), user["id"], input_id))
            db.audit(user["id"], "input_remove_file" if removed else "input_restore_file", "input", input_id, f["filename"])
            self.reindex(input_id)

    # ---------------------------------------------------------------- update --
    def update(self, user: dict, input_id: int, data: dict) -> None:
        self._require_edit(user)
        inp = self.get(user, input_id)
        vals = self._clean({k: v for k, v in data.items() if k in (
            "title", "kind", "source", "date", "input_date", "url", "note", "comment", "confidentiality")}, new=False)
        access = self._access(data)
        if ("confidentiality" in vals and vals["confidentiality"] != inp["confidentiality"]) or access is not None:
            self._require_manage(user, inp, "change who may see it")
        tags = clean_tags(data["tags"]) if "tags" in data else None
        items = id_list(data["item_ids"], "item_ids") if "item_ids" in data else None
        projects_ = id_list(data["project_ids"], "project_ids") if "project_ids" in data else None
        self._check_links(items or [], projects_ or [])
        after = {**inp, **vals}
        if not inp_has_content(after, input_id):
            raise bad("An input needs at least one file, a web link or a note.")
        with db.tx():
            if vals:
                db.update("inputs", input_id, {**vals, "updated_at": db.now(), "updated_by": user["id"]})
            if access is not None:
                self._set_access(input_id, access)
            if after["confidentiality"] == "normal":
                db.execute("DELETE FROM input_access WHERE input_id = ?", (input_id,))
            if tags is not None:
                db.execute("DELETE FROM input_tags WHERE input_id = ?", (input_id,))
                self._add_tags(input_id, tags)
            if items is not None:
                db.execute("DELETE FROM input_items WHERE input_id = ?", (input_id,))
            if projects_ is not None:
                db.execute("DELETE FROM input_projects WHERE input_id = ?", (input_id,))
            self._link(user, input_id, items or [], projects_ or [])
            changed = list(vals) + [name for name, value in (("tags", tags), ("items", items), ("projects", projects_),
                                                             ("access", access)) if value is not None]
            db.audit(user["id"], "input_update", "input", input_id, ",".join(changed))
            self.reindex(input_id)

    def link(self, user: dict, input_id: int, item_ids: list[int], project_ids: list[int], unlink: bool = False) -> int:
        self._require_edit(user)
        self.get(user, input_id)
        self._check_links(item_ids, project_ids)
        with db.tx():
            n = self._unlink(input_id, item_ids, project_ids) if unlink else self._link(user, input_id, item_ids, project_ids)
            if n:
                db.audit(user["id"], "input_unlink" if unlink else "input_link", "input", input_id,
                         f"items={item_ids} projects={project_ids}")
                self.reindex(input_id)
        return n

    def set_deleted(self, user: dict, input_id: int, deleted: bool) -> None:
        self._require_edit(user)
        inp = self.get(user, input_id, include_deleted=True)
        self._require_manage(user, inp, "delete or restore it")
        with db.tx():
            db.update("inputs", input_id, {"deleted": 1 if deleted else 0, "deleted_by": user["id"] if deleted else None,
                                           "deleted_at": db.now() if deleted else None, "updated_at": db.now()})
            db.audit(user["id"], "input_delete" if deleted else "input_restore", "input", input_id, inp["title"][:200])

    # ------------------------------------------------------------------ bulk --
    BULK_ACTIONS = ("add_tags", "remove_tags", "link", "unlink", "set_kind", "delete", "restore")

    def bulk(self, user: dict, ids: list[int], action: str, data: dict) -> dict:
        self._require_edit(user)
        if action not in self.BULK_ACTIONS:
            raise bad(f"Unknown action. Use one of: {', '.join(self.BULK_ACTIONS)}.")
        ids = list(dict.fromkeys(ids))
        if not ids:
            raise bad("Select at least one input.")
        if len(ids) > 500:
            raise bad("Change at most 500 inputs at a time.")
        tags = clean_tags(data.get("tags")) if action in ("add_tags", "remove_tags") else []
        if action in ("add_tags", "remove_tags") and not tags:
            raise bad("Give at least one tag.")
        item_ids = id_list(data.get("item_ids"), "item_ids")
        project_ids = id_list(data.get("project_ids"), "project_ids")
        if action in ("link", "unlink"):
            if not item_ids and not project_ids:
                raise bad("Choose at least one part or project.")
            self._check_links(item_ids, project_ids)
        kind = _clip(data.get("kind"), 40)
        if action == "set_kind" and kind not in self.kind_labels():
            raise bad("Choose a kind.")
        links = f"items={item_ids} projects={project_ids}"
        detail = {"add_tags": ", ".join(tags), "remove_tags": ", ".join(tags), "link": links, "unlink": links,
                  "set_kind": kind}.get(action, "")
        done, skipped, touched = 0, [], []
        with db.tx():
            for iid in ids:
                try:
                    inp = self.get(user, iid, include_deleted=action == "restore")
                except ApiError:
                    skipped.append({"id": iid, "reason": "Input not found"})
                    continue
                if action in ("delete", "restore") and not self.is_manager(user, inp):
                    skipped.append({"id": iid, "reason": "Only the person who added it, an approver or an admin "
                                                         "can delete or restore it."})
                    continue
                if action == "add_tags":
                    current = db.scalar("SELECT COUNT(*) FROM input_tags WHERE input_id = ?", (iid,))
                    new = [t for t in tags if not db.scalar(
                        "SELECT 1 FROM input_tags WHERE input_id = ? AND tag_key = ?", (iid, t.lower()))]
                    if current + len(new) > MAX_TAGS:
                        skipped.append({"id": iid, "reason": f"It would have more than {MAX_TAGS} tags."})
                        continue
                    self._add_tags(iid, new)
                elif action == "remove_tags":
                    for t in tags:
                        db.execute("DELETE FROM input_tags WHERE input_id = ? AND tag_key = ?", (iid, t.lower()))
                elif action == "link":
                    self._link(user, iid, item_ids, project_ids)
                elif action == "unlink":
                    self._unlink(iid, item_ids, project_ids)
                elif action == "set_kind":
                    db.update("inputs", iid, {"kind": kind})
                elif action in ("delete", "restore"):
                    d = action == "delete"
                    db.update("inputs", iid, {"deleted": 1 if d else 0, "deleted_by": user["id"] if d else None,
                                              "deleted_at": db.now() if d else None})
                db.update("inputs", iid, {"updated_at": db.now(), "updated_by": user["id"]})
                db.audit(user["id"], f"input_bulk_{action}", "input", iid, detail)
                touched.append(iid)
                done += 1
            for iid in touched:
                self.reindex(iid)
        return {"done": done, "skipped": skipped}

    # ---------------------------------------------------------------- index --
    def _meta_fields(self, input_id: int) -> list[tuple[str, str]]:
        inp = db.one("SELECT * FROM inputs WHERE id = ?", (input_id,))
        if not inp:
            return []
        labels = self.kind_labels()
        fields = [("title", inp["title"]), ("source", inp["source"]), ("kind", labels.get(inp["kind"], inp["kind"])),
                  ("comment", inp["comment"]), ("note", inp["note"]), ("url", inp["url"])]
        fields += [("tag", r["tag"]) for r in db.rows("SELECT tag FROM input_tags WHERE input_id = ?", (input_id,))]
        fields += [("file", r["filename"]) for r in db.rows(
            "SELECT filename FROM input_files WHERE input_id = ? AND removed = 0", (input_id,))]
        for r in db.rows("SELECT it.part_number, it.description FROM input_items x JOIN items it ON it.id = x.item_id "
                         "WHERE x.input_id = ?", (input_id,)):
            fields += [("part", r["part_number"]), ("item_desc", r["description"])]
        for r in db.rows("SELECT p.code, p.name, p.customer FROM input_projects x JOIN projects p ON p.id = x.project_id "
                         "WHERE x.input_id = ?", (input_id,)):
            fields += [("project", r["code"]), ("project_name", r["name"]), ("customer", r["customer"])]
        return fields

    def reindex(self, input_id: int) -> None:
        """Rebuild the index rows of an input's own fields (cheap; file text is indexed per fingerprint)."""
        with db.tx():
            self.index.write_input(input_id, self.index.weigh(self._meta_fields(input_id)))

    def index_text(self, sha: str, text: str | None = None) -> None:
        """Index the text read from one file (by fingerprint)."""
        if text is None:
            row = db.one("SELECT text_sha FROM input_content WHERE sha = ?", (sha,))
            text = ""
            if row and row["text_sha"]:
                try:
                    text = store.read_bytes(row["text_sha"]).decode("utf-8", "replace")
                except OSError:
                    log.warning("Extracted text of %s is missing; read the file again to rebuild it", sha)
        self.index.write_text(sha, self.index.weigh([("text", text)], limit=int(self.settings()["max_chars"])))

    def reindex_where(self, sql: str, params: Iterable) -> int:
        ids = [r[0] for r in db.connect().execute(sql, tuple(params)).fetchall()]
        for iid in ids:
            self.reindex(iid)
        return len(ids)

    def reindex_all(self) -> dict:
        """Rebuild the whole search index from the records and the stored text (safe to run any time)."""
        n = self.reindex_where("SELECT id FROM inputs ORDER BY id", ())
        shas = [r["sha"] for r in db.rows("SELECT sha FROM input_content WHERE text_sha IS NOT NULL")]
        for sha in shas:
            self.index_text(sha)
        with db.tx():
            db.execute("DELETE FROM input_text_terms WHERE sha NOT IN (SELECT sha FROM input_content "
                       "WHERE text_sha IS NOT NULL)")
        return {"inputs": n, "files": len(shas)}

    # ------------------------------------------------------------ extraction --
    def extract_options(self) -> dict:
        s = self.settings()
        return {k: s[k] for k in ("ocr", "ocr_lang", "ocr_max_pages", "time_limit_s", "max_chars")}

    def process_content(self, sha: str, isolate: bool = True) -> None:
        f = db.one("SELECT filename FROM input_files WHERE blob_sha = ? ORDER BY removed, id LIMIT 1", (sha,))
        s = self.settings()
        if not f or not store.exists(sha):
            res = Extracted("failed", note="The stored file is missing on the server. Restore it from the backup.")
        else:
            size = store.path_for(sha).stat().st_size
            if size > int(s["max_mb"]) * 1024 * 1024:
                res = Extracted("skipped", note=f"Larger than {s['max_mb']} MB, so its text was not read. "
                                                "The file itself is stored safely.")
            elif isolate:
                res = run_isolated(str(store.path_for(sha)), f["filename"], self.extract_options(),
                                   float(s["time_limit_s"]))
            else:
                res = TextExtractor(self.extract_options()).run(store.read_bytes(sha), f["filename"])
        text_sha = store.put_bytes(res.text.encode("utf-8"))[0] if res.text else None
        thumb_sha = thumb_kind = None
        if res.thumb_original and f:
            thumb_sha, thumb_kind = sha, _ext(f["filename"])
        elif res.thumb:
            thumb_kind = res.thumb[0]
            thumb_sha = store.put_bytes(res.thumb[1])[0]
        excerpt = re.sub(r"\n{3,}", "\n\n", res.text[:1500]).strip()
        weights = self.index.weigh([("text", res.text)], limit=int(s["max_chars"]))
        with db.tx():
            db.execute("UPDATE input_content SET file_type = ?, status = ?, method = ?, note = ?, text_sha = ?, excerpt = ?, "
                       "chars = ?, pages = ?, thumb_sha = ?, thumb_kind = ?, updated_at = ? WHERE sha = ?",
                       (res.file_type, res.status, res.method, res.note[:1000], text_sha, excerpt, len(res.text),
                        res.pages, thumb_sha, thumb_kind, db.now(), sha))
            self.index.write_text(sha, weights)

    def reread(self, user: dict, input_id: int | None = None, n: int | None = None, only_unread: bool = False) -> int:
        """Queue text reading again (after OCR was switched on, or a reader was improved)."""
        if input_id is not None:
            self._require_edit(user)
            self.get(user, input_id)
            sql = "SELECT DISTINCT blob_sha FROM input_files WHERE input_id = ?" + (" AND position = ?" if n else "")
            shas = [r[0] for r in db.connect().execute(sql, (input_id, n) if n else (input_id,)).fetchall()]
        else:
            where = " WHERE status IN ('empty', 'skipped', 'failed')" if only_unread else ""
            shas = [r[0] for r in db.connect().execute("SELECT sha FROM input_content" + where).fetchall()]
        with db.tx():
            for sha in shas:
                db.execute("UPDATE input_content SET status = 'pending', attempts = 0, updated_at = ? "
                           "WHERE sha = ? AND status != 'running'", (db.now(), sha))
        return len(shas)

    # ---------------------------------------------------------------- read --
    def _filters(self, user: dict, p: dict, skip: str | None = None) -> tuple[list[str], list]:
        vis, params = self.visible_sql(user)
        where = [vis]
        args = list(params)
        if str(p.get("deleted") or "") in ("1", "true"):
            where.append("i.deleted = 1")
            if ROLE_LEVEL.get(user.get("role"), 0) < ROLE_LEVEL["approver"]:
                where.append("i.created_by = ?")
                args.append(user["id"])
        else:
            where.append("i.deleted = 0")
        kinds = [k for k in str(p.get("kind") or "").split(",") if k]
        if kinds and skip != "kind":
            where.append(f"i.kind IN ({','.join('?' * len(kinds))})")
            args += kinds
        types = [t for t in str(p.get("type") or "").split(",") if t]
        if types and skip != "type":
            ors = []
            for t in types:
                if t == "link":
                    ors.append("i.url != ''")
                elif t == "note":
                    ors.append("(i.note != '' AND NOT EXISTS (SELECT 1 FROM input_files tf WHERE tf.input_id = i.id "
                               "AND tf.removed = 0))")
                elif t == "other":
                    ors.append(f"EXISTS (SELECT 1 FROM input_files tf WHERE tf.input_id = i.id AND tf.removed = 0 AND "
                               f"tf.ext NOT IN ({','.join('?' * len(KNOWN_EXTS))}))")
                    args += list(KNOWN_EXTS)
                elif t in TYPE_GROUPS:
                    exts = TYPE_GROUPS[t][1]
                    ors.append(f"EXISTS (SELECT 1 FROM input_files tf WHERE tf.input_id = i.id AND tf.removed = 0 AND "
                               f"tf.ext IN ({','.join('?' * len(exts))}))")
                    args += list(exts)
            if ors:
                where.append("(" + " OR ".join(ors) + ")")
        tags = [t.strip().lower() for t in str(p.get("tag") or "").split(",") if t.strip()]
        if tags and skip != "tag":
            for t in tags:
                where.append("EXISTS (SELECT 1 FROM input_tags tt WHERE tt.input_id = i.id AND tt.tag_key = ?)")
                args.append(t)

        def number(key: str) -> int:
            try:
                return int(p[key])
            except (TypeError, ValueError):
                raise bad(f"{key} must be a number.")

        for key, table, col in (("project", "input_projects", "project_id"), ("item", "input_items", "item_id")):
            if p.get(key) and skip != key:
                where.append(f"EXISTS (SELECT 1 FROM {table} lt WHERE lt.input_id = i.id AND lt.{col} = ?)")
                args.append(number(key))
        if p.get("by") and skip != "by":
            where.append("i.created_by = ?")
            args.append(number("by"))
        from .projects import clean_date
        if p.get("from"):
            where.append("i.input_date >= ?")
            args.append(clean_date(p["from"], "From date"))
        if p.get("to"):
            where.append("i.input_date <= ?")
            args.append(clean_date(p["to"], "To date"))
        if p.get("restricted") in ("1", "true"):
            where.append("i.confidentiality = 'restricted'")
        terms = self.index.query_terms(p.get("q") or "")
        if terms:
            w, wp, _, _ = self.index.match_sql(terms)
            where += w
            args += wp
        return where, args

    SORTS = {"date": "COALESCE(i.input_date, substr(i.created_at, 1, 10))", "title": "i.title COLLATE NOCASE",
             "kind": "i.kind", "added": "i.created_at", "updated": "i.updated_at"}

    def search(self, user: dict, p: dict) -> dict:
        where, args = self._filters(user, p)
        terms = self.index.query_terms(p.get("q") or "")
        sort = p.get("sort") or ("relevance" if terms else "date")
        desc = (p.get("dir") or ("asc" if sort in ("title", "kind") else "desc")).lower() != "asc"
        try:
            limit = max(1, min(int(p.get("limit") or 60), 500))
            offset = max(0, int(p.get("offset") or 0))
        except ValueError:
            raise bad("limit and offset must be numbers.")
        score_sql, score_args = "0", []
        if terms:
            _, _, score_sql, score_args = self.index.match_sql(terms)
        if sort == "relevance" and terms:
            order = f"score {'DESC' if desc else 'ASC'}, COALESCE(i.input_date, '') DESC, i.id DESC"
        else:
            col = self.SORTS.get(sort, self.SORTS["date"])
            order = f"{col} {'DESC' if desc else 'ASC'}, i.id DESC"
        wsql = " AND ".join(where)
        rows = db.rows(f"SELECT i.*, ({score_sql}) AS score FROM inputs i WHERE {wsql} ORDER BY {order} LIMIT ? OFFSET ?",
                       [*score_args, *args, limit, offset])
        total = db.scalar(f"SELECT COUNT(*) FROM inputs i WHERE {wsql}", args)
        out = self._summaries(rows, terms)
        result = {"inputs": out, "total": total, "terms": terms, "sort": sort, "dir": "desc" if desc else "asc",
                  "limit": limit, "offset": offset}
        if str(p.get("facets", "1")) not in ("0", "false"):
            result["facets"] = self.facets(user, p)
        return result

    def facets(self, user: dict, p: dict) -> dict:
        out: dict = {}
        labels = self.kind_labels()
        w, a = self._filters(user, p, skip="kind")
        counts = {r["kind"]: r["n"] for r in db.rows(
            f"SELECT i.kind, COUNT(*) AS n FROM inputs i WHERE {' AND '.join(w)} GROUP BY i.kind", a)}
        out["kind"] = [{"key": k, "label": lab, "n": counts.get(k, 0)} for k, lab in labels.items()] + \
            [{"key": k, "label": k, "n": n} for k, n in counts.items() if k not in labels]
        w, a = self._filters(user, p, skip="type")
        case = " ".join(f"WHEN f.ext IN ({','.join(repr(e) for e in exts)}) THEN '{key}'"
                        for key, (_, exts) in TYPE_GROUPS.items())
        tcounts = {r["g"]: r["n"] for r in db.rows(
            f"SELECT g, COUNT(DISTINCT input_id) AS n FROM (SELECT f.input_id, CASE {case} ELSE 'other' END AS g "
            f"FROM input_files f JOIN inputs i ON i.id = f.input_id WHERE f.removed = 0 AND {' AND '.join(w)}) t GROUP BY g", a)}
        tcounts["link"] = db.scalar(f"SELECT COUNT(*) FROM inputs i WHERE i.url != '' AND {' AND '.join(w)}", a)
        tcounts["note"] = db.scalar(f"SELECT COUNT(*) FROM inputs i WHERE i.note != '' AND NOT EXISTS (SELECT 1 FROM "
                                    f"input_files f WHERE f.input_id = i.id AND f.removed = 0) AND {' AND '.join(w)}", a)
        out["type"] = [{"key": k, "label": TYPE_LABEL[k], "n": tcounts.get(k, 0)} for k in TYPE_LABEL if tcounts.get(k)]
        w, a = self._filters(user, p, skip="tag")
        out["tag"] = db.rows(f"SELECT t.tag_key AS key, MIN(t.tag) AS label, COUNT(*) AS n FROM input_tags t "
                             f"JOIN inputs i ON i.id = t.input_id WHERE {' AND '.join(w)} GROUP BY t.tag_key "
                             f"ORDER BY n DESC, t.tag_key LIMIT 40", a)
        w, a = self._filters(user, p, skip="project")
        out["project"] = db.rows(f"SELECT p.id AS key, p.code AS label, p.name, p.status, COUNT(*) AS n FROM input_projects x "
                                 f"JOIN projects p ON p.id = x.project_id JOIN inputs i ON i.id = x.input_id "
                                 f"WHERE {' AND '.join(w)} GROUP BY p.id, p.code, p.name, p.status "
                                 f"ORDER BY n DESC, p.code LIMIT 40", a)
        w, a = self._filters(user, p, skip="by")
        out["by"] = db.rows(f"SELECT u.id AS key, u.full_name AS label, COUNT(*) AS n FROM inputs i JOIN users u "
                            f"ON u.id = i.created_by WHERE {' AND '.join(w)} GROUP BY u.id, u.full_name "
                            f"ORDER BY n DESC LIMIT 40", a)
        return out

    def _summaries(self, rows: list[dict], terms: list[str] | None = None) -> list[dict]:
        if not rows:
            return []
        ids = [r["id"] for r in rows]
        marks = ",".join("?" * len(ids))
        labels = self.kind_labels()
        names = {r["id"]: r["full_name"] for r in db.rows("SELECT id, full_name FROM users")}
        tags: dict[int, list] = {}
        for r in db.rows(f"SELECT input_id, tag FROM input_tags WHERE input_id IN ({marks}) ORDER BY tag", ids):
            tags.setdefault(r["input_id"], []).append(r["tag"])
        items: dict[int, list] = {}
        for r in db.rows(f"SELECT x.input_id, it.id, it.part_number, it.description FROM input_items x JOIN items it "
                         f"ON it.id = x.item_id WHERE x.input_id IN ({marks}) ORDER BY it.part_number", ids):
            items.setdefault(r.pop("input_id"), []).append(r)
        projs: dict[int, list] = {}
        for r in db.rows(f"SELECT x.input_id, p.id, p.code, p.name, p.status FROM input_projects x JOIN projects p "
                         f"ON p.id = x.project_id WHERE x.input_id IN ({marks}) ORDER BY p.code", ids):
            projs.setdefault(r.pop("input_id"), []).append(r)
        files: dict[int, list] = {}
        for r in db.rows(f"SELECT f.input_id, f.position AS n, f.filename, f.ext, f.size, c.status, c.thumb_sha, "
                         f"c.text_sha, c.excerpt FROM input_files f LEFT JOIN input_content c ON c.sha = f.blob_sha "
                         f"WHERE f.input_id IN ({marks}) AND f.removed = 0 ORDER BY f.position", ids):
            files.setdefault(r.pop("input_id"), []).append(r)
        out = []
        for r in rows:
            fl = files.get(r["id"], [])
            statuses = {f["status"] for f in fl}
            text_status = "pending" if statuses & {"pending", "running", None} else \
                ("done" if "done" in statuses else (sorted(statuses)[0] if statuses else ""))
            thumb = next((f["n"] for f in fl if f["thumb_sha"]), None)
            s = {
                "id": r["id"], "title": r["title"], "kind": r["kind"], "kind_label": labels.get(r["kind"], r["kind"]),
                "source": r["source"], "date": r["input_date"], "url": r["url"], "has_note": bool(r["note"]),
                "comment": r["comment"], "confidentiality": r["confidentiality"], "deleted": bool(r["deleted"]),
                "created_by": r["created_by"], "created_by_name": names.get(r["created_by"], ""),
                "created_at": r["created_at"], "updated_at": r["updated_at"],
                "tags": tags.get(r["id"], []), "items": items.get(r["id"], []), "projects": projs.get(r["id"], []),
                "files": [{k: f[k] for k in ("n", "filename", "ext", "size", "status")} for f in fl],
                "file_count": len(fl), "types": sorted({type_group(f["ext"]) for f in fl}),
                "thumb": thumb, "text_status": text_status,
            }
            if terms:
                s["score"] = round(r.get("score") or 0, 2)
                s["title_hl"] = self.index.highlight(r["title"], terms)
                s["snippet"], s["snippet_file"] = self._snippet(r, fl, terms)
            out.append(s)
        return out

    def _snippet(self, inp: dict, files: list[dict], terms: list[str]):
        for text, label in ((inp["comment"], "Comment"), (inp["note"], "Note"), (inp["source"], "Source")):
            sn = self.index.snippet(text, terms)
            if sn:
                return sn, label
        for f in files:
            if not f.get("text_sha"):
                continue
            try:
                text = store.read_bytes(f["text_sha"]).decode("utf-8", "replace")
            except OSError:
                continue
            sn = self.index.snippet(text, terms)
            if sn:
                return sn, f["filename"]
        return None, None

    def detail(self, user: dict, input_id: int) -> dict:
        inp = self.get(user, input_id, include_deleted=True)
        manager = self.is_manager(user, inp)
        out = self._summaries([inp])[0]
        out.update({"note": inp["note"], "updated_by_name": db.scalar("SELECT full_name FROM users WHERE id = ?",
                                                                      (inp["updated_by"],)) or ""})
        files = db.rows("SELECT f.position AS n, f.filename, f.ext, f.size, f.removed, f.blob_sha, f.created_at, "
                        "f.created_by, c.status, c.note, c.method, c.chars, c.pages, c.excerpt, c.thumb_sha, c.file_type "
                        "FROM input_files f LEFT JOIN input_content c ON c.sha = f.blob_sha WHERE f.input_id = ? "
                        "ORDER BY f.position", (input_id,))
        out["files"] = []
        for f in files:
            if f["removed"] and not manager:
                continue
            sha = f.pop("blob_sha")
            f["sha"] = sha[:12]
            f["has_thumb"] = bool(f.pop("thumb_sha"))
            f["removed"] = bool(f["removed"])
            f["status"] = f["status"] or "pending"
            f["status_label"] = STATUS_LABEL.get(f["status"], f["status"])
            f["type"] = type_group(f["ext"])
            f["inline"] = f["ext"] in INLINE_MIME and f["ext"] != "svg"
            f["duplicates"] = self.duplicates_of(user, sha, input_id)
            out["files"].append(f)
        out["file_count"] = sum(1 for f in out["files"] if not f["removed"])
        out["access"] = {
            "users": db.rows("SELECT u.id, u.full_name, u.username FROM input_access a JOIN users u ON u.id = a.user_id "
                             "WHERE a.input_id = ? ORDER BY u.full_name", (input_id,)),
            "roles": [r["role"] for r in db.rows("SELECT role FROM input_access WHERE input_id = ? AND role IS NOT NULL "
                                                 "ORDER BY role", (input_id,))],
        }
        lvl = ROLE_LEVEL.get(user.get("role"), 0)

        def a(ok, why=""):
            return {"ok": bool(ok), "why": "" if ok else why}

        eng = "You have view-only access." if lvl < ROLE_LEVEL["engineer"] else ""
        mgr = "Only the person who added this input, an approver or an admin can do this."
        out["actions"] = {
            "edit": a(lvl >= 1 and not inp["deleted"], eng or "Restore the input first."),
            "restrict": a(lvl >= 1 and manager, eng or mgr),
            "delete": a(lvl >= 1 and manager and not inp["deleted"], eng or (mgr if not manager else "Already deleted.")),
            "restore": a(lvl >= 1 and manager and inp["deleted"], eng or (mgr if not manager else "Not deleted.")),
        }
        return out

    def file_row(self, user: dict, input_id: int, n: int) -> dict:
        inp = self.get(user, input_id, include_deleted=True)
        f = db.one("SELECT f.*, c.thumb_sha, c.thumb_kind, c.text_sha FROM input_files f LEFT JOIN input_content c "
                   "ON c.sha = f.blob_sha WHERE f.input_id = ? AND f.position = ?", (input_id, n))
        if not f or ((f["removed"] or inp["deleted"]) and not self.is_manager(user, inp)):
            raise not_found("File not found")
        return f

    def tags(self, user: dict, q: str = "", limit: int = 20) -> list[dict]:
        vis, params = self.visible_sql(user)
        where, args = [vis, "i.deleted = 0"], list(params)
        if q:
            esc = q.strip().lower().replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")
            where.append("t.tag_key LIKE ? ESCAPE '\\'")
            args.append(f"%{esc}%")
        rows = db.rows(f"SELECT t.tag_key AS key, MIN(t.tag) AS tag, COUNT(*) AS n FROM input_tags t JOIN inputs i "
                       f"ON i.id = t.input_id WHERE {' AND '.join(where)} GROUP BY t.tag_key", args)
        ql = (q or "").strip().lower()
        rows.sort(key=lambda r: (not r["key"].startswith(ql), -r["n"], r["key"]))
        return rows[:max(1, min(int(limit), 200))]

    def count_by(self, user: dict, link: str, ids: list[int] | None = None) -> dict[int, int]:
        """Visible inputs per project or per item (``link`` = 'project' or 'item')."""
        table, col = ("input_projects", "project_id") if link == "project" else ("input_items", "item_id")
        vis, params = self.visible_sql(user)
        sql = (f"SELECT x.{col} AS k, COUNT(*) AS n FROM {table} x JOIN inputs i ON i.id = x.input_id "
               f"WHERE i.deleted = 0 AND {vis}")
        args = list(params)
        if ids is not None:
            if not ids:
                return {}
            sql += f" AND x.{col} IN ({','.join('?' * len(ids))})"
            args += ids
        return {r["k"]: r["n"] for r in db.rows(sql + f" GROUP BY x.{col}", args)}


def inp_has_content(inp: dict, input_id: int) -> bool:
    if (inp.get("url") or "").strip() or (inp.get("note") or "").strip():
        return True
    return bool(db.scalar("SELECT 1 FROM input_files WHERE input_id = ? AND removed = 0", (input_id,)))


library = InputLibrary()


def start_worker() -> None:
    """Start the background text reader (called when the server starts)."""
    library.worker.start()


# ===================================================================== demo data ==

def _simple_pdf(title: str, lines: list[str]) -> bytes:
    """A small, valid one-page PDF with a text layer (Helvetica), made without extra libraries."""
    def esc(s: str) -> str:
        return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    content = ["BT", "/F1 18 Tf", "56 780 Td", f"({esc(title)}) Tj", "/F1 11 Tf", "0 -30 Td"]
    for line in lines:
        content += [f"({esc(line)}) Tj", "0 -16 Td"]
    content.append("ET")
    stream = "\n".join(content).encode("latin-1", "replace")
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Resources << /Font << /F1 4 0 R >> >> "
            b"/Contents 5 0 R >>",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
            b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = []
    for i, o in enumerate(objs, 1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode() + o + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def _benchmark_sheet() -> tuple[str, bytes]:
    rows = [["Model", "Maker", "Motor kW", "Shaft Ø mm", "Speed rpm", "Seal", "Price index"],
            ["VA-1200 (ours)", "Demo Engineering", 15, 110, 84, "Mechanical seal", 100],
            ["Agitator AX-12", "Competitor A", 18.5, 120, 90, "Packed gland", 108],
            ["TopMix 1250", "Competitor B", 15, 100, 75, "Mechanical seal", 94],
            ["Note", "Competitor A uses 42CrMo4 shafts with a Ø120 hub seat; Competitor B has no bottom bearing."]]
    try:
        import openpyxl
    except ImportError:
        csv = "\n".join(",".join(f'"{c}"' if isinstance(c, str) and "," in c else str(c) for c in r) for r in rows)
        return "Agitator benchmark 2026.csv", csv.encode("utf-8")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Agitators"
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return "Agitator benchmark 2026.xlsx", buf.getvalue()


def _demo_photo() -> bytes | None:
    try:
        from PIL import Image, ImageDraw, ImageFont
    except ImportError:
        return None
    try:
        font = ImageFont.load_default(size=30)
    except (TypeError, OSError, ImportError):    # older Pillow or no FreeType: small bitmap font
        font = ImageFont.load_default()
    img = Image.new("RGB", (640, 420), (214, 220, 226))
    d = ImageDraw.Draw(img)
    d.rectangle((60, 60, 580, 360), fill=(236, 238, 240), outline=(90, 98, 106), width=6)
    for k, line in enumerate(("CENTRIFUGAL PUMP CP-80", "7.5 kW   2900 rpm", "Q 80 m3/h   H 32 m")):
        d.text((100, 110 + 70 * k), line, fill=(20, 24, 28), font=font)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def seed_demo() -> None:
    """Demo projects and inputs for the invented Demo Engineering. Every name here is made up."""
    if db.scalar("SELECT COUNT(*) FROM inputs") or db.scalar("SELECT COUNT(*) FROM projects"):
        return
    try:
        from . import projects
        users = {r["username"]: r for r in db.rows("SELECT id, username, role FROM users")}
        asha, daniel = users.get("asha"), users.get("daniel")
        if not asha or not daniel:
            return
        item = {r["part_number"]: r["id"] for r in db.rows("SELECT id, part_number FROM items")}
        p1 = projects.create(asha, {"code": "PRJ-2026-014", "name": "Vertical agitator for a dairy plant",
                                    "customer": "Customer A (dairy)", "start_date": "2026-03-02",
                                    "due_date": "2026-12-18",
                                    "description": "Two MX-VA-1200 agitators for milk storage tanks."})
        p2 = projects.create(asha, {"code": "PRJ-2025-031", "name": "Screw conveyor upgrade",
                                    "customer": "Customer B (cement)", "start_date": "2025-04-01",
                                    "due_date": "2025-11-30", "status": "closed"})
        for pn in ("MX-100-000", "MX-100-020"):
            if pn in item:
                projects.link_item(asha, p1["id"], item[pn])
        if "SC-300-021" in item:
            projects.link_item(asha, p2["id"], item["SC-300-021"])

        def upload(name: str, data: bytes) -> dict:
            return InputLibrary.store_upload(Upload(name, io.BytesIO(data)))

        lib = library
        name, data = _benchmark_sheet()
        lib.create(asha, {"title": "Agitator benchmark: Competitor A and Competitor B", "kind": "competitor_benchmark",
                          "source": "Trade fair visit, hall 4", "date": "2026-03-18",
                          "tags": ["agitator", "benchmark", "shaft"],
                          "item_ids": [item[x] for x in ("MX-100-000", "MX-100-020") if x in item],
                          "project_ids": [p1["id"]],
                          "comment": "Competitor A's shaft is bigger; check our Ø110 hub seat against it."},
                   [upload(name, data)])
        pdf = _simple_pdf("Spherical roller bearing 22220 EK", [
            "Supplier C - product datasheet (example data)",
            "Bore d = 100 mm, outside diameter D = 180 mm, width B = 46 mm",
            "Basic dynamic load rating C = 425 kN; static C0 = 490 kN",
            "Reference speed 3200 rpm; limiting speed 4300 rpm",
            "Tapered bore 1:12, use with adapter sleeve H 320 and lock nut KM 20",
            "Mass 4.9 kg"])
        lib.create(asha, {"title": "Spherical roller bearing 22220 EK datasheet", "kind": "supplier_datasheet",
                          "source": "Supplier C catalogue", "date": "2026-02-05", "tags": ["bearing", "22220"],
                          "item_ids": [item["STD-BRG-22220"]] if "STD-BRG-22220" in item else []},
                   [upload("22220 EK datasheet.pdf", pdf)])
        lib.create(asha, {"title": "Kick-off meeting with Customer A", "kind": "meeting_note",
                          "source": "Customer A, project kick-off", "date": "2026-03-04", "tags": ["noise", "CIP"],
                          "project_ids": [p1["id"]],
                          "note": "Noise limit 75 dB(A) at 1 m. Tanks are cleaned in place (CIP) at 85 °C with "
                                  "caustic; all wetted parts 1.4404 (316L), surface Ra 0.8 or better. Shaft seals must "
                                  "survive CIP. Delivery wanted before the plant shutdown in December."})
        lib.create(asha, {"title": "Hygienic design: surface finish guidance", "kind": "standard",
                          "source": "Industry guideline (summary page)", "date": "2026-03-10",
                          "tags": ["hygienic", "surface finish"], "project_ids": [p1["id"]],
                          "url": "https://www.example.com/hygienic-design/surface-finish",
                          "comment": "Summary of the surface finish rules for food-contact parts."})
        lib.create(daniel, {"title": "Customer A target price", "kind": "customer_requirement",
                            "source": "Sales, after the kick-off", "date": "2026-03-06", "tags": ["price"],
                            "project_ids": [p1["id"]], "confidentiality": "restricted", "access_roles": ["approver"],
                            "note": "Target price per agitator agreed with sales; do not share outside the approvers."})
        photo = _demo_photo()
        if photo:
            lib.create(asha, {"title": "CP-80 pump nameplate (field photo)", "kind": "photo",
                              "source": "Site visit, Customer B", "date": "2025-06-12", "tags": ["nameplate", "pump"],
                              "project_ids": [p2["id"]]},
                       [upload("CP-80 nameplate.png", photo)])
        log.info("Demo inputs and projects loaded")
    except Exception:  # demo extras must never stop the demo from loading
        log.exception("Demo inputs could not be loaded")


if __name__ == "__main__":
    sys.exit(_child_main())
