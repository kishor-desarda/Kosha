# SPDX-License-Identifier: AGPL-3.0-or-later
"""Duplicate / common-part detection.

Four independent checks, each producing review groups of a different kind:

  identical    the same file content stored more than once (copies on different drives)
  part_number  one part number used on two different drawings or models
  geometry     drawings or models whose geometry is (nearly) the same: the same part
               drawn again under another number (signature from cad/dxf.py, cad/mesh.py)
  description  bought-out and standard parts described the same way with the same sizes,
               e.g. "HEX BOLT M16X60 GR8.8" and "BLT HEX M16 X 60 8.8"
"""
from __future__ import annotations

import json
import logging
import math
import re
import threading
from collections import Counter

from . import db, lifecycle

log = logging.getLogger("kosha.similarity")

try:  # optional speed-up
    import numpy as _np
except Exception:  # pragma: no cover
    _np = None

GRID = 16


def _flip_x(grid: list[float]) -> list[float]:
    return [grid[r * GRID + (GRID - 1 - c)] for r in range(GRID) for c in range(GRID)]


def _flip_y(grid: list[float]) -> list[float]:
    return [grid[(GRID - 1 - r) * GRID + c] for r in range(GRID) for c in range(GRID)]


def _cos(a, b) -> float:
    if _np is not None:
        a, b = _np.asarray(a), _np.asarray(b)
        na, nb = float(_np.linalg.norm(a)), float(_np.linalg.norm(b))
        return float(a.dot(b) / (na * nb)) if na and nb else 0.0
    dot = sum(x * y for x, y in zip(a, b))
    na = math.sqrt(sum(x * x for x in a))
    nb = math.sqrt(sum(y * y for y in b))
    return dot / (na * nb) if na and nb else 0.0


def _ratio(a: float, b: float) -> float:
    if a <= 0 and b <= 0:
        return 1.0
    if a <= 0 or b <= 0:
        return 0.0
    return min(a, b) / max(a, b)


def _list_sim(a: list[float], b: list[float]) -> float:
    if not a and not b:
        return 1.0
    if not a or not b:
        return 0.0
    used = [False] * len(b)
    total = 0.0
    for x in a:
        best, bi = 0.0, -1
        for j, y in enumerate(b):
            if used[j]:
                continue
            r = _ratio(x, y)
            if r > best:
                best, bi = r, j
        if bi >= 0:
            used[bi] = True
            total += best ** 4  # punish real size differences
    return total / max(len(a), len(b))


def _tokens(s: str) -> set[str]:
    return {t for t in re.split(r"[^A-Z0-9]+", (s or "").upper()) if len(t) > 1}


def score(a: dict, b: dict, desc_a: str = "", desc_b: str = "") -> float:
    if a.get("kind") == "mesh" or b.get("kind") == "mesh":
        if a.get("kind") != b.get("kind"):
            return 0.0
        dims = sum(_ratio(x, y) for x, y in zip(a["bbox"], b["bbox"])) / 3
        geo = 0.45 * dims + 0.3 * _ratio(a["volume"], b["volume"]) + 0.25 * _ratio(a["area"], b["area"])
    else:
        size = _ratio(a["w"], b["w"]) * _ratio(a["h"], b["h"])
        ga = a["grid"]
        gb = b["grid"]
        grid = max(_cos(ga, gb), _cos(ga, _flip_x(gb)), _cos(ga, _flip_y(gb)))
        diam = _list_sim(a.get("diam", [])[:12], b.get("diam", [])[:12])
        ca, cb = a.get("counts", {}), b.get("counts", {})
        counts = sum(1 - abs(ca.get(k, 0) - cb.get(k, 0)) / max(ca.get(k, 0), cb.get(k, 0), 1)
                     for k in ("line", "poly", "circle", "arc")) / 4
        length = _ratio(a.get("len_ratio", 0), b.get("len_ratio", 0))
        geo = 0.40 * grid + 0.20 * size + 0.15 * diam + 0.15 * counts + 0.10 * length
    ta, tb = _tokens(desc_a), _tokens(desc_b)
    if ta and tb:
        # descriptions nudge the score by at most ±2 %; geometry decides
        jac = len(ta & tb) / len(ta | tb)
        return round(max(0.0, min(1.0, geo + 0.04 * (jac - 0.5))), 4)
    return round(geo, 4)


def _prefilter(a: dict, b: dict) -> bool:
    if (a.get("kind") == "mesh") != (b.get("kind") == "mesh"):
        return False
    if a.get("kind") == "mesh":
        return _ratio(a["bbox"][0], b["bbox"][0]) > 0.7
    if _ratio(a["w"], b["w"]) < 0.75 or _ratio(a["h"], b["h"]) < 0.6:
        return False
    ca, cb = a.get("counts", {}).get("circle", 0), b.get("counts", {}).get("circle", 0)
    return abs(ca - cb) <= max(2, 0.4 * max(ca, cb))



KINDS = ("identical", "part_number", "geometry", "description")
KIND_LABEL = {"identical": "Identical files", "part_number": "Same part number, different files",
              "geometry": "Similar geometry", "description": "Similar description"}

# ---------------------------------------------------------------- scope

def _scope(folder_id: int | None) -> set | None:
    if not folder_id:
        return None
    ids = {folder_id}
    frontier = [folder_id]
    while frontier:
        kids = [r["id"] for r in db.rows(
            f"SELECT id FROM folders WHERE parent_id IN ({','.join('?' * len(frontier))})", frontier)]
        frontier = [k for k in kids if k not in ids]
        ids.update(frontier)
    return ids


def _docs(scope: set | None) -> list[dict]:
    rows = db.rows(
        "SELECT d.id, d.folder_id, d.filename, d.ext, d.item_id, d.description, d.part_number, d.state, d.created_at, "
        "v.blob_sha, v.signature_json FROM documents d JOIN versions v ON v.document_id = d.id AND "
        "v.version = d.latest_version WHERE d.deleted = 0 AND d.state != 'obsolete'")
    return [r for r in rows if scope is None or r["folder_id"] in scope]


_ORDER = {"released": 0, "in_review": 1, "in_work": 2}


def _doc_rank(d: dict):
    return (0 if _is_standard(d) else 1, _ORDER.get(d["state"], 3), d["created_at"], d["id"])


def _is_standard(d: dict) -> bool:
    pn = (d.get("part_number") or "").upper()
    if pn.startswith("STD"):
        return True
    path = " ".join(f["name"].lower() for f in lifecycle.folder_path(d["folder_id"]))
    return "standard" in path


# ---------------------------------------------------------------- 1 identical

def identical_groups(docs: list[dict]) -> list[dict]:
    by_sha: dict[str, list] = {}
    for d in docs:
        by_sha.setdefault(d["blob_sha"], []).append(d)
    out = []
    for sha, ds in by_sha.items():
        if len(ds) < 2:
            continue
        ds.sort(key=_doc_rank)
        out.append({"kind": "identical", "label": f"{ds[0]['filename']} ({len(ds)} copies)",
                    "members": [(d, 1.0) for d in ds]})
    return out


# ---------------------------------------------------------------- 2 part number used twice

def part_number_groups(docs: list[dict]) -> list[dict]:
    by_key: dict[tuple, list] = {}
    for d in docs:
        if d["item_id"] and d["ext"] in ("dwg", "dxf", "slddrw", "sldprt", "sldasm"):
            by_key.setdefault((d["item_id"], d["ext"]), []).append(d)
    out = []
    for (_, ext), ds in by_key.items():
        if len({d["blob_sha"] for d in ds}) < 2:
            continue
        ds.sort(key=_doc_rank)
        out.append({"kind": "part_number", "label": f"{ds[0]['part_number']} on {len(ds)} different .{ext} files",
                    "members": [(d, 1.0) for d in ds]})
    return out


# ---------------------------------------------------------------- 3 geometry

def geometry_groups(docs: list[dict], threshold: float) -> tuple[list[dict], int]:
    seen_sha = set()
    cands = []
    for d in sorted(docs, key=_doc_rank):
        if not d["signature_json"] or d["blob_sha"] in seen_sha:
            continue  # identical copies are reported by the identical-file check
        seen_sha.add(d["blob_sha"])
        try:
            sig = json.loads(d["signature_json"])
        except ValueError:
            continue
        if sig:
            cands.append({**d, "sig": sig})
    key = lambda c: c["sig"]["bbox"][0] if c["sig"].get("kind") == "mesh" else c["sig"]["w"]
    cands.sort(key=key)
    n = len(cands)
    parent = list(range(n))

    def find(i):
        while parent[i] != i:
            parent[i] = parent[parent[i]]
            i = parent[i]
        return i

    for i in range(n):
        wi = key(cands[i])
        for j in range(i + 1, n):
            if key(cands[j]) > wi * 1.34 + 1e-9:
                break
            a, b = cands[i]["sig"], cands[j]["sig"]
            if cands[i]["item_id"] and cands[i]["item_id"] == cands[j]["item_id"]:
                continue  # drawing and model of the same part
            if not _prefilter(a, b):
                continue
            if score(a, b, cands[i]["description"], cands[j]["description"]) >= threshold:
                ri, rj = find(i), find(j)
                if ri != rj:
                    parent[ri] = rj
    clusters: dict[int, list[int]] = {}
    for i in range(n):
        clusters.setdefault(find(i), []).append(i)
    groups = []
    for members in clusters.values():
        ds = [cands[i] for i in members]
        if len({d["item_id"] or -d["id"] for d in ds}) < 2:
            continue
        ds.sort(key=_doc_rank)
        ref = ds[0]
        scored = [(ref, 1.0)] + [(d, score(ref["sig"], d["sig"], ref["description"], d["description"])) for d in ds[1:]]
        groups.append({"kind": "geometry", "label": _label([d["description"] for d in ds], ref["description"] or ref["filename"]),
                       "members": scored})
    return groups, n


def _label(descs: list[str], fallback: str) -> str:
    words = Counter(w for d in descs for w in _tokens(d) if not w.isdigit())
    common = {w for w, c in words.most_common(3) if c >= max(2, len(descs) // 2)}
    ref_order = [w for w in re.split(r"[^A-Z0-9]+", (descs[0] or "").upper()) if w in common]
    ordered = list(dict.fromkeys(ref_order)) + sorted(common - set(ref_order))
    return (" ".join(ordered).title() if ordered else fallback)[:80]


# ---------------------------------------------------------------- 4 description

ABBR = {
    "BRG": "BEARING", "BRGS": "BEARING", "BEARINGS": "BEARING", "HSG": "HOUSING", "HSNG": "HOUSING",
    "ASSY": "ASSEMBLY", "ASM": "ASSEMBLY", "PLT": "PLATE", "PL": "PLATE", "SHFT": "SHAFT", "BLT": "BOLT",
    "BOLTS": "BOLT", "SCR": "SCREW", "SCREWS": "SCREW", "WSHR": "WASHER", "WSH": "WASHER", "WASHERS": "WASHER",
    "FLG": "FLANGE", "FLNG": "FLANGE", "FLANGES": "FLANGE", "CVR": "COVER", "BRKT": "BRACKET", "SPKT": "SPROCKET",
    "HEXAGONAL": "HEX", "HEXAGON": "HEX", "HH": "HEX", "SOC": "SOCKET", "SKT": "SOCKET", "SHCS": "SOCKET",
    "CSK": "COUNTERSUNK", "SS": "STAINLESS", "NUTS": "NUT", "GR": "", "GRD": "", "GRADE": "", "CL": "",
    "HD": "", "HEAD": "", "PCS": "", "NOS": "", "DIA": "", "NO": "", "OF": "", "FOR": "", "AND": "", "WITH": "", "THE": "", "TYPE": "",
    "X": "", "SIZE": "", "MM": "", "LG": "", "LONG": "", "L": "", "THK": "", "STD": "",
}


def normalise(desc: str) -> tuple[frozenset, tuple]:
    s = (desc or "").upper().replace("Ø", " ").replace("%%C", " ").replace("⌀", " ")
    words, nums = set(), []
    for t in re.findall(r"\d+(?:[.,]\d+)?|[A-Z]+", s):
        if t[0].isdigit():
            v = float(t.replace(",", "."))
            nums.append(str(int(v)) if v == int(v) else str(v))
        else:
            w = ABBR.get(t, t)
            if w and (len(w) > 1 or w == "M"):
                words.add(w)
    return frozenset(words), tuple(sorted(nums))


def description_groups(scope: set | None, threshold: float) -> list[dict]:
    items = db.rows("SELECT id, part_number, description, make_buy, item_group, state, created_at FROM items "
                    "WHERE state != 'obsolete' AND standard_item_id IS NULL")
    if scope is not None:
        in_scope = {r["item_id"] for r in db.rows("SELECT DISTINCT item_id FROM documents WHERE deleted = 0 AND "
                                                  f"folder_id IN ({','.join('?' * len(scope))})", list(scope))}
        items = [i for i in items if i["id"] in in_scope]
    with_geometry = {r["item_id"] for r in db.rows(
        "SELECT DISTINCT d.item_id FROM documents d JOIN versions v ON v.document_id = d.id AND v.version = "
        "d.latest_version WHERE d.deleted = 0 AND v.signature_json IS NOT NULL AND d.item_id IS NOT NULL")}
    buckets: dict[tuple, list] = {}
    for it in items:
        words, nums = normalise(it["description"])
        if not nums or not words:
            continue  # without sizes in the text, a description match means nothing
        it["_w"] = words
        buckets.setdefault(nums, []).append(it)
    limit = max(0.5, threshold - 0.3)  # word overlap needed once all sizes match exactly
    groups = []
    for nums, its in buckets.items():
        if len(its) < 2:
            continue
        n = len(its)
        parent = list(range(n))

        def find(i):
            while parent[i] != i:
                parent[i] = parent[parent[i]]
                i = parent[i]
            return i

        for i in range(n):
            for j in range(i + 1, n):
                a, b = its[i], its[j]
                if a["id"] in with_geometry and b["id"] in with_geometry:
                    continue  # both have drawings: geometry decides
                jac = len(a["_w"] & b["_w"]) / len(a["_w"] | b["_w"])
                if jac >= limit:
                    ri, rj = find(i), find(j)
                    if ri != rj:
                        parent[ri] = rj
        clusters: dict[int, list] = {}
        for i in range(n):
            clusters.setdefault(find(i), []).append(its[i])
        for members in clusters.values():
            if len(members) < 2:
                continue
            members.sort(key=lambda i: (0 if i["part_number"].upper().startswith("STD") or
                                        i["item_group"] == "Standard Parts" else 1,
                                        _ORDER.get(i["state"], 3), i["created_at"], i["id"]))
            ref = members[0]
            scored = [(ref, 1.0)] + [(m, round(len(ref["_w"] & m["_w"]) / len(ref["_w"] | m["_w"]), 4))
                                     for m in members[1:]]
            groups.append({"kind": "description", "label": (ref["description"] or ref["part_number"])[:80],
                           "items": scored})
    return groups


# ---------------------------------------------------------------- run

def analyse(folder_id: int | None, threshold: float, kinds=KINDS) -> tuple[list[dict], int]:
    scope = _scope(folder_id)
    docs = _docs(scope)
    groups: list[dict] = []
    if "identical" in kinds:
        groups += identical_groups(docs)
    if "part_number" in kinds:
        groups += part_number_groups(docs)
    scanned = len(docs)
    if "geometry" in kinds:
        g, _ = geometry_groups(docs, threshold)
        groups += g
    if "description" in kinds:
        groups += description_groups(scope, threshold)
    return groups, scanned


def _key(g: dict) -> frozenset:
    if "items" in g:
        return frozenset([g["kind"]] + [f"i{i['id']}" for i, _ in g["items"]])
    return frozenset([g["kind"]] + [f"d{d['id']}" for d, _ in g["members"]])


def start_run(user_id: int, folder_id: int | None, threshold: float, kinds=None) -> int:
    threshold = max(0.5, min(0.995, float(threshold)))
    kinds = [k for k in (kinds or KINDS) if k in KINDS] or list(KINDS)
    run_id = db.insert("cp_runs", {"folder_id": folder_id, "threshold": threshold, "status": "running",
                                   "created_by": user_id, "started_at": db.now(), "message": ",".join(kinds)})

    def work():
        try:
            groups, scanned = analyse(folder_id, threshold, kinds)
            with db.tx():
                decided = set()
                for g in db.rows("SELECT id, kind FROM cp_groups WHERE status IN ('standardized','rejected','resolved')"):
                    ids = [f"d{r['document_id']}" for r in db.rows("SELECT document_id FROM cp_members WHERE group_id = ?", (g["id"],))]
                    ids += [f"i{r['item_id']}" for r in db.rows("SELECT item_id FROM cp_item_members WHERE group_id = ?", (g["id"],))]
                    decided.add(frozenset([g["kind"]] + ids))
                made = 0
                for g in groups:
                    if _key(g) in decided:
                        continue
                    members = g.get("members") or g.get("items")
                    scores = [s for _, s in members[1:]] or [1.0]
                    gid = db.insert("cp_groups", {"run_id": run_id, "label": g["label"], "kind": g["kind"],
                                                  "score_min": min(scores), "score_max": max(scores), "status": "open"})
                    for k, (m, sc) in enumerate(members):
                        if "items" in g:
                            db.insert("cp_item_members", {"group_id": gid, "item_id": m["id"], "score": sc,
                                                          "is_reference": 1 if k == 0 else 0})
                        else:
                            db.insert("cp_members", {"group_id": gid, "document_id": m["id"], "score": sc,
                                                     "is_reference": 1 if k == 0 else 0})
                    made += 1
                # earlier open groups of the analysed kinds are replaced by this run
                db.execute(f"DELETE FROM cp_groups WHERE status IN ('open','later') AND run_id != ? AND kind IN "
                           f"({','.join('?' * len(kinds))})", [run_id, *kinds])
                db.update("cp_runs", run_id, {"status": "done", "scanned": scanned, "groups_found": made,
                                              "finished_at": db.now()})
        except Exception as e:  # pragma: no cover
            log.exception("Duplicate analysis failed")
            db.update("cp_runs", run_id, {"status": "failed", "message": str(e), "finished_at": db.now()})
        finally:
            db.close_thread_connection()

    threading.Thread(target=work, name=f"cp-run-{run_id}", daemon=True).start()
    return run_id


# ---------------------------------------------------------------- review

def _where_used(item_id):
    return db.rows("SELECT DISTINCT b.item_id FROM bom_lines bl JOIN boms b ON b.id = bl.bom_id WHERE bl.child_item_id = ?",
                   (item_id,))


def group_detail(group_id: int) -> dict | None:
    g = db.one("SELECT * FROM cp_groups WHERE id = ?", (group_id,))
    if not g:
        return None
    names = lifecycle.user_names()
    members = []
    for m in db.rows("SELECT * FROM cp_members WHERE group_id = ? ORDER BY is_reference DESC, score DESC", (group_id,)):
        d = db.one("SELECT * FROM documents WHERE id = ?", (m["document_id"],))
        if not d:
            continue
        v = db.one("SELECT signature_json, meta_json, blob_sha FROM versions WHERE document_id = ? AND version = ?",
                   (d["id"], d["latest_version"]))
        sig = json.loads(v["signature_json"]) if v and v["signature_json"] else {}
        meta = json.loads(v["meta_json"]) if v else {}
        item = db.one("SELECT material, standard_item_id FROM items WHERE id = ?", (d["item_id"],)) if d["item_id"] else None
        s = lifecycle.doc_summary(d, names)
        s["preview_kind"] = db.scalar("SELECT preview_kind FROM versions WHERE document_id = ? AND version = ?",
                                      (d["id"], d["latest_version"]))
        members.append({
            "document": s, "score": m["score"], "is_reference": bool(m["is_reference"]),
            "folder": " › ".join(f["name"] for f in lifecycle.folder_path(d["folder_id"])),
            "sha": (v or {}).get("blob_sha", "")[:12], "source_path": meta.get("source_path", ""),
            "size": [sig.get("w"), sig.get("h")] if "w" in sig else sig.get("bbox"),
            "diameters": sig.get("diam", [])[:8], "counts": sig.get("counts"),
            "material": (item or {}).get("material") or (meta.get("title_block") or {}).get("material", ""),
            "where_used": len(_where_used(d["item_id"])) if d["item_id"] else 0,
        })
    items = []
    for m in db.rows("SELECT * FROM cp_item_members WHERE group_id = ? ORDER BY is_reference DESC, score DESC", (group_id,)):
        it = db.one("SELECT * FROM items WHERE id = ?", (m["item_id"],))
        if not it:
            continue
        items.append({"item": {k: it[k] for k in ("id", "part_number", "description", "material", "make_buy", "uom",
                                                  "state", "revision", "erp_stock", "erp_synced_at", "item_group")},
                      "score": m["score"], "is_reference": bool(m["is_reference"]),
                      "where_used": len(_where_used(it["id"])),
                      "files": db.scalar("SELECT COUNT(*) FROM documents WHERE item_id = ? AND deleted = 0", (it["id"],))})
    return {**g, "kind_label": KIND_LABEL.get(g["kind"], g["kind"]), "members": members, "items": items}


def decide(user: dict, group_id: int, decision: str, standard_document_id: int | None = None, note: str = "",
           standard_item_id: int | None = None) -> dict:
    from .web import ROLE_LEVEL, bad, forbidden, not_found
    g = db.one("SELECT * FROM cp_groups WHERE id = ?", (group_id,))
    if not g:
        raise not_found("Group not found")
    if decision not in ("standardized", "rejected", "later", "open", "resolved", "obsolete_others"):
        raise bad("Unknown decision")
    affected: list[int] = []
    done: list[str] = []
    if decision == "obsolete_others":
        if ROLE_LEVEL[user["role"]] < ROLE_LEVEL["approver"]:
            raise forbidden("Making files obsolete needs the approver role.")
        if not standard_document_id:
            raise bad("Pick the file to keep.")
        keep = int(standard_document_id)
        for m in db.rows("SELECT document_id FROM cp_members WHERE group_id = ?", (group_id,)):
            if m["document_id"] == keep:
                continue
            d = db.one("SELECT * FROM documents WHERE id = ? AND deleted = 0", (m["document_id"],))
            if not d:
                continue
            if d["state"] == "released":
                try:       # the duplicate decision is the recorded decision, so no change order is needed
                    lifecycle.set_state(user, d["id"], "obsolete", f"Duplicate of document #{keep} ({g['label']})",
                                        via_change=True)
                    done.append(f"{d['filename']} made obsolete")
                except Exception as e:  # still used in a released assembly: leave it
                    done.append(f"{d['filename']} kept ({getattr(e, 'message', e)})")
            elif d["state"] == "in_work" and not d["checked_out_by"]:
                try:
                    lifecycle.delete(user, d["id"], True)
                    done.append(f"{d['filename']} deleted")
                except Exception as e:  # released before: leave it
                    done.append(f"{d['filename']} kept ({getattr(e, 'message', e)})")
        decision = "resolved"
    with db.tx():
        if decision == "standardized":
            std_item = None
            if standard_item_id:
                std_item = int(standard_item_id)
            elif standard_document_id:
                std_doc = lifecycle.doc_or_404(int(standard_document_id))
                std_item = std_doc["item_id"]
                if not std_item:
                    raise bad("The standard file needs a part number first.")
            else:
                raise bad("Pick which part becomes the standard part.")
            db.update("items", std_item, {"item_group": "Standard Parts", "updated_at": db.now()})
            others = [r["item_id"] for r in db.rows("SELECT item_id FROM cp_item_members WHERE group_id = ?", (group_id,))]
            for m in db.rows("SELECT document_id FROM cp_members WHERE group_id = ?", (group_id,)):
                d = db.one("SELECT item_id FROM documents WHERE id = ?", (m["document_id"],))
                if d and d["item_id"]:
                    others.append(d["item_id"])
            for iid in set(others):
                if iid != std_item:
                    db.update("items", iid, {"standard_item_id": std_item, "updated_at": db.now()})
                    affected.extend(r["item_id"] for r in _where_used(iid))
            g_upd = {"standard_item_id": std_item}
        else:
            g_upd = {}
        db.update("cp_groups", group_id, {"status": decision, "standard_document_id": standard_document_id,
                                          "decided_by": user["id"], "decided_at": db.now(), "note": note or "", **g_upd})
        db.audit(user["id"], f"duplicate_{decision}", "cp_group", group_id, g["label"])
    assemblies = db.rows(f"SELECT id, part_number, description FROM items WHERE id IN ({','.join('?' * len(set(affected)))})",
                         list(set(affected))) if affected else []
    return {"assemblies_to_update": assemblies, "done": done}


def report_bytes(status: str = "open") -> tuple[bytes, str]:
    """Excel workbook of duplicate groups for offline review."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    groups = db.rows("SELECT * FROM cp_groups" + ("" if status == "all" else " WHERE status = ?") + " ORDER BY kind, id",
                     () if status == "all" else (status,))
    ws.append(["Duplicate review", f"status: {status}", db.now()])
    ws.append([])
    ws.append(["Check", "Groups", "What it means"])
    meaning = {"identical": "The same file is stored more than once. Keep one copy.",
               "part_number": "One part number is used on different drawings/models. One of them needs a new number.",
               "geometry": "Different part numbers, (nearly) the same part. Choose one standard part.",
               "description": "Bought-out/standard parts described the same way with the same sizes."}
    for k in KINDS:
        ws.append([KIND_LABEL[k], sum(1 for g in groups if g["kind"] == k), meaning[k]])
    for c in ws[3]:
        c.font = Font(bold=True)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["C"].width = 90
    sh = wb.create_sheet("Groups")
    head = ["Group", "Check", "Group label", "Status", "Reference", "Part number", "Description", "File / folder",
            "Similarity", "State", "Used in assemblies", "Original location"]
    sh.append(head)
    for c in sh[1]:
        c.font = Font(bold=True)
    fill = PatternFill("solid", fgColor="EEF3FB")
    for g in groups:
        d = group_detail(g["id"])
        first = True
        for m in d["members"]:
            doc = m["document"]
            sh.append([g["id"], KIND_LABEL.get(g["kind"]), g["label"], g["status"], "yes" if m["is_reference"] else "",
                       doc["part_number"], doc["description"], f"{m['folder']} › {doc['filename']}",
                       round(m["score"] * 100), doc["state_label"], m["where_used"], m["source_path"]])
            if first:
                for c in sh[sh.max_row]:
                    c.fill = fill
                first = False
        for m in d["items"]:
            it = m["item"]
            sh.append([g["id"], KIND_LABEL.get(g["kind"]), g["label"], g["status"], "yes" if m["is_reference"] else "",
                       it["part_number"], it["description"], f"{m['files']} file(s)", round(m["score"] * 100),
                       it["state"], m["where_used"], ""])
            if first:
                for c in sh[sh.max_row]:
                    c.fill = fill
                first = False
    for col, w in zip("ABCDEFGHIJKL", (8, 28, 34, 12, 10, 18, 36, 60, 10, 11, 10, 60)):
        sh.column_dimensions[col].width = w
    sh.freeze_panes = "A2"
    import io
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), "duplicate-review.xlsx"
