# SPDX-License-Identifier: AGPL-3.0-or-later
"""Quality: nonconformance reports, corrective and preventive actions, and FMEA.

**Nonconformance report (NCR)** ``NCR-YY-NNNN``: something does not match its drawing (found at incoming
inspection, in production, at assembly, at the customer or at a supplier). It names the part and revision, the batch
or serials and the quantity, and describes what is wrong, with photos. Lifecycle::

    open -> dispositioned -> closed            (cancelled from open)

An approver or someone with the quality role decides the **disposition**: use as is, rework (back to the drawing),
repair (not back to the drawing), scrap or return to the supplier. *Use as is* and *repair* leave the part different
from its drawing, so Kosha raises a draft deviation for it (the departure is then approved there). Closing records
what was done; an NCR can be linked to a corrective action.

**Corrective and preventive action (CAPA)** ``CAPA-YY-NNNN``: the cause of a problem found and removed, so it does not
come back (corrective) or does not happen at all (preventive). Lifecycle::

    open -> actions -> verifying -> closed     (cancelled from open or actions)

- *open*: the problem and the immediate containment. Moving on needs the **root cause** (for example five whys).
- *actions*: the actions, each with an owner and a due date; overdue ones are marked. Moving on needs every action
  done. A change request can be raised from a CAPA when the design has to change.
- *verifying*: someone other than the owner (four eyes), an approver or quality, checks that it worked and closes
  it with the **effectiveness** check written down.

**FMEA** (design or process): for a part or assembly, rows of function, failure mode, effect, severity, cause,
occurrence, prevention and detection controls and detection, each scored 1 to 10. Kosha works out the risk priority
number (RPN = S × O × D), marks rows over the company's threshold and every row with severity 9 or 10 (safety or
the law), which needs an action whatever its RPN, and keeps the scores after the action next to the ones before.

Records are never deleted: cancelled ones stay in the register. Numbers are allocated inside the write transaction and
never reused.
"""
from __future__ import annotations

import io
import logging
from datetime import date, datetime, timezone
from typing import BinaryIO

from . import access, db, lifecycle, store
from .web import ROLE_LEVEL, bad, conflict, forbidden, not_found

log = logging.getLogger("kosha.quality")

FOUND_AT = {"incoming": "Incoming inspection", "production": "Production", "assembly": "Assembly",
            "customer": "At the customer", "supplier": "At the supplier"}
DISPOSITIONS = {"use_as_is": "Use as is", "rework": "Rework to drawing", "repair": "Repair",
                "scrap": "Scrap", "return": "Return to supplier"}
DEPARTS = ("use_as_is", "repair")
NCR_STATES = {"open": "Open", "dispositioned": "Dispositioned", "closed": "Closed", "cancelled": "Cancelled"}
CAPA_STATES = {"open": "Open: finding the cause", "actions": "Actions under way", "verifying": "Checking it worked",
               "closed": "Closed", "cancelled": "Cancelled"}
CAPA_KINDS = {"corrective": "Corrective", "preventive": "Preventive"}
FMEA_KINDS = {"design": "Design FMEA", "process": "Process FMEA"}
DEFAULT_RPN_LIMIT = 100
FMEA_FIELDS = ("function", "failure_mode", "effect", "severity", "cause", "occurrence", "prevention",
               "detection_control", "detection", "action", "owner_id", "due_date", "action_taken", "new_severity",
               "new_occurrence", "new_detection")
SCORES = ("severity", "occurrence", "detection", "new_severity", "new_occurrence", "new_detection")


def today() -> date:
    """The workshop's calendar day. Tests replace this."""
    return date.today()


def _level(user: dict) -> int:
    return ROLE_LEVEL[user["role"]]


def _may_decide(user: dict) -> bool:
    return _level(user) >= ROLE_LEVEL["approver"] or access.has_role(user, "quality")


def _may_raise(user: dict) -> bool:
    return _level(user) >= ROLE_LEVEL["engineer"] or access.has_role(user, "quality")


def _need(ok: bool, why: str) -> None:
    if not ok:
        raise forbidden(why)


def _text(data: dict, key: str, limit: int = 4000) -> str:
    return str(data.get(key) or "").strip()[:limit]


def _line(data: dict, key: str, limit: int = 200) -> str:
    return " ".join(str(data.get(key) or "").split())[:limit]


def _date(value, label: str) -> str | None:
    if value in (None, ""):
        return None
    try:
        return date.fromisoformat(str(value).strip()[:10]).isoformat()
    except ValueError:
        raise bad(f"{label} must be a date like 2026-11-30.") from None


def _qty(value) -> float | None:
    if value in (None, ""):
        return None
    try:
        q = float(value)
    except (TypeError, ValueError):
        raise bad("The quantity must be a number.") from None
    if q <= 0:
        raise bad("The quantity must be more than zero.")
    return q


def _user_id(value, label: str) -> int | None:
    if value in (None, ""):
        return None
    try:
        uid = int(value)
    except (TypeError, ValueError):
        raise bad(f"Choose the {label}.") from None
    if not db.scalar("SELECT 1 FROM users WHERE id = ? AND active = 1", (uid,)):
        raise bad(f"The {label} is not an active user.")
    return uid


def _item(item_id) -> dict | None:
    if item_id in (None, ""):
        return None
    it = db.one("SELECT id, part_number, description, revision FROM items WHERE id = ?", (int(item_id),))
    if not it:
        raise bad("The part does not exist.")
    access.check_item(it["id"])
    return it


def _allocate(kind: str) -> str:
    """Next NCR-YY-NNNN or CAPA-YY-NNNN. Inside db.tx() (BEGIN IMMEDIATE), which makes it atomic."""
    year = datetime.now(timezone.utc).year
    prefix = f"{kind}-{year % 100:02d}-"
    table = {"NCR": "ncrs", "CAPA": "capas"}[kind]
    db.execute("INSERT INTO quality_numbers (kind, year, last_value) VALUES (?, ?, 0) ON CONFLICT(kind, year) "
               "DO NOTHING", (kind, year))
    n = int(db.scalar("SELECT last_value FROM quality_numbers WHERE kind = ? AND year = ?", (kind, year)) or 0) + 1
    highest = db.scalar(f"SELECT MAX(CAST(SUBSTR(number, ?) AS INTEGER)) FROM {table} WHERE number LIKE ?",
                        (len(prefix) + 1, prefix + "%"))
    if highest is not None and int(highest) >= n:
        n = int(highest) + 1
    db.execute("UPDATE quality_numbers SET last_value = ? WHERE kind = ? AND year = ?", (n, kind, year))
    return f"{prefix}{n:04d}"


# ===================================================================== NCRs ==

def ncr_or_404(ncr_id: int) -> dict:
    n = db.one("SELECT * FROM ncrs WHERE id = ?", (int(ncr_id),))
    if not n:
        raise not_found("Nonconformance report not found")
    if n["item_id"]:
        access.check_item(n["item_id"])
    return n


def create_ncr(user: dict, data: dict) -> int:
    _need(_may_raise(user), "Raising a nonconformance report needs the engineer or quality role.")
    title = _line(data, "title")
    if not title:
        raise bad("Say in a few words what is wrong, for example \"Bore Ø110 measured 110.08, above h6\".")
    found = data.get("found_at") or "production"
    if found not in FOUND_AT:
        raise bad("Say where it was found: " + ", ".join(FOUND_AT.values()) + ".")
    it = _item(data.get("item_id"))
    now = db.now()
    with db.tx():
        nid = db.insert("ncrs", {
            "number": _allocate("NCR"), "title": title, "description": _text(data, "description"),
            "item_id": it["id"] if it else None, "revision": _line(data, "revision", 10) or (it["revision"] if it else ""),
            "found_at": found, "batch": _line(data, "batch", 200), "qty": _qty(data.get("qty")),
            "supplier": _line(data, "supplier", 120), "state": "open", "raised_by": user["id"],
            "created_at": now, "updated_at": now})
        n = ncr_or_404(nid)
        db.audit(user["id"], "ncr_create", "ncr", nid, f"{n['number']} {title}")
    return nid


def update_ncr(user: dict, ncr_id: int, data: dict) -> None:
    n = ncr_or_404(ncr_id)
    if n["state"] != "open":
        raise conflict("Only open reports can be changed; this one is " + NCR_STATES[n["state"]].lower() + ".",
                       "ncr_not_open")
    _need(n["raised_by"] == user["id"] or _may_decide(user), "Only the person who raised it, an approver or quality "
                                                             "can change it.")
    upd: dict = {}
    for k, limit in (("title", 200), ("batch", 200), ("supplier", 120), ("revision", 10)):
        if k in data:
            upd[k] = _line(data, k, limit)
    if "description" in data:
        upd["description"] = _text(data, "description")
    if "qty" in data:
        upd["qty"] = _qty(data.get("qty"))
    if "found_at" in data:
        if data["found_at"] not in FOUND_AT:
            raise bad("Say where it was found: " + ", ".join(FOUND_AT.values()) + ".")
        upd["found_at"] = data["found_at"]
    if "item_id" in data:
        it = _item(data.get("item_id"))
        upd["item_id"] = it["id"] if it else None
    if upd.get("title") == "":
        raise bad("The title can't be empty.")
    with db.tx():
        db.update("ncrs", n["id"], {**upd, "updated_at": db.now()})
        db.audit(user["id"], "ncr_edit", "ncr", n["id"], ", ".join(upd))


def disposition(user: dict, ncr_id: int, data: dict) -> dict:
    """Decide what happens to the parts. *Use as is* and *repair* raise a draft deviation for the part, unless one
    is named (``deviation_id``)."""
    n = ncr_or_404(ncr_id)
    _need(_may_decide(user), "Deciding what happens to the parts needs the approver or quality role.")
    if n["state"] != "open":
        raise conflict("The disposition is decided once, on an open report.", "ncr_not_open")
    what = data.get("disposition")
    if what not in DISPOSITIONS:
        raise bad("Choose the disposition: " + ", ".join(DISPOSITIONS.values()) + ".")
    note = _text(data, "note", 2000)
    if what in ("use_as_is", "repair") and not note:
        raise bad("Say why the part can be used although it differs from the drawing.")
    if what == "return" and not (n["supplier"] or _line(data, "supplier", 120)):
        raise bad("Name the supplier the parts go back to.")
    dev_id = data.get("deviation_id")
    now = db.now()
    with db.tx():
        if what in DEPARTS and not dev_id:
            from . import deviations
            if not n["item_id"]:
                raise bad("Name the part first: a part used although it differs from its drawing needs a deviation.")
            dev_id = deviations.create(user | {"role": "engineer" if _level(user) < ROLE_LEVEL["engineer"]
                                               else user["role"]}, {
                "title": f"{n['number']}: {DISPOSITIONS[what].lower()} — {n['title']}"[:200],
                "reason": f"Nonconformance {n['number']}: {n['description'] or n['title']}",
                "departure": note, "serials": n["batch"], "qty_limit": n["qty"],
                "affected": [{"item_id": n["item_id"], "revision": n["revision"]}]})
        elif dev_id:
            if not db.scalar("SELECT 1 FROM deviations WHERE id = ?", (int(dev_id),)):
                raise bad("The deviation does not exist.")
        upd = {"state": "dispositioned", "disposition": what, "disposition_note": note,
               "deviation_id": int(dev_id) if dev_id else None, "dispositioned_by": user["id"],
               "dispositioned_at": now, "updated_at": now}
        if data.get("supplier"):
            upd["supplier"] = _line(data, "supplier", 120)
        db.update("ncrs", n["id"], upd)
        db.audit(user["id"], "ncr_disposition", "ncr", n["id"], f"{n['number']}: {DISPOSITIONS[what]}")
    return ncr_detail(n["id"], user)


def close_ncr(user: dict, ncr_id: int, data: dict) -> dict:
    n = ncr_or_404(ncr_id)
    _need(_may_decide(user), "Closing a report needs the approver or quality role.")
    if n["state"] != "dispositioned":
        raise conflict("Decide the disposition before closing the report.", "ncr_not_dispositioned")
    done = _text(data, "note", 2000)
    if not done:
        raise bad("Say what was done with the parts (reworked and re-inspected, scrapped, sent back …).")
    with db.tx():
        db.update("ncrs", n["id"], {"state": "closed", "close_note": done, "closed_by": user["id"],
                                    "closed_at": db.now(), "updated_at": db.now()})
        db.audit(user["id"], "ncr_close", "ncr", n["id"], n["number"])
    return ncr_detail(n["id"], user)


def cancel_ncr(user: dict, ncr_id: int, reason: str) -> dict:
    n = ncr_or_404(ncr_id)
    _need(n["raised_by"] == user["id"] or _may_decide(user), "Only the person who raised it, an approver or "
                                                             "quality can cancel it.")
    if n["state"] != "open":
        raise conflict("Only open reports can be cancelled.", "ncr_not_open")
    reason = _line({"r": reason}, "r", 300)
    if not reason:
        raise bad("Say why the report is cancelled (raised twice, the part was fine after all …).")
    with db.tx():
        db.update("ncrs", n["id"], {"state": "cancelled", "close_note": reason, "closed_by": user["id"],
                                    "closed_at": db.now(), "updated_at": db.now()})
        db.audit(user["id"], "ncr_cancel", "ncr", n["id"], f"{n['number']}: {reason}")
    return ncr_detail(n["id"], user)


def _names() -> dict:
    return lifecycle.user_names()


def ncr_detail(ncr_id: int, user: dict) -> dict:
    n = ncr_or_404(ncr_id)
    names = _names()
    it = db.one("SELECT id, part_number, description FROM items WHERE id = ?", (n["item_id"],)) if n["item_id"] else None
    dev = db.one("SELECT id, number, state FROM deviations WHERE id = ?", (n["deviation_id"],)) if n["deviation_id"] else None
    capa = db.one("SELECT id, number, title, state FROM capas WHERE id = ?", (n["capa_id"],)) if n["capa_id"] else None
    mine = n["raised_by"] == user["id"]
    return {**n, "state_label": NCR_STATES[n["state"]], "found_at_label": FOUND_AT.get(n["found_at"], n["found_at"]),
            "disposition_label": DISPOSITIONS.get(n["disposition"] or "", ""), "item": it, "deviation": dev,
            "capa": capa, "raised_by_name": names.get(n["raised_by"]),
            "dispositioned_by_name": names.get(n["dispositioned_by"]), "closed_by_name": names.get(n["closed_by"]),
            "attachments": attachments("ncr", n["id"]),
            "can": {"edit": n["state"] == "open" and (mine or _may_decide(user)),
                    "disposition": n["state"] == "open" and _may_decide(user),
                    "close": n["state"] == "dispositioned" and _may_decide(user),
                    "cancel": n["state"] == "open" and (mine or _may_decide(user)),
                    "attach": n["state"] not in ("closed", "cancelled") and _may_raise(user),
                    "capa": n["state"] != "cancelled" and not n["capa_id"] and _may_raise(user)},
            "dispositions": DISPOSITIONS, "found_at_choices": FOUND_AT}


def ncr_listing(f: dict) -> list[dict]:
    sql = ("SELECT n.id, n.number, n.title, n.state, n.found_at, n.disposition, n.qty, n.batch, n.supplier, "
           "n.created_at, n.item_id, i.part_number, i.description, c.number AS capa_number FROM ncrs n "
           "LEFT JOIN items i ON i.id = n.item_id LEFT JOIN capas c ON c.id = n.capa_id WHERE 1 = 1")
    args: list = []
    sc = access.scope()
    if f.get("state") == "open":
        sql += " AND n.state IN ('open', 'dispositioned')"
    elif f.get("state"):
        sql += " AND n.state = ?"
        args.append(f["state"])
    if f.get("item_id"):
        sql += " AND n.item_id = ?"
        args.append(int(f["item_id"]))
    if f.get("q"):
        sql += " AND (n.number LIKE ? OR n.title LIKE ? OR i.part_number LIKE ? OR n.batch LIKE ? OR n.supplier LIKE ?)"
        args += [f"%{f['q']}%"] * 5
    out = []
    for r in db.rows(sql + " ORDER BY n.id DESC LIMIT 1000", args):
        if r["item_id"] and not sc.can_see_item(r["item_id"]):
            continue
        r["state_label"] = NCR_STATES[r["state"]]
        r["found_at_label"] = FOUND_AT.get(r["found_at"], r["found_at"])
        r["disposition_label"] = DISPOSITIONS.get(r["disposition"] or "", "")
        out.append(r)
    return out


# ===================================================================== CAPAs ==

def capa_or_404(capa_id: int) -> dict:
    c = db.one("SELECT * FROM capas WHERE id = ?", (int(capa_id),))
    if not c:
        raise not_found("Corrective action not found")
    if c["item_id"]:
        access.check_item(c["item_id"])
    return c


def create_capa(user: dict, data: dict) -> int:
    _need(_may_raise(user), "Starting a corrective action needs the engineer or quality role.")
    title = _line(data, "title")
    if not title:
        raise bad("Say in a few words which problem this removes, for example \"Hub seat fretting on agitator shafts\".")
    kind = data.get("kind") or "corrective"
    if kind not in CAPA_KINDS:
        raise bad("A corrective action removes the cause of a problem that happened; a preventive one stops it "
                  "happening. Choose one.")
    it = _item(data.get("item_id"))
    owner = _user_id(data.get("owner_id"), "owner") or user["id"]
    ncr_ids = [int(x) for x in data.get("ncr_ids") or []]
    now = db.now()
    with db.tx():
        cid = db.insert("capas", {
            "number": _allocate("CAPA"), "title": title, "kind": kind, "problem": _text(data, "problem"),
            "containment": _text(data, "containment"), "item_id": it["id"] if it else None, "owner_id": owner,
            "due_date": _date(data.get("due_date"), "The due date"), "state": "open", "raised_by": user["id"],
            "created_at": now, "updated_at": now})
        for nid in ncr_ids:
            link_ncr(user, cid, nid, tx=False)
        c = capa_or_404(cid)
        db.audit(user["id"], "capa_create", "capa", cid, f"{c['number']} {title}")
    return cid


def link_ncr(user: dict, capa_id: int, ncr_id: int, tx: bool = True) -> None:
    n = ncr_or_404(ncr_id)
    if n["capa_id"] and n["capa_id"] != int(capa_id):
        raise conflict(f"{n['number']} already belongs to another corrective action.", "ncr_linked")

    def go():
        db.update("ncrs", n["id"], {"capa_id": int(capa_id), "updated_at": db.now()})
        db.audit(user["id"], "capa_link_ncr", "capa", int(capa_id), n["number"])
    if tx:
        _need(_may_raise(user), "Linking reports needs the engineer or quality role.")
        with db.tx():
            go()
    else:
        go()


def _may_edit_capa(c: dict, user: dict) -> bool:
    return c["owner_id"] == user["id"] or c["raised_by"] == user["id"] or _may_decide(user)


def update_capa(user: dict, capa_id: int, data: dict) -> None:
    c = capa_or_404(capa_id)
    if c["state"] in ("closed", "cancelled"):
        raise conflict("Closed and cancelled corrective actions are kept as they are.", "capa_closed")
    _need(_may_edit_capa(c, user), "Only the owner, the person who started it, an approver or quality can change it.")
    upd: dict = {}
    if "title" in data:
        upd["title"] = _line(data, "title")
        if not upd["title"]:
            raise bad("The title can't be empty.")
    for k in ("problem", "containment", "root_cause"):
        if k in data:
            upd[k] = _text(data, k)
    if "owner_id" in data:
        upd["owner_id"] = _user_id(data.get("owner_id"), "owner") or c["owner_id"]
    if "due_date" in data:
        upd["due_date"] = _date(data.get("due_date"), "The due date")
    if "item_id" in data:
        it = _item(data.get("item_id"))
        upd["item_id"] = it["id"] if it else None
    with db.tx():
        db.update("capas", c["id"], {**upd, "updated_at": db.now()})
        db.audit(user["id"], "capa_edit", "capa", c["id"], ", ".join(upd))


def actions(capa_id: int) -> list[dict]:
    names = _names()
    t = today().isoformat()
    out = db.rows("SELECT * FROM capa_actions WHERE capa_id = ? AND deleted = 0 ORDER BY id", (capa_id,))
    for a in out:
        a["owner_name"] = names.get(a["owner_id"])
        a["done_by_name"] = names.get(a["done_by"])
        a["overdue"] = not a["done_at"] and bool(a["due_date"]) and a["due_date"] < t
    return out


def add_action(user: dict, capa_id: int, data: dict) -> None:
    c = capa_or_404(capa_id)
    if c["state"] not in ("open", "actions"):
        raise conflict("Actions are added while the cause is found and the actions are under way.", "capa_state")
    _need(_may_edit_capa(c, user), "Only the owner, the person who started it, an approver or quality can add actions.")
    text = _line(data, "text", 500)
    if not text:
        raise bad("Say what will be done, for example \"Change the hub seat tolerance to h6 on MX-100-020\".")
    owner = _user_id(data.get("owner_id"), "person who does it")
    if not owner:
        raise bad("Choose who does it.")
    due = _date(data.get("due_date"), "The due date")
    if not due:
        raise bad("Give the date it should be done by.")
    with db.tx():
        db.insert("capa_actions", {"capa_id": c["id"], "text": text, "owner_id": owner, "due_date": due,
                                   "created_by": user["id"], "created_at": db.now()})
        db.audit(user["id"], "capa_action_add", "capa", c["id"], text)


def finish_action(user: dict, capa_id: int, action_id: int, note: str, done: bool = True) -> None:
    c = capa_or_404(capa_id)
    a = db.one("SELECT * FROM capa_actions WHERE id = ? AND capa_id = ? AND deleted = 0", (int(action_id), c["id"]))
    if not a:
        raise not_found("Action not found")
    if c["state"] not in ("open", "actions"):
        raise conflict("Actions are ticked off while the actions are under way.", "capa_state")
    _need(a["owner_id"] == user["id"] or _may_edit_capa(c, user), "Only the person who does it, the owner, an "
                                                                    "approver or quality can tick it off.")
    with db.tx():
        db.update("capa_actions", a["id"], {"done_at": db.now() if done else None, "done_by": user["id"] if done else None,
                                            "note": str(note or "").strip()[:1000]})
        db.audit(user["id"], "capa_action_done" if done else "capa_action_reopen", "capa", c["id"], a["text"])


def remove_action(user: dict, capa_id: int, action_id: int) -> None:
    c = capa_or_404(capa_id)
    if c["state"] not in ("open", "actions"):
        raise conflict("Actions are kept once the corrective action is being checked.", "capa_state")
    _need(_may_edit_capa(c, user), "Only the owner, the person who started it, an approver or quality can remove "
                                   "actions.")
    with db.tx():
        db.execute("UPDATE capa_actions SET deleted = 1 WHERE id = ? AND capa_id = ?", (int(action_id), c["id"]))
        db.audit(user["id"], "capa_action_remove", "capa", c["id"], str(action_id))


def advance(user: dict, capa_id: int, data: dict) -> dict:
    """Move a CAPA on: open -> actions (root cause known) -> verifying (every action done) -> closed (checked by
    someone else, with the effectiveness written down)."""
    c = capa_or_404(capa_id)
    if c["state"] == "open":
        _need(_may_edit_capa(c, user), "Only the owner, an approver or quality can move it on.")
        root = _text(data, "root_cause") or c["root_cause"]
        if not root.strip():
            raise conflict("Write down the root cause first (for example with five whys).", "no_root_cause")
        nxt, upd = "actions", {"root_cause": root}
    elif c["state"] == "actions":
        _need(_may_edit_capa(c, user), "Only the owner, an approver or quality can move it on.")
        acts = actions(c["id"])
        if not acts:
            raise conflict("Add at least one action that removes the cause.", "no_actions")
        open_ = [a["text"] for a in acts if not a["done_at"]]
        if open_:
            raise conflict(f"{len(open_)} action(s) not done yet: " + "; ".join(open_[:3]), "actions_open")
        nxt, upd = "verifying", {}
    elif c["state"] == "verifying":
        _need(_may_decide(user), "Checking that it worked needs the approver or quality role.")
        if user["id"] == c["owner_id"]:
            raise forbidden("Someone other than the owner checks that it worked (four eyes).")
        eff = _text(data, "effectiveness")
        if not eff:
            raise bad("Write down how you checked that it worked (for example: 3 months, 40 shafts, no fretting).")
        if data.get("effective") in (False, "false", "0", 0):
            with db.tx():   # it did not work: back to finding the cause, with what was learned
                db.update("capas", c["id"], {"state": "open", "effectiveness": eff, "updated_at": db.now()})
                db.audit(user["id"], "capa_not_effective", "capa", c["id"], eff[:200])
            return capa_detail(c["id"], user)
        nxt, upd = "closed", {"effectiveness": eff, "verified_by": user["id"], "closed_at": db.now()}
    else:
        raise conflict(f"This corrective action is {CAPA_STATES[c['state']].lower()}.", "capa_state")
    with db.tx():
        db.update("capas", c["id"], {**upd, "state": nxt, "updated_at": db.now()})
        db.audit(user["id"], f"capa_{nxt}", "capa", c["id"], c["number"])
    return capa_detail(c["id"], user)


def cancel_capa(user: dict, capa_id: int, reason: str) -> dict:
    c = capa_or_404(capa_id)
    _need(_may_decide(user) or c["raised_by"] == user["id"], "Only the person who started it, an approver or quality "
                                                             "can cancel it.")
    if c["state"] not in ("open", "actions"):
        raise conflict("Only corrective actions still open or under way can be cancelled.", "capa_state")
    reason = " ".join(str(reason or "").split())[:300]
    if not reason:
        raise bad("Say why it is cancelled.")
    with db.tx():
        db.update("capas", c["id"], {"state": "cancelled", "effectiveness": reason, "closed_at": db.now(),
                                     "updated_at": db.now()})
        db.audit(user["id"], "capa_cancel", "capa", c["id"], f"{c['number']}: {reason}")
    return capa_detail(c["id"], user)


def raise_change_request(user: dict, capa_id: int, data: dict) -> dict:
    """The design has to change: a change request from the CAPA, linked to it."""
    c = capa_or_404(capa_id)
    _need(_level(user) >= ROLE_LEVEL["engineer"], "Raising a change request needs the engineer role.")
    if c["ecr_id"]:
        raise conflict("A change request was already raised from this corrective action.", "ecr_exists")
    from . import change_requests
    title = _line(data, "title") or c["title"]
    reason = f"Corrective action {c['number']}: {c['problem'] or c['title']}" + (
        f"\nRoot cause: {c['root_cause']}" if c["root_cause"] else "")
    ecr = change_requests.create(user, title, reason[:4000], item_id=data.get("item_id") or c["item_id"])
    with db.tx():
        db.update("capas", c["id"], {"ecr_id": ecr, "updated_at": db.now()})
        db.audit(user["id"], "capa_raise_ecr", "capa", c["id"], str(ecr))
    return capa_detail(c["id"], user)


def capa_detail(capa_id: int, user: dict) -> dict:
    c = capa_or_404(capa_id)
    names = _names()
    acts = actions(c["id"])
    mine = _may_edit_capa(c, user)
    it = db.one("SELECT id, part_number, description FROM items WHERE id = ?", (c["item_id"],)) if c["item_id"] else None
    ecr = db.one("SELECT id, number, title, state FROM change_requests WHERE id = ?", (c["ecr_id"],)) if c["ecr_id"] else None
    nxt = {"open": "Root cause found: plan the actions", "actions": "Actions done: check it worked",
           "verifying": "It worked: close"}.get(c["state"])
    why = ""
    if c["state"] == "open" and not (c["root_cause"] or "").strip():
        why = "Write down the root cause first."
    elif c["state"] == "actions":
        if not acts:
            why = "Add at least one action."
        elif any(not a["done_at"] for a in acts):
            why = "Every action must be done first."
    elif c["state"] == "verifying" and user["id"] == c["owner_id"]:
        why = "Someone other than the owner checks that it worked."
    elif c["state"] == "verifying" and not _may_decide(user):
        why = "Checking it worked needs the approver or quality role."
    return {**c, "state_label": CAPA_STATES[c["state"]], "kind_label": CAPA_KINDS[c["kind"]], "item": it, "ecr": ecr,
            "owner_name": names.get(c["owner_id"]), "raised_by_name": names.get(c["raised_by"]),
            "verified_by_name": names.get(c["verified_by"]), "actions": acts,
            "overdue": (not c["closed_at"] and bool(c["due_date"]) and c["due_date"] < today().isoformat()),
            "ncrs": db.rows("SELECT id, number, title, state FROM ncrs WHERE capa_id = ? ORDER BY id", (c["id"],)),
            "attachments": attachments("capa", c["id"]),
            "next": {"label": nxt, "ok": bool(nxt) and not why and (mine or c["state"] == "verifying"), "why": why},
            "can": {"edit": c["state"] not in ("closed", "cancelled") and mine,
                    "actions": c["state"] in ("open", "actions") and mine,
                    "cancel": c["state"] in ("open", "actions") and (_may_decide(user) or c["raised_by"] == user["id"]),
                    "ecr": not c["ecr_id"] and _level(user) >= ROLE_LEVEL["engineer"] and c["state"] != "cancelled",
                    "attach": c["state"] not in ("closed", "cancelled") and _may_raise(user)}}


def capa_listing(f: dict) -> list[dict]:
    sql = ("SELECT c.id, c.number, c.title, c.kind, c.state, c.owner_id, c.due_date, c.created_at, c.closed_at, "
           "c.item_id, i.part_number, (SELECT COUNT(*) FROM capa_actions a WHERE a.capa_id = c.id AND a.deleted = 0 "
           "AND a.done_at IS NULL) AS open_actions, (SELECT COUNT(*) FROM ncrs n WHERE n.capa_id = c.id) AS ncrs "
           "FROM capas c LEFT JOIN items i ON i.id = c.item_id WHERE 1 = 1")
    args: list = []
    if f.get("state") == "open":
        sql += " AND c.state IN ('open', 'actions', 'verifying')"
    elif f.get("state"):
        sql += " AND c.state = ?"
        args.append(f["state"])
    if f.get("mine"):
        sql += " AND (c.owner_id = ? OR c.id IN (SELECT capa_id FROM capa_actions WHERE owner_id = ? AND deleted = 0 " \
               "AND done_at IS NULL))"
        args += [int(f["mine"])] * 2
    if f.get("q"):
        sql += " AND (c.number LIKE ? OR c.title LIKE ? OR i.part_number LIKE ?)"
        args += [f"%{f['q']}%"] * 3
    names = _names()
    t = today().isoformat()
    sc = access.scope()
    out = []
    for r in db.rows(sql + " ORDER BY c.id DESC LIMIT 1000", args):
        if r["item_id"] and not sc.can_see_item(r["item_id"]):
            continue
        r["state_label"] = CAPA_STATES[r["state"]]
        r["kind_label"] = CAPA_KINDS[r["kind"]]
        r["owner_name"] = names.get(r["owner_id"])
        r["overdue"] = not r["closed_at"] and bool(r["due_date"]) and r["due_date"] < t
        out.append(r)
    return out


# ===================================================================== FMEA ==

def rpn_limit() -> int:
    v = db.get_setting("fmea_rpn_limit")
    return int(v) if v else DEFAULT_RPN_LIMIT


def set_rpn_limit(user: dict, value) -> int:
    try:
        v = int(value)
    except (TypeError, ValueError):
        raise bad("The limit is a whole number from 1 to 1000.") from None
    if not 1 <= v <= 1000:
        raise bad("The limit is a whole number from 1 to 1000.")
    db.set_setting("fmea_rpn_limit", v)
    db.audit(user["id"], "fmea_rpn_limit", "settings", None, str(v))
    return v


def fmea_or_404(fmea_id: int) -> dict:
    f = db.one("SELECT * FROM fmeas WHERE id = ?", (int(fmea_id),))
    if not f:
        raise not_found("FMEA not found")
    if f["item_id"]:
        access.check_item(f["item_id"])
    return f


def create_fmea(user: dict, data: dict) -> int:
    _need(_may_raise(user), "Starting an FMEA needs the engineer or quality role.")
    kind = data.get("kind") or "design"
    if kind not in FMEA_KINDS:
        raise bad("Choose a design FMEA (how the part can fail) or a process FMEA (how making it can go wrong).")
    it = _item(data.get("item_id"))
    if not it:
        raise bad("Choose the part or assembly the FMEA is about.")
    title = _line(data, "title") or f"{FMEA_KINDS[kind]} {it['part_number']}"
    now = db.now()
    with db.tx():
        fid = db.insert("fmeas", {"item_id": it["id"], "kind": kind, "title": title, "team": _line(data, "team", 300),
                                  "created_by": user["id"], "created_at": now, "updated_at": now})
        db.audit(user["id"], "fmea_create", "fmea", fid, title)
    return fid


def _score(value, label: str, k: int) -> int | None:
    if value in (None, ""):
        return None
    try:
        v = int(value)
    except (TypeError, ValueError):
        raise bad(f"Row {k}: {label} is a whole number from 1 to 10.") from None
    if not 1 <= v <= 10:
        raise bad(f"Row {k}: {label} is a whole number from 1 to 10.")
    return v


def save_fmea_rows(user: dict, fmea_id: int, rows, title: str | None = None, team: str | None = None) -> None:
    f = fmea_or_404(fmea_id)
    _need(_may_raise(user), "Changing an FMEA needs the engineer or quality role.")
    if not isinstance(rows, list):
        raise bad("Send the rows as a list.")
    if len(rows) > 500:
        raise bad("An FMEA can have at most 500 rows; split it by function.")
    clean = []
    for k, r in enumerate(rows, start=1):
        if not isinstance(r, dict):
            raise bad(f"Row {k} is not valid.")
        mode = _line(r, "failure_mode", 300)
        if not mode:
            raise bad(f"Row {k}: say how it can fail (the failure mode).")
        row = {"function": _line(r, "function", 300), "failure_mode": mode, "effect": _line(r, "effect", 300),
               "cause": _line(r, "cause", 300), "prevention": _line(r, "prevention", 300),
               "detection_control": _line(r, "detection_control", 300), "action": _line(r, "action", 300),
               "action_taken": _line(r, "action_taken", 300), "owner_id": _user_id(r.get("owner_id"), "owner"),
               "due_date": _date(r.get("due_date"), f"Row {k}: the due date")}
        labels = {"severity": "severity", "occurrence": "occurrence", "detection": "detection",
                  "new_severity": "severity after the action", "new_occurrence": "occurrence after the action",
                  "new_detection": "detection after the action"}
        for s in SCORES:
            row[s] = _score(r.get(s), labels[s], k)
        for s in ("severity", "occurrence", "detection"):
            if row[s] is None:
                raise bad(f"Row {k}: score the {s} from 1 to 10.")
        clean.append(row)
    with db.tx():
        db.execute("DELETE FROM fmea_rows WHERE fmea_id = ?", (f["id"],))
        for p, row in enumerate(clean):
            db.insert("fmea_rows", {"fmea_id": f["id"], "position": p, **row})
        upd = {"updated_at": db.now()}
        if title is not None and title.strip():
            upd["title"] = " ".join(title.split())[:200]
        if team is not None:
            upd["team"] = " ".join(team.split())[:300]
        db.update("fmeas", f["id"], upd)
        db.audit(user["id"], "fmea_edit", "fmea", f["id"], f"{len(clean)} rows")


def _rated(r: dict, limit: int, t: str) -> dict:
    r["rpn"] = r["severity"] * r["occurrence"] * r["detection"]
    after = all(r[s] is not None for s in ("new_severity", "new_occurrence", "new_detection"))
    r["new_rpn"] = r["new_severity"] * r["new_occurrence"] * r["new_detection"] if after else None
    worst_s = r["new_severity"] if after else r["severity"]
    worst = r["new_rpn"] if after else r["rpn"]
    flags = []
    if worst_s >= 9:
        flags.append("Severity 9 or 10 (safety or the law): needs an action whatever the RPN.")
    if worst > limit:
        flags.append(f"RPN {worst} is above the limit of {limit}.")
    if (flags or r["action"]) and not after and not r["action"]:
        flags.append("No action planned yet.")
    r["flags"] = flags
    r["needs_action"] = bool(flags) and not (after and worst <= limit and worst_s < 9)
    r["overdue"] = bool(r["action"]) and not after and bool(r["due_date"]) and r["due_date"] < t
    return r


def fmea_detail(fmea_id: int) -> dict:
    f = fmea_or_404(fmea_id)
    names = _names()
    limit = rpn_limit()
    t = today().isoformat()
    rows = [_rated(r, limit, t) for r in db.rows("SELECT * FROM fmea_rows WHERE fmea_id = ? ORDER BY position, id",
                                                 (f["id"],))]
    for r in rows:
        r["owner_name"] = names.get(r["owner_id"])
    it = db.one("SELECT id, part_number, description, revision FROM items WHERE id = ?", (f["item_id"],))
    top = sorted(rows, key=lambda r: -(r["new_rpn"] if r["new_rpn"] is not None else r["rpn"]))[:5]
    return {**f, "kind_label": FMEA_KINDS[f["kind"]], "item": it, "rows": rows, "rpn_limit": limit,
            "created_by_name": names.get(f["created_by"]),
            "summary": {"rows": len(rows), "need_action": sum(1 for r in rows if r["needs_action"]),
                        "high_severity": sum(1 for r in rows if r["severity"] >= 9),
                        "max_rpn": max([r["rpn"] for r in rows] or [0]),
                        "top": [{"failure_mode": r["failure_mode"], "rpn": r["new_rpn"] or r["rpn"]} for r in top]}}


def fmea_listing(item_id: int | None = None) -> list[dict]:
    sql = ("SELECT f.id, f.title, f.kind, f.item_id, f.updated_at, i.part_number, i.description, "
           "(SELECT COUNT(*) FROM fmea_rows r WHERE r.fmea_id = f.id) AS rows, "
           "(SELECT MAX(r.severity * r.occurrence * r.detection) FROM fmea_rows r WHERE r.fmea_id = f.id) AS max_rpn "
           "FROM fmeas f JOIN items i ON i.id = f.item_id WHERE 1 = 1" + access.scope().items("i"))
    args: list = []
    if item_id:
        sql += " AND f.item_id = ?"
        args.append(int(item_id))
    out = db.rows(sql + " ORDER BY f.updated_at DESC LIMIT 500", args)
    for r in out:
        r["kind_label"] = FMEA_KINDS[r["kind"]]
    return out


def export_fmea(fmea_id: int) -> tuple[bytes, str]:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    d = fmea_detail(fmea_id)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "FMEA"
    ws.append([d["title"], f"{d['item']['part_number']} {d['item']['description']}", d["kind_label"],
               f"RPN limit {d['rpn_limit']}"])
    ws.append([f"Team: {d['team']}" if d["team"] else ""])
    ws.append([])
    head = ["Function", "Failure mode", "Effect", "S", "Cause", "O", "Prevention", "Detection control", "D", "RPN",
            "Action", "Owner", "Due", "Action taken", "S", "O", "D", "RPN after", "Needs attention"]
    ws.append(head)
    for c in ws[4]:
        c.font = Font(bold=True)
    mark = PatternFill("solid", fgColor="FFE0B2")
    for r in d["rows"]:
        ws.append([r["function"], r["failure_mode"], r["effect"], r["severity"], r["cause"], r["occurrence"],
                   r["prevention"], r["detection_control"], r["detection"], r["rpn"], r["action"], r["owner_name"] or "",
                   r["due_date"] or "", r["action_taken"], r["new_severity"], r["new_occurrence"], r["new_detection"],
                   r["new_rpn"], " ".join(r["flags"])])
        for c in ws[ws.max_row]:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        if r["needs_action"]:
            for c in ws[ws.max_row]:
                c.fill = mark
    for col, width in zip("ABCDEFGHIJKLMNOPQRS", (20, 24, 24, 4, 24, 4, 22, 22, 4, 6, 26, 14, 11, 24, 4, 4, 4, 9, 36)):
        ws.column_dimensions[col].width = width
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), f"FMEA_{d['item']['part_number']}_{d['id']}.xlsx"


# ============================================================ attachments ==

KINDS = {"ncr": ncr_or_404, "capa": capa_or_404}


def attachments(kind: str, record_id: int) -> list[dict]:
    names = _names()
    out = db.rows("SELECT id, filename, size, uploaded_by, uploaded_at FROM quality_attachments WHERE kind = ? AND "
                  "record_id = ? AND deleted = 0 ORDER BY id", (kind, record_id))
    for a in out:
        a["uploaded_by_name"] = names.get(a["uploaded_by"])
    return out


def add_attachment(user: dict, kind: str, record_id: int, filename: str, stream: BinaryIO) -> int:
    rec = KINDS[kind](record_id)
    _need(_may_raise(user), "Adding photos and files needs the engineer or quality role.")
    if rec["state"] in ("closed", "cancelled"):
        raise conflict("Closed and cancelled records are kept as they are.", "closed")
    name = lifecycle.clean_filename(filename or "")
    from .filesafety import refuse_blocked
    refuse_blocked(name)
    sha, size = store.put_stream(stream)    # safely on disk before the database knows about it
    with db.tx():
        aid = db.insert("quality_attachments", {"kind": kind, "record_id": rec["id"], "filename": name,
                                                "blob_sha": sha, "size": size, "uploaded_by": user["id"],
                                                "uploaded_at": db.now()})
        db.audit(user["id"], f"{kind}_attach", kind, rec["id"], f"{name} ({sha[:12]})")
    return aid


def attachment_or_404(kind: str, record_id: int, attachment_id: int) -> dict:
    KINDS[kind](record_id)
    a = db.one("SELECT * FROM quality_attachments WHERE id = ? AND kind = ? AND record_id = ? AND deleted = 0",
               (int(attachment_id), kind, int(record_id)))
    if not a:
        raise not_found("Attachment not found")
    return a


def remove_attachment(user: dict, kind: str, record_id: int, attachment_id: int) -> None:
    """Soft delete: the file stays in the store and the row stays in the database."""
    rec = KINDS[kind](record_id)
    a = attachment_or_404(kind, record_id, attachment_id)
    _need(a["uploaded_by"] == user["id"] or _may_decide(user), "Only the person who added it, an approver or quality "
                                                               "can remove it.")
    if rec["state"] in ("closed", "cancelled"):
        raise conflict("Closed and cancelled records are kept as they are.", "closed")
    with db.tx():
        db.update("quality_attachments", a["id"], {"deleted": 1, "deleted_by": user["id"], "deleted_at": db.now()})
        db.audit(user["id"], f"{kind}_attachment_remove", kind, rec["id"], a["filename"])


# ================================================================ overview ==

def overview(user: dict) -> dict:
    """Counts for the Quality screen and the home page."""
    t = today().isoformat()
    return {
        "ncr_open": db.scalar("SELECT COUNT(*) FROM ncrs WHERE state = 'open'") or 0,
        "ncr_dispositioned": db.scalar("SELECT COUNT(*) FROM ncrs WHERE state = 'dispositioned'") or 0,
        "capa_open": db.scalar("SELECT COUNT(*) FROM capas WHERE state IN ('open', 'actions', 'verifying')") or 0,
        "capa_overdue": db.scalar("SELECT COUNT(*) FROM capas WHERE state IN ('open', 'actions', 'verifying') AND "
                                  "due_date IS NOT NULL AND due_date < ?", (t,)) or 0,
        "my_actions": db.scalar("SELECT COUNT(*) FROM capa_actions a JOIN capas c ON c.id = a.capa_id WHERE "
                                "a.owner_id = ? AND a.deleted = 0 AND a.done_at IS NULL AND c.state IN "
                                "('open', 'actions')", (user["id"],)) or 0,
        "verify": db.scalar("SELECT COUNT(*) FROM capas WHERE state = 'verifying' AND owner_id != ?",
                            (user["id"],)) if _may_decide(user) else 0,
        "fmeas": db.scalar("SELECT COUNT(*) FROM fmeas") or 0,
        "rpn_limit": rpn_limit(),
    }
