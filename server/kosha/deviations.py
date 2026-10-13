# SPDX-License-Identifier: AGPL-3.0-or-later
"""Deviations: temporary, approved departures from the released design.

    "Use 42CrMo4 instead of C45 for batch 12, because C45 is out of stock."

A deviation names what it affects (items and/or files at a revision), describes the departure in plain words and
is limited in scope (a quantity, serial or batch numbers, and/or a date range).

Lifecycle::

    draft -> submitted -> approved -> active -> closed
                      \\-> rejected (edit and submit again)
    draft / submitted / rejected / approved -> cancelled

- An engineer raises it; the approvers named on it decide. Four-eyes rule: the requester never approves their own
  deviation, whatever their role.
- When the last approver agrees it becomes active straight away, or on its start date. Add-ins hear
  ``deviation.approved``.
- It closes by itself when its date range ends or the allowed quantity is used up; it can be closed by hand, recording
  the quantity actually used. Closed and cancelled deviations are kept for the register.
- Numbers are ``DEV-YY-NNNN`` (year, running number), allocated inside the write transaction and never reused.
"""
from __future__ import annotations

import io
import logging
from datetime import date, datetime, timezone
from typing import BinaryIO

from . import changes, db, lifecycle, notifications, store
from .web import ROLE_LEVEL, bad, conflict, forbidden, not_found

log = logging.getLogger("kosha.deviations")

STATES = ("draft", "submitted", "approved", "rejected", "active", "closed", "cancelled")
STATE_LABEL = {
    "draft": "Draft", "submitted": "Waiting for approval", "approved": "Approved, starts later",
    "rejected": "Rejected", "active": "Active", "closed": "Closed", "cancelled": "Cancelled",
}
CLOSE_LABEL = {"expired": "Date range ended", "used_up": "Quantity used up", "closed": "Closed by hand"}
OPEN_STATES = ("draft", "submitted", "approved", "rejected", "active")
MONTHS = ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec")
EDITABLE_FIELDS = ("title", "reason", "departure", "qty_limit", "serials", "valid_from", "valid_until")


def today() -> date:
    """The workshop's calendar day (server local time). Tests replace this."""
    return date.today()


def _ok(ok: bool, why: str = "") -> dict:
    return {"ok": bool(ok), "why": "" if ok else why}


def _level(user: dict) -> int:
    return ROLE_LEVEL[user["role"]]


def dev_or_404(dev_id: int) -> dict:
    d = db.one("SELECT * FROM deviations WHERE id = ?", (int(dev_id),))
    if not d:
        raise not_found("Deviation not found")
    return d


# ------------------------------------------------------------------ values --

def _date(value, label: str) -> str | None:
    if value in (None, ""):
        return None
    try:
        return date.fromisoformat(str(value).strip()[:10]).isoformat()
    except ValueError:
        raise bad(f"{label} must be a date like 2026-11-30.")


def _qty(value, label: str = "Quantity") -> float | None:
    if value in (None, ""):
        return None
    try:
        q = float(value)
    except (TypeError, ValueError):
        raise bad(f"{label} must be a number.")
    if q < 0:
        raise bad(f"{label} can't be negative.")
    return q


def _clean(data: dict) -> dict:
    out = {}
    if "title" in data:
        out["title"] = " ".join(str(data["title"] or "").split())[:200]
    for k, limit in (("reason", 4000), ("departure", 4000), ("serials", 500)):
        if k in data:
            out[k] = str(data[k] or "").strip()[:limit]
    if "qty_limit" in data:
        q = _qty(data["qty_limit"], "The quantity limit")
        if q is not None and q <= 0:
            raise bad("The quantity limit must be more than zero, or left empty for no limit.")
        out["qty_limit"] = q
    if "valid_from" in data:
        out["valid_from"] = _date(data["valid_from"], "The start date")
    if "valid_until" in data:
        out["valid_until"] = _date(data["valid_until"], "The end date")
    return out


def _check_dates(valid_from: str | None, valid_until: str | None) -> None:
    if valid_from and valid_until and valid_from > valid_until:
        raise bad("The end date is before the start date.")


def _fmt_day(iso: str | None) -> str:
    if not iso:
        return ""
    d = date.fromisoformat(iso)
    return f"{d.day} {MONTHS[d.month - 1]}" + ("" if d.year == today().year else f" {d.year}")


def _fmt_qty(q: float | None) -> str:
    if q is None:
        return ""
    return str(int(q)) if float(q).is_integer() else f"{q:g}"


def scope_text(d: dict) -> str:
    """The limits in plain words: "batch 12, 4 of 10 used, until 30 Nov"."""
    parts = []
    if d.get("serials"):
        parts.append(d["serials"])
    if d.get("qty_limit") is not None:
        parts.append(f"{_fmt_qty(d.get('qty_used') or 0)} of {_fmt_qty(d['qty_limit'])} used")
    if d.get("valid_from") and d["valid_from"] > today().isoformat():
        parts.append(f"from {_fmt_day(d['valid_from'])}")
    if d.get("valid_until"):
        parts.append(f"until {_fmt_day(d['valid_until'])}")
    return ", ".join(parts)


def banner_text(d: dict) -> str:
    """One sentence for the record pages: "Deviation DEV-26-0003 active: material 42CrMo4 allowed for batch 12, until 30 Nov"."""
    what = (d.get("departure") or d.get("title") or "").strip().rstrip(".")
    text = f"Deviation {d['number']} {'active' if d['state'] == 'active' else STATE_LABEL[d['state']].lower()}: {what}"
    if d.get("serials"):
        text += f" for {d['serials']}"
    rest = []
    if d.get("qty_limit") is not None:
        rest.append(f"{_fmt_qty(max(d['qty_limit'] - (d.get('qty_used') or 0), 0))} of {_fmt_qty(d['qty_limit'])} left")
    if d.get("valid_until"):
        rest.append(f"until {_fmt_day(d['valid_until'])}")
    return text + (", " + ", ".join(rest) if rest else "")


# ---------------------------------------------------------------- numbering --

def _allocate_number() -> str:
    """Next DEV-YY-NNNN. Must run inside db.tx() (BEGIN IMMEDIATE), which makes it atomic."""
    year = datetime.now(timezone.utc).year
    prefix = f"DEV-{year % 100:02d}-"
    db.execute("INSERT INTO deviation_numbers (year, last_value) VALUES (?, 0) ON CONFLICT(year) DO NOTHING", (year,))
    n = int(db.scalar("SELECT last_value FROM deviation_numbers WHERE year = ?", (year,)) or 0) + 1
    # never reuse a number, even if the counter were ever restored from an older backup
    highest = db.scalar("SELECT MAX(CAST(SUBSTR(number, ?) AS INTEGER)) FROM deviations WHERE number LIKE ?",
                        (len(prefix) + 1, prefix + "%"))
    if highest is not None and int(highest) >= n:
        n = int(highest) + 1
    db.execute("UPDATE deviation_numbers SET last_value = ? WHERE year = ?", (n, year))
    return f"{prefix}{n:04d}"


# ----------------------------------------------------------------- editing --

def create(user: dict, data: dict) -> int:
    if _level(user) < ROLE_LEVEL["engineer"]:
        raise forbidden("Raising a deviation needs the engineer role.")
    vals = _clean(data)
    if not vals.get("title"):
        raise bad("Give the deviation a short title, for example \"42CrMo4 instead of C45 for batch 12\".")
    _check_dates(vals.get("valid_from"), vals.get("valid_until"))
    now = db.now()
    with db.tx():
        dev_id = db.insert("deviations", {
            "number": _allocate_number(), "title": vals["title"], "reason": vals.get("reason", ""),
            "departure": vals.get("departure", ""), "state": "draft", "qty_limit": vals.get("qty_limit"),
            "qty_used": 0, "serials": vals.get("serials", ""), "valid_from": vals.get("valid_from"),
            "valid_until": vals.get("valid_until"), "requested_by": user["id"], "created_at": now, "updated_at": now,
        })
        for a in data.get("affected") or []:
            add_affected(user, dev_id, a.get("item_id"), a.get("document_id"), a.get("revision"), a.get("note", ""))
        for uid in data.get("approver_ids") or []:
            add_approver(user, dev_id, int(uid))
        d = dev_or_404(dev_id)
        db.audit(user["id"], "deviation_create", "deviation", dev_id, f"{d['number']} {d['title']}")
    return dev_id


def _can_edit(d: dict, user: dict) -> dict:
    if d["state"] not in ("draft", "rejected"):
        return _ok(False, f"Only draft or rejected deviations can be changed; this one is {STATE_LABEL[d['state']].lower()}.")
    if d["requested_by"] != user["id"] and _level(user) < ROLE_LEVEL["approver"]:
        return _ok(False, "Only the requester or an approver can change this deviation.")
    return _ok(True)


def _editable(d: dict, user: dict) -> None:
    c = _can_edit(d, user)
    if not c["ok"]:
        if d["state"] not in ("draft", "rejected"):
            raise conflict(c["why"], "not_draft")
        raise forbidden(c["why"])


def update(user: dict, dev_id: int, data: dict) -> None:
    d = dev_or_404(dev_id)
    _editable(d, user)
    vals = _clean({k: v for k, v in data.items() if k in EDITABLE_FIELDS})
    if "title" in vals and not vals["title"]:
        raise bad("The title can't be empty.")
    if not vals:
        return
    _check_dates(vals.get("valid_from", d["valid_from"]), vals.get("valid_until", d["valid_until"]))
    with db.tx():
        db.update("deviations", d["id"], {**vals, "updated_at": db.now()})
        db.audit(user["id"], "deviation_update", "deviation", d["id"], ",".join(sorted(vals)))


def add_affected(user: dict, dev_id: int, item_id=None, document_id=None, revision: str | None = None,
                 note: str = "") -> int:
    d = dev_or_404(dev_id)
    _editable(d, user)
    if bool(item_id) == bool(document_id):
        raise bad("Choose either a part (item) or a file.")
    if item_id:
        it = db.one("SELECT id, part_number, revision FROM items WHERE id = ?", (int(item_id),))
        if not it:
            raise not_found("Item not found")
        rev = (revision or it["revision"] or "").strip().upper()[:20]
        if db.scalar("SELECT 1 FROM deviation_affected WHERE deviation_id = ? AND item_id = ? AND revision = ?",
                     (d["id"], it["id"], rev)):
            raise conflict(f"{it['part_number']} rev {rev} is already on {d['number']}.", "duplicate")
        row = {"item_id": it["id"], "document_id": None}
        label = f"{it['part_number']} rev {rev}"
    else:
        doc = lifecycle.doc_or_404(int(document_id))
        rev = (revision or doc["revision"] or "").strip().upper()[:20]
        if db.scalar("SELECT 1 FROM deviation_affected WHERE deviation_id = ? AND document_id = ? AND revision = ?",
                     (d["id"], doc["id"], rev)):
            raise conflict(f"{doc['filename']} rev {rev} is already on {d['number']}.", "duplicate")
        row = {"item_id": None, "document_id": doc["id"]}
        label = f"{doc['filename']} rev {rev}"
    with db.tx():
        aid = db.insert("deviation_affected", {"deviation_id": d["id"], **row, "revision": rev,
                                               "note": (note or "").strip()[:500]})
        db.update("deviations", d["id"], {"updated_at": db.now()})
        db.audit(user["id"], "deviation_affected_add", "deviation", d["id"], label)
    return aid


def remove_affected(user: dict, dev_id: int, affected_id: int) -> None:
    d = dev_or_404(dev_id)
    _editable(d, user)
    row = db.one("SELECT * FROM deviation_affected WHERE id = ? AND deviation_id = ?", (int(affected_id), d["id"]))
    if not row:
        raise not_found("That record is not on this deviation.")
    with db.tx():
        db.execute("DELETE FROM deviation_affected WHERE id = ?", (row["id"],))
        db.audit(user["id"], "deviation_affected_remove", "deviation", d["id"],
                 f"item {row['item_id']}" if row["item_id"] else f"document {row['document_id']} rev {row['revision']}")


def add_approver(user: dict, dev_id: int, user_id: int) -> None:
    d = dev_or_404(dev_id)
    _editable(d, user)
    u = db.one("SELECT id, role, active, full_name FROM users WHERE id = ?", (int(user_id),))
    if not u or not u["active"]:
        raise bad("Unknown user")
    if u["id"] == d["requested_by"]:
        raise bad("The requester can't approve their own deviation (four-eyes rule). Choose someone else.",
                  "four_eyes")
    if ROLE_LEVEL[u["role"]] < ROLE_LEVEL["approver"]:
        raise bad(f"{u['full_name']} can't approve: approvers need the approver or admin role.")
    with db.tx():
        db.execute("INSERT INTO deviation_approvers (deviation_id, user_id) VALUES (?, ?) "
                   "ON CONFLICT(deviation_id, user_id) DO NOTHING", (d["id"], u["id"]))
        db.audit(user["id"], "deviation_approver_add", "deviation", d["id"], u["full_name"])


def remove_approver(user: dict, dev_id: int, user_id: int) -> None:
    d = dev_or_404(dev_id)
    _editable(d, user)
    with db.tx():
        db.execute("DELETE FROM deviation_approvers WHERE deviation_id = ? AND user_id = ?", (d["id"], int(user_id)))
        db.audit(user["id"], "deviation_approver_remove", "deviation", d["id"], f"user {user_id}")


# --------------------------------------------------------------- lifecycle --

def readiness(d: dict) -> list[str]:
    """What still stops this deviation from being submitted, in plain words (empty when ready)."""
    problems = []
    if not db.scalar("SELECT COUNT(*) FROM deviation_affected WHERE deviation_id = ?", (d["id"],)):
        problems.append("add at least one affected part or file")
    if not (d["departure"] or "").strip():
        problems.append("describe the departure from the released design")
    if d["qty_limit"] is None and not (d["serials"] or "").strip() and not d["valid_until"]:
        problems.append("limit it by a quantity, serial or batch numbers, or an end date (deviations are temporary)")
    if d["valid_until"] and d["valid_until"] < today().isoformat():
        problems.append("the end date has already passed")
    if not db.scalar("SELECT COUNT(*) FROM deviation_approvers a JOIN users u ON u.id = a.user_id WHERE "
                     "a.deviation_id = ? AND a.user_id != ? AND u.active = 1", (d["id"], d["requested_by"])):
        problems.append("add at least one approver other than the requester")
    return problems


def submit(user: dict, dev_id: int) -> None:
    d = dev_or_404(dev_id)
    _editable(d, user)
    problems = readiness(d)
    if problems:
        raise conflict("Can't submit yet: " + "; ".join(problems) + ".", "not_ready")
    now = db.now()
    with db.tx():
        db.execute("UPDATE deviation_approvers SET decision = 'pending', comment = '', decided_at = NULL "
                   "WHERE deviation_id = ?", (d["id"],))
        db.update("deviations", d["id"], {"state": "submitted", "submitted_at": now, "decided_at": None,
                                          "updated_at": now})
        approvers = [r["user_id"] for r in db.rows("SELECT user_id FROM deviation_approvers WHERE deviation_id = ?",
                                                    (d["id"],))]
        notifications.notify_many(approvers, "deviation_approval",
                                  f"{d['number']} is waiting for your approval: {d['title']}",
                                  f"#/deviation/{d['id']}", user["id"], skip=[user["id"]])
        db.audit(user["id"], "deviation_submit", "deviation", d["id"], d["number"])
    from . import notify_outbox   # e-mail and chat, for the approvers who chose them (the bell is written above)
    notify_outbox.enqueue([a for a in approvers if a != user["id"]], f"{d['number']} needs your approval: {d['title']}",
                          f"{user['full_name']} asks you to approve the deviation {d['number']}: {d['title']}.\n"
                          f"Departure from the released design: {d['departure']}", f"#/deviation/{d['id']}",
                          "approvals", key=f"deviation.submitted:{d['id']}:{now}", source=f"deviation:{d['id']}",
                          in_app=False)


def can_decide(d: dict, user: dict) -> dict:
    if d["state"] != "submitted":
        return _ok(False, "This deviation is not waiting for approval.")
    if d["requested_by"] == user["id"]:
        return _ok(False, "You raised this deviation, so someone else must approve it (four-eyes rule).")
    if _level(user) < ROLE_LEVEL["approver"]:
        return _ok(False, "Deciding needs the approver role.")
    row = db.one("SELECT decision FROM deviation_approvers WHERE deviation_id = ? AND user_id = ?", (d["id"], user["id"]))
    if not row and user["role"] != "admin":
        return _ok(False, "You are not an approver on this deviation.")
    if row and row["decision"] != "pending":
        return _ok(False, f"You have already {row['decision']} it.")
    return _ok(True)


def decide(user: dict, dev_id: int, decision: str, comment: str = "") -> dict:
    """Record one approver's decision. Returns the deviation; emits ``deviation.approved`` on the final approval."""
    d = dev_or_404(dev_id)
    if decision not in ("approved", "rejected"):
        raise bad("Decision must be approved or rejected")
    allowed = can_decide(d, user)
    if not allowed["ok"]:
        if d["state"] != "submitted":
            raise conflict(allowed["why"], "not_submitted")
        raise forbidden(allowed["why"])
    comment = (comment or "").strip()[:2000]
    if decision == "rejected" and not comment:
        raise bad("Please say why you are rejecting it, so the requester knows what to change.")
    if decision == "approved" and d["valid_until"] and d["valid_until"] < today().isoformat():
        raise conflict("The end date of this deviation has already passed. Reject it so the requester can update the "
                       "dates.", "expired")
    now = db.now()
    final = False
    with db.tx():
        row = db.one("SELECT id FROM deviation_approvers WHERE deviation_id = ? AND user_id = ?", (d["id"], user["id"]))
        if row:
            db.execute("UPDATE deviation_approvers SET decision = ?, comment = ?, decided_at = ? WHERE id = ?",
                       (decision, comment, now, row["id"]))
        else:   # an admin deciding without being on the list is recorded as an approver too
            db.execute("INSERT INTO deviation_approvers (deviation_id, user_id, decision, comment, decided_at) "
                       "VALUES (?, ?, ?, ?, ?)", (d["id"], user["id"], decision, comment, now))
        if decision == "rejected":
            db.update("deviations", d["id"], {"state": "rejected", "decided_at": now, "updated_at": now})
            notifications.notify(d["requested_by"], "deviation_rejected",
                                 f"{d['number']} was rejected by {user['full_name']}: {comment}",
                                 f"#/deviation/{d['id']}", user["id"])
        else:
            pending = db.scalar("SELECT COUNT(*) FROM deviation_approvers WHERE deviation_id = ? AND decision != 'approved'",
                                (d["id"],))
            if not pending:
                final = True
                starts_now = not d["valid_from"] or d["valid_from"] <= today().isoformat()
                db.update("deviations", d["id"], {
                    "state": "active" if starts_now else "approved", "decided_at": now, "updated_at": now,
                    "activated_at": now if starts_now else None,
                })
                notifications.notify(d["requested_by"], "deviation_approved",
                                     f"{d['number']} was approved" + (" and is active now" if starts_now else
                                                                      f"; it starts on {_fmt_day(d['valid_from'])}"),
                                     f"#/deviation/{d['id']}", user["id"])
        db.audit(user["id"], f"deviation_{decision}", "deviation", d["id"], d["number"] + (f": {comment}" if comment else ""))
    d = dev_or_404(dev_id)
    if final:
        _emit_approved(d, user)
    return d


def _emit_approved(d: dict, user: dict) -> None:
    from .addins import registry
    try:
        registry.emit("deviation.approved", event_payload(d, user["id"]))
    except Exception:   # an add-in problem must never undo an approval
        log.exception("Could not hand deviation.approved for %s to add-ins", d["number"])


def event_payload(d: dict, user_id: int | None) -> dict:
    return {
        "deviation_id": d["id"], "number": d["number"], "title": d["title"], "departure": d["departure"],
        "reason": d["reason"], "state": d["state"], "qty_limit": d["qty_limit"], "serials": d["serials"],
        "valid_from": d["valid_from"], "valid_until": d["valid_until"], "user_id": user_id,
        "affected": [{"part_number": a["part_number"], "item_id": a["item_id"], "document_id": a["document_id"],
                      "filename": a["filename"], "revision": a["revision"]} for a in affected(d["id"])],
    }


def record_usage(user: dict, dev_id: int, qty, reference: str = "", note: str = "") -> dict:
    d = dev_or_404(dev_id)
    if d["state"] != "active":
        raise conflict("Usage can only be recorded while the deviation is active.", "not_active")
    if _level(user) < ROLE_LEVEL["engineer"]:
        raise forbidden("Recording usage needs the engineer role.")
    q = _qty(qty)
    if not q:
        raise bad("Enter the quantity made or fitted under this deviation.")
    used = d["qty_used"] or 0
    if d["qty_limit"] is not None and used + q > d["qty_limit"] + 1e-9:
        left = max(d["qty_limit"] - used, 0)
        raise conflict(f"Only {_fmt_qty(left)} of {_fmt_qty(d['qty_limit'])} are left on {d['number']}. Record at most "
                       f"{_fmt_qty(left)}, or raise a new deviation for the rest.", "over_limit")
    now = db.now()
    with db.tx():
        db.insert("deviation_usage", {"deviation_id": d["id"], "qty": q, "reference": (reference or "").strip()[:200],
                                      "note": (note or "").strip()[:1000], "user_id": user["id"], "at": now})
        new_used = used + q
        db.update("deviations", d["id"], {"qty_used": new_used, "updated_at": now})
        db.audit(user["id"], "deviation_usage", "deviation", d["id"], f"{_fmt_qty(q)} {reference or ''}".strip())
        if d["qty_limit"] is not None and new_used >= d["qty_limit"] - 1e-9:
            _close(d, None, "used_up", "The allowed quantity is used up.")
    return dev_or_404(dev_id)


def _close(d: dict, user: dict | None, reason: str, note: str = "") -> None:
    now = db.now()
    db.update("deviations", d["id"], {"state": "closed", "closed_at": now, "closed_by": user["id"] if user else None,
                                      "close_reason": reason, "close_note": note[:2000], "updated_at": now})
    db.audit(user["id"] if user else None, f"deviation_close_{reason}", "deviation", d["id"], f"{d['number']}: {note}")
    if not user or user["id"] != d["requested_by"]:
        notifications.notify(d["requested_by"], "deviation_closed", f"{d['number']} has ended: {CLOSE_LABEL[reason].lower()}",
                             f"#/deviation/{d['id']}", user["id"] if user else None)


def can_close(d: dict, user: dict) -> dict:
    if d["state"] not in ("active", "approved"):
        return _ok(False, "Only approved or active deviations can be closed.")
    if d["requested_by"] != user["id"] and _level(user) < ROLE_LEVEL["approver"]:
        return _ok(False, "Only the requester or an approver can close it.")
    return _ok(True)


def close(user: dict, dev_id: int, qty_used=None, note: str = "") -> dict:
    """Close by hand. ``qty_used`` is the quantity actually used in total; the difference is booked as a usage line."""
    d = dev_or_404(dev_id)
    allowed = can_close(d, user)
    if not allowed["ok"]:
        if d["state"] not in ("active", "approved"):
            raise conflict(allowed["why"], "not_active")
        raise forbidden(allowed["why"])
    actual = _qty(qty_used, "The quantity used")
    with db.tx():
        if actual is not None and abs(actual - (d["qty_used"] or 0)) > 1e-9:
            db.insert("deviation_usage", {"deviation_id": d["id"], "qty": actual - (d["qty_used"] or 0),
                                          "reference": "Actual quantity recorded at closing", "note": "",
                                          "user_id": user["id"], "at": db.now()})
            db.update("deviations", d["id"], {"qty_used": actual})
        _close(d, user, "closed", (note or "").strip())
    return dev_or_404(dev_id)


def can_cancel(d: dict, user: dict) -> dict:
    if d["state"] not in ("draft", "submitted", "rejected", "approved"):
        return _ok(False, "Active deviations are closed, not cancelled; closed or cancelled ones stay as they are."
                   if d["state"] == "active" else "This deviation is already closed or cancelled.")
    if d["requested_by"] != user["id"] and _level(user) < ROLE_LEVEL["approver"]:
        return _ok(False, "Only the requester or an approver can cancel it.")
    return _ok(True)


def cancel(user: dict, dev_id: int, note: str = "") -> dict:
    d = dev_or_404(dev_id)
    allowed = can_cancel(d, user)
    if not allowed["ok"]:
        if d["state"] not in ("draft", "submitted", "rejected", "approved"):
            raise conflict(allowed["why"], "closed")
        raise forbidden(allowed["why"])
    now = db.now()
    with db.tx():
        db.update("deviations", d["id"], {"state": "cancelled", "closed_at": now, "closed_by": user["id"],
                                          "close_note": (note or "").strip()[:2000], "updated_at": now})
        db.audit(user["id"], "deviation_cancel", "deviation", d["id"], d["number"])
    return dev_or_404(dev_id)


def sweep(day: date | None = None) -> dict:
    """Start deviations whose start date has come and close those whose date range has ended or quantity is used up.

    Cheap when there is nothing to do (one read), so every page that shows deviations calls it first."""
    day_iso = (day or today()).isoformat()
    due = db.rows(
        "SELECT * FROM deviations WHERE "
        "(state IN ('approved','active') AND valid_until IS NOT NULL AND valid_until < ?) OR "
        "(state = 'approved' AND (valid_from IS NULL OR valid_from <= ?)) OR "
        "(state = 'active' AND qty_limit IS NOT NULL AND qty_used >= qty_limit)", (day_iso, day_iso))
    result = {"activated": [], "expired": [], "used_up": []}
    for d in due:
        with db.tx():
            d = dev_or_404(d["id"])   # re-read inside the transaction: another request may have done it
            if d["state"] in ("approved", "active") and d["valid_until"] and d["valid_until"] < day_iso:
                _close(d, None, "expired", f"The date range ended on {d['valid_until']}.")
                result["expired"].append(d["number"])
            elif d["state"] == "active" and d["qty_limit"] is not None and (d["qty_used"] or 0) >= d["qty_limit"]:
                _close(d, None, "used_up", "The allowed quantity is used up.")
                result["used_up"].append(d["number"])
            elif d["state"] == "approved" and (not d["valid_from"] or d["valid_from"] <= day_iso):
                now = db.now()
                db.update("deviations", d["id"], {"state": "active", "activated_at": now, "updated_at": now})
                db.audit(None, "deviation_activate", "deviation", d["id"], f"{d['number']} start date reached")
                result["activated"].append(d["number"])
    return result


# ------------------------------------------------------------ change orders --

def can_make_permanent(d: dict, user: dict) -> dict:
    if _level(user) < ROLE_LEVEL["engineer"]:
        return _ok(False, "Raising a change order needs the engineer role.")
    if d["state"] not in ("approved", "active", "closed"):
        return _ok(False, "Only approved, active or closed deviations can be made permanent.")
    done = db.one("SELECT c.number FROM deviation_changes x JOIN change_orders c ON c.id = x.change_order_id WHERE "
                  "x.deviation_id = ? AND x.relation = 'made_permanent' AND c.state != 'cancelled'", (d["id"],))
    if done:
        return _ok(False, f"Already being made permanent by {done['number']}.")
    return _ok(True)


def make_permanent(user: dict, dev_id: int) -> dict:
    """Create a draft change order pre-filled from the deviation and link the two. Returns {change_id, skipped}."""
    d = dev_or_404(dev_id)
    allowed = can_make_permanent(d, user)
    if not allowed["ok"]:
        raise conflict(allowed["why"], "not_allowed") if _level(user) >= ROLE_LEVEL["engineer"] else forbidden(allowed["why"])
    rows = affected(d["id"])
    doc_ids: list[int] = []
    for a in rows:
        if a["document_id"]:
            doc_ids.append(a["document_id"])
        elif a["item_id"]:
            doc_ids += [r["id"] for r in db.rows("SELECT id FROM documents WHERE item_id = ? AND deleted = 0 AND "
                                                  "state != 'obsolete' ORDER BY id", (a["item_id"],))]
    approvers = [r["user_id"] for r in db.rows("SELECT user_id FROM deviation_approvers WHERE deviation_id = ? "
                                                "ORDER BY id", (d["id"],)) if r["user_id"] != user["id"]]
    scope = scope_text(d)
    skipped = []
    with db.tx():
        co_id = changes.create(
            user, f"Make {d['number']} permanent: {d['title']}"[:200],
            f"Raised from deviation {d['number']}." + (f" {d['reason']}" if d["reason"] else ""),
            f"Departure allowed by {d['number']}: {d['departure']}" + (f"\nScope of the deviation: {scope}" if scope else ""),
            "normal")
        for doc_id in dict.fromkeys(doc_ids):
            try:
                changes.add_item(user, co_id, doc_id)
            except Exception as e:   # e.g. the file is already in another open change order
                skipped.append(getattr(e, "message", str(e)))
        for uid in approvers:
            try:
                changes.add_approver(user, co_id, uid)
            except Exception as e:
                skipped.append(getattr(e, "message", str(e)))
        db.insert("deviation_changes", {"deviation_id": d["id"], "change_order_id": co_id, "relation": "made_permanent",
                                        "created_by": user["id"], "created_at": db.now()})
        number = db.scalar("SELECT number FROM change_orders WHERE id = ?", (co_id,))
        db.audit(user["id"], "deviation_make_permanent", "deviation", d["id"], f"{d['number']} -> {number}")
    return {"change_id": co_id, "change_number": number, "skipped": skipped}


def link_change(user: dict, dev_id: int, co_id: int) -> None:
    d = dev_or_404(dev_id)
    if _level(user) < ROLE_LEVEL["engineer"]:
        raise forbidden("Linking change orders needs the engineer role.")
    co = changes.co_or_404(int(co_id))
    with db.tx():
        db.execute("INSERT INTO deviation_changes (deviation_id, change_order_id, relation, created_by, created_at) "
                   "VALUES (?, ?, 'related', ?, ?) ON CONFLICT(deviation_id, change_order_id) DO NOTHING",
                   (d["id"], co["id"], user["id"], db.now()))
        db.audit(user["id"], "deviation_link_change", "deviation", d["id"], co["number"])


def unlink_change(user: dict, dev_id: int, co_id: int) -> None:
    d = dev_or_404(dev_id)
    if _level(user) < ROLE_LEVEL["engineer"]:
        raise forbidden("Linking change orders needs the engineer role.")
    with db.tx():
        db.execute("DELETE FROM deviation_changes WHERE deviation_id = ? AND change_order_id = ?", (d["id"], int(co_id)))
        db.audit(user["id"], "deviation_unlink_change", "deviation", d["id"], f"change order {co_id}")


# -------------------------------------------------------------- attachments --

def can_attach(d: dict, user: dict) -> dict:
    if _level(user) < ROLE_LEVEL["engineer"]:
        return _ok(False, "Adding attachments needs the engineer role.")
    if d["state"] == "cancelled":
        return _ok(False, "This deviation was cancelled.")
    return _ok(True)


def add_attachment(user: dict, dev_id: int, filename: str, stream: BinaryIO) -> int:
    d = dev_or_404(dev_id)
    allowed = can_attach(d, user)
    if not allowed["ok"]:
        raise forbidden(allowed["why"])
    name = lifecycle.clean_filename(filename or "")
    sha, size = store.put_stream(stream)   # safely on disk before the database knows about it
    with db.tx():
        aid = db.insert("deviation_attachments", {"deviation_id": d["id"], "filename": name, "blob_sha": sha, "size": size,
                                                  "uploaded_by": user["id"], "uploaded_at": db.now()})
        db.audit(user["id"], "deviation_attach", "deviation", d["id"], f"{name} ({sha[:12]})")
    return aid


def attachment_or_404(dev_id: int, attachment_id: int) -> dict:
    a = db.one("SELECT * FROM deviation_attachments WHERE id = ? AND deviation_id = ? AND deleted = 0",
               (int(attachment_id), int(dev_id)))
    if not a:
        raise not_found("Attachment not found")
    return a


def remove_attachment(user: dict, dev_id: int, attachment_id: int) -> None:
    """Soft delete: the file stays in the store and the row stays in the database."""
    d = dev_or_404(dev_id)
    a = attachment_or_404(d["id"], attachment_id)
    if a["uploaded_by"] != user["id"] and _level(user) < ROLE_LEVEL["approver"]:
        raise forbidden("Only the person who added it, or an approver, can remove an attachment.")
    if d["state"] in ("closed", "cancelled"):
        raise conflict("Attachments of closed or cancelled deviations are kept as they are.", "closed")
    with db.tx():
        db.update("deviation_attachments", a["id"], {"deleted": 1, "deleted_by": user["id"], "deleted_at": db.now()})
        db.audit(user["id"], "deviation_attachment_remove", "deviation", d["id"], a["filename"])


# ------------------------------------------------------------------ reading --

def affected(dev_id: int) -> list[dict]:
    rows = db.rows(
        "SELECT a.*, i.part_number AS item_part_number, i.description AS item_description, d.filename, "
        "d.part_number AS doc_part_number, d.description AS doc_description, d.item_id AS doc_item_id, d.ext "
        "FROM deviation_affected a LEFT JOIN items i ON i.id = a.item_id "
        "LEFT JOIN documents d ON d.id = a.document_id WHERE a.deviation_id = ? ORDER BY a.id", (dev_id,))
    for r in rows:
        r["kind"] = "item" if r["item_id"] else "document"
        r["part_number"] = r["item_part_number"] or r["doc_part_number"] or ""
        r["description"] = r["item_description"] or r["doc_description"] or ""
        r["label"] = (f"{r['part_number']} rev {r['revision']}" if r["item_id"] else
                      f"{r['filename']} rev {r['revision']}")
    return rows


def _summary(d: dict, names: dict) -> dict:
    aff = affected(d["id"])
    return {**d, "state_label": STATE_LABEL.get(d["state"], d["state"]),
            "requested_by_name": names.get(d["requested_by"], "?"), "affected": aff,
            "affected_text": ", ".join(a["label"] for a in aff), "scope": scope_text(d),
            "banner": banner_text(d), "close_reason_label": CLOSE_LABEL.get(d["close_reason"], "")}


def detail(dev_id: int, user: dict) -> dict:
    d = dev_or_404(dev_id)
    names = lifecycle.user_names()
    out = _summary(d, names)
    out["approvers"] = [{**a, "name": names.get(a["user_id"], "?")} for a in
                        db.rows("SELECT * FROM deviation_approvers WHERE deviation_id = ? ORDER BY id", (d["id"],))]
    out["usage"] = [{**u, "by": names.get(u["user_id"], "?")} for u in
                    db.rows("SELECT * FROM deviation_usage WHERE deviation_id = ? ORDER BY id", (d["id"],))]
    out["attachments"] = []
    for a in db.rows("SELECT id, filename, size, uploaded_by, uploaded_at, blob_sha FROM deviation_attachments "
                     "WHERE deviation_id = ? AND deleted = 0 ORDER BY id", (d["id"],)):
        sha = a.pop("blob_sha")
        out["attachments"].append({
            **a, "sha": sha[:12], "by": names.get(a["uploaded_by"], "?"),
            "can_remove": _ok(d["state"] not in ("closed", "cancelled") and (
                a["uploaded_by"] == user["id"] or _level(user) >= ROLE_LEVEL["approver"]),
                "Attachments of closed or cancelled deviations are kept as they are."
                if d["state"] in ("closed", "cancelled") else
                "Only the person who added it, or an approver, can remove it.")})
    out["changes"] = db.rows("SELECT x.relation, x.created_at, c.id, c.number, c.title, c.state FROM deviation_changes x "
                             "JOIN change_orders c ON c.id = x.change_order_id WHERE x.deviation_id = ? ORDER BY x.id",
                             (d["id"],))
    out["closed_by_name"] = names.get(d["closed_by"]) if d["closed_by"] else None
    out["over_limit"] = d["qty_limit"] is not None and (d["qty_used"] or 0) > d["qty_limit"] + 1e-9
    edit = _can_edit(d, user)
    problems = readiness(d) if d["state"] in ("draft", "rejected") else []
    out["readiness"] = problems
    out["actions"] = {
        "edit": edit,
        "submit": edit if not edit["ok"] else _ok(not problems, "Before submitting: " + "; ".join(problems) + "."),
        "decide": can_decide(d, user),
        "record_usage": _ok(d["state"] == "active" and _level(user) >= ROLE_LEVEL["engineer"],
                            "Usage is recorded while the deviation is active." if d["state"] != "active" else
                            "Recording usage needs the engineer role."),
        "close": can_close(d, user),
        "cancel": can_cancel(d, user),
        "make_permanent": can_make_permanent(d, user),
        "link_change": _ok(_level(user) >= ROLE_LEVEL["engineer"], "Linking change orders needs the engineer role."),
        "attach": can_attach(d, user),
    }
    return out


def _filters(f: dict) -> tuple[str, list]:
    where, params = [], []
    state = (f.get("state") or f.get("status") or "").strip()
    if state:
        if state == "open":
            where.append("d.state IN ('draft','submitted','approved','rejected','active')")
        elif state not in STATES:
            raise bad(f"Unknown status {state!r}; use one of: open, {', '.join(STATES)}.")
        else:
            where.append("d.state = ?")
            params.append(state)
    if f.get("item_id"):
        iid = int(f["item_id"])
        where.append("d.id IN (SELECT a.deviation_id FROM deviation_affected a LEFT JOIN documents x ON "
                     "x.id = a.document_id WHERE a.item_id = ? OR x.item_id = ?)")
        params += [iid, iid]
    item = (f.get("item") or "").strip()
    if item:
        like = f"%{item}%"
        where.append("d.id IN (SELECT a.deviation_id FROM deviation_affected a LEFT JOIN items i ON i.id = a.item_id "
                     "LEFT JOIN documents x ON x.id = a.document_id WHERE i.part_number LIKE ? OR i.description LIKE ? "
                     "OR x.part_number LIKE ? OR x.filename LIKE ?)")
        params += [like] * 4
    q = (f.get("q") or "").strip()
    if q:
        where.append("(d.number LIKE ? OR d.title LIKE ? OR d.departure LIKE ?)")
        params += [f"%{q}%"] * 3
    for key, op, label in (("date_from", ">=", "From"), ("date_to", "<=", "To")):
        v = _date(f.get(key), f"The '{label}' date")
        if v:
            where.append(f"SUBSTR(d.created_at, 1, 10) {op} ?")
            params.append(v)
    return (" WHERE " + " AND ".join(where)) if where else "", params


def listing(f: dict, limit: int = 500) -> list[dict]:
    where, params = _filters(f)
    names = lifecycle.user_names()
    rows = db.rows(f"SELECT d.* FROM deviations d{where} ORDER BY d.id DESC LIMIT ?", [*params, int(limit)])
    return [_summary(r, names) for r in rows]


def _item_ids(dev_id: int) -> set[int]:
    return {r["iid"] for r in db.rows(
        "SELECT a.item_id AS iid FROM deviation_affected a WHERE a.deviation_id = ? AND a.item_id IS NOT NULL UNION "
        "SELECT x.item_id AS iid FROM deviation_affected a JOIN documents x ON x.id = a.document_id "
        "WHERE a.deviation_id = ? AND x.item_id IS NOT NULL", (dev_id, dev_id))}


def active_for(entity_type: str, entity_id: int) -> dict:
    """Deviations to show as a banner on a record page.

    item / document / bom: the active deviations affecting it (for a BOM: its assembly and every part in it, with
    ``by_item`` telling which BOM lines are affected). change_order: the deviations linked to it, in any state."""
    eid = int(entity_id)
    names = lifecycle.user_names()
    if entity_type == "change_order":
        rows = db.rows("SELECT d.*, x.relation FROM deviation_changes x JOIN deviations d ON d.id = x.deviation_id "
                       "WHERE x.change_order_id = ? ORDER BY d.id", (eid,))
        return {"deviations": [{**_summary(r, names), "relation": r["relation"]} for r in rows], "by_item": {}}
    active = db.rows("SELECT * FROM deviations WHERE state = 'active' ORDER BY id")
    if entity_type == "document":
        doc = db.one("SELECT id, item_id FROM documents WHERE id = ?", (eid,))
        if not doc:
            raise not_found("File not found")
        hits = [d for d in active if db.scalar("SELECT 1 FROM deviation_affected WHERE deviation_id = ? AND "
                                               "document_id = ?", (d["id"], eid))
                or (doc["item_id"] and doc["item_id"] in _item_ids(d["id"]))]
        return {"deviations": [_summary(d, names) for d in hits], "by_item": {}}
    if entity_type == "item":
        if not db.scalar("SELECT 1 FROM items WHERE id = ?", (eid,)):
            raise not_found("Item not found")
        hits = [d for d in active if eid in _item_ids(d["id"])]
        return {"deviations": [_summary(d, names) for d in hits], "by_item": {}}
    if entity_type == "bom":
        from . import boms
        b = boms.bom_or_404(eid)
        in_bom = {b["item_id"]} | {r["item_id"] for r in boms.explode(eid)}
        hits, by_item = [], {}
        for d in active:
            ids = _item_ids(d["id"]) & in_bom
            if not ids:
                continue
            s = _summary(d, names)
            hits.append(s)
            for iid in ids:
                by_item.setdefault(str(iid), []).append({"id": d["id"], "number": d["number"], "banner": s["banner"]})
        return {"deviations": hits, "by_item": by_item}
    raise bad("Deviation banners are shown on items, files, BOMs and change orders.")


def home(user: dict) -> dict:
    names = lifecycle.user_names()
    waiting = db.rows("SELECT d.* FROM deviations d JOIN deviation_approvers a ON a.deviation_id = d.id WHERE "
                      "a.user_id = ? AND a.decision = 'pending' AND d.state = 'submitted' ORDER BY d.submitted_at",
                      (user["id"],))
    active = db.rows("SELECT * FROM deviations WHERE state = 'active' ORDER BY valid_until IS NULL, valid_until, id")
    mine = db.rows("SELECT * FROM deviations WHERE requested_by = ? AND state IN ('draft','submitted','approved',"
                   "'rejected','active') ORDER BY id DESC LIMIT 8", (user["id"],))
    return {"waiting_for_me": [_summary(d, names) for d in waiting], "active": [_summary(d, names) for d in active],
            "mine": [_summary(d, names) for d in mine],
            "is_approver": _level(user) >= ROLE_LEVEL["approver"]}


def summary_counts(user: dict) -> dict:
    return {
        "waiting_for_me": int(db.scalar("SELECT COUNT(*) FROM deviations d JOIN deviation_approvers a ON "
                                        "a.deviation_id = d.id WHERE a.user_id = ? AND a.decision = 'pending' "
                                        "AND d.state = 'submitted'", (user["id"],)) or 0),
        "active": int(db.scalar("SELECT COUNT(*) FROM deviations WHERE state = 'active'") or 0),
    }


# ------------------------------------------------------------------- report --

def report_xlsx(f: dict) -> tuple[bytes, str]:
    """The deviations register as an Excel workbook (same filters as the list)."""
    import openpyxl
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    rows = listing(f, limit=100000)
    names = lifecycle.user_names()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Deviations"
    used = [f"{k.replace('_', ' ')}: {v}" for k, v in f.items() if v not in (None, "")]
    ws.append([f"Deviations register, exported {today().isoformat()}", "Filters: " + ("; ".join(used) or "none")])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    head = ["Number", "Title", "Status", "Affected (revision)", "Departure", "Reason", "Serial / batch numbers",
            "Quantity allowed", "Quantity used", "Valid from", "Valid until", "Requested by", "Raised on", "Approvers",
            "Decided on", "Active since", "Closed on", "How it ended", "Change orders"]
    ws.append(head)
    for c in ws[3]:
        c.font = Font(bold=True)
    for r in rows:
        appr = db.rows("SELECT user_id, decision FROM deviation_approvers WHERE deviation_id = ? ORDER BY id", (r["id"],))
        cos = db.rows("SELECT c.number FROM deviation_changes x JOIN change_orders c ON c.id = x.change_order_id "
                      "WHERE x.deviation_id = ? ORDER BY x.id", (r["id"],))
        ws.append([
            r["number"], r["title"], r["state_label"], r["affected_text"], r["departure"], r["reason"], r["serials"],
            r["qty_limit"], r["qty_used"], r["valid_from"] or "", r["valid_until"] or "", r["requested_by_name"],
            (r["created_at"] or "")[:10], ", ".join(f"{names.get(a['user_id'], '?')} ({a['decision']})" for a in appr),
            (r["decided_at"] or "")[:10], (r["activated_at"] or "")[:10], (r["closed_at"] or "")[:10],
            r["close_reason_label"] or ("Cancelled" if r["state"] == "cancelled" else ""),
            ", ".join(c["number"] for c in cos),
        ])
    widths = (14, 36, 20, 34, 44, 36, 20, 10, 10, 12, 12, 18, 12, 30, 12, 12, 12, 18, 16)
    for i, w in enumerate(widths):
        ws.column_dimensions[get_column_letter(i + 1)].width = w
    for row in ws.iter_rows(min_row=4):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "B4"
    ws.auto_filter.ref = f"A3:{get_column_letter(len(head))}{max(3, ws.max_row)}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), f"Deviations_register_{today().isoformat()}.xlsx"
