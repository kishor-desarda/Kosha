# SPDX-License-Identifier: AGPL-3.0-or-later
"""Bulk import of legacy data: folders of drawings, models and documents.

    python -m kosha.importer "\\\\fileserver\\Drawings" --to "Legacy Drawings/File server" --state released
    python -m kosha.importer D:\\Old --to "Legacy Drawings/Old D drive" --metadata parts.xlsx --dry-run

The same engine runs from Admin › Import legacy data in the browser.

What it does
  * walks the source folder and recreates its sub-folders under the target Kosha folder
  * skips junk (AutoCAD .bak/.dwl, Office ~$ lock files, Thumbs.db, …)
  * finds identical copies by content (SHA-256), both against Kosha and inside the
    import itself, and imports each distinct file only once
  * reads part number, description, material and revision from title blocks, file
    names and an optional part-master spreadsheet
  * keeps the legacy revision (Rev 3, R-2, B …) and can mark files Released at once
  * is resumable: files already in the target folder are skipped, so a broken run can be
    started again
  * unpacks ZIP archives into a staging folder and imports what is inside, under a folder named after the archive
  * writes an Excel/CSV report of every file and what happened to it
  * keeps a record of every file per batch (kosha/legacy.py): quality score and work queue, old paths and numbers
    as cross-references, the count check, and Undo batch
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import logging
import os
import re
import shutil
import sys
import threading
import uuid
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Callable

from . import config, db, legacy, lifecycle, store
from .cad import dwg as dwgmod
from .cad import process as cadproc

log = logging.getLogger("kosha.import")

JUNK = re.compile(r"(^~\$|^\.|\.(bak|dwl2?|tmp|sv\$|ac\$|lck|log|swp|db)$|^thumbs\.db$|^desktop\.ini$|^~.*\.tmp$)", re.I)
TYPE_GROUPS = {
    "solidworks": {"sldprt", "sldasm", "slddrw"},
    "cad2d": {"dwg", "dxf"},
    "pdf": {"pdf"},
    "neutral3d": {"step", "stp", "igs", "iges", "x_t", "stl"},
    "office": {"xlsx", "xls", "docx", "doc", "csv", "txt", "pptx"},
    "images": {"png", "jpg", "jpeg", "tif", "tiff", "bmp"},
}
DEFAULT_GROUPS = ["solidworks", "cad2d", "pdf", "neutral3d"]
# archives are unpacked within these limits (a damaged or hostile archive is reported, never half-imported)
ZIP_MAX_FILES = 20000
ZIP_MAX_BYTES = 4 * 1024 ** 3
ZIP_MAX_RATIO = 1000
META_ALIASES = {
    "filename": ("file", "filename", "file name", "drawing file", "dwg file"),
    "part_number": ("part number", "part no", "part no.", "partno", "item code", "item_code", "drawing no",
                    "drawing number", "dwg no", "code"),
    "description": ("description", "item name", "item_name", "name", "title"),
    "material": ("material", "moc", "matl", "grade"),
    "revision": ("revision", "rev", "rev no", "rev."),
    "uom": ("uom", "unit", "stock uom", "stock_uom"),
    "make_buy": ("make/buy", "make buy", "make_buy", "procurement", "type"),
}

META_EXTRA = ("item_type", "weight_kg")      # also accepted from importer add-ins


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


# ------------------------------------------------------------------ metadata

def read_metadata(path_or_bytes, filename: str = "") -> tuple[dict, list[dict]]:
    """Part master spreadsheet -> ({file name lower: props}, [item rows])."""
    name = (filename or str(path_or_bytes)).lower()
    data = path_or_bytes if isinstance(path_or_bytes, bytes) else Path(path_or_bytes).read_bytes()
    if name.endswith((".xlsx", ".xlsm")):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        table = [["" if c is None else str(c).strip() for c in r] for r in wb.worksheets[0].iter_rows(values_only=True)]
    else:
        text = data.decode("utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        table = [[c.strip() for c in r] for r in csv.reader(io.StringIO(text), dialect)]
    header, start = None, 0
    for k, row in enumerate(table[:20]):
        low = [c.lower() for c in row]
        found = {}
        for key, aliases in META_ALIASES.items():
            for i, c in enumerate(low):
                if c in aliases and i not in found.values():
                    found[key] = i
                    break
        if "part_number" in found or "filename" in found:
            header, start = found, k + 1
            break
    if not header:
        raise ValueError("The spreadsheet needs a 'Part number' or 'File name' column in its first rows.")
    return metadata_from_rows({k: (row[i].strip() if i < len(row) and row[i] else "") for k, i in header.items()}
                              for row in table[start:])


def metadata_from_rows(rows) -> tuple[dict, list[dict]]:
    """Part master rows ({filename, part_number, description, ...}) -> ({file name lower: props}, [item rows])."""
    by_file: dict = {}
    items: list = []
    for rec in rows:
        if not isinstance(rec, dict):
            continue
        rec = {k: str(v).strip()[:500] for k, v in rec.items()
               if v is not None and (k in META_ALIASES or k in META_EXTRA)}
        if not any(rec.values()):
            continue
        if rec.get("filename"):
            by_file[Path(rec["filename"]).name.lower()] = rec
            stem = Path(rec["filename"]).stem.lower()
            by_file.setdefault(stem, rec)
        if rec.get("part_number"):
            items.append(rec)
    return by_file, items


def upsert_items(rows: list[dict], user_id) -> int:
    n = 0
    now = db.now()
    with db.tx():
        for r in rows:
            pn = r["part_number"].strip().upper()
            if not pn:
                continue
            it = db.one("SELECT * FROM items WHERE part_number = ? COLLATE NOCASE", (pn,))
            mb = (r.get("make_buy") or "").lower()
            mb = "buy" if mb.startswith(("b", "p")) else "make" if mb.startswith(("m", "f")) else None
            if it:
                upd = {k: r[k] for k in ("description", "material", "uom") if r.get(k) and not it[k]}
                if mb and it["make_buy"] != mb:
                    upd["make_buy"] = mb
                if upd:
                    upd["updated_at"] = now
                    db.update("items", it["id"], upd)
            else:
                db.insert("items", {"part_number": pn, "description": r.get("description", ""),
                                    "material": r.get("material", ""), "uom": r.get("uom") or "Nos",
                                    "make_buy": mb or "make", "revision": lifecycle.clean_revision(r.get("revision")) or "A",
                                    "state": "released", "props_json": json.dumps({"source": "legacy part master"}),
                                    "created_by": user_id, "created_at": now, "updated_at": now})
            n += 1
    return n


# ------------------------------------------------------------------ engine

class Importer:
    def __init__(self, source: str, target: str, user: dict, state: str = "released", keep_revision: bool = True,
                 skip_identical: bool = True, groups: list[str] | None = None, extensions: set | None = None,
                 metadata: tuple | None = None, dry_run: bool = False, job_id: int | None = None, progress=None,
                 unpack_archives: bool = True):
        self.source = Path(source)
        self.target = target.strip().strip("/\\") or "Legacy Drawings"
        self.user = user
        self.state = state if state in ("released", "in_work") else "released"
        self.keep_revision = keep_revision
        self.skip_identical = skip_identical
        exts = set(extensions or set())
        for g in groups or DEFAULT_GROUPS:
            exts |= TYPE_GROUPS.get(g, set())
        self.extensions = exts
        self.meta_by_file, self.meta_items = metadata or ({}, [])
        self.dry_run = dry_run
        self.job_id = job_id
        self.progress = progress or (lambda **kw: None)
        self.report: list[dict] = []
        self.counts = {"imported": 0, "duplicates": 0, "skipped": 0, "errors": 0}
        self.unpack_archives = unpack_archives
        self.staging: Path | None = None
        self.shown_as: dict[Path, str] = {}      # files unpacked from an archive: "archive.zip!/inner/path"
        self.reldirs: dict[Path, str] = {}       # ... and the folder they go to, relative to the target
        self.doc_of_sha: dict[str, int] = {}
        self.found = self.junk = self.unpacked = 0

    # -- helpers
    def cancelled(self) -> bool:
        return bool(self.job_id and db.scalar("SELECT cancel FROM import_jobs WHERE id = ?", (self.job_id,)))

    def scan(self) -> tuple[list[Path], list[dict]]:
        files, ignored = [], []
        for root, dirs, names in os.walk(self.source):
            dirs[:] = sorted(d for d in dirs if not d.startswith((".", "~", "$")) and d.lower() not in ("_vti_cnf", "backup"))
            for n in sorted(names):
                p = Path(root) / n
                self.found += 1
                if self.unpack_archives and n.lower().endswith(".zip") and not JUNK.search(n):
                    files += self.unpack(p, ignored)
                else:
                    self.sort_one(p, n, files, ignored)
        return files, ignored

    def sort_one(self, p: Path, n: str, files: list, ignored: list) -> None:
        ext = n.rsplit(".", 1)[-1].lower() if "." in n else ""
        if JUNK.search(n):
            ignored.append({"path": p, "reason": "temporary or backup file"})
            self.junk += 1
        elif ext not in self.extensions:
            ignored.append({"path": p, "reason": f".{ext or '?'} files not selected"})
        else:
            files.append(p)

    def shown(self, p: Path) -> str:
        return self.shown_as.get(p) or str(p)

    def reldir(self, p: Path) -> str:
        if p in self.reldirs:
            return self.reldirs[p]
        rel = str(p.parent.relative_to(self.source)).replace("\\", "/")
        return "" if rel == "." else rel

    def unpack(self, zp: Path, ignored: list) -> list[Path]:
        """Unpack a ZIP archive into staging; its files go to a folder named after it. Returns the files to import."""
        if self.staging is None:
            self.staging = Path(config.WORK_DIR) / f"import-{self.job_id or 'run'}-{uuid.uuid4().hex[:8]}"
        dest = self.staging / uuid.uuid4().hex[:12]
        out: list[Path] = []
        try:
            with zipfile.ZipFile(zp) as z:
                infos = [i for i in z.infolist() if not i.is_dir()]
                if len(infos) > ZIP_MAX_FILES:
                    raise ValueError(f"more than {ZIP_MAX_FILES} files inside")
                if sum(i.file_size for i in infos) > ZIP_MAX_BYTES:
                    raise ValueError("more than 4 GB when unpacked")
                for i in infos:
                    name = i.filename.replace("\\", "/")
                    parts = [x for x in name.split("/") if x not in ("", ".")]
                    if name.startswith("/") or ".." in parts or (parts and ":" in parts[0]):
                        raise ValueError(f"an unsafe path inside ({i.filename})")
                    if i.flag_bits & 0x1:
                        raise ValueError("it is protected with a password")
                    if i.compress_size and i.file_size > 50 * 1024 ** 2 and i.file_size / i.compress_size > ZIP_MAX_RATIO:
                        raise ValueError(f"{i.filename} unpacks to far more than it should")
                staged = []
                for i in infos:
                    parts = [x for x in i.filename.replace("\\", "/").split("/") if x not in ("", ".")]
                    target = dest.joinpath(*parts)
                    target.parent.mkdir(parents=True, exist_ok=True)
                    written = 0
                    with z.open(i) as src, open(target, "wb") as f:
                        for chunk in iter(lambda: src.read(1024 * 1024), b""):
                            written += len(chunk)
                            if written > i.file_size:
                                raise ValueError(f"{i.filename} is bigger than the archive says")
                            f.write(chunk)
                    staged.append((target, parts))
        except (zipfile.BadZipFile, ValueError, OSError, RuntimeError, NotImplementedError) as e:
            shutil.rmtree(dest, ignore_errors=True)
            self.add(zp, "error", folder=self.target, msg=f"The archive could not be unpacked: {e}")
            return []
        base = "/".join(x for x in (self.reldir(zp), zp.stem) if x)
        for target, parts in staged:
            self.found += 1
            self.shown_as[target] = f"{zp}!/{'/'.join(parts)}"
            self.reldirs[target] = "/".join(x for x in (base, *parts[:-1]) if x)
            self.sort_one(target, parts[-1], out, ignored)
        self.unpacked += 1
        self.add(zp, "unpacked", folder="/".join(x for x in (self.target, base) if x),
                 msg=f"Archive unpacked: {len(staged)} files inside, imported from the folder {zp.stem}")
        return out

    def add(self, path: Path, status: str, **kw):
        row = {"source": self.shown(path), "status": status, "kosha_folder": kw.get("folder", ""),
               "file": self.shown(path).replace("\\", "/").rsplit("/", 1)[-1], "part_number": kw.get("pn", ""),
               "description": kw.get("desc", ""), "revision": kw.get("rev", ""), "duplicate_of": kw.get("dup", ""),
               "message": kw.get("msg", ""), "document_id": kw.get("doc_id")}
        self.report.append(row)
        key = {"imported": "imported", "would import": "imported", "identical copy": "duplicates",
               "already in Kosha": "skipped", "ignored": "skipped", "unpacked": "skipped"}.get(status, "errors")
        self.counts[key] += 1
        if self.job_id:
            legacy.add_record(self.job_id, row, kw.get("sha"), kw.get("scored"), kw.get("material", ""))
            if kw.get("doc_id") and not self.dry_run and status != "imported":
                legacy.add_ref(self.job_id, kw["doc_id"], "path", row["source"])

    # -- main
    def run(self) -> dict:
        if not self.source.is_dir():
            raise ValueError(f"Folder not found or not reachable from the server: {self.source}")
        self.progress(phase="Scanning folders")
        files, ignored = self.scan()
        for ig in ignored:
            self.add(ig["path"], "ignored", msg=ig["reason"])
        total = len(files)
        self.progress(phase="Loading part master", total=total)
        if self.meta_items and not self.dry_run:
            upsert_items(self.meta_items, self.user["id"])

        # identical content already in Kosha
        known: dict[str, str] = {}
        fpath: dict[int, str] = {}
        for r in db.rows("SELECT v.blob_sha, d.id, d.folder_id, d.filename FROM versions v JOIN documents d ON "
                         "d.id = v.document_id WHERE d.deleted = 0"):
            if r["folder_id"] not in fpath:
                fpath[r["folder_id"]] = "/".join(f["name"] for f in lifecycle.folder_path(r["folder_id"]))
            known.setdefault(r["blob_sha"], f"{fpath[r['folder_id']]}/{r['filename']}")
            self.doc_of_sha.setdefault(r["blob_sha"], r["id"])

        root_target = None if self.dry_run else lifecycle.ensure_folder_path("Kosha/" + self.target, self.user["id"])
        if self.job_id and root_target:
            db.update("import_jobs", self.job_id, {"target_folder_id": root_target})

        # hash everything first so identical copies can be grouped and the best-named copy kept
        self.progress(phase="Finding identical copies")
        hashes: dict[Path, str] = {}
        for k, p in enumerate(files):
            if k % 200 == 0 and self.cancelled():
                break
            try:
                hashes[p] = sha256_file(p)
            except OSError as e:
                self.add(p, "error", folder=self.target, msg=f"Cannot read file: {e}")
        clusters: dict[str, list[Path]] = {}
        for p, h in hashes.items():
            clusters.setdefault(h, []).append(p)
        todo: list[Path] = []
        copies: dict[Path, Path] = {}
        for h, ps in clusters.items():
            ps.sort(key=self.name_quality)
            if self.skip_identical:
                todo.append(ps[0])
                for other in ps[1:]:
                    copies[other] = ps[0]
            else:
                todo.extend(ps)
        todo.sort(key=lambda p: str(p).lower())
        seen_in_run: dict[str, str] = {}
        window = 100
        for w0 in range(0, len(todo), window):
            if self.cancelled():
                self.progress(phase="Cancelled")
                break
            batch = todo[w0:w0 + window]
            if not self.dry_run:
                dwgs = [p for p in batch if p.suffix.lower() == ".dwg" and hashes[p] not in known]
                if dwgs:
                    self.progress(phase="Converting DWG drawings")
                    for src, dxf_bytes in dwgmod.batch_to_dxf([str(p) for p in dwgs]).items():
                        cadproc.DWG_CACHE[hashes[Path(src)]] = dxf_bytes
                self.progress(phase="Importing files")
            for p in batch:
                self.one(p, hashes[p], known, seen_in_run, root_target)
                self.progress(done=len(self.report) - len(ignored), **self.counts)
            cadproc.DWG_CACHE.clear()
        for other, rep in sorted(copies.items(), key=lambda kv: self.shown(kv[0]).lower()):
            folder = "/".join(x for x in (self.target, self.reldir(other)) if x)
            self.add(other, "identical copy", folder=folder, sha=hashes[rep], doc_id=self.doc_of_sha.get(hashes[rep]),
                     dup=seen_in_run.get(hashes[rep]) or known.get(hashes[rep]) or self.shown(rep),
                     msg="Same content as another file; only one copy was imported")
        if self.staging:
            shutil.rmtree(self.staging, ignore_errors=True)
        if self.job_id:
            db.update("import_jobs", self.job_id, {"found": self.found, "junk": self.junk, "unpacked": self.unpacked})
        self.progress(phase="Finished", done=total, **self.counts)
        return {"total": total, **self.counts}

    def name_quality(self, p: Path):
        """Lower is better: avoid 'Copy of', '(1)', backup/old folders; prefer descriptive names."""
        name = p.name.lower()
        rel = "/".join(x for x in (self.reldir(p), p.name) if x).lower()
        bad = 0
        if name.startswith("copy of") or " - copy" in name or re.search(r"\(\d+\)", name):
            bad += 4
        if re.search(r"(^|/)(old|backup|bak|temp|tmp|archive|copy|recycle)[^/]*/", rel):
            bad += 2
        return (bad, rel.count("/"), -len(name), rel)

    def one(self, p: Path, sha: str, known: dict, seen: dict, root_target):
        rel_dir = self.reldir(p)
        kosha_folder = "/".join(x for x in (self.target, rel_dir) if x)
        dup = known.get(sha) or seen.get(sha)
        if dup and dup.lower() == f"kosha/{kosha_folder}/{p.name}".lower():
            self.add(p, "already in Kosha", folder=kosha_folder, sha=sha, doc_id=self.doc_of_sha.get(sha),
                     msg="Imported in an earlier run (unchanged)")
            return
        if dup and self.skip_identical:
            self.add(p, "identical copy", folder=kosha_folder, dup=dup, sha=sha, doc_id=self.doc_of_sha.get(sha),
                     msg="Same content as a file already imported; not imported again")
            return
        meta = self.meta_by_file.get(p.name.lower()) or self.meta_by_file.get(p.stem.lower()) or {}
        props = {k: meta.get(k) for k in ("part_number", "description", "material", "revision") if meta.get(k)}
        if self.keep_revision:
            props["use_title_block_revision"] = True
        if self.dry_run:
            pn, desc = lifecycle.parse_name(p.name)
            seen[sha] = f"{kosha_folder}/{p.name}"
            self.add(p, "would import", folder=kosha_folder, pn=props.get("part_number") or pn,
                     desc=props.get("description") or desc, rev=props.get("revision", ""), sha=sha)
            return
        try:
            fid = lifecycle.ensure_subfolder(root_target, rel_dir, self.user["id"]) if rel_dir else root_target
            existing = db.one("SELECT id FROM documents WHERE folder_id = ? AND filename = ? COLLATE NOCASE AND deleted = 0",
                              (fid, p.name))
            if existing:
                seen[sha] = f"{kosha_folder}/{p.name}"
                self.add(p, "already in Kosha", folder=kosha_folder, doc_id=existing["id"], sha=sha,
                         msg="A file with this name is already in the target folder (skipped; resumable import)")
                return
            mtime = datetime.fromtimestamp(p.stat().st_mtime, tz=timezone.utc).replace(microsecond=0).isoformat()
            with open(p, "rb") as f:
                d = lifecycle.checkin(self.user, fid, p.name, f, f"Imported from {self.shown(p)}", False, props,
                                      extra_meta={"source_path": self.shown(p), "source_modified": mtime,
                                                  "imported": True})
            if self.state == "released":
                self.release_quietly(d)
            seen[sha] = f"{kosha_folder}/{p.name}"
            self.doc_of_sha[sha] = d["id"]
            scored, material = self.score(d, props)
            self.add(p, "imported", folder=kosha_folder, pn=d["part_number"], desc=d["description"], rev=d["revision"],
                     doc_id=d["id"], msg=self.note_for(d), sha=sha, scored=scored, material=material)
            if self.job_id:
                self.cross_references(d)
        except Exception as e:
            log.warning("Import failed for %s: %s", p, e)
            self.add(p, "error", folder=kosha_folder, msg=getattr(e, "message", None) or str(e))

    def score(self, d: dict, props: dict) -> tuple[tuple[int, list[str]], str]:
        v = db.one("SELECT meta_json, preview_kind FROM versions WHERE document_id = ? AND version = 1", (d["id"],))
        tb = (json.loads(v["meta_json"] or "{}") if v else {}).get("title_block") or {}
        material = (props.get("material") or tb.get("material") or "").strip()
        return legacy.score(d, props, tb, bool(v and v["preview_kind"])), material

    def cross_references(self, d: dict) -> None:
        """Old paths and old numbers stay searchable for ever (specification 11.1)."""
        legacy.add_ref(self.job_id, d["id"], "path", self.report[-1]["source"])
        name_pn, _ = lifecycle.parse_name(d["filename"])
        if name_pn and name_pn.upper() != (d["part_number"] or "").upper():
            legacy.add_ref(self.job_id, d["id"], "old_number", name_pn.upper())

    @staticmethod
    def note_for(d: dict) -> str:
        v = db.one("SELECT meta_json, preview_kind FROM versions WHERE document_id = ? AND version = ?",
                   (d["id"], d["latest_version"]))
        meta = json.loads(v["meta_json"] or "{}") if v else {}
        notes = []
        if not d["part_number"]:
            notes.append("no part number found")
        if v and not v["preview_kind"]:
            notes.append(meta.get("preview_note") or "no preview")
        return "; ".join(notes)

    def release_quietly(self, d: dict) -> None:
        """Mark legacy files Released without triggering ERP pushes."""
        with db.tx():
            db.update("documents", d["id"], {"state": "released", "updated_at": db.now()})
            db.insert("lifecycle_events", {"document_id": d["id"], "revision": d["revision"], "from_state": "in_work",
                                           "to_state": "released", "user_id": self.user["id"],
                                           "comment": "Legacy import", "at": db.now()})
            lifecycle._sync_item_on_release({**d, "state": "released"})


# ------------------------------------------------------------------ report

REPORT_COLS = ["status", "source", "kosha_folder", "file", "part_number", "description", "revision", "duplicate_of",
               "message", "document_id"]


def report_bytes(rows: list[dict], summary: dict) -> tuple[bytes, str]:
    try:
        import openpyxl
        from openpyxl.styles import Font
    except ImportError:
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(REPORT_COLS)
        for r in rows:
            w.writerow([r.get(c, "") for c in REPORT_COLS])
        return buf.getvalue().encode("utf-8-sig"), "csv"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    for k, v in summary.items():
        ws.append([k, v])
    for c in ws["A"]:
        c.font = Font(bold=True)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 80
    order = ["error", "identical copy", "imported", "would import", "already in Kosha", "ignored"]
    for status in order:
        sel = [r for r in rows if r["status"] == status]
        if not sel:
            continue
        sh = wb.create_sheet(status.title()[:31])
        sh.append([c.replace("_", " ").title() for c in REPORT_COLS])
        for c in sh[1]:
            c.font = Font(bold=True)
        for r in sel:
            sh.append([r.get(c, "") if r.get(c) is not None else "" for c in REPORT_COLS])
        for col, wdt in zip("ABCDEFGHIJ", (16, 60, 40, 36, 18, 36, 9, 50, 50, 10)):
            sh.column_dimensions[col].width = wdt
        sh.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), "xlsx"


# ------------------------------------------------------------------ jobs

def start_job(user: dict, source: str, target: str, options: dict, metadata: tuple | None,
              collect: "Callable[[], list[dict]] | None" = None) -> int:
    """Run an import in the background. `collect`, for importer add-ins, first fills the source folder and returns
    the part master rows."""
    job_id = db.insert("import_jobs", {"source": source, "target": target, "options_json": json.dumps(options),
                                       "status": "running", "phase": "Starting", "created_by": user["id"],
                                       "started_at": db.now()})

    def progress(**kw):
        upd = {k: v for k, v in kw.items() if k in ("phase", "total", "done", "imported", "duplicates", "skipped", "errors")}
        if upd:
            db.update("import_jobs", job_id, upd)

    def work():
        nonlocal metadata
        imp = None
        try:
            if collect:
                progress(phase="Collecting the files from the add-in")
                rows = collect()
                if rows:
                    metadata = metadata_from_rows(rows)
                    options["metadata_rows"] = len(metadata[1])
                    db.update("import_jobs", job_id, {"options_json": json.dumps(options)})
            imp = Importer(source, target, user, options.get("state", "released"), options.get("keep_revision", True),
                           options.get("skip_identical", True), options.get("groups"), set(options.get("extensions") or []),
                           metadata, options.get("dry_run", False), job_id, progress,
                           options.get("unpack_archives", True))
            res = imp.run()
            cancelled = db.scalar("SELECT cancel FROM import_jobs WHERE id = ?", (job_id,))
            status = "cancelled" if cancelled else "done"
            msg = (f"{'Dry run: ' if imp.dry_run else ''}{res['imported']} {'would be ' if imp.dry_run else ''}imported, "
                   f"{res['duplicates']} identical copies, {res['skipped']} skipped, {res['errors']} errors")
        except Exception as e:
            log.exception("Import job %s failed", job_id)
            status, msg = "failed", str(e)
        try:
            rows = imp.report if imp else []
            summary = {"Source": source, "Kosha folder": target, "Started": db.scalar(
                "SELECT started_at FROM import_jobs WHERE id = ?", (job_id,)), "Finished": db.now(), "Result": msg,
                "Options": json.dumps(options)}
            data, _ = report_bytes(rows, summary)
            sha, _ = store.put_bytes(data)
            db.update("import_jobs", job_id, {"status": status, "message": msg, "report_sha": sha,
                                              "finished_at": db.now(), "phase": "Finished"})
            db.audit(user["id"], "legacy_import", "import_job", job_id, msg)
        finally:
            if collect and status == "done" and not options.get("dry_run"):
                shutil.rmtree(source, ignore_errors=True)     # the add-in's copies; the files are in the store now
            db.close_thread_connection()

    threading.Thread(target=work, name=f"import-{job_id}", daemon=True).start()
    return job_id


def preview(source: str, groups: list[str] | None = None, limit: int = 200000) -> dict:
    src = Path(source)
    if not src.is_dir():
        raise ValueError(f"Folder not found or not reachable from the server: {source}")
    by_ext: dict = {}
    folders = 0
    size = 0
    n = 0
    sample = []
    for root, dirs, names in os.walk(src):
        folders += 1
        if len(sample) < 12 and root != str(src):
            sample.append(str(Path(root).relative_to(src)))
        for nm in names:
            n += 1
            if n > limit:
                break
            ext = nm.rsplit(".", 1)[-1].lower() if "." in nm else ""
            try:
                sz = (Path(root) / nm).stat().st_size
            except OSError:
                sz = 0
            e = by_ext.setdefault(ext or "(none)", {"files": 0, "bytes": 0})
            e["files"] += 1
            e["bytes"] += sz
            size += sz
    selected = set()
    for g in groups or DEFAULT_GROUPS:
        selected |= TYPE_GROUPS.get(g, set())
    return {"files": n, "folders": folders, "bytes": size, "by_ext": by_ext,
            "selected_files": sum(v["files"] for k, v in by_ext.items() if k in selected),
            "sample_folders": sample, "dwg_converter": bool(dwgmod.converter_status()["oda"] or
                                                           dwgmod.converter_status()["libredwg"])}


# ------------------------------------------------------------------ CLI

def main(argv=None) -> int:
    ap = argparse.ArgumentParser(prog="python -m kosha.importer", description="Import legacy files into Kosha")
    ap.add_argument("source", help="folder or network share to import, e.g. \\\\fileserver\\Drawings")
    ap.add_argument("--to", default="", help="Kosha folder to import into (default: Legacy Drawings/<source name>)")
    ap.add_argument("--state", choices=["released", "in_work"], default="released")
    ap.add_argument("--types", default=",".join(DEFAULT_GROUPS),
                    help="comma list of: " + ", ".join(TYPE_GROUPS) + " (default: %(default)s)")
    ap.add_argument("--ext", default="", help="extra extensions, e.g. tif,txt")
    ap.add_argument("--metadata", help="part master spreadsheet (.xlsx or .csv)")
    ap.add_argument("--keep-duplicates", action="store_true", help="import identical copies too")
    ap.add_argument("--no-revision", action="store_true", help="ignore title-block revisions; start at A")
    ap.add_argument("--user", default="admin", help="Kosha user recorded as the importer")
    ap.add_argument("--dry-run", action="store_true", help="only report what would happen")
    ap.add_argument("--report", default="", help="write the report to this .xlsx/.csv file")
    a = ap.parse_args(argv)
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    db.init()
    user = db.one("SELECT id, username, full_name, role FROM users WHERE username = ?", (a.user,))
    if not user:
        print(f"User {a.user!r} not found")
        return 2
    target = a.to or f"Legacy Drawings/{Path(a.source).name or 'Import'}"
    metadata = read_metadata(a.metadata) if a.metadata else None
    last = {"t": 0}

    def progress(**kw):
        if "phase" in kw:
            print(f"-- {kw['phase']}")
        if "done" in kw and kw["done"] - last["t"] >= 50:
            last["t"] = kw["done"]
            print(f"   {kw['done']} files · imported {kw.get('imported', 0)} · identical {kw.get('duplicates', 0)} · "
                  f"errors {kw.get('errors', 0)}")

    imp = Importer(a.source, target, user, a.state, not a.no_revision, not a.keep_duplicates,
                   [t.strip() for t in a.types.split(",") if t.strip()],
                   {e.strip().lower().lstrip(".") for e in a.ext.split(",") if e.strip()}, metadata, a.dry_run,
                   progress=progress)
    res = imp.run()
    print(f"Done: {res}")
    if True:  # always leave a report behind
        data, kind = report_bytes(imp.report, {"Source": a.source, "Kosha folder": target, "Result": json.dumps(res)})
        out = Path(a.report or f"import-report-{datetime.now():%Y%m%d-%H%M}.{kind}")
        out.write_bytes(data)
        print(f"Report: {out.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
