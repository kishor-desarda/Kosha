# SPDX-License-Identifier: AGPL-3.0-or-later
"""Text-like files: Word paragraphs, Excel cells, PDF and plain-text lines.

Lists the counts and the first few changed paragraphs, lines or cells; a change log needs the
gist, not every comma.
"""
from __future__ import annotations

import csv
import difflib
import io

from .core import Change, Comparer, Side

FIRST_N = 8
MAX_LINES = 20_000
MAX_CELLS = 200_000


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


class _DocTypeRefused(ValueError):
    pass


def _refuse(*_a):
    raise _DocTypeRefused("The file contains an XML document type declaration, which Kosha does not read.")


MAX_ZIP_PART = 64 * 1024 * 1024     # a Word document.xml larger than this is not compared (zip bomb guard)


def docx_paragraphs(data: bytes) -> list[str]:
    """Paragraphs of a Word document, then one line per table row ("Table 1, row 2: a | b").

    Reads word/document.xml with a plain XML parser that refuses document type declarations (no entity expansion
    attacks) and stops at MAX_ZIP_PART bytes, so it needs no extra library."""
    import xml.parsers.expat
    import zipfile
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        info = zf.getinfo("word/document.xml")
        if info.file_size > MAX_ZIP_PART:
            raise ValueError("The Word document is too large to compare.")
        xml_bytes = zf.read(info)
    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main|"
    paras: list[str] = []
    tables: list[list[list[str]]] = []          # table -> rows -> cell texts
    state = {"in_t": 0, "tbl": 0}
    buf: list[str] = []
    cell: list[str] = []

    def start(name, _attrs):
        if name == W + "t":
            state["in_t"] += 1
        elif name == W + "tab":
            buf.append("\t")
        elif name in (W + "br", W + "cr"):
            buf.append("\n")
        elif name == W + "tbl":
            state["tbl"] += 1
            if state["tbl"] == 1:
                tables.append([])
        elif name == W + "tr" and state["tbl"] == 1:
            tables[-1].append([])
        elif name == W + "tc" and state["tbl"] == 1:
            cell.clear()

    def end(name):
        if name == W + "t":
            state["in_t"] -= 1
        elif name == W + "p":
            text = "".join(buf).strip()
            buf.clear()
            if state["tbl"]:
                if text:
                    cell.append(text)
            elif text:
                paras.append(text)
        elif name == W + "tc" and state["tbl"] == 1 and tables and tables[-1]:
            tables[-1][-1].append(" ".join(cell))
        elif name == W + "tbl":
            state["tbl"] -= 1

    def chars(text):
        if state["in_t"]:
            buf.append(text)

    p = xml.parsers.expat.ParserCreate(namespace_separator="|")
    p.StartDoctypeDeclHandler = _refuse
    p.EntityDeclHandler = _refuse
    p.StartElementHandler, p.EndElementHandler, p.CharacterDataHandler = start, end, chars
    p.Parse(xml_bytes, True)
    out = paras[:MAX_LINES]
    for ti, t in enumerate(tables, start=1):
        for ri, row in enumerate(t, start=1):
            if any(row):
                out.append(f"Table {ti}, row {ri}: " + " | ".join(row))
    return out[:MAX_LINES]


def pdf_lines(data: bytes) -> list[tuple[int, str]]:
    """(page, line) pairs of a PDF's text layer, with pypdfium2 (in requirements.txt) or else pypdf."""
    out: list[tuple[int, str]] = []
    try:
        import pypdfium2 as pdfium
    except ImportError:
        pdfium = None
    if pdfium is not None:
        pdf = pdfium.PdfDocument(data)
        try:
            for n in range(len(pdf)):
                page = pdf[n]
                try:
                    tp = page.get_textpage()
                    try:
                        text = tp.get_text_range()
                    finally:
                        tp.close()
                finally:
                    page.close()
                for line in text.splitlines():
                    if line.strip():
                        out.append((n + 1, " ".join(line.split())))
                    if len(out) >= MAX_LINES:
                        return out
        finally:
            pdf.close()
        return out
    from pypdf import PdfReader
    r = PdfReader(io.BytesIO(data))
    for n, page in enumerate(r.pages, start=1):
        for line in (page.extract_text() or "").splitlines():
            if line.strip():
                out.append((n, " ".join(line.split())))
            if len(out) >= MAX_LINES:
                return out
    return out


def xlsx_cells(data: bytes) -> dict[str, str]:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out: dict[str, str] = {}
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    v = getattr(c, "value", None)
                    if v is None or str(v).strip() == "":
                        continue
                    ref = getattr(c, "coordinate", None)
                    if ref is None:
                        continue
                    out[f"{ws.title}!{ref}"] = str(v).strip()
                    if len(out) >= MAX_CELLS:
                        return out
    finally:
        wb.close()
    return out


def diff_lines(old: list[str], new: list[str], unit: str, pages_old=None, pages_new=None) -> list[Change]:
    """Changed / added / removed lines (or paragraphs) with the first few as examples."""
    sm = difflib.SequenceMatcher(None, old, new, autojunk=False)
    changed = added = removed = 0
    examples: list[Change] = []

    def example(i, j, b, a):
        if len(examples) < FIRST_N:
            page = (pages_new[j] if a is not None and pages_new else pages_old[i] if pages_old else None)
            examples.append(Change("text_line", f"{unit.capitalize()} {(j if a is not None else i) + 1}", b, a, "minor",
                                   None, {"unit": unit, "index": (j if a is not None else i) + 1, "page": page}))

    for op, i1, i2, j1, j2 in sm.get_opcodes():
        if op == "equal":
            continue
        n_old, n_new = i2 - i1, j2 - j1
        common = min(n_old, n_new) if op == "replace" else 0
        for k in range(common):
            changed += 1
            example(i1 + k, j1 + k, old[i1 + k], new[j1 + k])
        for k in range(common, n_new):
            added += 1
            example(i1, j1 + k, None, new[j1 + k])
        for k in range(common, n_old):
            removed += 1
            example(i1 + k, j1, old[i1 + k], None)
    if not (changed or added or removed):
        return []
    summary = Change("text_summary", "Document", len(old), len(new), "minor", None,
                     {"unit": unit, "changed": changed, "added": added, "removed": removed})
    return [summary] + examples


def diff_cells(old: dict[str, str], new: dict[str, str]) -> list[Change]:
    keys = list(dict.fromkeys(list(old) + list(new)))
    changed = added = removed = 0
    examples = []
    for k in keys:
        a, b = old.get(k), new.get(k)
        if a == b:
            continue
        if a is None:
            added += 1
        elif b is None:
            removed += 1
        else:
            changed += 1
        if len(examples) < FIRST_N:
            examples.append(Change("cell", k, a, b, "minor"))
    if not examples:
        return []
    return [Change("text_summary", "Workbook", len(old), len(new), "minor", None,
                   {"unit": "cell", "changed": changed, "added": added, "removed": removed})] + examples


class TextComparer(Comparer):
    name = "text"
    label = "Document text comparison"
    exts = ("docx", "xlsx", "xlsm", "pdf", "txt", "csv")

    def applies(self, a: Side, b: Side) -> bool:
        return a.has_content and b.has_content

    def compare(self, a: Side, b: Side) -> list[Change]:
        da, db = a.data(), b.data()
        if da == db:
            return []
        ext = (b.ext or a.ext).lower()
        if ext == "docx":
            return diff_lines(docx_paragraphs(da), docx_paragraphs(db), "paragraph")
        if ext in ("xlsx", "xlsm"):
            return diff_cells(xlsx_cells(da), xlsx_cells(db))
        if ext == "pdf":
            la, lb = pdf_lines(da), pdf_lines(db)
            if not la and not lb:
                self.note("These PDFs have no text layer (scanned pages?), so only the file fingerprint was compared.")
                return []
            return diff_lines([t for _, t in la], [t for _, t in lb], "line", [p for p, _ in la], [p for p, _ in lb])
        if ext == "csv":
            rows_a = [" | ".join(r) for r in csv.reader(io.StringIO(_decode(da))) if any(x.strip() for x in r)]
            rows_b = [" | ".join(r) for r in csv.reader(io.StringIO(_decode(db))) if any(x.strip() for x in r)]
            return diff_lines(rows_a[:MAX_LINES], rows_b[:MAX_LINES], "row")
        la = [ln.rstrip() for ln in _decode(da).splitlines() if ln.strip()]
        lb = [ln.rstrip() for ln in _decode(db).splitlines() if ln.strip()]
        return diff_lines(la[:MAX_LINES], lb[:MAX_LINES], "line")
