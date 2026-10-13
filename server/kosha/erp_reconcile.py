# SPDX-License-Identifier: AGPL-3.0-or-later
"""ERP reconciliation: compare Kosha's released items and BOMs with what the ERP holds, in plain words.

A check runs in the background with progress. It only reads from the ERP, through the connector's
optional read-back methods (`fetch_items`, `fetch_bom`, `list_items_from_kosha`) and the add-in guard.
Nothing in the ERP changes unless a user presses "Send again" on a difference; that queues an ordinary,
idempotent ERP job, exactly like a manual push.

What is compared:
  * every released item: missing in the ERP, disabled there, description, unit, revision, material
    (a field the connector cannot read is not compared);
  * every obsolete item the ERP still has active;
  * the released BOM of every released item: missing, lines missing or extra, quantities;
  * items the ERP marks as created by Kosha that Kosha does not have.
"""
from __future__ import annotations

import io
import logging
import threading

from . import db, erp
from .addins import registry

from kosha_sdk import AddinError  # noqa: E402
from kosha_sdk.erp import ErpNotSupported  # noqa: E402

log = logging.getLogger("kosha.erp.reconcile")
BATCH = 50

CHECKS = {
    "missing_in_erp": "Missing in the ERP",
    "inactive_in_erp": "Disabled in the ERP",
    "description": "Description differs",
    "uom": "Unit differs",
    "revision": "Revision differs",
    "material": "Material differs",
    "obsolete_active": "Obsolete in Kosha, active in the ERP",
    "bom_missing": "BOM missing in the ERP",
    "bom_line_missing": "BOM line missing in the ERP",
    "bom_line_extra": "Extra BOM line in the ERP",
    "bom_qty": "BOM quantity differs",
    "erp_only": "Only in the ERP",
}
SCOPES = {"item": "Item", "bom": "BOM", "erp": "Only in the ERP"}

_running: set[int] = set()
_guard = threading.Lock()


class ReconcileError(AddinError):
    """The check cannot start (switched off, already running, no connector). The message says why."""


def _norm(text) -> str:
    return " ".join(str(text or "").split())


def _qty(q: float) -> str:
    q = round(float(q), 6)
    return str(int(q)) if q == int(q) else f"{q:g}"


def _connector() -> tuple[str, str]:
    cid = erp.connector_id()
    a = registry.get(cid)
    name = a.manifest.name.replace(" connector", "") if a and a.manifest else cid
    return cid, name


# -------------------------------------------------------------------- start

def start(user_id: int | None) -> int:
    cid, name = _connector()
    if not erp.settings().get("enabled"):
        raise ReconcileError(f"Tick 'Connect to {name}' in the connection settings first. The check only reads from "
                             f"{name}; it works while Dry run is on.")
    with _guard:
        db.execute("UPDATE erp_recon_runs SET status = 'interrupted', phase = 'Server restarted', finished_at = ? "
                   f"WHERE status = 'running' AND id NOT IN ({','.join(str(i) for i in _running) or '0'})", (db.now(),))
        if db.scalar("SELECT COUNT(*) FROM erp_recon_runs WHERE status = 'running'"):
            raise ReconcileError("A check is already running. Wait for it to finish, or stop it.")
        run_id = db.insert("erp_recon_runs", {"connector": cid, "connector_name": name, "status": "running",
                                              "phase": "Starting", "started_by": user_id, "started_at": db.now()})
        _running.add(run_id)
    threading.Thread(target=_work, args=(run_id, cid, name), name=f"erp-reconcile-{run_id}", daemon=True).start()
    return run_id


def nightly() -> int | None:
    """The nightly check (run by the health worker): only when the ERP is connected for real and the admin left it on.
    Returns the run's id, or None when it was not started. Differences show on the health page."""
    cfg = erp.settings()
    if not (cfg.get("enabled") and not cfg.get("dry_run") and cfg.get("nightly_reconcile", True)):
        return None
    try:
        return start(None)
    except (ReconcileError, erp.ErpError) as e:
        log.info("Nightly ERP check not started: %s", e)
        return None


def cancel(run_id: int) -> None:
    db.execute("UPDATE erp_recon_runs SET cancel = 1 WHERE id = ? AND status = 'running'", (run_id,))


class _Stopped(Exception):
    pass


# --------------------------------------------------------------------- work

class _Run:
    def __init__(self, run_id: int, cid: str, name: str):
        self.id, self.cid, self.name = run_id, cid, name
        self.cfg = erp.settings()
        self.done = 0
        self.notes: list[str] = []
        self.waiting = {(j["kind"], j["entity_id"]): j for j in db.rows(
            "SELECT id, kind, entity_id, status, error FROM erp_jobs WHERE status IN ('pending', 'running', 'failed') "
            "ORDER BY id")}

    def progress(self, phase: str | None = None, step: int = 0) -> None:
        self.done += step
        upd = {"done": self.done}
        if phase:
            upd["phase"] = phase
        db.update("erp_recon_runs", self.id, upd)
        if db.scalar("SELECT cancel FROM erp_recon_runs WHERE id = ?", (self.id,)):
            raise _Stopped()

    def call(self, method: str, *args):
        return registry.call(self.cid, method, *args)

    def job_note(self, kind: str, entity_id: int | None) -> str:
        j = self.waiting.get((kind, entity_id))
        if not j:
            return ""
        if j["status"] == "failed":
            return f" The last transfer failed (job #{j['id']}): {j['error'] or 'see the queue'}."
        return f" A transfer is waiting in the queue (job #{j['id']})."

    def add(self, *, part_number: str, scope: str, check: str, message: str, kosha="", erp_value="", advice="",
            item_id=None, bom_id=None, fix: tuple[str, int] | None = None) -> None:
        db.insert("erp_recon_rows", {
            "run_id": self.id, "part_number": part_number, "item_id": item_id, "bom_id": bom_id, "scope": scope,
            "check_code": check, "kosha_value": str(kosha)[:2000], "erp_value": str(erp_value)[:2000],
            "message": message[:2000], "advice": advice[:2000],
            "fix_kind": fix[0] if fix else "", "fix_entity_id": fix[1] if fix else None})


def _work(run_id: int, cid: str, name: str) -> None:
    run = _Run(run_id, cid, name)
    try:
        released = db.rows("SELECT * FROM items WHERE state = 'released' ORDER BY part_number")
        obsolete = db.rows("SELECT * FROM items WHERE state = 'obsolete' ORDER BY part_number")
        boms = _released_boms()
        db.update("erp_recon_runs", run_id, {"total": len(released) + len(obsolete) + len(boms) + 1})
        run.progress(f"Reading items from {name}")
        for start_at in range(0, len(released) + len(obsolete), BATCH):
            chunk = (released + obsolete)[start_at:start_at + BATCH]
            got = run.call("fetch_items", [r["part_number"] for r in chunk]) or {}
            found = {str(k).upper(): v for k, v in got.items()}
            for row in chunk:
                rec = found.get(row["part_number"].upper())
                if row["state"] == "released":
                    _compare_item(run, row, rec)
                elif rec is not None and rec.active is not False:
                    run.add(part_number=row["part_number"], scope="item", check="obsolete_active", item_id=row["id"],
                            kosha="Obsolete", erp_value="Active",
                            message=f"Obsolete in Kosha, but still active in {name}.",
                            advice=f"Disable or archive it in {name} by hand if it should no longer be bought or made. "
                                   "Kosha never disables items in the ERP itself.")
            run.progress(step=len(chunk))
        db.update("erp_recon_runs", run_id, {"items_checked": len(released) + len(obsolete)})
        run.progress(f"Reading BOMs from {name}")
        try:
            for i, b in enumerate(boms):
                _compare_bom(run, b, run.call("fetch_bom", b["part_number"]))
                run.progress(step=1)
                if i % 10 == 9:
                    db.update("erp_recon_runs", run_id, {"boms_checked": i + 1})
            db.update("erp_recon_runs", run_id, {"boms_checked": len(boms)})
        except ErpNotSupported as e:
            run.notes.append(f"BOMs were not compared: {e}")
            run.progress(step=len(released) + len(obsolete) + len(boms) - run.done)
        run.progress(f"Looking for items that are only in {name}")
        try:
            codes = run.call("list_items_from_kosha") or []
        except ErpNotSupported as e:
            run.notes.append(f"Items that exist only in {name} were not looked for: {e}")
        else:
            known = {r["pn"].upper() for r in db.rows("SELECT part_number AS pn FROM items")}
            for code in sorted({str(c) for c in codes if c and str(c).upper() not in known}):
                run.add(part_number=code, scope="erp", check="erp_only", kosha="", erp_value="Marked as from Kosha",
                        message=f"{name} has {code} marked as coming from Kosha, but Kosha has no such item.",
                        advice=f"Check whether it was renamed or came from a test system, then disable it in {name} by "
                               "hand. Kosha never deletes or disables items in the ERP.")
        run.progress(step=1)
        _finish(run, "done")
    except _Stopped:
        _finish(run, "cancelled", "Stopped by a user. The differences found until then are kept.")
    except ErpNotSupported as e:
        _finish(run, "failed", f"The {name} connector cannot read items back, so Kosha cannot compare: {e} "
                               "Update the add-in, or ask its publisher.")
    except AddinError as e:
        _finish(run, "failed", f"{name} could not be read: {e} Nothing was changed anywhere. "
                               "The differences found until then are kept; run the check again later.")
    except Exception as e:  # pragma: no cover - a bug in Kosha; logged, and the run is closed properly
        log.exception("Reconciliation %s crashed", run_id)
        _finish(run, "failed", f"The check stopped because of an error in Kosha ({e.__class__.__name__}). Nothing was "
                               "changed anywhere. The admin can see the details in the log.")
    finally:
        _running.discard(run_id)


def _released_boms() -> list[dict]:
    """The released BOM of every released item: the one for the item's current revision, else the newest."""
    out = {}
    for b in db.rows("SELECT b.id, b.item_id, b.revision, i.part_number, i.revision AS item_revision FROM boms b "
                     "JOIN items i ON i.id = b.item_id WHERE b.state = 'released' AND i.state = 'released' "
                     "ORDER BY b.id"):
        cur = out.get(b["item_id"])
        if cur is None or cur["revision"] != cur["item_revision"] or b["revision"] == b["item_revision"]:
            out[b["item_id"]] = b
    return sorted(out.values(), key=lambda b: b["part_number"])


def _compare_item(run: _Run, row: dict, rec) -> None:
    want = erp.to_erp_item(run.cfg, row)
    pn, name = row["part_number"], run.name
    note = run.job_note("item", row["id"])
    common = {"part_number": pn, "scope": "item", "item_id": row["id"]}
    if rec is None:
        run.add(**common, check="missing_in_erp", kosha=want.description, erp_value="",
                message=f"Released in Kosha (rev {want.revision}), but not in {name}.",
                advice=f"Send it again to create it in {name}." + note, fix=("item", row["id"]))
        return
    if rec.active is False:
        run.add(**common, check="inactive_in_erp", kosha="Released", erp_value="Disabled or archived",
                message=f"Released in Kosha, but disabled or archived in {name}.",
                advice=f"Switch it back on in {name} by hand if it is still used. Sending again does not change this.")
    if rec.description is not None and _norm(rec.description) != _norm(want.description):
        run.add(**common, check="description", kosha=want.description, erp_value=rec.description,
                message="The description differs.", advice=f"Send it again to give {name} Kosha's description." + note,
                fix=("item", row["id"]))
    if rec.uom is not None and _norm(rec.uom).lower() != _norm(want.uom).lower():
        run.add(**common, check="uom", kosha=want.uom, erp_value=rec.uom,
                message=f"The unit differs (Kosha's unit {row['uom'] or '(none)'} maps to {want.uom}).",
                advice=f"ERPs lock the unit once an item has stock or transactions, so sending again does not change it. "
                       f"Correct it in {name} by hand, or fix the unit mapping in the connection settings if Kosha's "
                       "side is wrong.")
    if rec.revision is not None and _norm(rec.revision) != _norm(want.revision):
        run.add(**common, check="revision", kosha=want.revision, erp_value=rec.revision or "(empty)",
                message="The revision differs.", advice=f"Send it again to give {name} Kosha's revision." + note,
                fix=("item", row["id"]))
    if rec.material is not None and _norm(rec.material) != _norm(want.material):
        fixable = bool(want.material)
        run.add(**common, check="material", kosha=want.material or "(empty)", erp_value=rec.material or "(empty)",
                message="The material differs.",
                advice=(f"Send it again to give {name} Kosha's material." + note) if fixable else
                f"Kosha has no material for this item and never clears one in {name}. Add it in Kosha, or clear it "
                f"in {name} by hand.", fix=("item", row["id"]) if fixable else None)


def _lines(pairs) -> dict[str, tuple[str, float]]:
    """{CODE: (code as written, total quantity)}: a part listed twice counts once with both quantities."""
    out: dict[str, tuple[str, float]] = {}
    for code, qty in pairs:
        key = str(code).upper()
        prev = out.get(key)
        out[key] = (prev[0] if prev else str(code), (prev[1] if prev else 0.0) + float(qty or 0))
    return out


def _compare_bom(run: _Run, b: dict, rec) -> None:
    name, pn = run.name, b["part_number"]
    kosha = _lines((ln["part_number"], ln["qty"]) for ln in erp._bom_lines(b["id"]))
    note = run.job_note("bom", b["id"])
    common = {"part_number": pn, "scope": "bom", "item_id": b["item_id"], "bom_id": b["id"], "fix": ("bom", b["id"])}
    resend = f"Send the BOM again to make {name}'s BOM match rev {b['revision']}." + note
    if rec is None:
        run.add(**common, check="bom_missing", kosha=f"rev {b['revision']}, {len(kosha)} lines", erp_value="",
                message=f"The released BOM (rev {b['revision']}) is not in {name}.", advice=resend)
        return
    theirs = _lines((ln.code, ln.qty) for ln in rec.lines)
    ref = f" ({rec.reference})" if rec.reference else ""
    for key in sorted(kosha.keys() - theirs.keys()):
        code, q = kosha[key]
        run.add(**common, check="bom_line_missing", kosha=f"{code} × {_qty(q)}", erp_value="",
                message=f"Line {code} × {_qty(q)} is missing in {name}'s BOM{ref}.", advice=resend)
    for key in sorted(theirs.keys() - kosha.keys()):
        code, q = theirs[key]
        run.add(**common, check="bom_line_extra", kosha="", erp_value=f"{code} × {_qty(q)}",
                message=f"{name}'s BOM{ref} has {code} × {_qty(q)}, which is not in Kosha's released BOM.",
                advice=resend)
    for key in sorted(kosha.keys() & theirs.keys()):
        (code, q1), (_, q2) = kosha[key], theirs[key]
        if abs(q1 - q2) > 1e-6:
            run.add(**common, check="bom_qty", kosha=f"{code} × {_qty(q1)}", erp_value=f"{code} × {_qty(q2)}",
                    message=f"The quantity of {code} differs: Kosha {_qty(q1)}, {name} {_qty(q2)}.", advice=resend)


def _finish(run: _Run, status: str, message: str = "") -> None:
    diffs = db.scalar("SELECT COUNT(*) FROM erp_recon_rows WHERE run_id = ?", (run.id,))
    fixable = db.scalar("SELECT COUNT(*) FROM erp_recon_rows WHERE run_id = ? AND fix_kind != ''", (run.id,))
    if status == "done":
        message = (f"Kosha and {run.name} match." if not diffs else
                   f"{diffs} difference{'s' if diffs != 1 else ''} found; {fixable} can be fixed by sending again.")
    db.update("erp_recon_runs", run.id, {"status": status, "phase": "", "message": message, "notes": "\n".join(run.notes),
                                         "differences": diffs, "fixable": fixable, "finished_at": db.now()})


# ------------------------------------------------------------------- reading

def run_detail(run_id: int | None = None) -> dict:
    run = db.one("SELECT * FROM erp_recon_runs WHERE id = ?", (run_id,)) if run_id else \
        db.one("SELECT * FROM erp_recon_runs ORDER BY id DESC LIMIT 1")
    if run and run["status"] == "running" and run["id"] not in _running:
        db.update("erp_recon_runs", run["id"], {"status": "interrupted", "phase": "Server restarted",
                                                "finished_at": db.now()})
        run = db.one("SELECT * FROM erp_recon_runs WHERE id = ?", (run["id"],))
    rows = db.rows("SELECT r.*, j.status AS job_status, j.error AS job_error FROM erp_recon_rows r LEFT JOIN erp_jobs j "
                   "ON j.id = r.resend_job_id WHERE r.run_id = ? ORDER BY r.part_number, r.scope DESC, r.id",
                   (run["id"],)) if run else []
    runs = db.rows("SELECT id, status, started_at, finished_at, differences, connector_name FROM erp_recon_runs "
                   "ORDER BY id DESC LIMIT 10")
    return {"run": run, "rows": rows, "runs": runs, "checks": CHECKS, "scopes": SCOPES}


# --------------------------------------------------------------- sending again

def _queue_fix(kind: str, entity_id: int, user_id: int | None, ref: str) -> int:
    if kind == "item":
        if not db.one("SELECT id FROM items WHERE id = ?", (entity_id,)):
            raise ReconcileError("The item no longer exists in Kosha.")
        return erp.enqueue("item", entity_id, user_id, ref)
    b = db.one("SELECT id, item_id FROM boms WHERE id = ?", (entity_id,))
    if not b:
        raise ReconcileError("The BOM no longer exists in Kosha.")
    erp.enqueue("item", b["item_id"], user_id, ref)      # the parent first, as a manual push does
    return erp.enqueue("bom", entity_id, user_id, ref)


def _mark(run_id: int, kind: str, entity_id: int, job_id: int, user_id: int | None) -> None:
    db.execute("UPDATE erp_recon_rows SET resend_job_id = ?, resent_by = ?, resent_at = ? WHERE run_id = ? AND "
               "fix_kind = ? AND fix_entity_id = ?", (job_id, user_id, db.now(), run_id, kind, entity_id))


def _kick() -> None:
    threading.Thread(target=erp.process_pending, name="erp-resend", daemon=True).start()


def resend_row(row_id: int, user_id: int | None) -> dict:
    row = db.one("SELECT * FROM erp_recon_rows WHERE id = ?", (row_id,))
    if not row:
        raise ReconcileError("That difference is not in the list any more. Run the check again.")
    if not row["fix_kind"]:
        raise ReconcileError("Sending again cannot fix this difference. " + row["advice"])
    with db.tx():
        jid = _queue_fix(row["fix_kind"], row["fix_entity_id"], user_id, f"Reconciliation #{row['run_id']}")
        _mark(row["run_id"], row["fix_kind"], row["fix_entity_id"], jid, user_id)
        db.audit(user_id, "erp_reconcile_resend", "item", row["item_id"], f"{row['part_number']}: {row['check_code']}")
    _kick()
    return db.one("SELECT * FROM erp_jobs WHERE id = ?", (jid,))


def resend_all(run_id: int, user_id: int | None) -> int:
    fixes = db.rows("SELECT DISTINCT fix_kind, fix_entity_id FROM erp_recon_rows WHERE run_id = ? AND fix_kind != '' "
                    "AND resend_job_id IS NULL ORDER BY fix_kind DESC", (run_id,))       # items before BOMs
    n = 0
    with db.tx():
        for f in fixes:
            try:
                jid = _queue_fix(f["fix_kind"], f["fix_entity_id"], user_id, f"Reconciliation #{run_id}")
            except ReconcileError:
                continue
            _mark(run_id, f["fix_kind"], f["fix_entity_id"], jid, user_id)
            n += 1
        if n:
            db.audit(user_id, "erp_reconcile_resend", "erp", None, f"run {run_id}: {n} transfer(s) queued")
    if n:
        _kick()
    return n


# -------------------------------------------------------------------- export

def export_xlsx(run_id: int) -> tuple[bytes, str]:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    d = run_detail(run_id)
    run = d["run"]
    if not run:
        raise ReconcileError("That check does not exist.")
    name = run["connector_name"] or "ERP"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    bold = Font(bold=True)
    ws.append([f"Kosha and {name}: reconciliation #{run['id']}"])
    ws["A1"].font = Font(bold=True, size=14)
    for label, value in (("Started", run["started_at"]), ("Finished", run["finished_at"] or ""), ("Status", run["status"]),
                         ("Result", run["message"]), ("Items checked", run["items_checked"]),
                         ("BOMs checked", run["boms_checked"]), ("Differences", run["differences"]),
                         ("Can be fixed by sending again", run["fixable"])):
        ws.append([label, value])
    for note in (run["notes"] or "").splitlines():
        ws.append(["Note", note])
    ws.append([])
    ws.append(["Difference", "How many"])
    ws.cell(ws.max_row, 1).font = ws.cell(ws.max_row, 2).font = bold
    counts: dict[str, int] = {}
    for r in d["rows"]:
        counts[r["check_code"]] = counts.get(r["check_code"], 0) + 1
    for code, label in CHECKS.items():
        if counts.get(code):
            ws.append([label.replace("the ERP", name), counts[code]])
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 70

    ws2 = wb.create_sheet("Differences")
    head = ["Part number", "Where", "Difference", "In Kosha", f"In {name}", "Details", "What to do",
            "Fixed by sending again", "Sent again", "Transfer status"]
    ws2.append(head)
    fill = PatternFill("solid", fgColor="E7F0FA")
    for c in ws2[1]:
        c.font, c.fill = bold, fill
    for r in d["rows"]:
        ws2.append([r["part_number"], SCOPES.get(r["scope"], r["scope"]), CHECKS.get(r["check_code"], r["check_code"])
                    .replace("the ERP", name), r["kosha_value"], r["erp_value"], r["message"], r["advice"],
                    "yes" if r["fix_kind"] else "no", r["resent_at"] or "",
                    (f"job #{r['resend_job_id']}: {r['job_status']}" if r["resend_job_id"] else "")])
    for col, width in zip("ABCDEFGHIJ", (20, 14, 28, 34, 34, 60, 60, 12, 22, 22)):
        ws2.column_dimensions[col].width = width
    for row in ws2.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = f"A1:J{max(1, ws2.max_row)}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), f"kosha-{name.lower().replace(' ', '-')}-reconciliation-{run['id']}.xlsx"
