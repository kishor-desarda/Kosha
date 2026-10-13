# SPDX-License-Identifier: AGPL-3.0-or-later
"""Legacy migration batches (specification 11): what happens after the importer has read the files.

Each import job is a *batch*. Every file it saw becomes a record with a status (imported, identical copy, already in
Kosha, ignored, unpacked archive, error), so the batch can be reconciled: *files found = the sum of the statuses*.

Imported files get a quality score from 0 to 100 with its reasons, and go to a work queue: **Auto** (90 and above),
**Review** (60 to 89) or **Attention** (below 60). People confirm them on the review grid or in Excel, and fix the
part number, description or material on the way. Old paths and old numbers stay searchable as cross-references.

A batch is signed off when the counts add up, every stored file reads back with its fingerprint, and a random
sample was checked by an engineer at 98 % or better. Before that, *Undo batch* takes back what it made, as long as
nobody has used it since; what was used is listed instead.
"""
from __future__ import annotations

import io
import json
import random
import re
import sqlite3

from . import db, lifecycle, store
from .web import bad, conflict, not_found

QUEUES = ("auto", "review", "attention")
QUEUE_LABEL = {"auto": "Auto", "review": "Review", "attention": "Attention"}
EDITABLE = ("part_number", "description", "material")
PN = re.compile(r"^[A-Z0-9][A-Z0-9 ._/-]{0,39}$")
SAMPLE_PASS = 98.0
# how each status counts in the reconciliation (specification 11.1: found = imported + copies + junk + errors)
STATUS_GROUP = {"imported": "imported", "would import": "imported", "identical copy": "identical copies",
                "already in Kosha": "already in Kosha", "ignored": "set aside", "unpacked": "archives unpacked",
                "error": "errors"}


def queue_for(score: int) -> str:
    return "auto" if score >= 90 else "review" if score >= 60 else "attention"


# ------------------------------------------------------------------------------------------- scoring --

def score(doc: dict, props: dict, title_block: dict, preview: bool) -> tuple[int, list[str]]:
    """How complete the record is and how sure the values are, with the reason for every point."""
    pn_name, desc_name = lifecycle.parse_name(doc["filename"])
    tb = title_block or {}
    pts, why = 0, []
    pn = doc["part_number"]
    if not pn:
        why.append("No part number found (0 of 30)")
    elif props.get("part_number"):
        pts += 30
        why.append(f"Part number {pn}: from the part master (30)")
    elif tb.get("part_number"):
        pts += 30
        why.append(f"Part number {pn}: from the title block (30)")
    else:
        pts += 20
        why.append(f"Part number {pn}: from the file name only (20 of 30)")
    tb_pn = (tb.get("part_number") or "").strip().upper()
    if pn and tb_pn and pn_name and tb_pn != pn_name.upper():
        why.append(f"The title block says {tb_pn} but the file name says {pn_name.upper()} (0 of 10)")
    elif pn:
        pts += 10
        why.append("Part number agrees everywhere it was found (10)")
    if not doc["description"]:
        why.append("No description found (0 of 20)")
    elif props.get("description") or tb.get("description"):
        pts += 20
        why.append(f"Description from the {'part master' if props.get('description') else 'title block'} (20)")
    else:
        pts += 12
        why.append("Description from the file name only (12 of 20)")
    material = props.get("material") or tb.get("material")
    if material:
        pts += 15
        why.append(f"Material {material} (15)")
    else:
        why.append("No material found (0 of 15)")
    if props.get("revision") or tb.get("revision"):
        pts += 15
        why.append(f"Revision {doc['revision']} read from the {'part master' if props.get('revision') else 'title block'} (15)")
    else:
        why.append(f"No revision mark; {doc['revision']} assumed (0 of 15)")
    if preview:
        pts += 10
        why.append("Picture made (10)")
    else:
        why.append("No picture could be made (0 of 10)")
    return pts, why


# ------------------------------------------------------------------------------------- records and refs --

def add_record(batch_id: int, row: dict, sha: str | None = None, scored: tuple[int, list[str]] | None = None,
               material: str = "") -> int:
    sc, why = scored if scored else (None, [])
    return db.insert("legacy_records", {
        "batch_id": batch_id, "source_path": row["source"][:1000], "status": row["status"], "sha": sha,
        "document_id": row.get("document_id"), "part_number": row.get("part_number") or "",
        "description": row.get("description") or "", "material": material or "", "revision": row.get("revision") or "",
        "score": sc, "queue": queue_for(sc) if sc is not None else None, "reasons_json": json.dumps(why),
        "message": (row.get("message") or "")[:500], "created_at": db.now()})


def add_ref(batch_id: int | None, document_id: int, kind: str, value: str) -> None:
    value = " ".join(str(value or "").split())[:1000]
    if not value or not document_id:
        return
    if db.one("SELECT 1 FROM legacy_refs WHERE document_id = ? AND kind = ? AND value = ?", (document_id, kind, value)):
        return
    db.insert("legacy_refs", {"batch_id": batch_id, "document_id": document_id, "kind": kind, "value": value,
                              "created_at": db.now()})


def refs(document_id: int) -> list[dict]:
    return db.rows("SELECT kind, value, batch_id, created_at FROM legacy_refs WHERE document_id = ? ORDER BY id",
                   (document_id,))


def batch_or_404(batch_id: int) -> dict:
    j = db.one("SELECT * FROM import_jobs WHERE id = ?", (batch_id,))
    if not j:
        raise not_found("No such import batch.")
    return j


def _record(rid: int) -> dict:
    r = db.one("SELECT * FROM legacy_records WHERE id = ?", (rid,))
    if not r:
        raise not_found("No such record.")
    return r


def _out(r: dict, names: dict) -> dict:
    r = dict(r)
    r["reasons"] = json.loads(r.pop("reasons_json") or "[]")
    r["queue_label"] = QUEUE_LABEL.get(r["queue"] or "", "")
    r["confirmed_by_name"] = names.get(r["confirmed_by"]) if r["confirmed_by"] else None
    return r


def records(batch_id: int, queue: str = "", q: str = "", unconfirmed: bool = False, status: str = "",
            offset: int = 0, limit: int = 50) -> dict:
    batch_or_404(batch_id)
    where, params = ["batch_id = ?"], [batch_id]
    if queue:
        if queue not in QUEUES:
            raise bad("Choose the Auto, Review or Attention queue.")
        where.append("queue = ?")
        params.append(queue)
    if status:
        where.append("status = ?")
        params.append(status)
    if unconfirmed:
        where.append("confirmed_at IS NULL")
    if q:
        where.append("(source_path LIKE ? OR part_number LIKE ? OR description LIKE ?)")
        params += [f"%{q}%"] * 3
    sql = " WHERE " + " AND ".join(where)
    total = db.scalar("SELECT COUNT(*) FROM legacy_records" + sql, params)
    rows = db.rows("SELECT * FROM legacy_records" + sql + " ORDER BY COALESCE(score, 999), id LIMIT ? OFFSET ?",
                   [*params, min(max(limit, 1), 500), max(offset, 0)])
    counts = {qn: db.scalar("SELECT COUNT(*) FROM legacy_records WHERE batch_id = ? AND queue = ?", (batch_id, qn))
              for qn in QUEUES}
    counts["confirmed"] = db.scalar("SELECT COUNT(*) FROM legacy_records WHERE batch_id = ? AND confirmed_at IS NOT NULL",
                                    (batch_id,))
    previews = {}
    ids = [r["document_id"] for r in rows if r["document_id"]]
    if ids:
        for v in db.rows("SELECT d.id, d.latest_version, v.preview_kind FROM documents d JOIN versions v ON "
                         "v.document_id = d.id AND v.version = d.latest_version WHERE d.id IN (" +
                         ",".join("?" * len(ids)) + ")", ids):
            previews[v["id"]] = (v["preview_kind"], v["latest_version"])
    names = lifecycle.user_names()
    out = []
    for r in rows:
        x = _out(r, names)
        x["preview_kind"], x["version"] = previews.get(r["document_id"], (None, None))
        out.append(x)
    return {"records": out, "total": total, "counts": counts}


# ------------------------------------------------------------------------------------------- review --

def _clean(field: str, value) -> str:
    v = " ".join(str(value if value is not None else "").split())
    if field == "part_number":
        v = v.upper()
        if not PN.match(v):
            raise bad(f"{v or 'An empty part number'} is not a part number: use letters, digits and - . / _ "
                      "(at most 40).")
    elif field == "description":
        if not v:
            raise bad("The description cannot be empty.")
        v = v[:140]
    else:
        v = v[:60]
    return v


def _check_changes(changes: dict) -> dict:
    out = {}
    for k, v in changes.items():
        if k not in EDITABLE:
            raise bad(f"Only {', '.join(EDITABLE)} can be changed here.")
        out[k] = _clean(k, v)
    return out


def _apply(user: dict, r: dict, changes: dict, confirm: bool) -> None:
    """Inside a transaction: write the corrections to the file and its part, and confirm the record."""
    upd = {k: v for k, v in changes.items() if v != r[k]}
    d = db.one("SELECT * FROM documents WHERE id = ?", (r["document_id"],)) if r["document_id"] else None
    if upd and not d:
        raise conflict("This record has no file in Kosha, so there is nothing to correct.", "no_file")
    if d and d["deleted"]:
        raise conflict(f"{d['filename']} is in the recycle bin. Put it back first.", "in_bin")
    if upd:
        doc_upd = {"updated_at": d["updated_at"]}         # a correction in review is not "use" of the file
        if "part_number" in upd and upd["part_number"] != d["part_number"]:
            if d["part_number"]:
                add_ref(r["batch_id"], d["id"], "old_number", d["part_number"])
            doc_upd["part_number"] = upd["part_number"]
            doc_upd["item_id"] = lifecycle._link_item(upd["part_number"], upd.get("description", d["description"]),
                                                      d["doc_type"], user["id"], upd.get("material", r["material"]))
        if "description" in upd:
            doc_upd["description"] = upd["description"]
        db.update("documents", d["id"], doc_upd)
        item_id = doc_upd.get("item_id", d["item_id"])
        if item_id and ("description" in upd or "material" in upd):
            alone = db.scalar("SELECT COUNT(*) FROM documents WHERE item_id = ? AND deleted = 0", (item_id,)) <= 1
            it = db.one("SELECT description, material FROM items WHERE id = ?", (item_id,))
            iu = {k: upd[k] for k in ("description", "material") if k in upd and (alone or not it[k])}
            if iu:
                iu["updated_at"] = db.now()
                db.update("items", item_id, iu)
        db.audit(user["id"], "legacy_correct", "document", d["id"],
                 "; ".join(f"{k} {r[k] or '(empty)'} → {v}" for k, v in upd.items()))
    rec = dict(upd)
    if confirm and not r["confirmed_at"]:
        rec.update({"confirmed_by": user["id"], "confirmed_at": db.now()})
    if rec:
        db.update("legacy_records", r["id"], rec)


def correct(user: dict, rid: int, changes: dict, confirm: bool = False) -> dict:
    changes = _check_changes(changes)
    with db.tx():
        r = _record(rid)
        _not_closed(batch_or_404(r["batch_id"]))
        _apply(user, r, changes, confirm)
    return _out(_record(rid), lifecycle.user_names())


def _not_closed(batch: dict) -> None:
    if batch["undone_at"]:
        raise conflict("This batch was undone.", "batch_undone")
    if batch["status"] == "running":
        raise conflict("This batch is still importing. Wait for it to finish.", "batch_running")


def bulk(user: dict, batch_id: int, ids: list, changes: dict | None = None, confirm: bool = True) -> int:
    """Confirm many records at once, optionally setting one field across all of them."""
    changes = _check_changes(changes or {})
    if "part_number" in changes:
        raise bad("A part number is different for every part; correct it one row at a time.")
    try:
        ids = [int(i) for i in ids]
    except (TypeError, ValueError):
        raise bad("ids must be record numbers.") from None
    with db.tx():
        _not_closed(batch_or_404(batch_id))
        rows = db.rows("SELECT * FROM legacy_records WHERE batch_id = ? AND id IN (" + ",".join("?" * len(ids)) + ")",
                       [batch_id, *ids]) if ids else []
        if len(rows) != len(set(ids)):
            raise bad("Some of these records are not in this batch.")
        for r in rows:
            _apply(user, r, changes, confirm)
    return len(rows)


# --------------------------------------------------------------------------------- Excel round-trip --

XLSX_COLS = ["Record", "File", "Old path", "Part number", "Description", "Material", "Revision", "Score", "Queue",
             "Why", "Confirm"]


def export_xlsx(batch_id: int, queue: str = "") -> bytes:
    import openpyxl
    from openpyxl.styles import Font
    data = records(batch_id, queue=queue, limit=500)
    rows, offset = data["records"], 500
    while offset < data["total"]:
        rows += records(batch_id, queue=queue, offset=offset, limit=500)["records"]
        offset += 500
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = (QUEUE_LABEL.get(queue) or "All records")[:31]
    ws.append(XLSX_COLS)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        if r["status"] != "imported":
            continue
        ws.append([r["id"], r["source_path"].replace("\\", "/").rsplit("/", 1)[-1], r["source_path"], r["part_number"],
                   r["description"], r["material"], r["revision"], r["score"], r["queue_label"],
                   "; ".join(r["reasons"]), "yes" if r["confirmed_at"] else ""])
    for col, w in zip("ABCDEFGHIJK", (9, 34, 60, 18, 40, 16, 9, 7, 10, 80, 9)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    note = wb.create_sheet("How to use")
    for line in ("Correct Part number, Description and Material in the first sheet; write yes under Confirm for the "
                 "rows that are right.", "Do not change the Record column. Rows you leave out stay as they are.",
                 "Then import this file again on the batch's Review screen. Nothing is changed unless every row is "
                 "valid; the problems are listed by row."):
        note.append([line])
    note.column_dimensions["A"].width = 120
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def import_xlsx(user: dict, batch_id: int, data: bytes) -> dict:
    """Corrections made offline come back: every row is checked first; nothing changes unless all rows are valid."""
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise bad("This is not an Excel file (.xlsx). Export the queue again and edit that file.") from None
    ws = wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    head = [str(h or "").strip() for h in next(it, [])]
    need = ["Record", "Part number", "Description", "Material", "Confirm"]
    missing = [h for h in need if h not in head]
    if missing:
        raise bad(f"The sheet has no {', '.join(missing)} column. Export the queue again and edit that file.")
    col = {h: head.index(h) for h in need}
    plan, problems = [], []
    for n, row in enumerate(it, start=2):
        if not row or all(v in (None, "") for v in row):
            continue
        try:
            rid = int(row[col["Record"]])
        except (TypeError, ValueError):
            problems.append({"row": n, "message": "The Record column was changed; it must stay as exported."})
            continue
        r = db.one("SELECT * FROM legacy_records WHERE id = ? AND batch_id = ?", (rid, batch_id))
        if not r:
            problems.append({"row": n, "message": f"Record {rid} is not in this batch."})
            continue
        changes = {}
        for field, head_name in (("part_number", "Part number"), ("description", "Description"), ("material", "Material")):
            v = row[col[head_name]]
            v = "" if v is None else str(v)
            if " ".join(v.split()) == r[field]:
                continue
            try:
                changes[field] = _clean(field, v)
            except Exception as e:
                problems.append({"row": n, "message": getattr(e, "message", str(e))})
        confirm = str(row[col["Confirm"]] or "").strip().lower() in ("yes", "y", "x", "ok", "1", "true", "haan", "ha")
        if changes and not r["document_id"]:
            problems.append({"row": n, "message": "This record has no file in Kosha, so there is nothing to correct."})
        plan.append((r, changes, confirm))
    if problems:
        return {"applied": 0, "confirmed": 0, "problems": problems}
    applied = confirmed = 0
    with db.tx():
        _not_closed(batch_or_404(batch_id))
        for r, changes, confirm in plan:
            fresh = _record(r["id"])
            _apply(user, fresh, changes, confirm)
            applied += bool(changes)
            confirmed += bool(confirm and not fresh["confirmed_at"])
        db.audit(user["id"], "legacy_excel", "import_job", batch_id, f"{applied} corrected, {confirmed} confirmed")
    return {"applied": applied, "confirmed": confirmed, "problems": []}


# ------------------------------------------------------------------------- reconcile, sample, sign off --

def reconcile(batch_id: int) -> dict:
    j = batch_or_404(batch_id)
    by = {}
    for r in db.rows("SELECT status, COUNT(*) AS n FROM legacy_records WHERE batch_id = ? GROUP BY status", (batch_id,)):
        key = STATUS_GROUP.get(r["status"], "errors")
        by[key] = by.get(key, 0) + r["n"]
    accounted = sum(by.values())
    found = j["found"] or 0
    verify = json.loads(j["verify_json"]) if j.get("verify_json") else None
    sample = db.one("SELECT COUNT(*) AS size, SUM(CASE WHEN sample_result = 'ok' THEN 1 ELSE 0 END) AS ok, "
                    "SUM(CASE WHEN sample_result = 'wrong' THEN 1 ELSE 0 END) AS wrong FROM legacy_records "
                    "WHERE batch_id = ? AND sample_result IS NOT NULL", (batch_id,))
    size, ok, wrong = sample["size"] or 0, sample["ok"] or 0, sample["wrong"] or 0
    checked = ok + wrong
    accuracy = round(100.0 * ok / checked, 1) if checked else None
    blockers = []
    if found != accounted:
        blockers.append(f"{found} files were found but {accounted} are accounted for"
                        + (" (the import was cancelled)" if j["status"] == "cancelled" else "") + ".")
    if not verify:
        blockers.append("The stored files have not been read back yet.")
    elif verify["bad"]:
        blockers.append(f"{len(verify['bad'])} stored files do not read back correctly.")
    if not size:
        blockers.append("No sample has been drawn for checking.")
    elif checked < size:
        blockers.append(f"{size - checked} of the {size} sample files still need checking.")
    elif accuracy is not None and accuracy < SAMPLE_PASS:
        blockers.append(f"The sample is {accuracy} % right; {SAMPLE_PASS:g} % is needed. Correct the batch and draw a "
                        "new sample.")
    if j["status"] in ("running", "failed") or j["undone_at"]:
        blockers.insert(0, "Only a finished batch can be signed off.")
    if json.loads(j["options_json"] or "{}").get("dry_run"):
        blockers.insert(0, "A dry run loads nothing, so there is nothing to sign off.")
    return {"found": found, "by_status": by, "accounted": accounted, "adds_up": found == accounted,
            "verify": verify, "sample": {"size": size, "checked": checked, "ok": ok, "wrong": wrong, "accuracy": accuracy},
            "can_sign_off": not blockers and not j["signed_off_at"], "blockers": blockers,
            "signed_off_at": j["signed_off_at"], "signed_off_by": j["signed_off_by"]}


def verify_files(batch_id: int) -> dict:
    """Read back every stored file of the batch and compare it with its fingerprint (SHA-256)."""
    batch_or_404(batch_id)
    rows = db.rows("SELECT r.id, r.source_path, v.blob_sha FROM legacy_records r JOIN versions v ON "
                   "v.document_id = r.document_id AND v.version = 1 WHERE r.batch_id = ? AND r.status = 'imported'",
                   (batch_id,))
    bad_files = [{"record": r["id"], "source": r["source_path"]} for r in rows if not store.verify(r["blob_sha"])]
    res = {"checked": len(rows), "bad": bad_files, "at": db.now()}
    db.update("import_jobs", batch_id, {"verify_json": json.dumps(res)})
    return res


def draw_sample(user: dict, batch_id: int, size: int | None = None) -> dict:
    """A random sample of imported records for an engineer to check (2 %, at least 5, at most 50)."""
    with db.tx():
        _not_closed(batch_or_404(batch_id))
        ids = [r["id"] for r in db.rows("SELECT id FROM legacy_records WHERE batch_id = ? AND status = 'imported'",
                                        (batch_id,))]
        if not ids:
            raise conflict("This batch imported no files, so there is nothing to sample.", "nothing_imported")
        n = min(len(ids), size or max(5, min(50, -(-len(ids) * 2 // 100))))
        db.execute("UPDATE legacy_records SET sample_result = NULL, sample_by = NULL WHERE batch_id = ?", (batch_id,))
        picked = random.SystemRandom().sample(ids, n)
        db.execute("UPDATE legacy_records SET sample_result = 'pending' WHERE id IN (" + ",".join("?" * n) + ")", picked)
        db.audit(user["id"], "legacy_sample", "import_job", batch_id, f"{n} files drawn for checking")
    return reconcile(batch_id)["sample"]


def mark_sample(user: dict, rid: int, result: str) -> dict:
    if result not in ("ok", "wrong"):
        raise bad("Mark the file right (ok) or wrong.")
    with db.tx():
        r = _record(rid)
        if r["sample_result"] is None:
            raise conflict("This file is not in the sample.", "not_in_sample")
        db.update("legacy_records", rid, {"sample_result": result, "sample_by": user["id"]})
    return _out(_record(rid), lifecycle.user_names())


def sign_off(user: dict, batch_id: int) -> dict:
    with db.tx():
        rec = reconcile(batch_id)
        if rec["signed_off_at"]:
            raise conflict("This batch is already signed off.", "signed_off")
        if not rec["can_sign_off"]:
            raise conflict("The batch cannot be signed off yet: " + " ".join(rec["blockers"]), "not_ready",
                           blockers=rec["blockers"])
        db.update("import_jobs", batch_id, {"signed_off_by": user["id"], "signed_off_at": db.now()})
        db.audit(user["id"], "legacy_sign_off", "import_job", batch_id,
                 f"{rec['found']} files accounted for; sample {rec['sample']['accuracy']} % right")
    return reconcile(batch_id)


# ------------------------------------------------------------------------------------------ undo --

def _used(d: dict, batch: dict) -> str:
    """Why this file may not be taken back ("" when nobody has used it since the batch)."""
    if d["latest_version"] > 1:
        return "a newer version was checked in"
    if d["checked_out_by"]:
        return "it is checked out"
    if batch["finished_at"] and d["updated_at"] > batch["finished_at"]:
        return "it was changed after the import"
    if db.scalar("SELECT 1 FROM co_items WHERE document_id = ? LIMIT 1", (d["id"],)):
        return "it is on a change order"
    if d["item_id"] and db.scalar("SELECT 1 FROM bom_lines WHERE child_item_id = ? LIMIT 1", (d["item_id"],)):
        return "its part is used in a BOM"
    if d["item_id"] and db.scalar("SELECT 1 FROM boms WHERE item_id = ? LIMIT 1", (d["item_id"],)):
        return "its part has a BOM"
    return ""


def undo_plan(batch_id: int) -> dict:
    j = batch_or_404(batch_id)
    remove, keep = [], []
    for d in db.rows("SELECT r.id AS record_id, r.source_path, d.* FROM legacy_records r JOIN documents d ON "
                     "d.id = r.document_id WHERE r.batch_id = ? AND r.status = 'imported' AND d.deleted = 0", (batch_id,)):
        entry = {"record": d["record_id"], "document_id": d["id"], "filename": d["filename"], "source": d["source_path"]}
        why = _used(d, j)
        if why:
            keep.append({**entry, "why": why})
        else:
            remove.append(entry)
    return {"remove": remove, "keep": keep, "undone_at": j["undone_at"], "signed_off_at": j["signed_off_at"]}


def undo(user: dict, batch_id: int) -> dict:
    """Take back what the batch made and nobody used: its files go to the recycle bin (restorable), the parts and
    folders it created and nothing else uses are removed, its cross-references go. Used files are listed instead."""
    from .filesafety import to_bin
    with db.tx():
        j = batch_or_404(batch_id)
        if j["undone_at"]:
            raise conflict("This batch was undone already.", "batch_undone")
        if j["signed_off_at"]:
            raise conflict("This batch is signed off; its files are in use as the company's drawings. Delete single "
                           "files instead.", "signed_off")
        if j["status"] == "running":
            raise conflict("This batch is still importing. Cancel it first, then undo it.", "batch_running")
        plan = undo_plan(batch_id)
        root = lifecycle.root_folder_id()
        removed_ids = [x["document_id"] for x in plan["remove"]]
        for did in removed_ids:
            d = db.one("SELECT * FROM documents WHERE id = ?", (did,))
            to_bin(user["id"], d)
            db.update("documents", did, {"folder_id": root})          # its folder path stays in deleted_path
        if removed_ids:
            db.execute("DELETE FROM legacy_refs WHERE document_id IN (" + ",".join("?" * len(removed_ids)) + ")",
                       removed_ids)
        # parts the batch created (from its files or its part master) that nothing else uses
        parts = 0
        end = j["finished_at"] or db.now()
        marks = ",".join("?" * len(removed_ids)) or "NULL"
        candidates = db.rows("SELECT id FROM items WHERE (created_by = ? AND created_at >= ? AND created_at <= ?) OR "
                             f"(created_at >= ? AND id IN (SELECT item_id FROM documents WHERE id IN ({marks})))",
                             [j["created_by"], j["started_at"], end, j["started_at"], *removed_ids])
        for it in candidates:
            if db.scalar("SELECT 1 FROM documents WHERE item_id = ? AND deleted = 0 LIMIT 1", (it["id"],)):
                continue
            if db.scalar("SELECT 1 FROM erp_jobs WHERE kind IN ('item','refresh') AND entity_id = ? LIMIT 1", (it["id"],)):
                continue
            binned = [r["id"] for r in db.rows("SELECT id FROM documents WHERE item_id = ?", (it["id"],))]
            if any(b not in removed_ids for b in binned):
                continue
            for b in binned:
                db.update("documents", b, {"item_id": None})          # put back from the bin, it is linked again
            try:
                db.execute("DELETE FROM items WHERE id = ?", (it["id"],))
                parts += 1
            except sqlite3.IntegrityError:                           # something else refers to it: it stays
                for b in binned:
                    db.update("documents", b, {"item_id": it["id"]})
        # folders the batch created that are now empty
        folders = 0
        if j["target_folder_id"]:
            todo, ids = [j["target_folder_id"]], []
            while todo:
                f = todo.pop()
                ids.append(f)
                todo += [r["id"] for r in db.rows("SELECT id FROM folders WHERE parent_id = ?", (f,))]
            for fid in reversed(ids):
                f = db.one("SELECT * FROM folders WHERE id = ?", (fid,))
                if (f and f["parent_id"] and f["created_at"] >= j["started_at"]
                        and not db.scalar("SELECT 1 FROM documents WHERE folder_id = ? LIMIT 1", (fid,))
                        and not db.scalar("SELECT 1 FROM folders WHERE parent_id = ? LIMIT 1", (fid,))):
                    try:
                        db.execute("DELETE FROM folders WHERE id = ?", (fid,))
                        folders += 1
                    except sqlite3.IntegrityError:
                        pass
        db.update("import_jobs", batch_id, {"undone_by": user["id"], "undone_at": db.now()})
        summary = (f"{len(removed_ids)} files to the recycle bin, {parts} parts and {folders} folders removed; "
                   f"{len(plan['keep'])} files kept because they were used")
        db.audit(user["id"], "legacy_undo", "import_job", batch_id, summary)
    return {"removed": len(removed_ids), "parts_removed": parts, "folders_removed": folders, "kept": plan["keep"],
            "summary": summary}
