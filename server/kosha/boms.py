# SPDX-License-Identifier: AGPL-3.0-or-later
"""Engineering BOMs: generate, edit, validate, compare, import and export."""
from __future__ import annotations

import csv
import io
import json
import re

from . import db, lifecycle
from .web import bad, conflict, not_found

BUY_HINT = re.compile(r"\b(BOLT|NUT|WASHER|SCREW|BEARING|SEAL|GASKET|O-RING|ORING|MOTOR|GEARBOX|COUPLING|"
                      r"CIRCLIP|KEY|PIN|SPRING|VALVE|SENSOR|GREASE|PAINT)\b", re.I)


def bom_or_404(bom_id: int) -> dict:
    b = db.one("SELECT b.*, i.part_number, i.description, i.revision AS item_revision, i.state AS item_state, "
               "i.item_type FROM boms b JOIN items i ON i.id = b.item_id WHERE b.id = ?", (bom_id,))
    if not b:
        raise not_found("BOM not found")
    return b


def bom_for_item(item_id: int, revision: str | None = None) -> dict | None:
    if revision:
        b = db.one("SELECT id FROM boms WHERE item_id = ? AND revision = ?", (item_id, revision))
        if b:
            return b
    return db.one("SELECT id FROM boms WHERE item_id = ? ORDER BY CASE state WHEN 'released' THEN 0 ELSE 1 END, "
                  "updated_at DESC LIMIT 1", (item_id,))


def lines(bom_id: int) -> list[dict]:
    return db.rows(
        "SELECT bl.id, bl.line_no, bl.qty, bl.uom, bl.find_number, bl.note, bl.child_item_id AS item_id, "
        "i.part_number, i.description, i.material, i.make_buy, i.uom AS item_uom, i.revision, i.state, "
        "i.item_type, i.erp_synced_at, i.erp_stock, i.standard_item_id, i.weight_kg "
        "FROM bom_lines bl JOIN items i ON i.id = bl.child_item_id WHERE bl.bom_id = ? ORDER BY bl.line_no",
        (bom_id,))


def explode(bom_id: int, max_depth: int = 12) -> list[dict]:
    out: list[dict] = []

    def walk(bid: int, level: int, mult: float, path: tuple):
        for ln in lines(bid):
            child_bom = bom_for_item(ln["item_id"], ln["revision"])
            row = dict(ln)
            row.update({"level": level, "ext_qty": round(ln["qty"] * mult, 6), "path": "/".join(map(str, path + (ln["id"],))),
                        "has_children": bool(child_bom), "child_bom_id": child_bom["id"] if child_bom else None,
                        "cycle": ln["item_id"] in path})
            out.append(row)
            if child_bom and level < max_depth and ln["item_id"] not in path:
                walk(child_bom["id"], level + 1, mult * ln["qty"], path + (ln["item_id"],))

    root = bom_or_404(bom_id)
    walk(bom_id, 1, 1.0, (root["item_id"],))
    return out


def validate(bom_id: int) -> list[dict]:
    issues: list[dict] = []
    seen: dict[int, int] = {}
    for row in explode(bom_id):
        pn = row["part_number"]
        if row["cycle"]:
            issues.append({"severity": "error", "line": row["path"], "part_number": pn,
                           "message": f"{pn} contains itself (circular BOM)."})
        if row["level"] == 1:
            seen[row["item_id"]] = seen.get(row["item_id"], 0) + 1
        if row["qty"] <= 0:
            issues.append({"severity": "error", "line": row["path"], "part_number": pn, "message": "Quantity must be above zero."})
        if not (row["description"] or "").strip():
            issues.append({"severity": "warning", "line": row["path"], "part_number": pn, "message": "Missing description."})
        if row["make_buy"] == "make" and not (row["material"] or "").strip() and row["item_type"] != "assembly":
            issues.append({"severity": "warning", "line": row["path"], "part_number": pn,
                           "message": "Manufactured part has no material."})
        if row["state"] != "released":
            issues.append({"severity": "warning", "line": row["path"], "part_number": pn,
                           "message": f"Not released yet (revision {row['revision']} is {lifecycle.STATE_LABEL.get(row['state'], row['state']).lower()})."})
        if row["standard_item_id"]:
            std = db.scalar("SELECT part_number FROM items WHERE id = ?", (row["standard_item_id"],))
            issues.append({"severity": "info", "line": row["path"], "part_number": pn,
                           "message": f"Duplicate of standard part {std}. Consider using {std} instead."})
    for item_id, count in seen.items():
        if count > 1:
            pn = db.scalar("SELECT part_number FROM items WHERE id = ?", (item_id,))
            issues.append({"severity": "warning", "line": "", "part_number": pn,
                           "message": f"{pn} appears on {count} lines; combine them into one line."})
    return issues


def ensure_item(part_number: str, description: str = "", user_id=None, uom: str = "", material: str = "",
                make_buy: str | None = None) -> int:
    pn = part_number.strip().upper()
    if not pn:
        raise bad("Part number is required")
    item = db.one("SELECT id FROM items WHERE part_number = ? COLLATE NOCASE", (pn,))
    if item:
        return item["id"]
    now = db.now()
    mb = make_buy or ("buy" if BUY_HINT.search(description or "") else "make")
    return db.insert("items", {"part_number": pn, "description": description or "", "item_type": "part",
                               "material": material or "", "uom": uom or "Nos", "make_buy": mb,
                               "revision": db.get_setting("first_revision") or "A", "state": "in_work",
                               "props_json": "{}", "created_by": user_id, "created_at": now, "updated_at": now})


def create(user_id: int, item_id: int, revision: str | None = None, source: str = "manual", note: str = "") -> int:
    item = db.one("SELECT * FROM items WHERE id = ?", (item_id,))
    if not item:
        raise not_found("Item not found")
    rev = revision or _working_revision(item_id) or item["revision"]
    existing = db.one("SELECT * FROM boms WHERE item_id = ? AND revision = ?", (item_id, rev))
    if existing:
        return existing["id"]
    now = db.now()
    with db.tx():
        if item["item_type"] != "assembly":
            db.update("items", item_id, {"item_type": "assembly", "updated_at": now})
        return db.insert("boms", {"item_id": item_id, "revision": rev, "source": source, "state": "draft",
                                  "note": note, "created_by": user_id, "created_at": now, "updated_at": now})


def _working_revision(item_id: int) -> str | None:
    d = db.one("SELECT revision FROM documents WHERE item_id = ? AND deleted = 0 AND doc_type IN ('assembly','part','model') "
               "ORDER BY updated_at DESC LIMIT 1", (item_id,))
    return d["revision"] if d else None


def _check_editable(bom: dict) -> None:
    if bom["state"] == "released":
        raise conflict(f"BOM revision {bom['revision']} is released. Revise the assembly to change it.", "bom_released")


def set_lines(user_id: int, bom_id: int, new_lines: list[dict], source: str | None = None) -> None:
    bom = bom_or_404(bom_id)
    _check_editable(bom)
    with db.tx():
        db.execute("DELETE FROM bom_lines WHERE bom_id = ?", (bom_id,))
        for k, ln in enumerate(new_lines, start=1):
            item_id = ln.get("item_id")
            if not item_id:
                item_id = ensure_item(str(ln.get("part_number") or ""), ln.get("description", ""), user_id,
                                      ln.get("uom", ""), ln.get("material", ""))
            if item_id == bom["item_id"]:
                raise bad(f"{bom['part_number']} cannot contain itself.")
            try:
                qty = float(ln.get("qty", 1))
            except (TypeError, ValueError):
                raise bad(f"Line {k}: quantity must be a number")
            db.insert("bom_lines", {"bom_id": bom_id, "line_no": k, "child_item_id": item_id, "qty": qty,
                                    "uom": ln.get("uom") or "", "find_number": str(ln.get("find_number") or k),
                                    "note": ln.get("note") or ""})
        upd = {"updated_at": db.now()}
        if source:
            upd["source"] = source
        db.update("boms", bom_id, upd)
        db.audit(user_id, "bom_edit", "bom", bom_id, f"{len(new_lines)} lines")


def generate_from_document(user: dict, doc_id: int) -> dict:
    """Build the EBOM of an assembly from what Kosha knows about the model."""
    d = lifecycle.doc_or_404(doc_id)
    if d["doc_type"] != "assembly":
        raise bad("Choose a SolidWorks assembly (.SLDASM) to generate a BOM.")
    if not d["item_id"]:
        raise bad("The assembly has no part number yet. Add one in its properties first.")
    ver = db.one("SELECT meta_json FROM versions WHERE document_id = ? AND version = ?", (doc_id, d["latest_version"]))
    meta = json.loads(ver["meta_json"]) if ver else {}
    structure = meta.get("structure")
    unresolved: list[str] = []
    new_lines: list[dict] = []
    source = "solidworks"
    if structure:
        for k, comp in enumerate(structure, start=1):
            pn = (comp.get("part_number") or "").strip()
            if not pn:
                pn, _ = lifecycle.parse_name(comp.get("file", ""))
            if not pn:
                unresolved.append(comp.get("file") or "?")
                continue
            new_lines.append({"part_number": pn, "description": comp.get("description", ""), "qty": comp.get("qty", 1),
                              "material": comp.get("material", ""), "find_number": str(k)})
    else:
        refs = meta.get("references") or []
        source = "file_references"
        for name in refs:
            doc = db.one("SELECT item_id, part_number, description FROM documents WHERE filename = ? COLLATE NOCASE "
                         "AND deleted = 0 ORDER BY updated_at DESC LIMIT 1", (name,))
            pn = doc["part_number"] if doc and doc["part_number"] else lifecycle.parse_name(name)[0]
            if not pn:
                unresolved.append(name)
                continue
            new_lines.append({"part_number": pn, "description": (doc or {}).get("description", "") or
                              lifecycle.parse_name(name)[1], "qty": 1})
        if not refs:
            raise bad("Kosha could not read the component list from this file. Check it in from SolidWorks with the "
                      "Kosha add-in, or import the BOM table exported from SolidWorks.")
    # merge duplicate part numbers
    merged: dict[str, dict] = {}
    for ln in new_lines:
        key = ln["part_number"].upper()
        if key in merged:
            merged[key]["qty"] = float(merged[key]["qty"]) + float(ln["qty"])
        else:
            merged[key] = dict(ln)
    bom_id = create(user["id"], d["item_id"], d["revision"], source)
    set_lines(user["id"], bom_id, list(merged.values()), source)
    return {"bom_id": bom_id, "lines": len(merged), "unresolved": unresolved, "source": source,
            "quantities_verified": bool(structure)}


# ---------------------------------------------------------------- import --

HEADER_ALIASES = {
    "item": ("item no.", "item no", "item", "itemno", "find no", "find", "pos", "position", "sr no", "s.no", "no."),
    "part_number": ("part number", "part no.", "part no", "partno", "pn", "item code", "drawing no", "dwg no",
                    "component", "part"),
    "description": ("description", "desc", "name", "title"),
    "qty": ("qty.", "qty", "quantity", "qty per", "no. off", "nos"),
    "uom": ("uom", "unit", "units"),
    "material": ("material", "matl", "moc"),
}


def _read_table(filename: str, data: bytes) -> list[list[str]]:
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        try:
            import openpyxl
        except ImportError:
            raise bad("Excel import needs the openpyxl package on the server. Save the BOM as CSV instead.")
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        return [["" if c is None else str(c).strip() for c in row] for row in ws.iter_rows(values_only=True)]
    text = data.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), dialect)]


def _map_header(row: list[str]) -> dict[str, int] | None:
    found: dict[str, int] = {}
    low = [c.lower().strip() for c in row]
    for key, aliases in HEADER_ALIASES.items():
        for idx, cell in enumerate(low):
            if cell in aliases and key not in found and idx not in found.values():
                found[key] = idx
                break
    if "part_number" in found and "qty" in found:
        return found
    return None


def import_table(user_id: int, bom_id: int, filename: str, data: bytes) -> dict:
    """Import a BOM table (SolidWorks 'Save as Excel/CSV', indented or top-level only)."""
    bom = bom_or_404(bom_id)
    _check_editable(bom)
    table = _read_table(filename, data)
    header = None
    start = 0
    for k, row in enumerate(table[:30]):
        header = _map_header(row)
        if header:
            start = k + 1
            break
    if not header:
        raise bad("Couldn't find the BOM columns. The file needs at least 'Part Number' and 'Qty' columns.")
    nodes = []
    for row in table[start:]:
        def cell(key):
            i = header.get(key)
            return row[i].strip() if i is not None and i < len(row) and row[i] is not None else ""
        pn = cell("part_number")
        if not pn:
            continue
        try:
            qty = float(cell("qty").replace(",", ".") or 1)
        except ValueError:
            qty = 1.0
        item_no = cell("item") or str(len(nodes) + 1)
        nodes.append({"item": item_no, "part_number": pn, "description": cell("description"), "qty": qty,
                      "uom": cell("uom"), "material": cell("material")})
    if not nodes:
        raise bad("The table has no rows with a part number.")
    # build tree from item numbers like 1, 1.1, 1.2.3
    children: dict[str, list[dict]] = {"": []}
    for n in nodes:
        parent = n["item"].rsplit(".", 1)[0] if "." in n["item"] else ""
        children.setdefault(parent, []).append(n)
        children.setdefault(n["item"], [])
    created_sub = []
    skipped_sub = []
    with db.tx():
        for n in nodes:
            ensure_item(n["part_number"], n["description"], user_id, n["uom"], n["material"])
            if n["material"]:
                db.execute("UPDATE items SET material = ? WHERE part_number = ? COLLATE NOCASE AND material = ''",
                           (n["material"], n["part_number"].upper()))
        set_lines(user_id, bom_id, [{"part_number": n["part_number"], "qty": n["qty"], "uom": n["uom"],
                                      "find_number": n["item"]} for n in children[""]], "import")
        for n in nodes:
            kids = children.get(n["item"]) or []
            if not kids:
                continue
            item_id = db.scalar("SELECT id FROM items WHERE part_number = ? COLLATE NOCASE", (n["part_number"],))
            sub_id = create(user_id, item_id, None, "import")
            sub = bom_or_404(sub_id)
            if sub["state"] == "released":
                skipped_sub.append(n["part_number"])
                continue
            set_lines(user_id, sub_id, [{"part_number": k["part_number"], "qty": k["qty"], "uom": k["uom"],
                                         "find_number": k["item"].rsplit(".", 1)[-1]} for k in kids], "import")
            created_sub.append(n["part_number"])
    return {"lines": len(children[""]), "rows": len(nodes), "sub_assemblies": created_sub,
            "skipped_released": skipped_sub}


# ------------------------------------------------------------ compare/io --

def compare(bom_a: int, bom_b: int) -> dict:
    a = {ln["part_number"]: ln for ln in lines(bom_a)}
    b = {ln["part_number"]: ln for ln in lines(bom_b)}
    added = [b[k] for k in b if k not in a]
    removed = [a[k] for k in a if k not in b]
    changed = [{"part_number": k, "description": b[k]["description"], "from_qty": a[k]["qty"], "to_qty": b[k]["qty"]}
               for k in b if k in a and abs(a[k]["qty"] - b[k]["qty"]) > 1e-9]
    return {"added": added, "removed": removed, "changed": changed}


def export(bom_id: int, fmt: str = "csv", indented: bool = True) -> tuple[bytes, str, str]:
    bom = bom_or_404(bom_id)
    rows = explode(bom_id) if indented else lines(bom_id)
    head = ["Level", "Find no.", "Part number", "Description", "Qty", "Total qty", "UoM", "Material", "Make/Buy",
            "Revision", "State"]
    data = []
    for r in rows:
        data.append([r.get("level", 1), r["find_number"], r["part_number"], r["description"], r["qty"],
                     r.get("ext_qty", r["qty"]), r["uom"] or r["item_uom"], r["material"], r["make_buy"],
                     r["revision"], lifecycle.STATE_LABEL.get(r["state"], r["state"])])
    base = f"BOM_{bom['part_number']}_rev{bom['revision']}"
    if fmt == "xlsx":
        try:
            import openpyxl
            from openpyxl.styles import Font
        except ImportError:
            fmt = "csv"
        else:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "EBOM"
            ws.append([f"{bom['part_number']} {bom['description']}", f"Revision {bom['revision']}",
                       f"State: {bom['state']}"])
            ws.append([])
            ws.append(head)
            for c in ws[3]:
                c.font = Font(bold=True)
            for r in data:
                ws.append(r)
            for col, width in zip("ABCDEFGHIJK", (7, 9, 20, 40, 8, 10, 8, 16, 10, 9, 11)):
                ws.column_dimensions[col].width = width
            buf = io.BytesIO()
            wb.save(buf)
            return buf.getvalue(), base + ".xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(head)
    w.writerows(data)
    return buf.getvalue().encode("utf-8-sig"), base + ".csv", "text/csv"


def where_used(item_id: int) -> list[dict]:
    return db.rows(
        "SELECT b.id AS bom_id, b.revision AS bom_revision, b.state AS bom_state, i.id AS item_id, i.part_number, "
        "i.description, bl.qty FROM bom_lines bl JOIN boms b ON b.id = bl.bom_id JOIN items i ON i.id = b.item_id "
        "WHERE bl.child_item_id = ? ORDER BY i.part_number, b.revision", (item_id,))
