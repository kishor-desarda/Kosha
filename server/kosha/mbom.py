# SPDX-License-Identifier: AGPL-3.0-or-later
"""Manufacturing BOMs and process plans.

The engineering BOM (EBOM) says what a machine is made of, the way the designer structured it. The manufacturing BOM
(MBOM) says how the workshop builds it: which parts are issued at which operation, which sub-assemblies are built in
one go with their parent (*flattened*), and what the workshop adds that the design does not list (paint, grease,
thread-locker). The **process plan** (routing) is the list of numbered operations, each at a work centre with its
set-up and run times and instructions.

- An MBOM belongs to one assembly and one engineering revision, and is made from that revision's EBOM. A new
  revision carries the operations, the parts' operations, the flattened sub-assemblies and the added parts forward.
- **Checks** compare what the two BOMs finally consume, part by part down to the bought and made parts, so a
  flattened sub-assembly counts the same as one built separately. A part the design needs and the workshop does not
  issue (or issues too few of) is an error and stops the release; extra quantities are warnings; parts added for
  manufacturing are listed so a reviewer sees them.
- An approver releases it once the EBOM of the same revision is released; it is then locked. Add-ins hear
  ``mbom.released``, and when the ERP settings say the ERP should get manufacturing BOMs, it is queued for the ERP
  with its operations.
"""
from __future__ import annotations

import io
import json
import logging

from . import access, boms, db
from .web import bad, conflict, not_found

log = logging.getLogger("kosha.mbom")

STATE_LABEL = {"draft": "Draft", "released": "Released"}
MAX_OPS = 200
MAX_LINES = 2000


def _num(value, label: str, minimum: float = 0.0, allow_equal: bool = True) -> float:
    try:
        v = float(value)
    except (TypeError, ValueError):
        raise bad(f"{label} must be a number.") from None
    if v < minimum or (not allow_equal and v == minimum):
        raise bad(f"{label} must be {'at least' if allow_equal else 'more than'} {minimum:g}.")
    return v


def _fmt(q: float) -> str:
    return str(int(q)) if float(q).is_integer() else f"{q:g}"


# ------------------------------------------------------------------ reading --

def mbom_or_404(mbom_id: int) -> dict:
    m = db.one("SELECT m.*, i.part_number, i.description, i.revision AS item_revision FROM mboms m "
               "JOIN items i ON i.id = m.item_id WHERE m.id = ?", (int(mbom_id),))
    if not m:
        raise not_found("Manufacturing BOM not found")
    access.check_item(m["item_id"])
    m["flatten"] = json.loads(m.pop("flatten_json") or "[]")
    return m


def lines(mbom_id: int) -> list[dict]:
    return db.rows(
        "SELECT ml.id, ml.line_no, ml.qty, ml.uom, ml.op_no, ml.note, ml.child_item_id AS item_id, i.part_number, "
        "i.description, i.material, i.make_buy, i.uom AS item_uom, i.revision, i.state, i.item_type "
        "FROM mbom_lines ml JOIN items i ON i.id = ml.child_item_id WHERE ml.mbom_id = ? ORDER BY ml.line_no",
        (mbom_id,))


def ops(mbom_id: int) -> list[dict]:
    return db.rows("SELECT op_no, work_centre, description, setup_min, run_min, instructions FROM mbom_ops "
                   "WHERE mbom_id = ? ORDER BY op_no", (mbom_id,))


def listing(item_id: int | None = None, q: str = "") -> list[dict]:
    sql = ("SELECT m.id, m.item_id, m.revision, m.state, m.updated_at, m.released_at, i.part_number, i.description, "
           "(SELECT COUNT(*) FROM mbom_lines WHERE mbom_id = m.id) AS lines, "
           "(SELECT COUNT(*) FROM mbom_ops WHERE mbom_id = m.id) AS ops "
           "FROM mboms m JOIN items i ON i.id = m.item_id WHERE 1 = 1" + access.scope().items("i"))
    args: list = []
    if item_id:
        sql += " AND m.item_id = ?"
        args.append(int(item_id))
    if q.strip():
        sql += " AND (i.part_number LIKE ? OR i.description LIKE ?)"
        args += [f"%{q.strip()}%"] * 2
    return db.rows(sql + " ORDER BY m.updated_at DESC, m.id DESC LIMIT 500", args)


def _ebom(m: dict) -> dict | None:
    """The engineering BOM this MBOM follows: the same revision of the same assembly."""
    b = db.one("SELECT id, revision, state FROM boms WHERE item_id = ? AND revision = ?", (m["item_id"], m["revision"]))
    return b or (db.one("SELECT id, revision, state FROM boms WHERE id = ?", (m["ebom_id"],)) if m.get("ebom_id") else None)


def _child_bom(item_id: int, revision: str | None) -> int | None:
    b = boms.bom_for_item(item_id, revision)
    return b["id"] if b else None


def _leaves(item_lines: list[tuple[int, float, str | None]]) -> dict[int, float]:
    """What a list of (item, qty, revision) lines finally consumes: assemblies with a BOM are opened down to the
    parts without one, with quantities multiplied out."""
    out: dict[int, float] = {}
    for item_id, qty, rev in item_lines:
        sub = _child_bom(item_id, rev)
        if not sub:
            out[item_id] = out.get(item_id, 0.0) + qty
            continue
        for r in boms.explode(sub):
            if not r["has_children"] and not r["cycle"]:
                out[r["item_id"]] = out.get(r["item_id"], 0.0) + r["ext_qty"] * qty
    return out


def _ebom_items(ebom_id: int) -> set[int]:
    return {r["item_id"] for r in boms.explode(ebom_id)}


def checks(mbom_id: int) -> list[dict]:
    """The MBOM against its EBOM and its own process plan, worst first."""
    m = mbom_or_404(mbom_id)
    issues: list[dict] = []
    ls = lines(mbom_id)
    plan = ops(mbom_id)
    eb = _ebom(m)

    def add(severity, part_number, message):
        issues.append({"severity": severity, "part_number": part_number, "message": message})

    if not eb:
        add("error", m["part_number"], f"There is no engineering BOM for revision {m['revision']} to follow.")
    else:
        need = _leaves([(ln["item_id"], ln["qty"], ln["revision"]) for ln in boms.lines(eb["id"])])
        have = _leaves([(ln["item_id"], ln["qty"], ln["revision"]) for ln in ls])
        names = {r["id"]: r["part_number"] for r in db.rows(
            "SELECT id, part_number FROM items WHERE id IN (%s)" % ",".join(str(int(i)) for i in (set(need) | set(have)) or {0}))}
        in_ebom = _ebom_items(eb["id"])
        for item_id, q in sorted(need.items(), key=lambda kv: names.get(kv[0], "")):
            pn = names.get(item_id, "?")
            got = have.get(item_id, 0.0)
            if got <= 1e-9:
                add("error", pn, f"The engineering BOM needs {_fmt(round(q, 6))}, but the manufacturing BOM does not "
                                 "issue it at any operation.")
            elif got < q - 1e-9:
                add("error", pn, f"The engineering BOM needs {_fmt(round(q, 6))}; the manufacturing BOM issues only "
                                 f"{_fmt(round(got, 6))}.")
            elif got > q + 1e-9:
                add("warning", pn, f"The manufacturing BOM issues {_fmt(round(got, 6))}; the engineering BOM needs "
                                   f"{_fmt(round(q, 6))}. Fine for a scrap allowance; otherwise correct it.")
        for ln in ls:
            if ln["item_id"] not in in_ebom:
                add("info", ln["part_number"], "Added for manufacturing (not in the engineering BOM).")
        if eb["state"] != "released":
            add("warning", m["part_number"], f"The engineering BOM revision {eb['revision']} is not released yet; "
                                              "the manufacturing BOM can be released after it.")
    newer = db.one("SELECT revision FROM boms WHERE item_id = ? AND state = 'released' AND id > ? "
                   "ORDER BY id DESC LIMIT 1", (m["item_id"], (eb or {}).get("id") or 0))
    if newer and m["state"] == "released":
        add("warning", m["part_number"], f"Engineering revision {newer['revision']} has been released since. Make "
                                          f"the manufacturing BOM for revision {newer['revision']}.")
    if not plan:
        add("warning", m["part_number"], "No process plan yet: add the operations the workshop follows.")
    else:
        for ln in ls:
            if ln["op_no"] is None:
                add("warning", ln["part_number"], "Not issued at an operation. Choose the operation that uses it.")
        used = {ln["op_no"] for ln in ls}
        for op in plan:
            if not op["work_centre"]:
                add("warning", "", f"Operation {op['op_no']} has no work centre.")
            if op["run_min"] <= 0:
                add("info", "", f"Operation {op['op_no']} has no run time yet.")
            if op["op_no"] not in used:
                add("info", "", f"No parts are issued at operation {op['op_no']}.")
    if not ls:
        add("error", m["part_number"], "The manufacturing BOM is empty.")
    order = {"error": 0, "warning": 1, "info": 2}
    issues.sort(key=lambda i: order[i["severity"]])
    return issues


def _row_origin(m: dict, ls: list[dict]) -> list[dict]:
    eb = _ebom(m)
    in_ebom = _ebom_items(eb["id"]) if eb else set()
    for ln in ls:
        ln["origin"] = "ebom" if ln["item_id"] in in_ebom else "added"
    return ls


def detail(mbom_id: int) -> dict:
    m = mbom_or_404(mbom_id)
    eb = _ebom(m)
    ls = _row_origin(m, lines(mbom_id))
    plan = ops(mbom_id)
    for op in plan:
        op["parts"] = sum(1 for ln in ls if ln["op_no"] == op["op_no"])
    issues = checks(mbom_id)
    errors = [i for i in issues if i["severity"] == "error"]
    release_why = ""
    if m["state"] == "released":
        release_why = "Already released."
    elif not eb or eb["state"] != "released":
        release_why = f"Release the engineering BOM revision {m['revision']} first."
    elif errors:
        release_why = "Fix the errors under Checks first."
    flattened = db.rows("SELECT id, part_number, description FROM items WHERE id IN (%s)"
                        % ",".join(str(int(i)) for i in m["flatten"] or [0]))
    newer = db.one("SELECT revision FROM boms WHERE item_id = ? AND id > ? ORDER BY id DESC LIMIT 1",
                   (m["item_id"], (eb or {}).get("id") or 0))
    return {
        "mbom": {**m, "state_label": STATE_LABEL.get(m["state"], m["state"])}, "lines": ls, "ops": plan,
        "issues": issues, "ebom": eb, "flattened": flattened,
        "revisions": db.rows("SELECT id, revision, state FROM mboms WHERE item_id = ? ORDER BY id", (m["item_id"],)),
        "work_centres": [r["work_centre"] for r in db.rows(
            "SELECT DISTINCT work_centre FROM mbom_ops WHERE work_centre <> '' ORDER BY work_centre LIMIT 200")],
        "sub_assemblies": [{"item_id": ln["item_id"], "part_number": ln["part_number"]} for ln in ls
                           if _child_bom(ln["item_id"], ln["revision"])],
        "totals": {"setup_min": round(sum(o["setup_min"] for o in plan), 2),
                   "run_min": round(sum(o["run_min"] for o in plan), 2)},
        "can_release": {"ok": not release_why, "why": release_why},
        "newer_ebom": newer["revision"] if newer else None,
    }


# ------------------------------------------------------------------ writing --

def _from_ebom(ebom_id: int, flatten: set[int]) -> list[dict]:
    """The EBOM's lines, with the flattened sub-assemblies replaced by their own lines (and theirs, when also
    flattened), merged by part."""
    merged: dict[int, dict] = {}

    def walk(bom_id: int, mult: float, path: tuple):
        for ln in boms.lines(bom_id):
            q = ln["qty"] * mult
            sub = _child_bom(ln["item_id"], ln["revision"]) if ln["item_id"] in flatten else None
            if sub and ln["item_id"] not in path:
                walk(sub, q, path + (ln["item_id"],))
                continue
            row = merged.setdefault(ln["item_id"], {"item_id": ln["item_id"], "qty": 0.0, "uom": ln["uom"] or ""})
            row["qty"] = round(row["qty"] + q, 6)

    walk(ebom_id, 1.0, ())
    return list(merged.values())


def _write_lines(mbom_id: int, new_lines: list[dict]) -> None:
    db.execute("DELETE FROM mbom_lines WHERE mbom_id = ?", (mbom_id,))
    for k, ln in enumerate(new_lines, start=1):
        db.insert("mbom_lines", {"mbom_id": mbom_id, "line_no": k, "child_item_id": ln["item_id"], "qty": ln["qty"],
                                 "uom": ln.get("uom") or "", "op_no": ln.get("op_no"), "note": ln.get("note") or ""})


def _write_ops(mbom_id: int, plan: list[dict]) -> None:
    db.execute("DELETE FROM mbom_ops WHERE mbom_id = ?", (mbom_id,))
    for op in plan:
        db.insert("mbom_ops", {"mbom_id": mbom_id, **{k: op[k] for k in (
            "op_no", "work_centre", "description", "setup_min", "run_min", "instructions")}})


def create(user: dict, item_id: int, revision: str | None = None) -> dict:
    """The MBOM of an assembly's engineering revision (the newest when none is given), made from its EBOM and carrying
    the previous MBOM's process plan forward. Returns {"mbom_id", "created", "carried_from"}."""
    access.check_item(item_id)
    item = db.one("SELECT id, part_number FROM items WHERE id = ?", (int(item_id),))
    if not item:
        raise not_found("Item not found")
    if revision:
        eb = db.one("SELECT id, revision FROM boms WHERE item_id = ? AND revision = ?", (item_id, revision))
    else:
        eb = db.one("SELECT id, revision FROM boms WHERE item_id = ? ORDER BY id DESC LIMIT 1", (item_id,))
    if not eb:
        raise conflict(f"{item['part_number']} has no engineering BOM{' revision ' + revision if revision else ''} "
                       "yet. Make the engineering BOM first; the manufacturing BOM is made from it.", "no_ebom")
    have = db.one("SELECT id FROM mboms WHERE item_id = ? AND revision = ?", (item_id, eb["revision"]))
    if have:
        return {"mbom_id": have["id"], "created": False, "carried_from": None}
    prev = db.one("SELECT * FROM mboms WHERE item_id = ? ORDER BY id DESC LIMIT 1", (item_id,))
    flatten = set(json.loads(prev["flatten_json"] or "[]")) if prev else set()
    new_lines = _from_ebom(eb["id"], flatten)
    plan: list[dict] = []
    if prev:
        plan = ops(prev["id"])
        before = {ln["item_id"]: ln for ln in lines(prev["id"])}
        in_prev_ebom = _ebom_items(prev["ebom_id"]) if prev["ebom_id"] else set()
        for ln in new_lines:
            if ln["item_id"] in before:
                ln["op_no"] = before[ln["item_id"]]["op_no"]
                ln["note"] = before[ln["item_id"]]["note"]
        present = {ln["item_id"] for ln in new_lines}
        for item_id_, old in before.items():   # parts the workshop added: paint, grease, thread-locker
            if item_id_ not in in_prev_ebom and item_id_ not in present:
                new_lines.append({"item_id": item_id_, "qty": old["qty"], "uom": old["uom"], "op_no": old["op_no"],
                                  "note": old["note"]})
    now = db.now()
    with db.tx():
        mid = db.insert("mboms", {"item_id": item_id, "revision": eb["revision"], "ebom_id": eb["id"], "state": "draft",
                                  "flatten_json": json.dumps(sorted(flatten)), "created_by": user["id"],
                                  "created_at": now, "updated_at": now})
        _write_lines(mid, new_lines)
        _write_ops(mid, plan)
        db.audit(user["id"], "mbom_create", "mbom", mid, f"{item['part_number']} rev {eb['revision']}"
                 + (f", carried from rev {prev['revision']}" if prev else ""))
    return {"mbom_id": mid, "created": True, "carried_from": prev["revision"] if prev else None}


def _check_editable(m: dict) -> None:
    if m["state"] == "released":
        raise conflict(f"The manufacturing BOM for revision {m['revision']} is released and locked. Make the one for "
                       "the next engineering revision to change it.", "mbom_released")


def _clean_ops(raw) -> list[dict]:
    if not isinstance(raw, list):
        raise bad("Send the operations as a list.")
    if len(raw) > MAX_OPS:
        raise bad(f"A process plan can have at most {MAX_OPS} operations.")
    out, seen = [], set()
    for k, op in enumerate(raw, start=1):
        if not isinstance(op, dict):
            raise bad(f"Operation {k} is not valid.")
        try:
            no = int(op.get("op_no"))
        except (TypeError, ValueError):
            raise bad(f"Operation {k} needs a number, for example 10, 20, 30.") from None
        if no <= 0 or no > 99999:
            raise bad(f"Operation numbers run from 1 to 99999; {no} is outside.")
        if no in seen:
            raise bad(f"Operation {no} is listed twice. Give each operation its own number.")
        seen.add(no)
        desc = " ".join(str(op.get("description") or "").split())[:200]
        if not desc:
            raise bad(f"Say what happens at operation {no}, for example \"Fit bearings on the shaft\".")
        out.append({"op_no": no, "work_centre": " ".join(str(op.get("work_centre") or "").split())[:80],
                    "description": desc,
                    "setup_min": round(_num(op.get("setup_min") or 0, f"The set-up time of operation {no}"), 2),
                    "run_min": round(_num(op.get("run_min") or 0, f"The run time of operation {no}"), 2),
                    "instructions": str(op.get("instructions") or "").strip()[:4000]})
    return sorted(out, key=lambda o: o["op_no"])


def _clean_lines(user: dict, raw, op_numbers: set[int], self_item: int) -> list[dict]:
    if not isinstance(raw, list):
        raise bad("Send the parts as a list.")
    if len(raw) > MAX_LINES:
        raise bad(f"A manufacturing BOM can have at most {MAX_LINES} lines.")
    out: list[dict] = []
    seen: dict[tuple, int] = {}
    for k, ln in enumerate(raw, start=1):
        if not isinstance(ln, dict):
            raise bad(f"Line {k} is not valid.")
        item_id = ln.get("item_id")
        if not item_id:
            pn = str(ln.get("part_number") or "").strip()
            if not pn:
                raise bad(f"Line {k} has no part.")
            item_id = boms.ensure_item(pn, ln.get("description") or "", user["id"], ln.get("uom") or "")
        item_id = int(item_id)
        if not db.scalar("SELECT 1 FROM items WHERE id = ?", (item_id,)):
            raise bad(f"Line {k}: the part does not exist.")
        if item_id == self_item:
            raise bad("An assembly cannot contain itself.")
        qty = _num(ln.get("qty", 1), f"The quantity on line {k}", 0, allow_equal=False)
        op_no = ln.get("op_no")
        if op_no in (None, ""):
            op_no = None
        else:
            try:
                op_no = int(op_no)
            except (TypeError, ValueError):
                raise bad(f"Line {k}: the operation must be one of the plan's operation numbers.") from None
            if op_no not in op_numbers:
                raise bad(f"Line {k} is issued at operation {op_no}, which is not in the process plan.")
        key = (item_id, op_no)
        if key in seen:                       # the same part twice at the same operation: one line
            out[seen[key]]["qty"] = round(out[seen[key]]["qty"] + qty, 6)
            continue
        seen[key] = len(out)
        out.append({"item_id": item_id, "qty": round(qty, 6), "uom": str(ln.get("uom") or "")[:20], "op_no": op_no,
                    "note": str(ln.get("note") or "").strip()[:500]})
    return out


def save_plan(user: dict, mbom_id: int, body: dict) -> None:
    """Replace the parts and/or the operations. Both are checked together, so a part never points at a missing
    operation; nothing is saved when anything is wrong."""
    m = mbom_or_404(mbom_id)
    _check_editable(m)
    plan = _clean_ops(body["ops"]) if "ops" in body else ops(mbom_id)
    numbers = {o["op_no"] for o in plan}
    with db.tx():
        if "lines" in body:
            new_lines = _clean_lines(user, body["lines"], numbers, m["item_id"])
        else:
            new_lines = [dict(ln) for ln in lines(mbom_id)]
            orphans = [ln["part_number"] for ln in new_lines if ln["op_no"] is not None and ln["op_no"] not in numbers]
            if orphans:
                raise conflict("These parts are issued at an operation you removed: " + ", ".join(orphans[:10])
                               + ". Move them to another operation first.", "op_in_use")
        if "ops" in body:
            _write_ops(mbom_id, plan)
        _write_lines(mbom_id, new_lines)
        if "note" in body:
            db.update("mboms", mbom_id, {"note": str(body.get("note") or "").strip()[:2000]})
        db.update("mboms", mbom_id, {"updated_at": db.now()})
        db.audit(user["id"], "mbom_edit", "mbom", mbom_id, f"{len(new_lines)} parts, {len(plan)} operations")


def flatten(user: dict, mbom_id: int, item_id: int, on: bool = True) -> None:
    """Build a sub-assembly in one go with its parent (``on``): its line is replaced by its own parts, issued at the
    same operation. ``on=False`` puts the sub-assembly back as one line, at the operation most of its parts used."""
    m = mbom_or_404(mbom_id)
    _check_editable(m)
    item_id = int(item_id)
    pn = db.scalar("SELECT part_number FROM items WHERE id = ?", (item_id,)) or "?"
    current = lines(mbom_id)
    flat = set(m["flatten"])
    if on:
        line = next((ln for ln in current if ln["item_id"] == item_id), None)
        sub = _child_bom(item_id, line["revision"]) if line else None
        if not line or not sub:
            raise bad(f"{pn} is not a sub-assembly with its own BOM on this manufacturing BOM.")
        new_lines = []
        for ln in current:
            if ln["id"] != line["id"]:
                new_lines.append(dict(ln))
                continue
            for child in boms.lines(sub):
                new_lines.append({"item_id": child["item_id"], "qty": round(child["qty"] * ln["qty"], 6),
                                  "uom": child["uom"] or "", "op_no": ln["op_no"], "note": ""})
        flat.add(item_id)
    else:
        if item_id not in flat:
            raise bad(f"{pn} is not flattened.")
        eb = _ebom(m)
        if not eb:
            raise conflict("There is no engineering BOM to rebuild from.")
        flat.discard(item_id)
        target = {ln["item_id"]: ln for ln in _from_ebom(eb["id"], flat)}
        before = {ln["item_id"]: ln for ln in current}
        sub = _child_bom(item_id, None)
        kids = {c["item_id"] for c in boms.lines(sub)} if sub else set()
        ops_used = [before[k]["op_no"] for k in kids if k in before and before[k]["op_no"] is not None]
        new_lines = []
        for iid, ln in target.items():
            old = before.get(iid)
            new_lines.append({"item_id": iid, "qty": ln["qty"], "uom": ln["uom"],
                              "op_no": (max(set(ops_used), key=ops_used.count) if ops_used else None) if iid == item_id
                              else (old["op_no"] if old else None), "note": old["note"] if old else ""})
        in_ebom = _ebom_items(eb["id"])
        new_lines += [dict(ln) for ln in current if ln["item_id"] not in in_ebom]
    merged: dict[tuple, dict] = {}
    for ln in new_lines:
        key = (ln["item_id"], ln.get("op_no"))
        if key in merged:
            merged[key]["qty"] = round(merged[key]["qty"] + ln["qty"], 6)
        else:
            merged[key] = {k: ln.get(k) for k in ("item_id", "qty", "uom", "op_no", "note")}
    with db.tx():
        _write_lines(mbom_id, list(merged.values()))
        db.update("mboms", mbom_id, {"flatten_json": json.dumps(sorted(flat)), "updated_at": db.now()})
        db.audit(user["id"], "mbom_flatten" if on else "mbom_unflatten", "mbom", mbom_id, pn)


def rebuild(user: dict, mbom_id: int) -> None:
    """Make the parts again from the engineering BOM (after it changed), keeping each part's operation and note,
    the flattened sub-assemblies and the parts added for manufacturing."""
    m = mbom_or_404(mbom_id)
    _check_editable(m)
    eb = _ebom(m)
    if not eb:
        raise conflict(f"There is no engineering BOM for revision {m['revision']} to rebuild from.", "no_ebom")
    before = lines(mbom_id)
    by_item = {ln["item_id"]: ln for ln in before}
    new_lines = _from_ebom(eb["id"], set(m["flatten"]))
    for ln in new_lines:
        old = by_item.get(ln["item_id"])
        if old:
            ln["op_no"], ln["note"] = old["op_no"], old["note"]
    in_ebom = _ebom_items(eb["id"])
    present = {ln["item_id"] for ln in new_lines}
    new_lines += [{"item_id": ln["item_id"], "qty": ln["qty"], "uom": ln["uom"], "op_no": ln["op_no"],
                   "note": ln["note"]} for ln in before if ln["item_id"] not in in_ebom and ln["item_id"] not in present]
    with db.tx():
        _write_lines(mbom_id, new_lines)
        db.update("mboms", mbom_id, {"ebom_id": eb["id"], "updated_at": db.now()})
        db.audit(user["id"], "mbom_rebuild", "mbom", mbom_id, f"from engineering BOM rev {eb['revision']}")


def release(user: dict, mbom_id: int) -> dict:
    m = mbom_or_404(mbom_id)
    if m["state"] == "released":
        return m
    d = detail(mbom_id)
    if not d["can_release"]["ok"]:
        errors = [i for i in d["issues"] if i["severity"] == "error"]
        raise conflict(d["can_release"]["why"] + ("" if not errors else " " + "; ".join(
            f"{e['part_number']}: {e['message']}" for e in errors[:5])), "mbom_not_ready")
    now = db.now()
    with db.tx():
        db.update("mboms", mbom_id, {"state": "released", "released_at": now, "released_by": user["id"],
                                     "updated_at": now})
        db.audit(user["id"], "mbom_release", "mbom", mbom_id, f"{m['part_number']} rev {m['revision']}")
    from . import erp
    if erp.settings().get("bom_source") == "mbom" and erp.settings().get("auto_push_on_release", True):
        erp.enqueue("item", m["item_id"], user["id"], f"Manufacturing BOM {m['part_number']} rev {m['revision']}")
        erp.enqueue("mbom", mbom_id, user["id"], f"Manufacturing BOM {m['part_number']} rev {m['revision']}")
    try:
        from .addins import registry
        registry.emit("mbom.released", {"mbom_id": mbom_id, "item_id": m["item_id"], "part_number": m["part_number"],
                                        "revision": m["revision"], "user_id": user["id"]})
    except Exception:  # pragma: no cover - add-ins never undo a release
        log.exception("Could not hand mbom.released for %s to add-ins", m["part_number"])
    return mbom_or_404(mbom_id)


def released_for(item_id: int) -> dict | None:
    """The newest released MBOM of an assembly (what the ERP gets when it takes manufacturing BOMs)."""
    return db.one("SELECT id, revision FROM mboms WHERE item_id = ? AND state = 'released' ORDER BY id DESC LIMIT 1",
                  (item_id,))


# ------------------------------------------------------------------- export --

def export_xlsx(mbom_id: int) -> tuple[bytes, str]:
    import openpyxl
    from openpyxl.styles import Alignment, Font

    d = detail(mbom_id)
    m = d["mbom"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Manufacturing BOM"
    ws.append([f"{m['part_number']} {m['description']}", f"Revision {m['revision']}", m["state_label"]])
    ws.append([])
    ws.append(["Operation", "Part number", "Description", "Qty", "Unit", "Material", "Make/Buy", "From", "Note"])
    for c in ws[3]:
        c.font = Font(bold=True)
    for ln in sorted(d["lines"], key=lambda r: (r["op_no"] is None, r["op_no"] or 0, r["line_no"])):
        ws.append([ln["op_no"] or "", ln["part_number"], ln["description"], ln["qty"], ln["uom"] or ln["item_uom"],
                   ln["material"], ln["make_buy"], "Engineering BOM" if ln["origin"] == "ebom" else "Added",
                   ln["note"]])
    for col, width in zip("ABCDEFGHI", (10, 20, 40, 8, 8, 16, 10, 16, 30)):
        ws.column_dimensions[col].width = width
    pp = wb.create_sheet("Process plan")
    pp.append([f"Process plan {m['part_number']} rev {m['revision']}"])
    pp.append([])
    pp.append(["Operation", "Work centre", "What happens", "Set-up (min)", "Run (min)", "Parts issued",
               "Instructions"])
    for c in pp[3]:
        c.font = Font(bold=True)
    for op in d["ops"]:
        parts = ", ".join(f"{ln['part_number']} × {_fmt(ln['qty'])}" for ln in d["lines"] if ln["op_no"] == op["op_no"])
        pp.append([op["op_no"], op["work_centre"], op["description"], op["setup_min"], op["run_min"], parts,
                   op["instructions"]])
        for c in pp[pp.max_row]:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    pp.append(["", "", "Total", d["totals"]["setup_min"], d["totals"]["run_min"]])
    for col, width in zip("ABCDEFG", (10, 18, 36, 12, 10, 40, 50)):
        pp.column_dimensions[col].width = width
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), f"MBOM_{m['part_number']}_rev{m['revision']}.xlsx"
